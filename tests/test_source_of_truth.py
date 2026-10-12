"""The decoded layouts hold on the real files, and the saved workbook matches the backups."""
import os
import unittest

import openpyxl

import _paths
import extract_vanilla


class DecodedLayouts(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.problems = []
        cls.skills = {(s['category'], s['name']): s for s in extract_vanilla.read_all_skills(cls.problems)}

    def test_header_sizes_and_effect_counts_hold_for_every_skill(self):
        self.assertEqual(self.problems, [])
        self.assertEqual(len(extract_vanilla.read_all_skills([])), 312)

    def test_enemy_attack_effect_counts_hold(self):
        problems = []
        rows, _ = extract_vanilla.build_enemy_attack_rows(problems)
        self.assertEqual(problems, [])
        self.assertEqual(len(rows), 2375)

    def test_equipment_enchant_counts_hold(self):
        problems = []
        _, names = extract_vanilla.build_enchant_rows()
        _, ranks = extract_vanilla.build_weapon_rank_rows()
        rows = extract_vanilla.build_equipment_rows(names, ranks, problems)
        self.assertEqual(problems, [])
        self.assertEqual(len(rows), 551)

    def test_values_cross_checked_with_the_guides(self):
        """Figures confirmed against GameFAQs; they pin the meaning of the offsets."""
        halving = self.skills[('Technique', 'Halving Slice')]
        self.assertEqual(halving['levels'], 10)
        self.assertEqual(halving['arrays']['tp_cost'][0], [38, 42, 46, 50, 54, 58, 62, 66, 70, 74])
        power = next(s for s in halving['slots'] if s['id'] == 26)['values']
        self.assertEqual((round(100 * (1 + power[0])), round(100 * (1 + power[-1]))), (80, 110))

        megaton = self.skills[('Technique', 'Megaton Press')]
        self.assertEqual(next(s for s in megaton['slots'] if s['id'] == 137)['values'], [2, 2, 2, 3, 3, 3, 4, 4, 5, 5])

        flare = self.skills[('Magic', 'Crimson Flare')]
        self.assertEqual(flare['levels'], 5)
        self.assertEqual(next(s for s in flare['slots'] if s['id'] == 1)['values'][:5], [620, 760, 920, 1100, 1300])

    def test_vanilla_weapon_delays(self):
        _, ranks = extract_vanilla.build_weapon_rank_rows()
        self.assertEqual([ranks[i][0] for i in range(12)],
                         ['A+', 'A', 'B+', 'B', 'C+', 'C', 'D+', 'D', 'E+', 'E', 'F+', 'F'])
        self.assertEqual([ranks[i][1] for i in range(12)],
                         [1.0, 1.2, 1.4, 1.6, 1.8, 2.0, 2.15, 2.3, 2.45, 2.6, 2.75, 2.9])


class SavedWorkbook(unittest.TestCase):
    def test_workbook_is_not_stale(self):
        path = os.path.join(_paths.TABLES_DIR, 'vanilla_source_of_truth.xlsx')
        self.assertTrue(os.path.exists(path), 'run tools/extract_vanilla.py')
        wb = openpyxl.load_workbook(path)
        saved = [list(r) for r in wb['SkillLevels'].iter_rows(min_row=2, values_only=True)]
        skills = extract_vanilla.read_all_skills([])
        _, fresh = extract_vanilla.build_skill_sheets(skills)
        normalised = [[('' if v is None else v) for v in row] for row in fresh]
        saved = [[('' if v is None else v) for v in row] for row in saved]
        self.assertEqual(saved, normalised)


if __name__ == '__main__':
    unittest.main()
