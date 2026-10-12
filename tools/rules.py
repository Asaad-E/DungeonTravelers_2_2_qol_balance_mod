"""
The editable rules of the QoL text: data_tables/qol_rules.xlsx.

Rule of the user: nothing that reaches the screen is written in the tools. Every wording and every
figure that is not read from the game files lives in that workbook, one sheet per table (its About
sheet lists them). The tools read it through this module; to change a word, a bound or a figure, edit
the workbook and run the generators again.
"""
import os
import functools
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RULES_PATH = os.path.join(BASE_DIR, 'data_tables', 'qol_rules.xlsx')


@functools.lru_cache(maxsize=None)
def _book():
    if not os.path.exists(RULES_PATH):
        raise SystemExit(f"Error: {RULES_PATH} not found. It holds the rules of the QoL text and is not generated.")
    return openpyxl.load_workbook(RULES_PATH, data_only=True)


@functools.lru_cache(maxsize=None)
def rows(sheet):
    """Rows of a sheet as dicts by header. Empty cells are '' and rows without a first cell are skipped."""
    ws = _book()[sheet]
    header = [c.value for c in ws[1]]
    found = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None or str(r[0]).strip() == '':
            continue
        found.append({h: ('' if v is None else v) for h, v in zip(header, r)})
    return tuple(found)


def setting(key):
    for r in rows('Settings'):
        if r['Key'] == key:
            return r['Value']
    raise KeyError(f"qol_rules.xlsx/Settings has no key '{key}'")


def pairs(sheet, key, value):
    """A two-column sheet as an ordered dict."""
    return {r[key]: r[value] for r in rows(sheet)}


def id_set(text):
    """'161, 162' -> {161, 162}"""
    return {int(part) for part in str(text).replace(';', ',').split(',') if part.strip()}


def yes(value):
    return str(value).strip().lower() == 'yes'
