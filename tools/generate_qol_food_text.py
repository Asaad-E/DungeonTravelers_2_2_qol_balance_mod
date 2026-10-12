"""
QoL text of the ice cream and ramen menus: proposes it and keeps it in a workbook edited by hand.

    python tools/generate_qol_food_text.py

Workbook: data_tables/qol_food_text.xlsx
  * A meal has two lines: a flavour line (never touched) and an effect line.
  * In the effect line the vague amount gives way to the stored one, and the rest of the sentence stays:
        "Heal a little HP and TP."  ->  "Heal 30% HP and 20% TP."
  * The amounts are effects 62 (HP), 63 (TP) and 64 (HP and TP), stored in tenths of the maximum: the
    DT2-2 wiki lists every meal with the same figure ("HP 2/10" for a stored 2).
"""
import os
import sys
import collections
import openpyxl
from pck_handler import PckArchive, FoodMenuHelper
from text_checks import food_text, check_food_text, food_effects

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKUP_PATH = os.path.join(BASE_DIR, 'backups', 'FoodMenuData.pck')
OUTPUT_PATH = os.path.join(BASE_DIR, 'data_tables', 'qol_food_text.xlsx')
HEADERS = ['Filename', 'Meal', 'Flavour_Line', 'Vanilla_Effect', 'QoL_Effect', 'Edited', 'Effects']


def load_edited():
    if not os.path.exists(OUTPUT_PATH):
        return {}
    ws = openpyxl.load_workbook(OUTPUT_PATH)['Meals']
    header = [c.value for c in ws[1]]
    fn, text, edited = header.index('Filename'), header.index('QoL_Effect'), header.index('Edited')
    return {r[fn]: r[text] or '' for r in ws.iter_rows(min_row=2, values_only=True)
            if str(r[edited] or '').strip().lower() == 'yes'}


def main():
    edited, stats, rows = load_edited(), collections.Counter(), []
    for fn, b in PckArchive(BACKUP_PATH).files:
        meal = FoodMenuHelper.read_food_menu(b)
        effect = edited.get(fn) or food_text(b, meal['effect'])
        problems, _ = check_food_text(b, effect, meal['effect'])
        if problems:
            raise SystemExit(f"{fn} ({meal['name']}): {problems}")
        stats['exact amount in place of the vague one' if effect != meal['effect'] else 'kept as written'] += 1
        rows.append([fn, meal['name'], meal['flavor'], meal['effect'], effect, 'yes' if fn in edited else '',
                     ', '.join(f'{i}:{v:g}' for i, v in food_effects(b))])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Meals'
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    for col, width in zip('ABCDEFG', (22, 24, 60, 60, 60, 8, 40)):
        ws.column_dimensions[col].width = width
    wb.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}: {len(rows)} meals")
    for key in sorted(stats):
        print(f"    {stats[key]:4d}  {key}")
    for r in rows:
        if r[3] != r[4]:
            print(f"        {r[1]}: {r[4]}")


if __name__ == '__main__':
    sys.exit(main())
