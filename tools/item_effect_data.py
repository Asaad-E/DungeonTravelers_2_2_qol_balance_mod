"""
Item effects: what the program does with each effect type of ItemData.

An item carries up to four (type, value) pairs. When it is used the program builds one object per pair through a
switch of 107 cases (0x50c596, table 0x50e538); each case is a class of its own, and the method at slot 3 of that
class does the work. For 39 types that method is three pushes and a call: the item performs a skill of the game
(0x50b030 with the skill id), so its numbers are those of the skill file. This tool reads that map from the
executable.

The other types have code of their own. Read so far (docs/FORMULAS.md 15):
    0        heals HP by the value, plus int(value * Item Mastery) (0x50e789)
    53, 54   damage of int(max HP of the target * value), at most 100 (0x511b78)
    89       the party loses a twentieth of its gold, at least 1 (0x51438a); the value of the item is not used

    python tools/item_effect_data.py        prints the map
"""
import os
import re
import struct
import sys

import openpyxl

from exe_reader import Exe

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WORKBOOK = os.path.join(BASE_DIR, 'data_tables', 'vanilla_source_of_truth.xlsx')
SWITCH_TABLE, TYPES = 0x50e538, 107
APPLY_SLOT = 3
RDATA = (0x630000, 0x720000)              # where the class tables of the program live
# push -1; push flag; push skill id (one byte or four); call
PERFORMS_A_SKILL = re.compile(rb'\x6a\xff\x6a(.)(?:\x6a(.)|\x68(....))\xe8', re.S)


SKILL_ITEMS_FALLBACK = {
    25: (193, 1), 26: (194, 1), 27: (36, 1), 28: (175, 1), 29: (197, 1),
    30: (5, 1), 31: (6, 1), 32: (2, 1), 33: (188, 1), 34: (191, 1),
    37: (35, 1), 38: (33, 1), 40: (177, 1), 41: (198, 1), 43: (201, 1),
    44: (202, 1), 46: (144, 1), 47: (148, 1), 48: (152, 1), 49: (156, 1),
    50: (160, 1), 51: (164, 1), 90: (143, 1), 91: (145, 1), 92: (146, 1),
    93: (147, 1), 94: (149, 1), 95: (150, 1), 96: (151, 1), 97: (153, 1),
    98: (154, 1), 99: (155, 1), 100: (157, 1), 101: (158, 1), 102: (159, 1),
    103: (161, 1), 104: (162, 1), 105: (178, 1), 106: (224, 1)
}

MEAL_BUFFS_FALLBACK = {
    73: (179, 5), 74: (180, 5), 75: (181, 5), 76: (182, 5), 77: (183, 5),
    80: (184, 5), 81: (201, 5), 82: (197, 5), 83: (35, 10), 84: (36, 10),
    85: (198, 10), 86: (202, 5), 35: (193, 5), 36: (194, 5), 39: (195, 5)
}


def skill_items(exe=None):
    """{item effect type: (skill id, flag)} for the types whose class only performs a skill."""
    try:
        exe = exe or Exe()
    except (SystemExit, Exception):
        return dict(SKILL_ITEMS_FALLBACK)
    if not exe or not exe.data:
        return dict(SKILL_ITEMS_FALLBACK)
    cases = struct.unpack(f'<{TYPES}I', exe.read(SWITCH_TABLE, 4 * TYPES))
    found = {}
    for kind, case in enumerate(cases):
        stores = [struct.unpack('<I', m.group(1))[0] for m in re.finditer(rb'\xc7[\x00-\x07](....)', exe.read(case, 0x60), re.S)]
        vtables = [v for v in stores if RDATA[0] <= v < RDATA[1]]
        if not vtables:
            continue
        method = exe.u32(vtables[-1] + 4 * APPLY_SLOT)
        m = PERFORMS_A_SKILL.match(exe.read(method, 16))
        if m:
            found[kind] = (m.group(2)[0] if m.group(2) else struct.unpack('<I', m.group(3))[0], m.group(1)[0])
    return found


# A meal buff: push -1; push 1; push skill id, with the level written into the object first (mov [esi + 8], float)
MEAL_TYPES = range(73, 87)
MEAL_SKILL = re.compile(rb'\x6a\xff.{0,3}?\x6a\x01(?:\x6a(.)|\x68(....))', re.S)
MEAL_LEVEL = re.compile(rb'\xc7\x46\x08(....)', re.S)


def meal_buffs(exe=None):
    """{item effect type: (skill id, level)} of the effect types the meals use for their buffs (0x513c30 on): the
    method stores the level as a float in the effect and then performs the skill. The game announces the same
    level when the meal is eaten (StringData 183 to 196)."""
    try:
        exe = exe or Exe()
    except (SystemExit, Exception):
        return dict(MEAL_BUFFS_FALLBACK)
    if not exe or not exe.data:
        return dict(MEAL_BUFFS_FALLBACK)
    cases = struct.unpack(f'<{TYPES}I', exe.read(SWITCH_TABLE, 4 * TYPES))
    found = {}
    for kind in MEAL_TYPES:
        stores = [struct.unpack('<I', m.group(1))[0] for m in re.finditer(rb'\xc7[\x00-\x07](....)', exe.read(cases[kind], 0x4c), re.S)]
        vtables = [v for v in stores if RDATA[0] <= v < RDATA[1]]
        body = exe.read(exe.u32(vtables[-1] + 4 * APPLY_SLOT), 0x30)
        skill, level = MEAL_SKILL.search(body), MEAL_LEVEL.search(body)
        if skill and level:
            found[kind] = (skill.group(1)[0] if skill.group(1) else struct.unpack('<I', skill.group(2))[0],
                           int(struct.unpack('<f', level.group(1))[0]))
    # The three party buffs (Brave, Protect, Sun Veil) perform their skill with no level given: the program then
    # asks the skill table how many levels the skill has and uses the last one (0x50b0ea, 0x56c690).
    levels = {r[2]: r[4] for r in openpyxl.load_workbook(WORKBOOK, read_only=True)['Skills'].iter_rows(min_row=2, values_only=True)}
    for kind in MAX_LEVEL_TYPES:
        stores = [struct.unpack('<I', m.group(1))[0] for m in re.finditer(rb'\xc7[\x00-\x07](....)', exe.read(cases[kind], 0x4c), re.S)]
        vtables = [v for v in stores if RDATA[0] <= v < RDATA[1]]
        skill = NO_LEVEL_SKILL.search(exe.read(exe.u32(vtables[-1] + 4 * APPLY_SLOT), 0x30))
        skill_id = skill.group(1)[0] if skill.group(1) else struct.unpack('<I', skill.group(2))[0]
        found[kind] = (skill_id, levels[skill_id])
    return found


NO_LEVEL_SKILL = re.compile(rb'\x6a\xff.{0,3}?\x6a\x00(?:\x6a(.)|\x68(....))', re.S)      # push -1; push 0; push skill id


MAX_LEVEL_TYPES = (35, 36, 39)


def skill_names():
    rows = openpyxl.load_workbook(WORKBOOK, read_only=True)['Skills'].iter_rows(min_row=2, values_only=True)
    return {r[2]: r[3] for r in rows}


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    names = skill_names()
    for kind, (skill, flag) in sorted(skill_items().items()):
        print(f'type {kind:3d}: performs skill {skill} ({names.get(skill)}), flag {flag}')


if __name__ == '__main__':
    main()
