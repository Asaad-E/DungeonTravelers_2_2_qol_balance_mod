"""The two protecting dances of the Etoile raised once more, the user's word of 2026-10-12 ("raise the Etoile more"):
the rows A084 - A093 (Fascinating Waltz) and A124 - A133 (Dumbing Mazurka) of data_tables/balance_changes.xlsx take
1.5 times the value they held, which was already 1.5 times vanilla.

    python raise_etoile_dances.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CHANGES = os.path.join(PROJECT, 'data_tables', 'balance_changes.xlsx')
RECORDS = ('Technique 69 Fascinating Waltz', 'Technique 74 Dumbing Mazurka')
FACTOR = 1.5
WHY = (' | 2026-10-12 (the user: raise the Etoile more, without making her a Diva): times 1.5 again. As the one support '
       'of the Meta party she won 50% / 88% / 67% at C30 / C34 / C38 and wins 68% / 98% / 82% with it; C36 (10%) and the '
       'last boss (0%) stay out of her reach, they ask for a heal she does not have, nor with Primevil. Beside a Diva '
       'the party gains at most 14 points (C38), nothing elsewhere: scratch/audit_2026-10-11/etoile_trial.py')


def main():
    wb = openpyxl.load_workbook(CHANGES)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    col = {name: head.index(name) for name in head}
    done = 0
    for row in ws.iter_rows(min_row=2):
        if row[col['Record']].value in RECORDS and row[col['State']].value != 'dropped' and \
                '2026-10-12' not in str(row[col['Why']].value):
            row[col['New']].value = round(float(row[col['New']].value) * FACTOR, 2)
            row[col['Why']].value = str(row[col['Why']].value) + WHY
            row[col['State']].value = 'decided'
            done += 1
    wb.save(CHANGES)
    print(done, 'rows')


if __name__ == '__main__':
    main()
