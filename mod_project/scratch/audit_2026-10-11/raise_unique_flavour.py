"""Three unique skills raised a little for the feel of them, the user's word of 2026-10-12 (rows U021 - U050 of
data_tables/balance_changes.xlsx): the share Arcane Shield sends back times 1.5, the chance of Alter Fate times 1.5,
the chance of Tranquil Song times 1.25 ("not so much, it is one of the good ones"). The box of Arcane Shield is no
longer worded by hand: its line of the roll comes from the sheet UniqueSkillNotes (column Balance_Line).

    python raise_unique_flavour.py
"""
import os
import sys
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, PROJECT)
sys.path.insert(0, os.path.join(PROJECT, 'tools'))
from sim.data import game_data                            # noqa: E402

TABLES = os.path.join(PROJECT, 'data_tables')
FIRST = 20
# skill, factor, the most a value may be (a chance stays under 1; a share sent back does not), row it serves, why
RAISES = (('Arcane Shield', 1.5, None, 'I18',
           'The user, 2026-10-12, for the feel of it: with its roll at 25 in 100 (X13) the share sent back times 1.5. '
           'Tried at 1.5 and at 2 with the roll at 15 and at 25: the wins of the reference party do not move (81 - 82% '
           'of five late bosses), the damage sent back is little beside the HP of a boss'),
          ('Alter Fate', 1.5, 1.0, 'I14',
           'The user, 2026-10-12, for the feel of it: an insurance that is seen now and then. Tried: the reference party '
           'wins 92.6% of five late bosses with it and 94.0% at 1.5 times the chance (90.0% without it)'),
          ('Tranquil Song', 1.25, 1.0, 'I15',
           'The user, 2026-10-12: a little, it is one of the good ones. The singer, whose song fired 7 - 13% of her turns'))


def main():
    data = game_data()
    wb = openpyxl.load_workbook(os.path.join(TABLES, 'balance_changes.xlsx'))
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    have = {r[0].value for r in ws.iter_rows(min_row=2)}
    number = FIRST
    for name, factor, most, serves, why in RAISES:
        uid, skill = next((k, s) for k, s in data.unique.items() if s['name'] == name)
        for rank, value in enumerate(skill['values'], 1):
            number += 1
            rid = f'U{number:03d}'
            if rid in have:
                continue
            new = round(value * factor + 1e-9, 2)
            ws.append([{'ID': rid, 'Table': 'unique', 'Record': f'Unique {uid} {name}', 'Field': 'value', 'Level': rank,
                        'Vanilla': value, 'New': new if most is None else min(most, new), 'Lever': 'L41',
                        'Serves': serves, 'Why': why, 'Decided': '2026-10-12', 'State': 'decided'}.get(c) for c in head])
    wb.save(os.path.join(TABLES, 'balance_changes.xlsx'))
    wb = openpyxl.load_workbook(os.path.join(TABLES, 'balance_skill_text.xlsx'))
    ws = wb['Skills']
    bhead = [c.value for c in ws[1]]
    for row in list(ws.iter_rows(min_row=2)):
        if row[bhead.index('Filename')].value == 'UniqueSkillData_0379.bin':
            ws.delete_rows(row[0].row)
    wb.save(os.path.join(TABLES, 'balance_skill_text.xlsx'))
    print(number - FIRST, 'rows')


if __name__ == '__main__':
    main()
