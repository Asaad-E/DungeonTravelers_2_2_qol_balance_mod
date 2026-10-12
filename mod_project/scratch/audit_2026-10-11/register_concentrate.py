"""The decision of the user of 2026-10-12 on Concentrate, written into the workbooks: one byte of the program makes
the interruption rate the product its own message tells (row X14 of the sheet ExePatches, lever L45), for every unit
that has the passive; the box of the Balance variant states what it then takes off; the box of the QoL variant goes
back to the text of the game, which states no figure. Run once.

    python register_concentrate.py
"""
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TABLES = os.path.join(os.path.dirname(os.path.dirname(HERE)), 'data_tables')
FILE = 'PassiveSkillData_0281.bin'
BALANCE_LINE = 'Interruption rate: -9% -> -17% -> -23% -> -29% -> -33% (L1-5 Max)'


def main():
    path = os.path.join(TABLES, 'balance_criteria.xlsx')
    wb = openpyxl.load_workbook(path)
    wb['ExePatches'].append(['X14', 'balance', 'L45', '0x44d044', '58', '59', None, None,
                             'Interruption rate of a blow: max HP * 1.5 times (1 + Concentrate) where the program adds the '
                             'two (addss -> mulss at 0x44d042; its debug message writes the product). The user, 2026-10-12: '
                             'mend it for every unit, enemies with the passive too'])
    wb['Levers'].append(['L45', 'Concentrate divides the interruption rate (the product the message of the program tells)',
                         'executable, 0x44d042', 'one byte', 'Executable', 'The byte is 58 (addss) and becomes 59 (mulss)',
                         'docs/FORMULAS.md 2.5: the only reader of effect 264 as a status is 0x44cfad. Not seen working in '
                         'the game yet'])
    ws = wb['EffectRules']
    for row in ws.iter_rows(min_row=2):
        if row[0].value == 264:
            row[2].value = str(row[2].value) + ('. On the Balance variant row X14 of the sheet ExePatches makes it the '
                                                'product: the rate is divided by 1 + value')
    wb.save(path)

    path = os.path.join(TABLES, 'qol_skill_text.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb['Skills']
    col = {c.value: k for k, c in enumerate(ws[1])}
    for row in ws.iter_rows(min_row=2):
        if row[col['Filename']].value == FILE:
            row[col['QoL_LvUp']].value = row[col['Vanilla_LvUp']].value
            row[col['Edited']].value = 'yes'
            row[col['Line1_Notes']].value = ('Left as the game words it (the user, 2026-10-12): the program adds the value '
                                             'where it means to multiply (docs/FORMULAS.md 2.5), so no figure is true of it')
    wb.save(path)

    path = os.path.join(TABLES, 'balance_skill_text.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb['Skills']
    head = [c.value for c in ws[1]]
    line = {'Category': 'Passive', 'Filename': FILE, 'Skill_Name': 'Concentrate', 'Balance_LvUp': BALANCE_LINE,
            'Note': 'Row X14 of the sheet ExePatches: the rate is divided by 1 + value (0.1 to 0.5), which takes off 1 - 1 / '
                    '(1 + value) of it'}
    ws.append([line.get(h) for h in head])
    wb.save(path)
    print('written')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    main()
