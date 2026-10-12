"""The measure of each unique skill (unique_skills.py beside this file) written into its row of the register (rows
I01 - I18 of the sheet Criticisms). The rows stay open: the user decides after reading them.

    python register_unique_figures.py
"""
import glob
import json
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(os.path.dirname(HERE))
WORKBOOK = os.path.join(PROJECT, 'data_tables', 'balance_criteria.xlsx')
NOTES = {
    'Forerunner': ' Read again on 2026-10-12: one roll below 100 under value * 100 for each chanted command (0x4176a0, '
                  'called at 0x41d69c for a spell and at 0x419c55 for the other chanted kind, each behind a flag that '
                  'keeps it to one roll): 7 - 15% by rank, 20% crowned. Both guides felt 40 - 50% with the crown',
    'Positive Trance': ' The simulator holds Pep at 100; she was played at Pep 130, the most it can be, so this is the '
                       'most the skill gives: +6% damage at rank 5 - 6, +15% at rank 9, +30% crowned (the guide measured '
                       'about +15% at level 90)',
    'Shadow Heart': ' Played since 2026-10-12 as read (scratch/qol_fields_reading.md 5, checked): the shadow has the '
                    'stats of its owner at that moment with twice the ATK, her passives, none of the statuses techniques '
                    'put, a life of 3. The fault the Japanese page tells (its attack falling back after a drain) is not '
                    'in the program by that reading',
    'Third Eye Wide': ' Its seventh result, the dispel, removes nothing: 0x47ee6f jumps over the routine '
                      '(scratch/qol_fields_reading.md 7, checked), as the Japanese page says; played so since 2026-10-12',
    'Did You Call for Me?': ' Played since 2026-10-12 as read: the second action belongs to the same turn, so her '
                            'statuses do not lose a turn between the two (Power Charge serves both blows)',
    'Usurper of Souls': ' It only acts on what an enemy summons; which of these five fights hold a summoner was not looked at',
    'Demon Wall': ' It cuts the damage of demons, undead and apparitions: Samael (C38) and Melek Taus (P27) are demons; the race of the others was not checked against its list',
}


def main():
    rows = []
    for name in glob.glob(os.path.join(HERE, 'unique_skills_*.json')):
        with open(name, encoding='utf-8') as f:
            rows += json.load(f)
    found = {}
    for r in rows:
        wins, lengths, hps, cells = [], [], [], []
        for c in r['cells']:
            a, b = c['without'], c['with']
            wins.append(b['won'] - a['won'])
            hps.append(b['hp_left'] - a['hp_left'])
            if a['time'] and b['time']:
                lengths.append(1 - b['time'] / a['time'])
            cells.append(f"{c['checkpoint']}{' crowned' if c['crown'] else ''} {a['won']:.0%} > {b['won']:.0%}")
        mean = lambda xs: sum(xs) / len(xs) if xs else 0.0          # noqa: E731
        found[r['skill']] = (f"scratch/audit_2026-10-11/UNIQUE_SKILLS.md (2026-10-12, variant, 100 battles a cell; the "
                             f"reference party with {r['girl']} as {r['role']}, without the skill > with it): "
                             + ', '.join(cells) + f". In the mean {mean(wins) * 100:+.0f} points of wins, fights "
                             f"{mean(lengths):+.0%} shorter, {mean(hps) * 100:+.0f} points of HP left (a cell moves by "
                             f"about 5 points by chance)." + NOTES.get(r['skill'], ''))
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    col = {name: head.index(name) for name in head}
    done = 0
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value).startswith('I') and str(row[col['Theme']].value).startswith('T15'):
            skill = str(row[col['Subject']].value).split(' (')[0]
            if skill in found:
                row[col['Balance_Result']].value = found[skill]
                row[col['Verdict']].value = 'measured; to be read with the user'
                done += 1
    wb.save(WORKBOOK)
    print(done, 'rows')


if __name__ == '__main__':
    main()
