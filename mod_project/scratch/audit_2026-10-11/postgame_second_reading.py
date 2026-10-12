"""The second reading of the Japanese pages (2026-10-11 and 12, notes in postgame_notes_atwiki.md beside this file)
written into data_tables/sim_setup.xlsx: what each page says, beside the GameFAQs note of the boss (sheet
PostgameBosses), and a class for the 25 fights the guide curve did not hold (sheet GuideDifficulty). Where no page
rates a fight the lowest floor is taken (hard) and the row says so. Also the two rules of the sheet EffectRules of
data_tables/balance_criteria.xlsx that the reading corrected.

    python postgame_second_reading.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SETUP = os.path.join(PROJECT, 'data_tables', 'sim_setup.xlsx')
CRITERIA = os.path.join(PROJECT, 'data_tables', 'balance_criteria.xlsx')
UNRATED = 'no page rates the fight, only how to play it: the lowest floor is taken'
NOTES = {
    'P01': 'atwiki page 29: only a first showing; a party that beat Demiurge wins with room',
    'P05': 'atwiki page 186: his physical blows after his self buff kill in one hit; remove the buff early',
    'P06': 'atwiki page 227: high DEF, physical resistance and Auto-Guard; under about 3,000 HP a chant of 2,000 to everyone',
    'P07': 'atwiki page 225: magic forbidden; fire resistance first, then water; Siegdrad in the back row is the threat, '
           'the fight eases once one of the three falls',
    'P10': 'atwiki page 29: water zone and physical attack zone; the boss kills in one or two hits, a short fight that '
           'takes several tries',
    'P11': 'atwiki page 99: stronger than before, but the water zone lets Frost Slash and Enchant Ice hit hard',
    'P14': 'atwiki page 227: physical defence zone, two Beard Knights in front',
    'P16': 'atwiki page 207: Murder License, before each action of Aradia one member may die at once; without instant '
           'death immunity or high earth resistance the fight cannot be played',
    'P32': 'atwiki page 317: no race and no attribute; her plain attack drains TP; at about half HP a slow chant of 1,000 or more',
    'P34': 'atwiki page 231: three physical blows in a row; sleep, Cover or Provoke on a member of high defence; dispel '
           'her self buff each time',
    'P35': 'atwiki page 231: her Parry voids physical blows and comes back at once: magic or the Salsa; under about '
           '26,000 HP a party wide physical skill that needs Parry, clones or evasion',
    'P37': 'atwiki page 231: his Warning Shot and Chain Spike are water, not stab; most of his skills are stab',
    'P38': 'atwiki page 231: a Geass that lowers the stats, then strong spells; a critical hit breaks the Geass',
    'P39': 'atwiki page 231: three enemies of 25,000 HP or more',
    'P40': 'atwiki page 231: five enemies; the Princess (songs that raise speed and evasion) first, then the Prince '
           '(party dispel, Salsa); at half HP Trump Card and a fixed 2,000 to everyone after a long chant',
    'P41': 'atwiki page 333: mostly single blows, a Valkyrie tank makes it stable; her Carol heals 2,000 a turn; fight from range',
    'P42': 'atwiki page 333: the hardest of the four; DEF 2,450, a Salsa of 250 to all each turn, a full self buff; '
           'players report not winning without crowns',
    'P43': 'atwiki page 333: opens with a wall against short and mid range; with a tank the easiest of the four',
    'P44': 'atwiki page 333: party lightning with paralysis, eight clones with a speed and attack buff; the lowest '
           'defence of the four',
    'P45': 'atwiki page 334: two escorts with Warning Shot; after his magic buff a spell deals 2,000 or more',
    'P46': 'atwiki page 334: the three hit hard physically; Trump Card of an escort at half HP',
    'P47': 'atwiki page 334: magic and songs forbidden; Murder License every turn; the three open to slow',
    'P48': 'atwiki page 334: Dark Blast deals a fixed 1,111 to all; high RES on the three; a Geass lowers every resistance',
    'P52': 'atwiki page 317: uses Trump Card often (ATK 1,160, DEF 0 for 4 turns)',
    'P53': 'atwiki page 317: acts first for sure and opens with the Salsa; critical hits to break it',
    'P54': 'atwiki page 317: strong blows in a row, slow him; shoots ranged attacks down often',
    'P55': 'atwiki page 317: the Peggy Sentinel is hard and its Parry hard to remove; start with the boss to stop its songs',
}
GAMEFAQS = 'GameFAQs faqs/80834 (Hentaicheg), '
CLASSES = {
    'P31': ('hard', 'a weakened superboss', GAMEFAQS + 'Postgame Quick Summary'),
    'P32': ('hard', UNRATED, 'atwiki dt2-2 page 317'),
    'P33': ('normal', 'doable; much weaker than the superbosses', GAMEFAQS + 'Postgame Quick Summary'),
    'P34': ('hard', UNRATED, 'atwiki dt2-2 page 231'),
    'P35': ('hard', UNRATED, 'atwiki dt2-2 page 231'),
    'P36': ('hard', UNRATED, 'atwiki dt2-2 page 231'),
    'P37': ('hard', UNRATED, 'atwiki dt2-2 page 231'),
    'P38': ('hard', UNRATED, 'atwiki dt2-2 page 231'),
    'P39': ('hard', UNRATED, 'atwiki dt2-2 page 231'),
    'P40': ('hard', UNRATED, 'atwiki dt2-2 page 231'),
    'P41': ('normal', 'stable with a Valkyrie tank and a way to stop her Carol', 'atwiki dt2-2 page 333'),
    'P42': ('hard', 'the hardest of the brigade; players report not winning without crowns', 'atwiki dt2-2 page 333'),
    'P43': ('normal', 'with a tank the easiest of the brigade', 'atwiki dt2-2 page 333'),
    'P44': ('normal', 'the lowest defence of the brigade; a rush with enchanted blows works', 'atwiki dt2-2 page 333'),
    'P45': ('hard', 'each king asks for its own answer', 'atwiki dt2-2 page 334'),
    'P46': ('hard', 'each king asks for its own answer', 'atwiki dt2-2 page 334'),
    'P47': ('hard', 'each king asks for its own answer', 'atwiki dt2-2 page 334'),
    'P48': ('hard', 'a fixed 1,111 to all; the last of the four', 'atwiki dt2-2 page 334'),
    'P49': ('hard', 'a superboss', GAMEFAQS + 'Postgame Dungeon Superbosses'),
    'P50': ('hard', 'the strongest enemy of the game, hard even with every member crowned', 'atwiki dt2-2 page 231'),
    'P51': ('hard', 'a superboss', GAMEFAQS + 'Postgame Dungeon Superbosses'),
    'P52': ('hard', UNRATED, 'atwiki dt2-2 page 317'),
    'P53': ('hard', UNRATED, 'atwiki dt2-2 page 317'),
    'P54': ('hard', UNRATED, 'atwiki dt2-2 page 317'),
    'P55': ('hard', UNRATED, 'atwiki dt2-2 page 317'),
}
RULES = {
    '291': 'Heretic Samsara (enemy): when damage takes its bearer below half its max HP its value is set to 1. From then '
           'on, at each battle update of its turn, each of its eight resistance states (effects 416 to 423) is put again '
           'with value (rand(2) + 1) * 50: 50 or 100 (0x438f15 - 0x438f1e). When the stats of the unit are recomputed '
           'the state of a type replaces that resistance, after its clamp (0x4fc76f - 0x4fc880). The Japanese page 317 '
           'adds that putting the bearer to sleep ends it (not read in the code, not played)',
    '319': 'Murder License (enemy): at step 12 of the turn of its bearer an action is built from the skill linked to the '
           'status, on one target drawn at random (0x43736a - 0x43743a). Whether the own command of the bearer still '
           'follows, which skill is linked and its chance are not read yet; the Japanese pages 207 and 334 say that '
           'before each action of the bearer one member may die at once unless immune to instant death or of high '
           'earth resistance. An earlier wording here, that the action of the bearer is replaced, did not fit them',
}


def main():
    wb = openpyxl.load_workbook(SETUP)
    ws = wb['PostgameBosses']
    head = [c.value for c in ws[1]]
    says = head.index('What the guide says')
    names = {}
    for row in ws.iter_rows(min_row=2):
        names[row[0].value] = row[head.index('Boss')].value
        note = NOTES.get(row[0].value)
        if note and note not in str(row[says].value or ''):
            old = str(row[says].value or '').replace(' | atwiki page 29 (not read yet)', '')
            row[says].value = (old + ' | ' if old else '') + note
    ws = wb['GuideDifficulty']
    have = {r[0].value for r in ws.iter_rows(min_row=2)}
    for cid, (cls, text, source) in CLASSES.items():
        if cid not in have:
            ws.append((cid, names[cid], cls, text, source))
    wb.save(SETUP)
    wb = openpyxl.load_workbook(CRITERIA)
    ws = wb['EffectRules']
    found = 0
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) in RULES:
            row[2].value = RULES[str(row[0].value)]
            found += 1
    assert found == len(RULES), found
    wb.save(CRITERIA)
    print(len(NOTES), 'notes,', len(CLASSES), 'classes,', found, 'rules')


if __name__ == '__main__':
    main()
