"""Row J02 of the register: the elemental resistances of the six superboss fights where the best option of the mage
is under that of the Dark Lord (magic_wall.py beside this file). Each elemental resistance of their enemies that is
under 100 goes down by a step; a resistance of 100 stays (what the boss shrugs off is its flavour).

    python j02_resistances.py trial 20 30       the mage beside the Dark Lord at each step, nothing written
    python j02_resistances.py write 20          rows of data_tables/balance_changes.xlsx, state decided
    python j02_resistances.py res 0.9           the same for the RES stat of those enemies times a factor: what was built

Measured on 2026-10-12: the step on the elemental resistances raises the earth and water strikes of the Dark Lord as
much as the spells, so the mage stays where she was beside him (50 - 84% at -20). What walls the spells off is the
magic defence of these enemies (RES 407 to 1,648 against INT 353 - 363 of the mage), which no physical strike reads.
"""
import os
import sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, PROJECT)
sys.path.insert(0, os.path.join(PROJECT, 'tools'))
sys.path.insert(0, HERE)
import magic_wall                                         # noqa: E402
from sim import scenario, setup as setup_mod              # noqa: E402
from sim import policy as P                               # noqa: E402
from sim.data import game_data                            # noqa: E402
from sim.variant import balance_data                      # noqa: E402

FIGHTS = ('P28', 'P29', 'P30', 'P33', 'P50', 'P51')
ELEMENTS = ('fire', 'water', 'wind', 'earth', 'lightning')
CHANGES = os.path.join(PROJECT, 'data_tables', 'balance_changes.xlsx')
WHY = ('Decided 2026-10-09 (row J02: lower a little the elemental resistances of the superbosses that wall magic off, '
       'keeping their flavour) and confirmed by the user on 2026-10-12 (without abuse; the postgame stays hard, the mages '
       'keep some of their power). The six fights where the best option of the mage was under that of the Dark Lord '
       '(50 - 84%): every elemental resistance under 100 of their enemies goes down by {step}; a resistance of 100 stays. '
       'The mage beside the Dark Lord there: {before} before, {after} after')


def enemies_of(data, setup):
    found = {}
    for cid in FIGHTS:
        cp = setup.checkpoint(cid)
        for number in scenario.battle_numbers(cp):
            fight = next(b for b in data.fixed if b['floor'] == cp['Floor'] and b['number'] == number)
            for eid, _ in fight['enemies']:
                found[eid] = cid
    return found


def lowered(data, setup, step):
    changes = []
    for eid in enemies_of(data, setup):
        for e in ELEMENTS:
            value = data.enemies[eid]['res_' + e]
            if value < 100:
                changes.append(('enemies', eid, 'res_' + e, value - step))
    return changes


def ratios(data, setup, policy):
    out = {}
    for cid in FIGHTS:
        spell, mage, _ = magic_wall.best(data, setup, 'mage', cid, policy)
        _, lord, _ = magic_wall.best(data, setup, 'dark lord', cid, policy)
        out[cid] = mage / lord
    return out


def main():
    setup = setup_mod.load()
    policy = P.Policy(P.load_rules(setup.policies), setup.settings)
    data = balance_data(game_data())
    before = ratios(data, setup, policy)
    steps = [int(x) for x in sys.argv[2:]]
    after = {step: ratios(data.with_changes(lowered(data, setup, step)), setup, policy) for step in steps}
    for cid in FIGHTS:
        print(cid, f'{before[cid]:.0%}', ' '.join(f'-{step}: {after[step][cid]:.0%}' for step in steps))
    if sys.argv[1] != 'write':
        return
    step = steps[0]

    def span(r):
        return f"{min(r.values()):.0%} - {max(r.values()):.0%}"

    why = WHY.format(step=step, before=span(before), after=span(after[step]))
    wb = openpyxl.load_workbook(CHANGES)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    number = max(int(str(r[0].value)[1:4]) for r in ws.iter_rows(min_row=2) if str(r[0].value).startswith('E'))
    for _, eid, field, value in lowered(data, setup, step):
        number += 1
        enemy = data.enemies[eid]
        ws.append([{'ID': f'E{number:03d}', 'Table': 'enemies', 'Record': f"{eid} {enemy['name']}",
                    'Field': 'resistance ' + field[4:], 'Vanilla': enemy[field], 'New': value, 'Lever': 'L02',
                    'Serves': 'J02', 'Why': why, 'Decided': '2026-10-12', 'State': 'decided'}.get(name) for name in head])
    wb.save(CHANGES)
    print(number, 'is the last E row')


WHY_RES = ('Row J02 (decided 2026-10-09: lower a little what walls magic off at the superbosses, keeping their flavour; '
           'the user, 2026-10-12: without abuse, the postgame stays hard, the mages keep some of their power). At the six '
           'fights where the best option of the mage was under that of the Dark Lord, RES of every enemy times {factor}; '
           'the elemental resistances are left alone (lowering them raises the elemental strikes of the Dark Lord as much '
           'as the spells: measured, the mage stays at 50 - 84%). The mage beside the Dark Lord there: {before} before, '
           '{after} after')


def res_rows():
    factor = float(sys.argv[2])
    setup = setup_mod.load()
    policy = P.Policy(P.load_rules(setup.policies), setup.settings)
    data = balance_data(game_data())
    ids = enemies_of(data, setup)
    changes = [('enemies', e, 'res', int(data.enemies[e]['res'] * factor)) for e in ids]
    before, after = ratios(data, setup, policy), ratios(data.with_changes(changes), setup, policy)
    for cid in FIGHTS:
        print(cid, f'{before[cid]:.0%} > {after[cid]:.0%}')

    def each(r):
        return ', '.join(f'{c} {r[c]:.0%}' for c in FIGHTS)

    why = WHY_RES.format(factor=factor, before=each(before), after=each(after))
    wb = openpyxl.load_workbook(CHANGES)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    number = max(int(str(r[0].value)[1:4]) for r in ws.iter_rows(min_row=2) if str(r[0].value).startswith('E'))
    for _, eid, field, value in changes:
        number += 1
        enemy = data.enemies[eid]
        ws.append([{'ID': f'E{number:03d}', 'Table': 'enemies', 'Record': f"{eid} {enemy['name']} ({ids[eid]})",
                    'Field': 'res', 'Vanilla': enemy['res'], 'New': value, 'Lever': 'L01', 'Serves': 'J02', 'Why': why,
                    'Decided': '2026-10-12', 'State': 'decided'}.get(name) for name in head])
    wb.save(CHANGES)
    print(len(changes), 'rows, the last is E%03d' % number)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    res_rows() if sys.argv[1] == 'res' else main()
