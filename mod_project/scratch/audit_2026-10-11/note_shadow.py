"""A note on the row of Shadow Heart in the register (sheet Criticisms): what was read on 2026-10-12 in the
constructor of the shadow and what the simulator does not play, so that its measure is not taken for more than it is.

    python note_shadow.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
WORKBOOK = os.path.join(PROJECT, 'data_tables', 'balance_criteria.xlsx')
NOTE = (' Read on 2026-10-12 (0x4fe2e0): the shadow is built by the doll constructor and then walks a list of statuses '
        'of its own, keeping some and dropping others, so it does carry statuses of its owner, as both guides say (the '
        'buffs she holds when it appears, Double and Triple Attack). The simulator gives it the base stats of its owner '
        'and no status at all: this measure is a floor, not the worth of the skill. Which statuses stay, and whether '
        'its attack is twice hers, are not read')


def main():
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if str(row[head.index('Subject')].value).startswith('Shadow Heart') and NOTE not in str(row[head.index('Balance_Result')].value):
            row[head.index('Balance_Result')].value = str(row[head.index('Balance_Result')].value) + NOTE
    wb.save(WORKBOOK)


if __name__ == '__main__':
    main()
