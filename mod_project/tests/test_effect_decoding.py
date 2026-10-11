"""Effect ids decoded from the game's own text: the evidence is re-read from the files on every run."""
import os
import re
import unittest

import openpyxl

import _paths
import magnitude_words as mw
from extract_vanilla import EFFECT_NAMES
from pck_handler import PckArchive, SkillHelper, EnemyAttackHelper
from extract_vanilla import read_game_effect_names
from text_checks import STATED_FIGURES, correct_prose, enemy_attack_slots, skill_tables

SKILL_TABLE = os.path.join(_paths.TABLES_DIR, 'qol_skill_text.xlsx')

# Places where the original text gives a figure that its own data does not hold. The game does what
# the data says; the QoL text corrects the figure and keeps the sentence (text_checks.correct_prose).
VANILLA_TEXT_ERRORS = {
    ('Modest Fever', 'Deal 555 fixed damage', 777.0),
    ('Scream of Despair', 'With 50% accuracy', 65.0),
    ('Debilitating Salsa', '(Cap: 500)', 300.0),
    ('Debilitating Geas', '(Cap: 999)', 500.0),
    ('Weathering', 'fixed 50% chance to hit', 85.0),
}


class GameTextAgainstStoredValues(unittest.TestCase):
    """Where an enemy attack text states a figure, the figure is the value of the effect named for it."""

    @classmethod
    def setUpClass(cls):
        cls.attacks = []
        for _, b in PckArchive(os.path.join(_paths.BACKUP_DIR, 'EnemyAttackData.pck')).files:
            cls.attacks.append((EnemyAttackHelper.read_enemy_attack(b), enemy_attack_slots(b)))

    def test_figures_in_the_text_are_the_stored_values(self):
        matches, errors = {i: 0 for i in STATED_FIGURES}, set()
        for attack, slots in self.attacks:
            for effect_id, (pattern, factor) in STATED_FIGURES.items():
                m = re.search(pattern, attack['desc'])
                if not m:
                    continue
                stored = [round(v * factor, 2) for i, v, _ in slots if i == effect_id]
                self.assertTrue(stored, f"{attack['name']}: says '{m.group(0)}' but has no effect {effect_id}")
                if float(m.group(1)) in stored:
                    matches[effect_id] += 1
                else:
                    errors.add((attack['name'], m.group(0), stored[0]))
        self.assertEqual(matches, {70: 5, 296: 24, 8: 18, 55: 3})
        self.assertEqual(errors, VANILLA_TEXT_ERRORS)

    def test_wrong_figures_are_corrected_and_the_sentence_stays(self):
        corrected = {}
        for _, b in PckArchive(os.path.join(_paths.BACKUP_DIR, 'EnemyAttackData.pck')).files:
            attack = EnemyAttackHelper.read_enemy_attack(b)
            fixed = correct_prose(b, attack['desc'])
            if fixed != attack['desc']:
                corrected.setdefault(attack['name'], set()).add(
                    tuple((a, c) for a, c in zip(attack['desc'].split(), fixed.split()) if a != c)
                    if len(attack['desc'].split()) == len(fixed.split()) else ('reworded',))
        self.assertEqual(corrected['Modest Fever'], {(('555', '777'),)})
        self.assertEqual(corrected['Scream of Despair'], {(('50%', '65%'),)})
        self.assertEqual(corrected['Debilitating Salsa'], {(('500)', '300)'),)})
        self.assertEqual(corrected['Debilitating Geas'], {(('999)', '500)'),)})
        self.assertEqual(corrected['Weathering'], {(('50%', '85%'),)})
        # The game's own table names effect 161 "Physical Attack"; the sentence said "defense".
        self.assertEqual(corrected['Heavy Stomp'], {('reworded',)})
        self.assertEqual(sorted(corrected), ['Debilitating Geas', 'Debilitating Salsa', 'Heavy Stomp', 'Mega Stomp',
                                             'Modest Fever', 'Scream of Despair', 'Weathering'])

    def test_hit_count_effects_match_their_wording(self):
        wording = {90: r'thrice|three', 92: r'2.{1,3}4', 94: r'0.{1,3}4', 95: r'1.{1,3}5', 96: r'3.{1,3}5', 97: r'twice|two|double'}
        for effect_id, pattern in wording.items():
            texts = [a['desc'] for a, slots in self.attacks
                     if any(i == effect_id for i, _, _ in slots) and a['desc'].startswith('Target:')]
            said = sum(1 for t in texts if re.search(pattern, t, re.I))
            with self.subTest(effect=effect_id):
                self.assertGreater(len(texts), 5)
                self.assertGreaterEqual(said / len(texts), 0.6, f'{said} of {len(texts)}')

    def test_ignore_defense_flag_matches_its_wording(self):
        texts = [a['desc'] for a, slots in self.attacks if any(i == 69 for i, _, _ in slots)]
        said = sum(1 for t in texts if re.search(r'defense|armor|defence', t, re.I))
        self.assertEqual(len(texts), 30)
        self.assertGreaterEqual(said, 24)


class GameEffectTable(unittest.TestCase):
    """CommandData.pck is the game's own list of effects: one entry per effect id, with its name."""

    def test_names_agree_with_what_was_decoded_from_the_text(self):
        names = read_game_effect_names()
        self.assertGreater(len(names), 300)
        expected = {
            0: 'Physical Attack', 1: 'Magic Attack', 8: 'Fixed-Damage Attack', 26: 'Weapon Attack', 47: 'RES-Based Damage',
            54: 'Turn Delay', 55: 'Fixed Accuracy', 69: 'Ignore Defense', 70: 'Crit Rate Fixed Modifier',
            73: 'Damage xn', 74: 'Strip Buffs', 75: 'Absorb HP', 76: 'Absorb TP', 79: 'Speed Change Chance',
            87: 'Only Damages Living', 88: 'Big Damage to Undead, Demons, and Apparitions', 90: '3 Hits', 92: '2~4 Hits',
            93: 'Hit Count', 94: '0~4 Hits', 95: '1~5 Hits', 96: '3~5 Hits', 97: '2 Hits', 99: '1~3 Hits',
            110: 'Crits on Beasts', 111: 'More Damage with Heavy Lance', 113: 'Poison', 118: 'Stun', 122: 'Slow',
            125: 'More Crits Later', 133: 'Max HP Down', 136: 'Power Drops', 137: 'Target Count',
            138: 'Percent HP Damage', 161: 'Physical Attack', 162: 'Physical Defense', 163: 'Magic Defense',
            164: 'Physical Evasion', 165: 'Magic Evasion', 166: 'Physical Accuracy', 167: 'Magic Accuracy',
            168: 'Speed', 295: 'Healing Cap', 296: 'Damage Cap',
        }
        for effect_id, name in expected.items():
            with self.subTest(effect=effect_id):
                self.assertEqual(names[effect_id], name)


class Registry(unittest.TestCase):
    def test_every_effect_treated_as_harmless_is_named_with_its_evidence(self):
        for effect_id in mw.SAFE_EFFECTS:
            with self.subTest(effect=effect_id):
                self.assertIn(effect_id, EFFECT_NAMES)
                self.assertTrue(EFFECT_NAMES[effect_id][2])

    def test_effects_that_change_the_damage_are_not_treated_as_harmless(self):
        for effect_id in (138, 199, 8, 146, 147):
            self.assertNotIn(effect_id, mw.SAFE_EFFECTS)


class Line1AfterDecoding(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.skills = {}
        for category, pck in (('Technique', 'TechniqueSkillData.pck'), ('Magic', 'MagicSkillData.pck')):
            for _, b in PckArchive(os.path.join(_paths.BACKUP_DIR, pck)).files:
                info = SkillHelper.read_skill(b, category)
                cls.skills.setdefault(info['name'], (info['desc'], skill_tables(category, b)['slots']))

    def propose(self, name):
        desc, slots = self.skills[name]
        return mw.propose(desc, slots, 85)[0]

    def test_skills_that_were_held_back_now_have_their_word(self):
        expected = {
            'Cross Slash': 'Sword/lance skill. Deal severe damage to one enemy (effective vs. unholy monsters).',
            'Armor Penetration': 'Katana skill. Deal minor defense-piercing damage and DEF down to one enemy.',
            'Sharp Stab': 'Dagger skill. Inflict a critical hit with minor stab damage.',
            'Moonlight Purgatory': 'Katana skill. Deal moderate slash damage 3 times to one enemy. Living targets only.',
            'Bible Thump': 'Tome skill. Severe RES-based bash damage to one enemy. May freeze/silence (20-50%).',
            'Guided Shot': 'Bow/gun skill. Fire a single shot certain to hit one enemy for minor damage.',
            'Soul Eater': 'Deal severe damage to one enemy that possesses life. Low chance of instakill (12%).',
        }
        for name, text in expected.items():
            with self.subTest(skill=name):
                self.assertEqual(self.propose(name), text)

    def test_chances_of_side_effects_are_stated(self):
        self.assertEqual(self.propose('Aura Eliminate'), 'Short-range skill. Attack one enemy. May cancel enemy buffs (68-85%).')
        self.assertIn('May delay enemy turn (25-43%).', self.propose('Tinker Squash'))

    def test_damage_that_the_power_table_does_not_measure_stays_as_written(self):
        for name in ('Star Dust Arrow', 'Fire Scroll', 'Debilitating Salsa'):
            with self.subTest(skill=name):
                self.assertEqual(self.propose(name), self.skills[name][0])

    def test_damage_xn_is_the_measure_of_the_hit(self):
        """Effect 73, "Damage xn": the DT2-2 wiki lists Spiral Arrow as ATK 330% -> 400%."""
        self.assertEqual(self.propose('Spiral Arrow'), 'Bow skill. Shoot a tremendous single shot at one enemy for stab damage.')
        self.assertEqual(self.propose('Holy Gavel'), 'Mace skill. Deal tremendous bash damage with an inexorable thud.')


class LevelUpLinesAfterDecoding(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ws = openpyxl.load_workbook(SKILL_TABLE)['Skills']
        header = [c.value for c in ws[1]]
        cls.rows = {r[header.index('Skill_Name')]: dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)}

    def test_phrases_resolved_with_the_decoded_effects(self):
        expected = {
            'Last Stand': ['Critical rate: 5% (+2.5%/Lv) -> 15%', 'Enemy critical rate: 4% (+2%/Lv)'],
            'Power Charge': ['Damage: 1.05x -> 1.1x -> 1.16x -> 1.22x -> 1.3x'],
            'Warm-Up': ['Physical evasion: +15%', 'Magic evasion: +10%'],
            'Fire Scroll': ['Max HP damage: 4% -> 5% -> 6% -> 7% -> 8%', 'Damage cap: 250 -> 500'],
            'Debilitating Salsa': ['Max HP damage: 3%', 'Total damage cap: 250'],
            'Carrot and Stick': ['DEF: -5%', 'Magic defense: -10% -> -15%', 'HP recovery: 35 -> 25'],
            'Generous Heart': ['Post-battle HP: 6%', 'Post-battle TP: 3%'],
            'Phoenix': ['Speed decrease chance: 30% -> 75%', 'Speed decrease: 10%'],
            'Spiral Arrow': ['Damage: 3.3x -> 3.45x -> 3.6x -> 3.8x -> 4x', 'ATK: 88%'],
            'Heal Mastery': ['HP recovery bonus: 8%', 'Heal Cap: +10%'],
        }
        for name, pieces in expected.items():
            for piece in pieces:
                with self.subTest(skill=name, piece=piece):
                    self.assertIn(piece, self.rows[name]['QoL_LvUp'] or self.rows[name]['QoL_Desc'])

    def test_values_that_never_change_are_in_the_box(self):
        """The user's rule: a value the skill has is shown, also when every level has the same."""
        expected = {
            'Phoenix': 'Speed decrease: 10%', 'Justice': 'DEF: -10%', 'Cross Slash': 'Damage vs. unholy: +65%',
            'Daedalus Attack': 'Heavy Lance bonus: +20%', 'Halving Slice': 'Power decrease: 10%',
            'Moonlight Purgatory': 'Critical rate: 10%', 'Blood Absorb': 'HP Absorb: 100%', 'Dispel': 'Buff cancel rate: 100%',
        }
        for name, piece in expected.items():
            with self.subTest(skill=name):
                self.assertIn(piece, self.rows[name]['QoL_LvUp'])

    def test_no_effect_is_left_out_of_a_box(self):
        """Rule B: only speed, chant, turns or TP may give way."""
        for name, row in self.rows.items():
            dropped = [x for x in (row['Left_Out'] or '').split(', ') if x and x not in ('Speed', 'Chant', 'Turns', 'TP Cost')]
            with self.subTest(skill=name):
                self.assertEqual(dropped, [])

    def test_summon_parameters_are_decoded(self):
        """Dolls now show their verified hidden action power rather than vague ATK UP."""
        expected = {
            'Spartan': 'ATK: 144% -> 147',
            'Oracles': 'ATK: 126% (+2-5%/Lv) -> 157%',
            'Mercidarian': 'ATK: 100% -> 104',
            'Catalina': 'Heal Power: 4 -> 6',
            'Musiy': 'Pwr: 180 -> 240',
            'Star Dust Arrow': 'ATK: 100% -> 102',
        }
        for name, piece in expected.items():
            with self.subTest(skill=name):
                self.assertIn(piece, self.rows[name]['QoL_LvUp'])


if __name__ == '__main__':
    unittest.main()
