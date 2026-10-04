"""
QoL enchant names: proposes the name of each enchantment for the QoL variant and keeps
them in a workbook that is edited by hand.

    python mod_project/tools/generate_qol_enchant_text.py

Workbook: mod_project/data_tables/qol_enchant_text.xlsx
  * QoL_Name is the name the build uses. The game lists the enchantments of an item by name
    inside the equipment box, so the name is the place where their effect is shown.
  * Put "yes" in the Edited column to keep your wording; this script then leaves that row alone.
  * The binary holds no magnitude for enchantments. Every figure comes from a player guide and
    names it (tools/enchant_values.py, Figure_Source column). A figure from a guide that is not in
    SHOWN_SOURCES stays in Candidate_Name and the vanilla name is kept.
  * The name keeps the game's own words; only the tier word gives way to the figure.
  * An enchantment without a figure keeps its vanilla name.
"""
import os
import re
import sys
import collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pck_handler import PckArchive, EnchantHelper
import enchant_values

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKUP_PATH = os.path.join(BASE_DIR, 'backups', 'EnchantData.pck')
TABLES_DIR = os.path.join(BASE_DIR, 'data_tables')
PREVIOUS_PATH = os.path.join(TABLES_DIR, 'enchants.xlsx')
OUTPUT_PATH = os.path.join(TABLES_DIR, 'qol_enchant_text.xlsx')

# Longest vanilla name, and longest vanilla name of an enchantment that random equipment can
# roll (those are drawn in two columns inside the equipment box). Not measured on screen.
MAX_NAME = 46
MAX_NAME_ON_EQUIPMENT = 40

HEADERS = ['Enchant_ID', 'Filename', 'Vanilla_Name', 'Prefix', 'Suffix', 'On_Random_Equipment', 'QoL_Name', 'Edited',
           'Figure_Source', 'Candidate_Name', 'Candidate_Source', 'Note', 'Previous_Mod_Name', 'Previous_Mod_Check']
NUM_RE = re.compile(r'\d+(?:\.\d+)?')


def read_enchants():
    rows = []
    for fn, b in PckArchive(BACKUP_PATH).files:
        e = EnchantHelper.read_enchant(b)
        rows.append({'id': int.from_bytes(b[:4], 'little'), 'filename': fn, 'name': e['name'],
                     'prefix': e['prefix'], 'suffix': e['suffix']})
    return rows


def name_limit(row):
    return MAX_NAME_ON_EQUIPMENT if row['prefix'] != '0' else MAX_NAME


def load_previous_mod():
    if not os.path.exists(PREVIOUS_PATH):
        return {}
    ws = openpyxl.load_workbook(PREVIOUS_PATH, read_only=True)['Enchants']
    return {r[0]: r[4] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}


def load_edited():
    """Rows the user marked as edited keep their wording."""
    if not os.path.exists(OUTPUT_PATH):
        return {}
    ws = openpyxl.load_workbook(OUTPUT_PATH)['Enchants']
    header = [c.value for c in ws[1]]
    fn, name, edited = header.index('Filename'), header.index('QoL_Name'), header.index('Edited')
    return {r[fn]: r[name] for r in ws.iter_rows(min_row=2, values_only=True)
            if str(r[edited] or '').strip().lower() == 'yes'}


def check_previous(previous, vanilla, proposed):
    """How the earlier build's name stands against the guides."""
    if not previous or previous == vanilla:
        return 'unchanged'
    figures = NUM_RE.findall(previous)
    if not figures or figures == NUM_RE.findall(vanilla):
        return 'reworded, no new figure'
    if proposed is None:
        return 'figure without a source'
    name, source, _ = proposed
    return f"agrees with {source}" if figures == NUM_RE.findall(name) else f"contradicts {source} ({name})"


def build_rows():
    previous, edited = load_previous_mod(), load_edited()
    rows, stats = [], collections.Counter()
    for e in read_enchants():
        proposed = enchant_values.proposal(e['id'], e['name'])
        shown = proposed is not None and proposed[1] in enchant_values.SHOWN_SOURCES
        name = e['name']
        if e['filename'] in edited:
            name = edited[e['filename']]
            stats['edited by hand'] += 1
        elif shown:
            name = proposed[0]
            stats[f'figure from {proposed[1]}'] += 1
        elif proposed is not None:
            stats[f'candidate from {proposed[1]}, not shown (vanilla name kept)'] += 1
        elif NUM_RE.search(e['name']):
            stats['figure already in the vanilla name'] += 1
        else:
            stats['no figure known (vanilla name kept)'] += 1
        if len(name) > name_limit(e):
            raise SystemExit(f"{e['filename']}: '{name}' is longer than {name_limit(e)} characters")
        name.encode('latin-1')
        check = check_previous(previous.get(e['filename']), e['name'], proposed)
        stats[f"earlier build: {check.split(' (')[0]}"] += 1
        rows.append([
            e['id'], e['filename'], e['name'], e['prefix'], e['suffix'], 'yes' if e['prefix'] != '0' else '',
            name, 'yes' if e['filename'] in edited else '',
            proposed[1] if shown else '',
            proposed[0] if proposed and not shown else '', proposed[1] if proposed and not shown else '',
            proposed[2] if proposed else '',
            previous.get(e['filename'], ''), check,
        ])
    return rows, stats


def write_workbook(rows, stats):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Enchants'
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='305496')
    ws.freeze_panes = 'D2'
    for i, width in enumerate([10, 22, 44, 14, 18, 10, 44, 8, 12, 44, 12, 60, 44, 44], 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    summary = wb.create_sheet('Summary')
    summary.append(['What', 'Enchantments'])
    for key in sorted(stats):
        summary.append([key, stats[key]])
    summary.column_dimensions['A'].width = 70

    sources = wb.create_sheet('Sources')
    sources.append(['Source', 'What it is', 'Shown by the QoL variant'])
    for key, text in enchant_values.SOURCES.items():
        sources.append([key, text, 'yes' if key in enchant_values.SHOWN_SOURCES else 'no'])
    sources.column_dimensions['B'].width = 120

    how = wb.create_sheet('How to edit')
    for line in [
        'QoL_Name is the name the build writes into EnchantData.pck.',
        'To word an enchantment yourself: change QoL_Name and put "yes" in Edited. The generator keeps that row.',
        'To accept a candidate from the first game: copy Candidate_Name into QoL_Name and put "yes" in Edited.',
        f'Limits: {MAX_NAME_ON_EQUIPMENT} characters when On_Random_Equipment is yes (drawn in two columns in the '
        f'equipment box), {MAX_NAME} otherwise. They are the longest vanilla names, not a measurement.',
        'The binary holds no magnitude for enchantments: a figure is only as good as its source (Sources sheet).',
    ]:
        how.append([line])
    how.column_dimensions['A'].width = 130
    for sheet in wb:
        for row in sheet.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical='top', wrap_text=False)
    wb.save(OUTPUT_PATH)


def main():
    rows, stats = build_rows()
    write_workbook(rows, stats)
    print(f"Wrote {OUTPUT_PATH}: {len(rows)} enchantments")
    for key in sorted(stats):
        print(f"    {stats[key]:4d}  {key}")


if __name__ == '__main__':
    sys.exit(main())
