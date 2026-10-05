"""
Vanilla Source-of-Truth extractor.

Reads ONLY the clean backups in mod_project/backups/ and dumps every numeric
value the mod needs to describe into one read-only workbook:

    mod_project/data_tables/vanilla_source_of_truth.xlsx

The workbook is generated, never hand-edited. Regenerate it with:

    python mod_project/tools/extract_vanilla.py
    python mod_project/tools/extract_vanilla.py --check-guides   # also cross-check vs Data/*.mhtml

Each run also asserts the structural assumptions of the decoded layouts
(header sizes, effect counts) and exits with code 1 if any of them fails.
"""
import os
import re
import sys
import math
import struct
import argparse
import collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pck_handler import (
    PckArchive,
    SkillHelper,
    ItemHelper,
    EquipmentHelper,
    EnchantHelper,
    TrapHelper,
    EnemyAttackHelper,
    FoodMenuHelper,
    JoinSkillHelper,
    PartySkillHelper,
    SealScrollHelper,
    ClassDataHelper,
)

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKUP_DIR = os.path.join(BASE_DIR, 'backups')
TABLES_DIR = os.path.join(BASE_DIR, 'data_tables')
GUIDES_DIR = os.path.join(os.path.dirname(BASE_DIR), 'Data')
OUTPUT_PATH = os.path.join(TABLES_DIR, 'vanilla_source_of_truth.xlsx')

LEVELS = 10
SLOT_SIZE = 0x30  # effect slot: id (4) + 10 values (40) + tail (4)

# ---------------------------------------------------------------------------
# Decoded layouts. Offsets are relative to the start of each .bin entry.
# ---------------------------------------------------------------------------
SKILL_LAYOUTS = {
    'Technique': {
        'pck': 'TechniqueSkillData.pck', 'header': 0x1F0,
        'target': 0x28, 'attr': 0x2C, 'group': 0x30,
        'sp': 0x34, 'tp': 0x5C, 'turns': 0x84, 'speed': 0xAC,
        'weapon': (0xD4, 0xD8, 0xDC),
        'count': 0xE0, 'slots': 0xE4, 'n_slots': 5, 'trailer': 0x1D4,
    },
    'Magic': {
        'pck': 'MagicSkillData.pck', 'header': 0x238,
        'target': 0x28, 'attr': 0x2C, 'group': 0x30,
        'sp': 0x34, 'tp': 0x5C, 'turns': 0x84, 'speed': 0xAC,
        'cast': 0xD4, 'arr_0xFC': 0xFC, 'misc': 0x124,
        'count': 0x128, 'slots': 0x12C, 'n_slots': 5, 'trailer': 0x21C,
    },
    'Passive': {
        'pck': 'PassiveSkillData.pck', 'header': 0xD8,
        'flag': 0x28, 'sp': 0x2C, 'weapon1': 0x54,
        'count': 0x58, 'slots': 0x5C, 'n_slots': 2, 'trailer': 0xBC,
    },
}

# Verified against the "Target: X / Type: Y" header line of all 2375 enemy attacks.
TARGET_NAMES = {
    0: 'Single enemy', 1: 'Random enemies', 2: 'Enemy row', 3: 'Piercing',
    4: 'All enemies', 5: 'Single ally', 7: 'Ally row', 8: 'All allies',
    9: 'Self', 14: 'Indiscriminate',
}
ATTR_NAMES = {
    0: "Weapon's", 1: 'Slash', 2: 'Stab', 3: 'Bash', 4: 'Fire', 5: 'Water',
    6: 'Wind', 7: 'Earth', 8: 'Lightning', 9: 'Neutral',
}
# Verified against the "<Weapon> skill." prefix of vanilla technique descriptions
# and the weapon mastery passives.
WEAPON_NAMES = {
    0: 'Short-range weapon', 1: 'Sword', 2: 'Dagger', 3: 'Axe', 4: 'Lance',
    5: 'Mace', 7: 'Tome', 8: 'Bow', 9: 'Katana', 10: 'Projectile', 11: 'Whip',
    12: 'Fan', 15: 'Gun', 16: 'Shield',
}
EQUIP_CATEGORIES = {0: 'Weapon', 1: 'Shield', 2: 'Helm', 3: 'Armor', 4: 'Accessory'}
ITEM_EFFECT_NAMES = {
    0: 'Restore HP (one ally)', 1: 'Restore TP (one ally)', 3: 'Restore HP (party)',
    6: 'Cure status ailment', 8: 'Revive', 9: 'Light (steps)', 10: 'Escape dungeon',
}

# effect id -> (name, status, evidence). Only ids with evidence are named here;
# every other id is listed in the Effects sheet with the skills that use it.
EFFECT_NAMES = {
    1: ('Magic power', 'Verified', 'GameFAQs: Fireball 70-520, Crimson Flare 620-1300'),
    26: ('Physical power modifier (damage = 100% + value)', 'Verified',
         'GameFAQs: Halving Slice 80-110%, Megaton Press 120-184%, 4 more'),
    29: ('Accuracy modifier', 'By vanilla text', 'Varies in every technique whose LvUp line says "Accuracy UP"'),
    74: ('Dispel', 'By usage', 'Aura Eliminate, Dispel'),
    75: ('HP drain', 'Verified', 'GameFAQs: Blood Drain 15-25% life steal'),
    76: ('TP drain', 'Verified', 'GameFAQs: Soul Drain 10-20% TP steal'),
    103: ('Heal power', 'By usage', 'Heal, Circle Heal'),
    113: ('Poison', 'Verified', 'TrapData: Poison Gas'),
    114: ('Freeze', 'Verified', 'TrapData: Freeze Gas'),
    115: ('Paralysis', 'Verified', 'TrapData: Paralysis Gas'),
    116: ('Sleep', 'Verified', 'TrapData: Sleep Gas'),
    117: ('Silence', 'Verified', 'TrapData: Silence Gas'),
    118: ('Stun', 'By usage', 'Stun Sting, Shield Bash'),
    119: ('Instant death', 'Verified', 'GameFAQs: Halving Slice 10-20% instant death'),
    120: ('Petrification', 'Verified', 'TrapData: Petrification Gas'),
    122: ('Slow', 'Verified', 'TrapData: Slow Gas'),
    137: ('Number of targets', 'Verified', 'GameFAQs: Megaton Press 2-5 targets'),
    162: ('DEF modifier', 'Verified', 'GameFAQs: Armor Penetration 16-25% PDEF down; Protect'),
    169: ('All-element resistance', 'By usage', 'Moon Curtain'),
    170: ('Aggro', 'By usage', 'Provoke, Attract'),
    295: ('Heal cap', 'Verified', 'GameFAQs: Heal 100-800 healing cap'),
    # Decoded from the game's own text: an enemy attack description states the effect, and where it
    # gives a figure the figure is the stored value.
    0: ('Plain hit that deals damage', 'By usage', 'Present in 1,179 enemy attacks; attacks that say "Deals no damage" carry 25 instead'),
    8: ('Fixed damage', 'By game text', 'Value 666 <-> "does 666 damage"; 1111 <-> "Deal 1111 fixed damage"'),
    25: ('Deals no damage', 'By game text', '"Deals no damage, but lowers physical accuracy" (Slapyon Lotion and 100 more)'),
    31: ('Critical rate bonus, in points', 'By game text', 'Last Stand and Divine Cannon say "Critical rate UP"; Moonlight Purgatory 10 <-> "high critical hit rate"'),
    32: ('Vulnerability to critical hits, in points', 'By vanilla text', 'Last Stand only: "Enemy critical rate UP", next to 31'),
    47: ('Damage based on RES', 'By game text', '"Righteous damage derived from RES, not ATK" (Bible Thump), "RES-based damage, not INT" (Prism Dawn)'),
    54: ('Chance to push the target back in the turn order', 'By game text', 'Value 0.5-1.0 <-> "Sends the target straight into next week... in the turn order"; Tinker Squash "May delay enemy turn"'),
    55: ('Fixed accuracy', 'By game text', 'Value 0.5 <-> "With 50% accuracy", 0.65 <-> "With 65% accuracy", 1.0 <-> "guaranteed to hit home"'),
    69: ('Ignores defense', 'By game text', '30 enemy attacks: "ignoring defense"; Armor Penetration "defense-piercing"'),
    70: ('Fixed critical hit rate', 'By game text', 'Value 0.35 / 0.5 / 0.75 / 1.0 <-> "fixed critical hit rate of 35% / 50% / 75% / 100%" (Divine Cannon)'),
    71: ('More damage to targets with buffs or debuffs', 'By game text', '"Deal more damage to targets with buffs or debuffs" (10 enemy attacks); Pursuit Stab'),
    72: ('Certain to hit', 'By vanilla text', 'Guided Shot only: "certain to hit"'),
    73: ('Damage multiplier', 'By game table', 'CommandData names it "Damage xn"; enemy text: "deal several times ATK in damage" (Beat the Odds 2.6). How it combines with the power table is not measured'),
    79: ('Chance of the speed decrease', 'By game text', 'Value 0.25-0.75 <-> "May reduce the targets\' turn speed" (Phoenix, Coin Toss)'),
    80: ('Pep change', 'By game text', '"bums everyone out", "Saps the enemy\'s pep as well" (value -5 to -15)'),
    87: ('Only works on living targets', 'By game text', '"they\'ll only cut living foes" (13 enemy attacks); Moonlight Purgatory, Soul Eater'),
    88: ('Bonus damage to undead, demons and apparitions', 'By game text', '"Does more damage to the undead, demons, and apparitions" (Cross Slash 0.35-0.5)'),
    90: ('Hits three times', 'By game text', '24 enemy attacks: "thrice", "three times"'),
    92: ('Hits 2-4 times', 'By game text', '58 enemy attacks: "2-4 flaps", "Attack all enemies 2-4 times"'),
    93: ('Number of hits', 'By game text', 'Value 5 <-> "Dispense five punches"; Chain Spike "Number of attacks UP" 3 -> 5'),
    94: ('Hits 0-4 times', 'By game text', '17 enemy attacks: "0-4 rapid thrusts", "Attack 0-4 times"'),
    95: ('Hits 1-5 times', 'By game text', '8 enemy attacks: "1-5 potshots", "Slash away at one foe 1-5 times"'),
    96: ('Hits 3-5 times', 'By game text', '29 enemy attacks: "Attack 3-5 times"'),
    97: ('Hits twice', 'By game text', '63 enemy attacks: "Attack twice with bared claws"'),
    99: ('Hits 1-3 times', 'By game text', '47 enemy attacks: "Rend foes 1-3 times"'),
    105: ('HP on revival, as a fraction of max HP', 'By vanilla text', 'Primevil: "Amount of HP on revival UP" 0.1 -> 0.5'),
    106: ('How many ailments of the listed order it cures', 'By vanilla text', 'Cleanse 1-5 <-> "Can cure poison > paralysis > silence > freeze > sleep"'),
    110: ('Crits against Beasts', 'By vanilla text', 'Beast Binder only'),
    111: ('Bonus with a Heavy Lance', 'By vanilla text', 'Daedalus Attack 0.2; a player test on the DT2-2 wiki measures about +20%'),
    123: ('Shadow Sew chance', 'By game text', 'Value 0.5 / 0.75 / 1.0 on attacks that "may inflict Shadow Sew"'),
    125: ('Critical rate grows with each hit', 'By game text', '"The later blows are weaker, but more likely to crit" (Chain Spike)'),
    133: ('Max HP reduction; unit not known', 'Partly', 'Life Bringer only: "Max HP reduction UP" 0.5 -> 1.0'),
    136: ('Damage change for each later target or hit', 'By game text', 'Value -0.05 to -0.25 <-> "more softly on later hits", "damage will decrease towards the end of the row"'),
    138: ('Damage as a fraction of the target max HP', 'By game text', '"Deals more damage to enemies with more max HP"; always next to a cap (296)'),
    146: ('Number of units summoned', 'By game text', 'Value 2 <-> "Summon not one but two magical helpers"'),
    147: ('Turns a summon lasts', 'By vanilla text', 'Rises where the summon skills say "Duration UP"'),
    161: ('ATK modifier', 'By game text', '"Lowers target\'s physical attack" (Dragon Claw); Fire Scroll "ATK decrease UP"'),
    163: ('Magic defense modifier', 'By game text', '"lower magical defense" (Anima Steal); Water Scroll "Magic defense decrease UP"'),
    164: ('Physical evasion modifier', 'By vanilla text', 'Evade: "Increase own physical evasion rate", its only effect'),
    165: ('Magic evasion modifier', 'By vanilla text', 'Warm-Up: "physical evasion rate and magical evasion rate" = 164 + 165'),
    166: ('Physical accuracy modifier', 'By game text', '"may rob them of their physical accuracy" (Spinning Eye, 166 alone)'),
    167: ('Magic accuracy modifier', 'By game text', '"lowering physical and magic accuracy" (Cherubeam: 166 + 167)'),
    168: ('Speed modifier', 'By game text', '"decrease their turn speed" (It\'s a Road Roller!); Fleet Foot "Speed increase UP"'),
    191: ('Damage multiplier of the next attack', 'By vanilla text', 'Power Charge only: "Damage UP" 1.05 -> 1.3'),
    199: ('Damage every turn as a fraction of the target max HP', 'By game text', '"More damage to higher-HP foes" (Debilitating Salsa), next to a cap (296)'),
    272: ('HP restored after a battle, as a fraction', 'By vanilla text', 'Generous Heart: "Post-battle HP & TP recovery UP", first table'),
    273: ('TP restored after a battle, as a fraction', 'By vanilla text', 'Generous Heart, second table'),
    296: ('Damage cap', 'By game text', 'Value 200 / 300 / 400 / 500 <-> "(Cap: 200)" ... "(Cap: 500)" (Fire Scroll)'),
}

# (Sheet, Column, PCK, Offset, Type, Meaning, Status, Evidence)
FIELD_DOCS = [
    ('SkillLevels', 'sp_cost', 'Technique/Magic 0x34, Passive 0x2C', '10 x int',
     'Skill points needed for each level. Number of non-zero entries = max level.', 'Inferred',
     '1..10 on 10-level skills, 2,4..10 on 5-level skills'),
    ('SkillLevels', 'tp_cost', 'Technique/Magic 0x5C', '10 x int', 'TP cost per level', 'Verified',
     'GameFAQs TP ranges'),
    ('SkillLevels', 'turns', 'Technique/Magic 0x84', '10 x int', 'Duration in turns per level', 'Verified',
     'GameFAQs: Attract 3-7 turns, Insurance 6-10 turns'),
    ('SkillLevels', 'speed', 'Technique/Magic 0xAC', '10 x float', 'Turn speed multiplier after acting', 'Verified',
     'GameFAQs speed values'),
    ('SkillLevels', 'cast', 'Magic 0xD4', '10 x float', 'Chant time multiplier', 'Verified',
     'GameFAQs cast values'),
    ('SkillLevels', 'arr_0xFC', 'Magic 0xFC', '10 x float', 'Magic accuracy', 'By vanilla text',
     'Rises with level in 30 of the 35 spells whose LvUp line says "Accuracy UP"'),
    ('SkillLevels', 'effect_N', 'Technique 0xE4, Magic 0x12C, Passive 0x5C (+0x30 per slot)',
     'id + 10 values + tail', 'Effect slot. Effect_ID is resolved in the Effects sheet.', 'Verified',
     'Effect count field equals used slots in all 312 skills'),
    ('Skills', 'Target_Code', 'Technique/Magic 0x28', 'int', 'Targeting', 'Verified',
     'Same codes as enemy attack header lines'),
    ('Skills', 'Attr_Code', 'Technique/Magic 0x2C', 'int', 'Damage type / element', 'Verified',
     'Same codes as enemy attack header lines'),
    ('Skills', 'Weapon_Codes', 'Technique 0xD4-0xDC, Passive 0x54', 'int', 'Required weapon type(s)', 'Verified',
     '"<Weapon> skill." prefix of vanilla descriptions'),
    ('Skills', 'Misc_0x124', 'Magic 0x124', 'float', 'Unknown single value', 'Unknown', ''),
    ('Skills', 'Trailer', 'last 7 words of the header', 'int', 'Unknown', 'Unknown', ''),
    ('UniqueSkills', 'Lv1..Lv9, EX', 'UniqueSkillData 0x1C', '10 x float', 'Value per unique skill level; 10th is EX',
     'Inferred', 'Lesson Sp03: unique skill reaches EX after Lv 9'),
    ('Items', 'Buy_Price', 'ItemData 0x1C', 'int', 'Shop price', 'Verified',
     'Vanilla Rin-Rin lines: Danish 450, Fairy Apple 900'),
    ('Items', 'Sell_Price', 'ItemData 0x20', 'int', 'Sell value (buy / 3)', 'Inferred', ''),
    ('Items', 'EffN_Type / EffN_Value', 'ItemData 0x24 (+8 per pair)', 'int + float',
     'Up to 4 effect pairs. Types: ' + ', '.join(f'{k}={v}' for k, v in ITEM_EFFECT_NAMES.items()),
     'Partly verified', 'Vanilla descriptions; pairs 3-4 not understood'),
    ('Equipment', 'Category_Code', 'EquipmentData 0x20', 'int', 'Equipment slot', 'Verified', 'Unidentified name string'),
    ('Equipment', 'Rank_Idx', 'EquipmentData 0x4C', 'int', 'Index into WeaponRanks (-1 for non-weapons)', 'Verified',
     'Range 0-11 on weapons only'),
    ('Equipment', 'Enchant_Count / EnchantN_ID', 'EquipmentData 0x5C / 0x60-0x6C', 'int',
     'Innate enchantments (ids into the Enchants sheet)', 'Verified', 'Count equals number of ids in all 551 entries'),
    ('Equipment', 'w_0x24..w_0x58', 'EquipmentData', 'int/float', 'Stats and price, not individually identified',
     'Unknown', '0x24 looks like ATK, 0x58 like price'),
    ('WeaponRanks', 'Delay_Mult', 'WeaponRankData 0x10', 'float', 'Turn delay multiplier of the weapon rank', 'Verified', ''),
    ('Traps', 'Status_ID / Rate', 'TrapData 0x10 / 0x14', 'int / float', 'Effect id and rate of the trap', 'Verified',
     'Trap names match effect ids'),
    ('QuestRanks', 'Reward_Type / Value', 'QuestRankData 0x10 / 0x24', 'int', 'Reward kind and amount', 'Verified',
     'Reward messages'),
    ('ExpCurve', 'Req_EXP / SP', 'LvExpSkptData 0x04 / 0x08', 'int', 'EXP needed for the level and SP granted', 'Verified', ''),
    ('EnemyAttacks', 'Power_0x44', 'EnemyAttackData 0x44', 'float', 'Physical power multiplier', 'Inferred', ''),
    ('EnemyAttacks', 'EffN_ID / Value / Tail', 'EnemyAttackData 0x4C (+12 per slot)', 'int + number + int',
     'Effect slots (same effect ids as skills)', 'Verified', 'Effect count equals used slots in all 2375 entries'),
    ('FoodMenu', 'EffN_ID / Value', 'FoodMenuData 0x20 (+8 per pair)', 'int + float', 'Meal effects', 'Inferred', ''),
    ('Raw_*', 'w_0x..', 'JoinSkill / PartySkill / SealScroll / Class', 'int/float',
     'Header words not decoded yet (string pointers included)', 'Unknown', ''),
]


# ---------------------------------------------------------------------------
# Low-level readers
# ---------------------------------------------------------------------------
def i32(b, off):
    return struct.unpack_from('<i', b, off)[0]


def f32(b, off):
    v = struct.unpack_from('<f', b, off)[0]
    return round(v, 4) if math.isfinite(v) else None


def looks_float(raw):
    # Integers stored in these tables are small; any larger word is an IEEE float.
    return abs(raw) > 0xFFFF


def word(b, off):
    """One header word as int, or as float when it only makes sense as one."""
    raw = i32(b, off)
    if looks_float(raw):
        f = f32(b, off)
        if f is not None and 1e-4 < abs(f) < 1e7:
            return f
    return raw


def read_array(b, off, n=LEVELS):
    raw = [i32(b, off + 4 * k) for k in range(n)]
    if any(looks_float(v) for v in raw):
        return [v if v == -1 else f32(b, off + 4 * k) for k, v in enumerate(raw)], 'float'
    return raw, 'int'


def cstr(b, off):
    return b[off:b.find(b'\x00', off)].decode('latin-1', 'replace')


def load(pck_name):
    return PckArchive(os.path.join(BACKUP_DIR, pck_name)).files


def leading_nonzero(values):
    n = 0
    for v in values:
        if not v:
            break
        n += 1
    return n


# ---------------------------------------------------------------------------
# Skills
# ---------------------------------------------------------------------------
def read_skill_entry(category, fn, b, problems):
    lay = SKILL_LAYOUTS[category]
    info = SkillHelper.read_skill(b, category)
    if i32(b, 8) != lay['header']:
        problems.append(f"{fn}: header size {i32(b, 8):#x}, expected {lay['header']:#x}")

    sp, _ = read_array(b, lay['sp'])
    arrays = collections.OrderedDict()
    arrays['sp_cost'] = (sp, 'int')
    for key, name in (('tp', 'tp_cost'), ('turns', 'turns'), ('speed', 'speed'),
                      ('cast', 'cast'), ('arr_0xFC', 'arr_0xFC')):
        if key in lay:
            arrays[name] = read_array(b, lay[key])

    count = i32(b, lay['count'])
    slots = []
    for k in range(lay['n_slots']):
        off = lay['slots'] + k * SLOT_SIZE
        values, vtype = read_array(b, off + 4)
        slots.append({'id': i32(b, off), 'values': values, 'type': vtype, 'tail': i32(b, off + 4 + 4 * LEVELS)})
    used = sum(1 for s in slots if not (s['id'] == -1 and all(v == -1 for v in s['values'])))
    if used != count:
        problems.append(f"{fn} ({info['name']}): effect count {count} but {used} slots in use")

    levels = leading_nonzero(sp) or leading_nonzero(arrays['speed'][0] if 'speed' in arrays else []) or 1

    weapon_codes = []
    if 'weapon' in lay:
        n_off, w1, w2 = lay['weapon']
        weapon_codes = [i32(b, o) for o in (w1, w2)][:max(i32(b, n_off), 0)]
    elif 'weapon1' in lay and i32(b, lay['weapon1']) != -1:
        weapon_codes = [i32(b, lay['weapon1'])]

    return {
        'category': category, 'fn': fn, 'id': i32(b, 0), 'name': info['name'],
        'desc': info['desc'], 'lvup': info['lvup'], 'levels': levels,
        'target': i32(b, lay['target']) if 'target' in lay else None,
        'attr': i32(b, lay['attr']) if 'attr' in lay else None,
        'group': i32(b, lay['group']) if 'group' in lay else None,
        'flag': i32(b, lay['flag']) if 'flag' in lay else None,
        'misc': f32(b, lay['misc']) if 'misc' in lay else None,
        'weapon_codes': weapon_codes,
        'arrays': arrays, 'count': count, 'slots': slots[:count],
        'trailer': [i32(b, lay['trailer'] + 4 * k) for k in range(7)],
    }


def read_all_skills(problems):
    skills = []
    for category, lay in SKILL_LAYOUTS.items():
        for fn, b in load(lay['pck']):
            skills.append(read_skill_entry(category, fn, b, problems))
    return skills


def effect_name(eff_id):
    return EFFECT_NAMES.get(eff_id, ('', '', ''))[0]


def pad_levels(values, levels):
    # Past the max level only padding (0 / -1) is hidden; real data is always shown.
    return [None if k >= levels and v in (0, -1, None) else v for k, v in enumerate(values)]


def build_skill_sheets(skills):
    skill_rows, level_rows = [], []
    for s in skills:
        skill_rows.append([
            s['category'], s['fn'], s['id'], s['name'], s['levels'],
            s['target'], TARGET_NAMES.get(s['target'], ''),
            s['attr'], ATTR_NAMES.get(s['attr'], ''),
            s['group'],
            ' / '.join(WEAPON_NAMES.get(c, f'#{c}') for c in s['weapon_codes']),
            ', '.join(str(c) for c in s['weapon_codes']),
            s['misc'], s['flag'], s['count'],
            ', '.join(str(sl['id']) for sl in s['slots']),
            ', '.join(str(v) for v in s['trailer']),
            s['desc'], s['lvup'],
        ])
        base = [s['category'], s['fn'], s['id'], s['name'], s['levels']]
        for field, (values, vtype) in s['arrays'].items():
            if field in ('turns', 'cast', 'arr_0xFC') and not any(values):
                continue
            level_rows.append(base + [field, None, '', None, vtype] + pad_levels(values, s['levels']))
        for k, sl in enumerate(s['slots'], 1):
            level_rows.append(base + [f'effect_{k}', sl['id'], effect_name(sl['id']), sl['tail'], sl['type']]
                              + pad_levels(sl['values'], s['levels']))
    return skill_rows, level_rows


def build_unique_rows():
    rows = []
    for fn, b in load('UniqueSkillData.pck'):
        info = SkillHelper.read_skill(b, 'Unique')
        values, _ = read_array(b, 0x1C)
        rows.append([fn, i32(b, 0), info['name'], i32(b, 0x18)] + values
                    + [i32(b, o) for o in range(0x44, 0x60, 4)] + [info['desc']])
    return rows


# ---------------------------------------------------------------------------
# Other tables
# ---------------------------------------------------------------------------
def build_enchant_rows():
    rows, names = [], {}
    for fn, b in load('EnchantData.pck'):
        e = EnchantHelper.read_enchant(b)
        names[i32(b, 0)] = e['name']
        rows.append([fn, i32(b, 0), e['name'], e['prefix'], e['suffix']] + [word(b, o) for o in range(0x20, 0x38, 4)])
    return rows, names


def build_weapon_rank_rows():
    rows, ranks = [], {}
    for fn, b in load('WeaponRankData.pck'):
        ranks[i32(b, 0)] = (cstr(b, i32(b, 8)), f32(b, 0x10))
        rows.append([fn, i32(b, 0), cstr(b, i32(b, 8)), f32(b, 0x10)])
    return rows, ranks


def build_item_rows():
    rows = []
    for fn, b in load('ItemData.pck'):
        it = ItemHelper.read_item(b)
        row = [fn, i32(b, 0), it['name'], i32(b, 0x18), i32(b, 0x1C), i32(b, 0x20)]
        for k in range(4):
            eff_type = i32(b, 0x24 + 8 * k)
            row += [eff_type, f32(b, 0x28 + 8 * k) if eff_type != -1 else None]
        row += [i32(b, o) for o in range(0x44, 0x60, 4)] + [it['desc']]
        rows.append(row)
    return rows


def build_equipment_rows(enchant_names, ranks, problems):
    rows = []
    for fn, b in load('EquipmentData.pck'):
        e = EquipmentHelper.read_equipment(b)
        rank_idx = i32(b, 0x4C)
        rank_label, rank_delay = ranks.get(rank_idx, ('', None))
        ids = [i32(b, o) for o in range(0x60, 0x70, 4)]
        if i32(b, 0x5C) != sum(1 for v in ids if v != -1):
            problems.append(f"{fn} ({e['name']}): enchant count {i32(b, 0x5C)} does not match ids {ids}")
        row = [fn, i32(b, 0), e['name'], e['s2'], i32(b, 0x20), EQUIP_CATEGORIES.get(i32(b, 0x20), '')]
        row += [i32(b, o) for o in range(0x24, 0x3C, 4)]
        row += [f32(b, 0x3C), f"{i32(b, 0x40) & 0xFFFFFFFF:#010x}", i32(b, 0x44), f32(b, 0x48)]
        row += [rank_idx, rank_label, rank_delay, i32(b, 0x50), i32(b, 0x54), i32(b, 0x58), i32(b, 0x5C)]
        for v in ids:
            row += [v if v != -1 else None, enchant_names.get(v, '') if v != -1 else '']
        row.append(e['desc'])
        rows.append(row)
    return rows


def build_trap_rows():
    rows = []
    for fn, b in load('TrapData.pck'):
        status = i32(b, 0x10)
        rows.append([fn, i32(b, 0), TrapHelper.read_trap(b), status if status != -1 else None,
                     effect_name(status), f32(b, 0x14), i32(b, 0x18), i32(b, 0x1C)])
    return rows


def build_quest_rank_rows():
    rows = []
    for fn, b in load('QuestRankData.pck'):
        w = struct.unpack_from('<10i', b, 0)
        rows.append([fn, w[0], w[1], w[4], w[5], w[6], w[7], w[8], w[9], cstr(b, w[2])])
    return rows


def build_exp_rows():
    return [[fn] + list(struct.unpack_from('<3I', b, 0)) for fn, b in load('LvExpSkptData.pck')]


def build_enemy_attack_rows(problems):
    rows, uses = [], collections.Counter()
    for fn, b in load('EnemyAttackData.pck'):
        e = EnemyAttackHelper.read_enemy_attack(b)
        count = i32(b, 0x48)
        slots = [(i32(b, 0x4C + 12 * k), word(b, 0x50 + 12 * k), i32(b, 0x54 + 12 * k)) for k in range(5)]
        if count != sum(1 for s in slots if s[0] != -1):
            problems.append(f"{fn} ({e['name']}): effect count {count} does not match slots")
        row = [fn, i32(b, 0), e['name'], e['desc'].split('\n')[0], i32(b, 0x28),
               i32(b, 0x2C), TARGET_NAMES.get(i32(b, 0x2C), ''), i32(b, 0x30), ATTR_NAMES.get(i32(b, 0x30), ''),
               f32(b, 0x34), f32(b, 0x38), word(b, 0x3C), word(b, 0x40), f32(b, 0x44), count]
        for k, (eff_id, value, tail) in enumerate(slots):
            if k < count:
                uses[eff_id] += 1
                row += [eff_id, effect_name(eff_id), value, tail]
            else:
                row += [None, '', None, None]
        row += [', '.join(str(i32(b, o)) for o in range(0x88, 0xA0, 4)), e['desc']]
        rows.append(row)
    return rows, uses


def build_food_rows():
    rows = []
    for fn, b in load('FoodMenuData.pck'):
        f = FoodMenuHelper.read_food_menu(b)
        row = [fn, i32(b, 0), f['name'], i32(b, 0x18), i32(b, 0x1C)]
        for k in range(5):
            eff_id = i32(b, 0x20 + 8 * k)
            row += [eff_id if eff_id != -1 else None, f32(b, 0x24 + 8 * k) if eff_id != -1 else None]
        row += [i32(b, 0x48), i32(b, 0x4C), f['flavor'], f['effect']]
        rows.append(row)
    return rows


def build_raw_rows(pck_name, header_size, name_reader):
    """Undecoded PCKs: one column per header word so nothing is hidden."""
    rows, width = [], 0
    for fn, b in load(pck_name):
        size = header_size if header_size else struct.unpack_from('<Q', b, 8)[0]
        words = [word(b, o) for o in range(0, size, 4)]
        width = max(width, len(words))
        rows.append([fn, name_reader(b)] + words)
    return ['Filename', 'Name'] + [f'w_{4 * k:#04x}' for k in range(width)], rows


def read_game_effect_names():
    """{effect id: name} from CommandData.pck, the game's own table of effects (one entry per effect id:
    int id, then four name pointers and the battle messages). 317 entries carry a name."""
    names = {}
    for _, b in load('CommandData.pck'):
        pointer = struct.unpack_from('<Q', b, 8)[0]
        name = bytes(b[pointer:b.find(b'\x00', pointer)]).decode('latin-1') if 0 < pointer < len(b) else ''
        if name:
            names[i32(b, 0)] = name
    return names


def build_effect_rows(skills, enemy_uses):
    used_by = collections.defaultdict(list)
    for s in skills:
        for sl in s['slots']:
            used_by[sl['id']].append(s['name'])
    game_names = read_game_effect_names()
    rows = []
    for eff_id in sorted(set(used_by) | set(enemy_uses) | set(EFFECT_NAMES)):
        name, status, evidence = EFFECT_NAMES.get(eff_id, ('', 'Unnamed', ''))
        if status == 'Unnamed' and eff_id in game_names:
            status = 'Named by the game only'
        skill_names = used_by.get(eff_id, [])
        rows.append([eff_id, game_names.get(eff_id, ''), name, status, evidence, len(skill_names),
                     enemy_uses.get(eff_id, 0), ', '.join(dict.fromkeys(skill_names))[:200]])
    return rows


# ---------------------------------------------------------------------------
# Workbook output
# ---------------------------------------------------------------------------
def add_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    hdr_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    hdr_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = 'A2'
    for col_num, header in enumerate(headers, 1):
        sample = [len(str(r[col_num - 1])) for r in rows[:200] if col_num <= len(r) and r[col_num - 1] is not None]
        ws.column_dimensions[get_column_letter(col_num)].width = min(max([len(str(header))] + sample) + 2, 60)
    return ws


def write_workbook(problems):
    skills = read_all_skills(problems)
    skill_rows, level_rows = build_skill_sheets(skills)
    enchant_rows, enchant_names = build_enchant_rows()
    rank_rows, ranks = build_weapon_rank_rows()
    enemy_rows, enemy_uses = build_enemy_attack_rows(problems)
    lv_cols = [f'Lv{k}' for k in range(1, LEVELS + 1)]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    add_sheet(wb, 'About', ['Topic', 'Detail'], [
        ['What', 'Numeric values of the unmodified game, extracted from mod_project/backups/*.pck'],
        ['Rule', 'Generated file. Do not edit by hand; run tools/extract_vanilla.py to regenerate.'],
        ['Use', 'QoL texts must quote these values. Balance texts quote these values plus their overrides.'],
        ['Confidence', 'See the Fields sheet: Verified / Inferred / Unknown for every decoded column.'],
        ['Effects', 'Effect ids are shared by skills, enemy attacks and traps; see the Effects sheet.'],
    ])
    add_sheet(wb, 'Fields', ['Sheet', 'Column', 'PCK / Offset', 'Type', 'Meaning', 'Status', 'Evidence'], FIELD_DOCS)
    add_sheet(wb, 'Skills', [
        'Category', 'Filename', 'Skill_ID', 'Name', 'Levels', 'Target_Code', 'Target', 'Attr_Code', 'Attr',
        'Group_0x30', 'Weapon_Req', 'Weapon_Codes', 'Misc_0x124', 'Flag_0x28', 'Effect_Count', 'Effect_IDs',
        'Trailer', 'Vanilla_Desc', 'Vanilla_LvUp'], skill_rows)
    add_sheet(wb, 'SkillLevels', [
        'Category', 'Filename', 'Skill_ID', 'Name', 'Levels', 'Field', 'Effect_ID', 'Effect_Name', 'Slot_Tail',
        'Value_Type'] + lv_cols, level_rows)
    add_sheet(wb, 'UniqueSkills', ['Filename', 'Skill_ID', 'Name', 'Kind_0x18'] + lv_cols[:9] + ['EX']
              + [f'w_{o:#04x}' for o in range(0x44, 0x60, 4)] + ['Vanilla_Desc'], build_unique_rows())
    add_sheet(wb, 'Effects', ['Effect_ID', 'Game_Name', 'Name', 'Status', 'Evidence', 'Skill_Uses', 'EnemyAttack_Uses', 'Used_By'],
              build_effect_rows(skills, enemy_uses))
    add_sheet(wb, 'Items', ['Filename', 'Item_ID', 'Name', 'Flags_0x18', 'Buy_Price', 'Sell_Price',
                            'Eff1_Type', 'Eff1_Value', 'Eff2_Type', 'Eff2_Value', 'Eff3_Type', 'Eff3_Value',
                            'Eff4_Type', 'Eff4_Value'] + [f'w_{o:#04x}' for o in range(0x44, 0x60, 4)]
              + ['Vanilla_Desc'], build_item_rows())
    add_sheet(wb, 'Equipment', ['Filename', 'Equip_ID', 'Name', 'Unidentified_Name', 'Category_Code', 'Category']
              + [f'w_{o:#04x}' for o in range(0x24, 0x3C, 4)]
              + ['f_0x3c', 'mask_0x40', 'w_0x44', 'f_0x48', 'Rank_Idx', 'Rank', 'Delay_Mult', 'Type_Code', 'w_0x54',
                 'w_0x58', 'Enchant_Count', 'Enchant1_ID', 'Enchant1_Name', 'Enchant2_ID', 'Enchant2_Name',
                 'Enchant3_ID', 'Enchant3_Name', 'Enchant4_ID', 'Enchant4_Name', 'Vanilla_Desc'],
              build_equipment_rows(enchant_names, ranks, problems))
    add_sheet(wb, 'Enchants', ['Filename', 'Enchant_ID', 'Name', 'Prefix', 'Suffix']
              + [f'w_{o:#04x}' for o in range(0x20, 0x38, 4)], enchant_rows)
    add_sheet(wb, 'WeaponRanks', ['Filename', 'Rank_Idx', 'Rank', 'Delay_Mult'], rank_rows)
    add_sheet(wb, 'Traps', ['Filename', 'Trap_ID', 'Name', 'Status_ID', 'Status_Name', 'Rate', 'w_0x18', 'w_0x1c'],
              build_trap_rows())
    add_sheet(wb, 'QuestRanks', ['Filename', 'Rank', 'Trust_Points', 'Reward_Type', 'Item_ID', 'w_0x18', 'w_0x1c',
                                 'w_0x20', 'Value', 'Message'], build_quest_rank_rows())
    add_sheet(wb, 'ExpCurve', ['Filename', 'Level', 'Req_EXP', 'SP'], build_exp_rows())
    enemy_headers = ['Filename', 'Skill_ID', 'Name', 'Header_Line', 'Anim_0x28', 'Target_Code', 'Target',
                     'Attr_Code', 'Attr', 'f_0x34', 'f_0x38', 'w_0x3c', 'w_0x40', 'Power_0x44', 'Effect_Count']
    for k in range(1, 6):
        enemy_headers += [f'Eff{k}_ID', f'Eff{k}_Name', f'Eff{k}_Value', f'Eff{k}_Tail']
    add_sheet(wb, 'EnemyAttacks', enemy_headers + ['Trailer', 'Vanilla_Desc'], enemy_rows)
    food_headers = ['Filename', 'Food_ID', 'Name', 'Price', 'Effect_Count']
    for k in range(1, 6):
        food_headers += [f'Eff{k}_ID', f'Eff{k}_Value']
    add_sheet(wb, 'FoodMenu', food_headers + ['Tier_0x48', 'Shop_0x4c', 'Flavor', 'Vanilla_Effect'], build_food_rows())
    for title, pck_name, header_size, reader in (
        ('Raw_JoinSkill', 'JoinSkillData.pck', 0x98, lambda b: JoinSkillHelper.read_joint(b)['name']),
        ('Raw_PartySkill', 'PartySkillData.pck', None, lambda b: PartySkillHelper.read_party_skill(b)['name']),
        ('Raw_SealScroll', 'SealScrollData.pck', 0x48, lambda b: SealScrollHelper.read_seal(b)['name']),
        ('Raw_Class', 'ClassData.pck', 0x28, lambda b: ClassDataHelper.read_class(b)['name']),
    ):
        headers, rows = build_raw_rows(pck_name, header_size, reader)
        add_sheet(wb, title, headers, rows)

    wb.save(OUTPUT_PATH)
    return wb, skills


# ---------------------------------------------------------------------------
# Cross-check against the GameFAQs guides saved in Data/*.mhtml
# ---------------------------------------------------------------------------
def guide_lines(path):
    import email
    import html
    with open(path, encoding='utf-8', errors='replace') as f:
        msg = email.message_from_file(f)
    text = ''
    for part in msg.walk():
        if part.get_content_type() == 'text/html':
            h = part.get_payload(decode=True).decode('utf-8', 'replace')
            h = re.sub(r'<(script|style)[\s\S]*?</\1>', '', h)
            h = re.sub(r'<br\s*/?>|</p>|</div>|</tr>|</h\d>|</pre>|</li>', '\n', h)
            text += html.unescape(re.sub(r'<[^>]+>', '', h)) + '\n'
    return [l.strip() for l in text.split('\n') if l.strip()]


def parse_range(pattern, line):
    m = re.search(pattern, line)
    if not m:
        return None
    lo = float(m.group(1))
    return lo, float(m.group(2)) if m.group(2) else lo


def check_guides(skills):
    files = ['Fighter.mhtml', 'Mage.mhtml', 'Scout.mhtml', 'Maid.mhtml', 'Spieler.mhtml']
    by_name = collections.defaultdict(list)
    for s in skills:
        by_name[s['name']].append(s)

    def ends(values, levels):
        return values[0], values[levels - 1]

    stats = collections.defaultdict(lambda: [0, 0])  # field -> [match, mismatch]
    mismatches = []
    for guide in files:
        path = os.path.join(GUIDES_DIR, guide)
        if not os.path.exists(path):
            print(f"    [SKIP] {guide} not found in {GUIDES_DIR}")
            continue
        lines = guide_lines(path)
        for idx, line in enumerate(lines):
            m = re.match(r"^([A-Z][A-Za-z' ,\-]+?):\s", line)
            if not m or m.group(1) not in by_name:
                continue
            # The stats line follows the entry; stop at the next "Name:" entry so a
            # skill without stats never borrows the numbers of its neighbour.
            stat_line = None
            for l in lines[idx + 1:idx + 6]:
                if re.match(r"^[A-Z][A-Za-z' ,\-/]+?:\s", l):
                    break
                if re.match(r'^\d', l) and re.search(r'TP|speed|% ATK|power', l):
                    stat_line = l
                    break
            if not stat_line:
                continue
            expected = {
                'tp_cost': parse_range(r'(\d+)(?:-(\d+))? ?TP', stat_line),
                'atk_pct': parse_range(r'(\d+)(?:-(\d+))?% ATK', stat_line),
                'power': parse_range(r'(\d+)(?:-(\d+))? (?:magic )?power', stat_line),
                'speed': parse_range(r'([\d.]+)(?:-([\d.]+))? speed', stat_line),
                'cast': parse_range(r'([\d.]+)(?:-([\d.]+))? cast(?! speed)', stat_line),
                'turns': parse_range(r'(\d+)(?:-(\d+))? turns', stat_line),
            }
            for field, exp in expected.items():
                if exp is None:
                    continue
                candidates = []
                for s in by_name[m.group(1)]:
                    if field in ('atk_pct', 'power'):
                        eff_id = 26 if field == 'atk_pct' else 1
                        for sl in s['slots']:
                            if sl['id'] == eff_id:
                                lo, hi = ends(sl['values'], s['levels'])
                                if field == 'atk_pct':
                                    lo, hi = round(100 * (1 + lo)), round(100 * (1 + hi))
                                candidates.append((lo, hi))
                    elif field in s['arrays']:
                        candidates.append(ends(s['arrays'][field][0], s['levels']))
                if not candidates:
                    continue
                tol = 0.011 if field in ('speed', 'cast') else 1.01 if field == 'atk_pct' else 0.01
                ok = any(abs(lo - exp[0]) <= tol and abs(hi - exp[1]) <= tol for lo, hi in candidates)
                stats[field][0 if ok else 1] += 1
                if not ok:
                    mismatches.append(f"{guide[:-6]}: {m.group(1)} {field}: guide {exp} vs binary {candidates}")

    print("\n[GUIDE CHECK] Binary Lv1/max values vs GameFAQs ranges:")
    for field, (ok, bad) in stats.items():
        print(f"    {field:<8}: {ok:3d} match, {bad:3d} differ")
    for line in mismatches:
        print(f"    [DIFF] {line}")
    return stats, mismatches


def main():
    parser = argparse.ArgumentParser(description='Extract vanilla numeric values into a read-only workbook.')
    parser.add_argument('--check-guides', action='store_true', help='cross-check skills against Data/*.mhtml')
    args = parser.parse_args()

    print("=== Extracting vanilla source of truth from backups ===")
    problems = []
    wb, skills = write_workbook(problems)
    for ws in wb.worksheets:
        print(f"    {ws.title:<16}: {ws.max_row - 1} rows")
    print(f"Saved: {OUTPUT_PATH}")

    if args.check_guides:
        check_guides(skills)

    if problems:
        print(f"\n[FAIL] {len(problems)} structural assumption(s) do not hold:")
        for p in problems:
            print(f"    {p}")
        sys.exit(1)
    print("\n[OK] All structural assumptions hold.")


if __name__ == "__main__":
    main()
