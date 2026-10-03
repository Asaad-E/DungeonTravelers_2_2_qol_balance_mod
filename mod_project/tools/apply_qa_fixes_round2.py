import openpyxl

# 1. Update skills.xlsx
wb_skills = openpyxl.load_workbook('mod_project/data_tables/skills.xlsx')
ws_skills = wb_skills['Skills']

for r in range(2, ws_skills.max_row + 1):
    name = ws_skills.cell(row=r, column=4).value
    if name == 'Mini Auto-Recovery':
        ws_skills.cell(row=r, column=9, value="Slightly recover HP every turn during battle.")
        ws_skills.cell(row=r, column=10, value="HP Regen (L1-5): 1% -> 2% -> 3% -> 4% -> 5% (+1%/Lv)\nTriggers at start of user's action in combat")
        ws_skills.cell(row=r, column=11, value="Approved")
        ws_skills.cell(row=r, column=12, value="QoL transparency: exact 1-5% per-turn HP regeneration scaling.")
        print(f"Updated skills.xlsx: Row {r} [Mini Auto-Recovery]")

wb_skills.save('mod_project/data_tables/skills.xlsx')

# 2. Update items.xlsx
wb_items = openpyxl.load_workbook('mod_project/data_tables/items.xlsx')
ws_items = wb_items['Items']

ITEM_UPDATES = {
    "ItemData_0043.bin": "Restore an ally's HP by 30 and TP by 40.",
    "ItemData_0044.bin": "Restore all allies' HP by 75 and TP by 25.",
    "ItemData_0045.bin": "Restore an ally's TP by 30. May inflict Sleep.",
    "ItemData_0046.bin": "Restore an ally's HP by 20 and TP by 15. May inflict Sleep.",
    "ItemData_0047.bin": "Restore an ally's HP by 20 and TP by 15. May inflict Sleep.",
    "ItemData_0048.bin": "Restore an ally's TP by 30. May inflict Sleep.",
    "ItemData_0052.bin": "Revive a fallen ally with 100% HP.",
    "ItemData_0053.bin": "Cure all status ailments for all allies.",
    "ItemData_0064.bin": "Restore all allies' HP by 150.",
    "ItemData_0081.bin": "Restore an ally's HP by 100.",
    "ItemData_0082.bin": "Restore an ally's HP by 75 and TP by 75."
}

updated_items = 0
for r in range(2, ws_items.max_row + 1):
    fn = ws_items.cell(row=r, column=2).value
    if fn in ITEM_UPDATES:
        ws_items.cell(row=r, column=6, value=ITEM_UPDATES[fn])
        ws_items.cell(row=r, column=7, value="Approved")
        ws_items.cell(row=r, column=8, value="QoL transparency: exact recovery numbers.")
        updated_items += 1
        print(f"Updated items.xlsx: Row {r} ({fn}) [{ws_items.cell(row=r, column=3).value}] -> {ITEM_UPDATES[fn]}")

wb_items.save('mod_project/data_tables/items.xlsx')
print(f"Successfully updated {updated_items} items in items.xlsx!")
