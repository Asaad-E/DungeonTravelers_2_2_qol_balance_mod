import openpyxl
import re

# 1. Update skills.xlsx
wb_skills = openpyxl.load_workbook('mod_project/data_tables/skills.xlsx')
ws_skills = wb_skills['Skills']

# Step corrections for 2-point ranges
step_corrections = {
    "Berserk": {
        "lvup": "ATK: +25% -> +65% (+10%/Lv) | DEF: -15% -> -35% (-5%/Lv)\nTurns: 4 -> 5 -> 6 -> 7 -> 8 | TP Cost: 8 -> 10 -> 12 -> 14 -> 16"
    },
    "Shield Bash": {
        "lvup": "ATK: 110% -> 150% (+10%/Lv) | Stun: 20% -> 40% (+5%/Lv)\nTP Cost: 6 -> 8 -> 10 -> 12 -> 14 | Requires equipped shield"
    },
    "Halving Slice": {
        "lvup": "ATK: 115% -> 175% (+15%/Lv) | Instakill: 10% -> 25% (+3.75%/Lv)\nTP Cost: 15 -> 18 -> 21 -> 24 -> 28"
    },
    "Armor Penetration": {
        "lvup": "ATK: 120% -> 180% (+15%/Lv) | DEF Debuff: -15% -> -30% (-3.75%/Lv)\nTP Cost: 12 -> 14 -> 16 -> 18 -> 20 | Bypasses enemy armor"
    },
    "Frost Slash": {
        "lvup": "ATK: 120% -> 180% (+15%/Lv) | Freeze: 20% -> 40% (+5%/Lv)\nTP Cost: 10 -> 12 -> 14 -> 16 -> 18"
    },
    "Geo Break": {
        "lvup": "ATK: 130% -> 210% (+20%/Lv) | Petrify: 15% -> 30% (+3.75%/Lv)\nTP Cost: 12 -> 15 -> 18 -> 21 -> 24"
    },
    "Shock Spiral": {
        "lvup": "ATK: 110% -> 170% (+15%/Lv) | Paralysis: 20% -> 40% (+5%/Lv)\nTP Cost: 14 -> 17 -> 20 -> 23 -> 26"
    },
    "Gust Thrust": {
        "lvup": "ATK: 125% -> 185% (+15%/Lv) | Silence: 20% -> 40% (+5%/Lv)\nTP Cost: 10 -> 12 -> 14 -> 16 -> 18"
    },
    "Pursuit Stab": {
        "lvup": "Base ATK: 120% -> 160% (+10%/Lv) | +25% dmg per active buff/debuff\nTP Cost: 10 -> 12 -> 14 -> 16 -> 18"
    },
    "Sleep Sting": {
        "lvup": "ATK: 110% -> 170% (+15%/Lv) | Sleep: 30% -> 60% (+7.5%/Lv)\nTP Cost: 6 -> 8 -> 10 -> 12 -> 14"
    },
    "Silence Sting": {
        "lvup": "ATK: 110% -> 170% (+15%/Lv) | Silence: 30% -> 60% (+7.5%/Lv)\nTP Cost: 6 -> 8 -> 10 -> 12 -> 14"
    },
    "Stun Sting": {
        "lvup": "ATK: 110% -> 170% (+15%/Lv) | Stun: 25% -> 50% (+6.25%/Lv)\nTP Cost: 6 -> 8 -> 10 -> 12 -> 14"
    },
    "Chain Spike": {
        "lvup": "Hits: 2 -> 3 -> 3 -> 4 -> 4 | ATK: 65% -> 85% per hit (+5%/Lv)\nTP Cost: 12 -> 15 -> 18 -> 21 -> 25"
    },
    "Poison Shot": {
        "lvup": "ATK: 110% -> 170% (+15%/Lv) | Poison: 30% -> 60% (+7.5%/Lv)\nTP Cost: 6 -> 8 -> 10 -> 12 -> 14"
    },
    "Freeze Shot": {
        "lvup": "ATK: 115% -> 175% (+15%/Lv) | Freeze: 25% -> 50% (+6.25%/Lv)\nTP Cost: 8 -> 10 -> 12 -> 14 -> 16"
    },
    "Paralysis Shot": {
        "lvup": "ATK: 115% -> 175% (+15%/Lv) | Paralysis: 25% -> 50% (+6.25%/Lv)\nTP Cost: 8 -> 10 -> 12 -> 14 -> 16"
    },
    "Star Dust Arrow": {
        "lvup": "Hits: 3-5 arrows per wave | ATK: 70% -> 95% per arrow (+6.25%/Lv)\nDuration: 2 -> 3 turns | TP Cost: 20 -> 24 -> 28 -> 32 -> 36"
    },
    "Gold Steal": {
        "lvup": "Success Rate: 40% -> 70% (+7.5%/Lv) | Gold: 1.0x -> 2.5x (+0.375x/Lv)\nTP Cost: 2 -> 3 -> 4 -> 5 -> 6"
    }
}

# Unique Skills: Merge mechanics directly into Modded_Desc (since UniqueSkillData has no LvUp pointer!)
unique_skills_fixed = {
    "Positive Trance": "Heroine: Luco | Damage scales with Pep.\nEffect: +1% damage per Pep point (up to +100% damage at Max Pep).",
    "Third Eye Wide": "Heroine: Mefmera | 15% trigger chance on enemy action.\nInflicts Poison, Paralyze, Silence, or cancels enemy buffs.",
    "Sugar Coma": "Heroine: Souffle | 25% counter chance on physical contact.\nInflicts Slow and -30% Speed on attacker.",
    "Elbow Grease": "Heroine: Grishna | Weapon Scaling Bonus.\nGrants +25% extra damage to all weapon-based techniques.",
    "Furious Strike": "Heroine: Yae | 20% activation rate on physical attacks.\nGuarantees 100% hit rate and completely ignores target DEF.",
    "Alter Fate": "Heroine: Lize | 10% nullification rate.\nChance to completely negate all incoming damage.",
    "Usurper of Souls": "Heroine: Lilian | 25% banish rate before acting.\nInstantly eliminates enemy summon entities, puppets, and familiars.",
    "Overstock Clearance": "Heroine: Monica | 20% trigger rate at start of turn.\nUses a free recovery item from personal stock without turn cost.",
    "Wreathed in Glory": "Heroine: Alisia | Action Recovery Acceleration.\nDecreases action recovery delay by 20% after using skills or magic.",
    "Maximum Maid": "Heroine: Conette | Maid Mastery Booster.\nIncreases potency and effectiveness of all Maid skills by +15%.",
    "Did You Call for Me?": "Heroine: Yuni | 15% double action chance.\nGrants an immediate second action in the same turn.",
    "Undaunted Spirit": "Heroine: Gina | Auto-Resurrection.\nResurrects with 50% Max HP 3 turns after falling in combat.",
    "Demon Wall": "Heroine: Ist | Row Protection Barrier.\nReduces incoming damage from Demons, Undead, and Apparitions by 20%.",
    "Tranquil Song": "Heroine: Fiora | 15% trigger rate at start of turn.\nInflicts -20% ATK and -20% INT on all enemies.",
    "Forerunner": "Heroine: Eltricia | 15% instant cast chance.\nOffensive and support spells activate immediately with 0 chant time.",
    "Arcane Shield": "Heroine: Melvy | 20% damage reflect rate.\nReflects 30% of received damage back to attacker as neutral damage.",
    "Shadow Heart": "Heroine: Tsurara | Shadow Clone Follow-up.\nNormal attacks trigger an extra strike dealing 75% damage.",
    "Yearning": "Heroine: Siena | Class Mastery Booster.\nIncreases all class stat bonuses and weapon masteries by +10%."
}

for row in ws_skills.iter_rows(min_row=2):
    cat = row[1].value
    name = row[3].value
    
    # Apply step corrections
    if name in step_corrections:
        row[9].value = step_corrections[name]["lvup"]
        
    # Apply unique skills fix
    if cat == "Unique" and name in unique_skills_fixed:
        row[8].value = unique_skills_fixed[name]
        row[9].value = "" # No LvUp in UniqueSkillData
        
    # Normalize pipe separators from '    |    ' to ' | '
    desc = row[8].value or ""
    lvup = row[9].value or ""
    if "    |    " in desc:
        row[8].value = desc.replace("    |    ", " | ")
    if "    |    " in lvup:
        row[9].value = lvup.replace("    |    ", " | ")

wb_skills.save('mod_project/data_tables/skills.xlsx')
print("Successfully normalized skills.xlsx.")

# 2. Update items.xlsx (clean up flavor text prefixes)
wb_items = openpyxl.load_workbook('mod_project/data_tables/items.xlsx')
ws_items = wb_items['Items']

clean_items = {
    "Refreshing Eye Drop": "Cure Blindness for an ally.",
    "Egg Sandwich": "Restore an ally's HP by 10.",
    "Yakisoba-Pan": "Restore an ally's HP by 20.",
    "Premium Steak Sandwich": "Restore an ally's HP by 30.",
    "Maharaja Curry": "Restore an ally's HP by 150.",
    "Killer Curry": "Restore an ally's HP by 1,000.",
    "Hipanera Tea": "Restore an ally's HP by 50 and TP by 50.",
    "Genghis Khan Hotpot": "Restore an ally's HP by 150.",
    "Hong Shao": "Restore an ally's HP by 300.",
    "Harihari Hotpot": "Restore an ally's HP by 400.",
    "Konpeito": "Restore an ally's HP by 60.",
    "Galactic Konpeito": "Restore an ally's HP by 120."
}

for row in ws_items.iter_rows(min_row=2):
    name = row[2].value
    if name in clean_items:
        row[5].value = clean_items[name]

wb_items.save('mod_project/data_tables/items.xlsx')
print("Successfully cleaned items.xlsx.")
