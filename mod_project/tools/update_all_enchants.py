import os
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TABLES_DIR = os.path.join(BASE_DIR, 'data_tables')
excel_path = os.path.join(TABLES_DIR, 'enchants.xlsx')

wb = openpyxl.load_workbook(excel_path)
ws = wb['Enchants']

all_target_enchants = {
    # HP / TP Regen
    'EnchantData_0144.bin': 'HP Regen +4%/turn',
    'EnchantData_0145.bin': 'HP Regen +2%/turn',
    'EnchantData_0148.bin': 'HP Regen +8%/turn',
    'EnchantData_0147.bin': 'TP Regen +2/turn',
    'EnchantData_0146.bin': 'TP Regen +4/turn',
    'EnchantData_0149.bin': 'TP Regen +6/turn',
    
    # HP / TP Drain on normal attack
    'EnchantData_0161.bin': 'Normal attacks: 50% HP Drain',
    'EnchantData_0162.bin': 'Normal attacks: 20% HP Drain',
    'EnchantData_0163.bin': 'Normal attacks: 20% TP Drain',
    'EnchantData_0164.bin': 'Normal attacks: 10% TP Drain',
    
    # Encounter Rates
    'EnchantData_0211.bin': 'Encounter Rate -25%',
    'EnchantData_0212.bin': 'Encounter Rate -50%',
    'EnchantData_0213.bin': 'Encounter Rate +25%',
    'EnchantData_0214.bin': 'Encounter Rate +50%',
    
    # Defending Regens
    'EnchantData_0257.bin': 'Defending HP Regen +10%',
    'EnchantData_0292.bin': 'Defending TP Regen +10',
    
    # TP Cost Reductions
    'EnchantData_0246.bin': 'Technique TP Cost -10%',
    'EnchantData_0247.bin': 'Technique TP Cost -25%',
    'EnchantData_0248.bin': 'Spell TP Cost -10%',
    'EnchantData_0249.bin': 'Spell TP Cost -25%',
    
    # Race Slayers (+20%, +50%, +100%)
    'EnchantData_0170.bin': 'Physical ATK vs. Demons +50%',
    'EnchantData_0171.bin': 'Physical ATK vs. Demons +20%',
    'EnchantData_0172.bin': 'Physical ATK vs. Apparitions +50%',
    'EnchantData_0173.bin': 'Physical ATK vs. Apparitions +20%',
    'EnchantData_0174.bin': 'Physical ATK vs. Magic Beings +50%',
    'EnchantData_0175.bin': 'Physical ATK vs. Magic Beings +20%',
    'EnchantData_0176.bin': 'Physical ATK vs. Humans +100%',
    'EnchantData_0177.bin': 'Physical ATK vs. Humans +50%',
    'EnchantData_0178.bin': 'Physical ATK vs. Humans +20%',
    'EnchantData_0179.bin': 'Physical ATK vs. Undead +50%',
    'EnchantData_0180.bin': 'Physical ATK vs. Undead +20%',
    'EnchantData_0181.bin': 'Physical ATK vs. Dragons +50%',
    'EnchantData_0182.bin': 'Physical ATK vs. Dragons +20%',
    'EnchantData_0183.bin': 'Physical ATK vs. Beasts +50%',
    'EnchantData_0184.bin': 'Physical ATK vs. Beasts +20%',
    'EnchantData_0185.bin': 'Physical ATK vs. Fruit +100%',
    'EnchantData_0186.bin': 'Physical ATK vs. Fruit +50%',
    'EnchantData_0196.bin': 'Physical ATK vs. Felines +100%',
    'EnchantData_0203.bin': 'Physical ATK vs. Demons +100%',
    'EnchantData_0204.bin': 'Physical ATK vs. Apparitions +100%',
    'EnchantData_0205.bin': 'Physical ATK vs. Magic Beings +100%',
    'EnchantData_0206.bin': 'Physical ATK vs. Undead +100%',
    'EnchantData_0207.bin': 'Physical ATK vs. Dragons +100%',
    'EnchantData_0208.bin': 'Physical ATK vs. Beasts +100%',
    'EnchantData_0209.bin': 'Physical ATK vs. Sea Creatures +100%',
    
    # Joint & Special
    'EnchantData_0261.bin': 'Joint Skill Damage +25%',
    'EnchantData_0290.bin': 'Joint Skill Damage +50%',
    'EnchantData_0276.bin': 'Equipment Drop Rate +25%',
    'EnchantData_0280.bin': 'Damage vs. Ailment Foes +25%',
    'EnchantData_0262.bin': 'Escape Chance +25%',
    'EnchantData_0295.bin': 'Mercenary Buddy Stats +10%',
    'EnchantData_0297.bin': 'Mercenary Buddy Stats +30',
}

updated = 0
for r in range(2, ws.max_row + 1):
    fn = str(ws.cell(row=r, column=1).value)
    if fn in all_target_enchants:
        new_val = all_target_enchants[fn]
        ws.cell(row=r, column=5).value = new_val
        ws.cell(row=r, column=6).value = 'Approved'
        updated += 1

wb.save(excel_path)
print(f"Successfully updated {updated} enchants in enchants.xlsx!")
