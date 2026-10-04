"""
QoL text of enemy attacks (also the abilities of summoned monsters, shown in "View the details of
your teammates' abilities"): proposes it and keeps it in a workbook that is edited by hand.

    python mod_project/tools/generate_qol_enemy_attack_text.py

Workbook: mod_project/data_tables/qol_enemy_attack_text.xlsx
  * QoL_Desc is the text the build uses. Put "yes" in Edited to keep your wording.
  * The game's own text has a data line and then prose:
        Target: Single / Type: Stab / Range: Long
        Fire an arrow doused in some real toxic stuff.
        If it hits, it may inflict poison. Also, ouch.
    The prose (jokes included) is never touched. The exact figures go on the data line, in its own
    format:  Target: Single / Type: Stab / Range: Long / Poison: 30%
  * An attack has no levels, so its figures are exact. Shown so far: the chance of each ailment.
"""
import os
import sys
import collections
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pck_handler import PckArchive, EnemyAttackHelper
from text_checks import (ENEMY_DATA_LINE, AILMENT_NAMES, enemy_attack_slots, enemy_ailments,
                         enemy_ailment_tokens, check_enemy_attack_text)

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKUP_PATH = os.path.join(BASE_DIR, 'backups', 'EnemyAttackData.pck')
OUTPUT_PATH = os.path.join(BASE_DIR, 'data_tables', 'qol_enemy_attack_text.xlsx')

MAX_LINE = 85     # measured on screen for description boxes
HEADERS = ['Filename', 'Attack_Name', 'Vanilla_Desc', 'QoL_Desc', 'Edited', 'Effect_Slots', 'Note']


def propose(b, desc):
    """Returns (text, note)."""
    lines = desc.split('\n')
    has_ailment = any(effect_id in AILMENT_NAMES for effect_id, _, _ in enemy_attack_slots(b))
    if not has_ailment:
        return desc, ''
    if not lines[0].startswith(ENEMY_DATA_LINE):
        return desc, 'has an ailment chance but no data line to put it on'
    if enemy_ailments(b) is None:
        return desc, 'ailment chance outside 0-100% or stated twice: left as written'
    line = ' / '.join([lines[0]] + enemy_ailment_tokens(b))
    if len(line) > MAX_LINE:
        return desc, f'data line would be {len(line)} characters'
    return '\n'.join([line] + lines[1:]), ''


def load_edited():
    if not os.path.exists(OUTPUT_PATH):
        return {}
    ws = openpyxl.load_workbook(OUTPUT_PATH)['EnemyAttacks']
    header = [c.value for c in ws[1]]
    fn, desc, edited = header.index('Filename'), header.index('QoL_Desc'), header.index('Edited')
    return {r[fn]: r[desc] or '' for r in ws.iter_rows(min_row=2, values_only=True)
            if str(r[edited] or '').strip().lower() == 'yes'}


def build_rows():
    edited = load_edited()
    rows, stats = [], collections.Counter()
    for fn, b in PckArchive(BACKUP_PATH).files:
        attack = EnemyAttackHelper.read_enemy_attack(b)
        desc, note = propose(b, attack['desc'])
        if fn in edited:
            desc, note = edited[fn], ''
        problems, _ = check_enemy_attack_text(b, desc, attack['desc'])
        if problems or any(len(line) > MAX_LINE for line in desc.split('\n')[:1] if desc != attack['desc']):
            raise SystemExit(f"{fn} ({attack['name']}): {problems or 'data line too long'}")
        stats['edited by hand' if fn in edited else 'ailment chance added to the data line' if desc != attack['desc']
              else 'could not be stated: ' + note.split(' characters')[0].rstrip('0123456789 ') if note
              else 'nothing to add'] += 1
        slots = ', '.join(f'{i}:{v:g}' for i, v, _ in enemy_attack_slots(b))
        rows.append([fn, attack['name'], attack['desc'], desc, 'yes' if fn in edited else '', slots, note])
    return rows, stats


def main():
    rows, stats = build_rows()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'EnemyAttacks'
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='305496')
    ws.freeze_panes = 'C2'
    for i, width in enumerate([26, 30, 90, 90, 8, 50, 50], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    how = wb.create_sheet('How to edit')
    for line in [
        'QoL_Desc is the text the build writes into EnemyAttackData.pck.',
        'Lines 2 and 3 are the game\'s prose and must stay exactly as they are; the build rejects any change to them.',
        'Line 1 is the data line. It may gain "Ailment: N%" items after " / ", and only the ones stored in the entry.',
        'Effect_Slots lists the raw (effect id:value) pairs. Ailments: 113 poison, 114 freeze, 115 paralysis, 116 sleep, '
        '117 silence, 118 stun, 119 instakill, 120 petrify, 122 slow.',
    ]:
        how.append([line])
    how.column_dimensions['A'].width = 130
    wb.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}: {len(rows)} enemy attacks")
    for key in sorted(stats):
        print(f"    {stats[key]:5d}  {key}")


if __name__ == '__main__':
    sys.exit(main())
