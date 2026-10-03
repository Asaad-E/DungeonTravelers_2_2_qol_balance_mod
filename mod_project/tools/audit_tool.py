import os
import openpyxl
from pck_handler import PckArchive, SkillHelper

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TABLES_DIR = os.path.join(BASE_DIR, 'data_tables')
BACKUP_DIR = os.path.join(BASE_DIR, 'backups')
GAME_PARAM_DIR = os.path.join(os.path.dirname(BASE_DIR), 'DungeonTravelers2-2_EN', 'Data', 'ENG', 'Parameter')

def run_audit():
    print("=" * 60)
    print("           DT2-2 MOD AUDIT & QA VERIFICATION REPORT")
    print("=" * 60)
    
    excel_path = os.path.join(TABLES_DIR, 'skills.xlsx')
    if not os.path.exists(excel_path):
        print(f"[FAIL] Missing {excel_path}")
        return False
        
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Skills"]
    
    total_skills = ws.max_row - 1
    modified_skills = 0
    errors = []
    warnings = []
    
    for row_idx in range(2, ws.max_row + 1):
        cat = ws.cell(row=row_idx, column=2).value
        name = ws.cell(row=row_idx, column=4).value
        status = ws.cell(row=row_idx, column=11).value
        mod_desc = ws.cell(row=row_idx, column=9).value or ""
        mod_lvup = ws.cell(row=row_idx, column=10).value or ""
        
        if status in ['Approved', 'Applied', 'Draft'] and (mod_desc or mod_lvup):
            modified_skills += 1
            
            # 1. Check line count
            lines = [mod_desc] if mod_desc else []
            if mod_lvup:
                lines.extend(mod_lvup.split('\n'))
                
            if len(lines) > 3:
                errors.append(f"[ROW {row_idx}] '{name}' ({cat}): Exceeds 3 lines ({len(lines)} lines)!")
            elif len(lines) == 3:
                # perfect max
                pass
                
            # 2. Check character length per line (max 85 chars, warn > 80)
            for line_idx, l in enumerate(lines):
                if len(l) > 85:
                    errors.append(f"[ROW {row_idx}] '{name}' Line {line_idx+1}: Too long ({len(l)} chars > 85 max): '{l}'")
                elif len(l) > 80:
                    warnings.append(f"[ROW {row_idx}] '{name}' Line {line_idx+1}: Close to limit ({len(l)} chars): '{l}'")
                    
    print(f"\n[1] EXCEL DATA TABLES CHECK:")
    print(f"    Total Skills Indexed:    {total_skills}")
    print(f"    Modified/Active Skills:  {modified_skills}")
    print(f"    Errors:   {len(errors)}")
    print(f"    Warnings: {len(warnings)}")
    
    for err in errors:
        print(f"    [FAIL] {err}")
    for warn in warnings:
        print(f"    [WARN] {warn}")
        
    # 3. Verify Game Parameter Files
    print(f"\n[2] GAME BINARY & PCK ARCHIVE INTEGRITY CHECK:")
    pck_files = ['TechniqueSkillData.pck', 'PassiveSkillData.pck', 'MagicSkillData.pck', 'ItemData.pck']
    for pck_file in pck_files:
        game_pck = os.path.join(GAME_PARAM_DIR, pck_file)
        if not os.path.exists(game_pck):
            print(f"    [FAIL] Missing in game directory: {pck_file}")
            continue
            
        try:
            archive = PckArchive(game_pck)
            print(f"    [PASS] {pck_file:<25}: Valid PCK archive ({len(archive.files)} entries, {os.path.getsize(game_pck)} bytes)")
        except Exception as e:
            print(f"    [FAIL] {pck_file} corrupted: {e}")
            errors.append(str(e))
            
    # 4. Git Status
    print(f"\n[3] VERSION CONTROL STATUS:")
    import subprocess
    res = subprocess.run(["git", "status", "--short"], cwd=os.path.dirname(BASE_DIR), capture_output=True, text=True)
    uncommitted = res.stdout.strip().splitlines()
    print(f"    Tracked/Changed files in repo: {len(uncommitted)}")
    
    print("\n" + "=" * 60)
    if not errors:
        print("          ALL CHECKS PASSED: READY FOR IN-GAME TESTING")
    else:
        print(f"          AUDIT FOUND {len(errors)} CRITICAL ISSUES TO FIX")
    print("=" * 60)
    return len(errors) == 0

if __name__ == "__main__":
    run_audit()
