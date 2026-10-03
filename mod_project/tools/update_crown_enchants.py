import os
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TABLES_DIR = os.path.join(BASE_DIR, 'data_tables')
excel_path = os.path.join(TABLES_DIR, 'enchants.xlsx')

wb = openpyxl.load_workbook(excel_path)
ws = wb['Enchants']

crown_enchants = {
    'EnchantData_0219.bin': 'ATK vs. Higher-Level Foes +25%',
    'EnchantData_0220.bin': 'ATK vs. Lower-Level Foes +25%',
    'EnchantData_0222.bin': 'Front Row ATK +8%',
    'EnchantData_0223.bin': 'Front Row ATK/DEF +4%',
    'EnchantData_0224.bin': 'Back Row INT +8%',
    'EnchantData_0225.bin': 'Back Row INT/RES +4%',
    'EnchantData_0258.bin': 'Treasure Chest Item Level +1',
    'EnchantData_0259.bin': 'Wandering Shop Encounter Rate +50%',
    'EnchantData_0260.bin': 'Non-Gas Trap Damage -90%',
    'EnchantData_0263.bin': 'Unarmed Action Speed: A-Rank',
    'EnchantData_0264.bin': 'Class Weapon Modifier +2%',
    'EnchantData_0265.bin': 'Ailment Natural Recovery Rate +50%',
    'EnchantData_0266.bin': 'Pep Decrease Rate -50%',
    'EnchantData_0267.bin': 'Damage +100%/empty party slot',
    'EnchantData_0270.bin': 'Higher-Level Enemy EXP +50%',
    'EnchantData_0271.bin': 'Damage at Max HP +10%',
    'EnchantData_0272.bin': 'Missed Spell Damage Penalty -50%',
    'EnchantData_0281.bin': 'Back Row DEX/AGI +2%',
    'EnchantData_0291.bin': 'Chest Mimic Appearance Rate UP',
    'EnchantData_0293.bin': 'Unique Skill Level +1',
    'EnchantData_0294.bin': 'Food Affinity Bonus/Penalty x2',
    'EnchantData_0296.bin': 'Turn Start Full HP Recovery (10%)',
    'EnchantData_0298.bin': 'Fast Spell Cast (Instant Cast chance)',
    'EnchantData_0217.bin': 'Recover 3 TP / 10 steps',
}

updated = 0
for r in range(2, ws.max_row + 1):
    fn = str(ws.cell(row=r, column=1).value)
    if fn in crown_enchants:
        old_val = ws.cell(row=r, column=5).value
        new_val = crown_enchants[fn]
        ws.cell(row=r, column=5).value = new_val
        ws.cell(row=r, column=6).value = 'Approved'
        print(f"Row {r:3d}: {fn} | '{old_val}' -> '{new_val}'")
        updated += 1

wb.save(excel_path)
print(f"Successfully updated {updated} crown enchants in enchants.xlsx!")
