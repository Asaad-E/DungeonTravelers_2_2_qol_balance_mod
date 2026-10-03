import os
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TABLES_DIR = os.path.join(BASE_DIR, 'data_tables')
excel_path = os.path.join(TABLES_DIR, 'enchants.xlsx')

wb = openpyxl.load_workbook(excel_path)
ws = wb['Enchants']

target_enchants = {
    'EnchantData_0144.bin': 'HP Regen +4%/turn',
    'EnchantData_0145.bin': 'HP Regen +2%/turn',
    'EnchantData_0148.bin': 'HP Regen +8%/turn',
    'EnchantData_0147.bin': 'TP Regen +2/turn',
    'EnchantData_0146.bin': 'TP Regen +4/turn',
    'EnchantData_0149.bin': 'TP Regen +6/turn',
    'EnchantData_0211.bin': 'Encounter Rate -25%',
    'EnchantData_0212.bin': 'Encounter Rate -50%',
    'EnchantData_0213.bin': 'Encounter Rate +25%',
    'EnchantData_0214.bin': 'Encounter Rate +50%',
    'EnchantData_0292.bin': 'Defending TP Regen +10',
    'EnchantData_0257.bin': 'Defending HP Regen +10%',
    'EnchantData_0246.bin': 'Technique TP Cost -10%',
    'EnchantData_0247.bin': 'Technique TP Cost -25%',
    'EnchantData_0248.bin': 'Spell TP Cost -10%',
    'EnchantData_0249.bin': 'Spell TP Cost -25%',
}

updated = 0
for r in range(2, ws.max_row + 1):
    fn = str(ws.cell(row=r, column=1).value)
    if fn in target_enchants:
        old_val = ws.cell(row=r, column=5).value
        new_val = target_enchants[fn]
        ws.cell(row=r, column=5).value = new_val
        ws.cell(row=r, column=6).value = 'Approved'
        print(f"Row {r:3d}: {fn} | '{old_val}' -> '{new_val}'")
        updated += 1

wb.save(excel_path)
print(f"Successfully updated {updated} enchants in enchants.xlsx!")
