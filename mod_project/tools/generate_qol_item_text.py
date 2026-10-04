"""
QoL item text: proposes the description of each item for the QoL variant and keeps it in a
workbook that is edited by hand.

    python mod_project/tools/generate_qol_item_text.py

Workbook: mod_project/data_tables/qol_item_text.xlsx
  * QoL_Desc is the text the build uses. Put "yes" in Edited to keep your wording.
  * Rule: where the vanilla text gives a vague size and the binary holds the exact amount,
    the amount replaces the word. "Restore an ally's HP (medium)." -> "Restore an ally's HP (+250)."
    A text never carries both.
  * Everything else stays vanilla. The same description is shown in battle, so it stays one line.
"""
import os
import re
import sys
import collections
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pck_handler import PckArchive, ItemHelper
from text_checks import ITEM_RESTORE, item_pairs, check_item_text

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKUP_PATH = os.path.join(BASE_DIR, 'backups', 'ItemData.pck')
TABLES_DIR = os.path.join(BASE_DIR, 'data_tables')
PREVIOUS_PATH = os.path.join(TABLES_DIR, 'items.xlsx')
OUTPUT_PATH = os.path.join(TABLES_DIR, 'qol_item_text.xlsx')

MAX_LINE = 85     # measured on screen
HEADERS = ['Filename', 'Item_Name', 'Vanilla_Desc', 'QoL_Desc', 'Edited', 'Effect_Pairs', 'Previous_Mod_Desc']
VAGUE_RE = re.compile(r"Restore (?P<what>an ally's HP|an ally's TP|all allies' HP) \((?:small|medium|large)\)")


def propose(b, desc):
    """Replaces the vague size by the amount stored in the first effect pair, when both exist."""
    effect_type, value = item_pairs(b)[0]
    m = VAGUE_RE.search(desc)
    if not m or ITEM_RESTORE[m.group('what')] != effect_type or value != int(value):
        return desc
    return desc[:m.start()] + f"Restore {m.group('what')} (+{int(value)})" + desc[m.end():]


def load_column(path, sheet, key, value, only_edited=False):
    if not os.path.exists(path):
        return {}
    wb = openpyxl.load_workbook(path)
    if sheet not in wb.sheetnames:
        return {}
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    if key not in header or value not in header:
        return {}
    k, v = header.index(key), header.index(value)
    e = header.index('Edited') if only_edited else None
    return {r[k]: r[v] for r in ws.iter_rows(min_row=2, values_only=True)
            if r[k] and (e is None or str(r[e] or '').strip().lower() == 'yes')}


def build_rows():
    edited = load_column(OUTPUT_PATH, 'Items', 'Filename', 'QoL_Desc', only_edited=True)
    previous = {}
    if os.path.exists(PREVIOUS_PATH):
        wb = openpyxl.load_workbook(PREVIOUS_PATH, read_only=True)
        ws = wb[wb.sheetnames[0]]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        if 'Filename' in header and 'Modded_Desc' in header:
            f, d = header.index('Filename'), header.index('Modded_Desc')
            previous = {r[f]: r[d] for r in ws.iter_rows(min_row=2, values_only=True) if r[f]}
    rows, stats = [], collections.Counter()
    for fn, b in PckArchive(BACKUP_PATH).files:
        item = ItemHelper.read_item(b)
        desc = edited[fn] if fn in edited else propose(b, item['desc'])
        problems, _ = check_item_text(b, desc)
        if problems or '\n' in desc or len(desc) > MAX_LINE:
            raise SystemExit(f"{fn} ({item['name']}): '{desc}' {problems or 'breaks the one-line limit'}")
        desc.encode('latin-1')
        stats['edited by hand' if fn in edited else
              'amount in place of the vague size' if desc != item['desc'] else 'vanilla text kept'] += 1
        pairs = ', '.join(f'{t}:{v:g}' for t, v in item_pairs(b) if t != -1)
        rows.append([fn, item['name'], item['desc'], desc, 'yes' if fn in edited else '', pairs, previous.get(fn, '')])
    return rows, stats


def main():
    rows, stats = build_rows()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Items'
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='305496')
    ws.freeze_panes = 'C2'
    for i, width in enumerate([20, 28, 80, 80, 8, 36, 80], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    how = wb.create_sheet('How to edit')
    for line in [
        'QoL_Desc is the text the build writes into ItemData.pck. One line, 85 characters at most.',
        'To word an item yourself: change QoL_Desc and put "yes" in Edited. The generator keeps that row.',
        'An amount is written as "Restore an ally\'s HP (+250)." and the build checks it against the binary.',
        'Effect_Pairs lists the (type:value) pairs of the item. Types read so far: 0 HP, 1 TP, 3 HP of all allies, '
        '6 cure one ailment, 7 cure all ailments, 8 revive, 9 light steps, 10 escape. The rest are not decoded.',
    ]:
        how.append([line])
    how.column_dimensions['A'].width = 130
    wb.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}: {len(rows)} items")
    for key in sorted(stats):
        print(f"    {stats[key]:4d}  {key}")
    for r in rows:
        if r[2] != r[3]:
            print(f"        {r[1]}: {r[3]}")


if __name__ == '__main__':
    sys.exit(main())
