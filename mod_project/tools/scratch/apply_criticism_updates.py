import openpyxl

# 1. Update balance_criteria.xlsx
wb_crit = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx')
ws_crit = wb_crit['Criticisms']
headers_crit = [c.value for c in ws_crit[1]]

updates = {
    'J04': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Vanilla_Result': 'Parties without a Valkyrie (using Paladin, debuffers like Etoile, or healers) win 92 - 100% of story bosses; in late postgame (P08 - P24) Valkyrie wins +0% to +6% on the median (CLASS_STAGES.md). Her party-wide guard is comfortable but not strictly mandatory',
        'Proposal': 'None: class identity kept without making her mandatory',
        'Decision': '2026-10-10: no change; alternative defensive setups are viable (docs/generated/CLASS_STAGES.md)',
        'Why': 'Valkyrie is the premier tank/guard class by identity, but not mandatory: Paladin and debuff builds clear all story checkpoints'
    },
    'J19': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Vanilla_Result': 'Intermediate stage (C10, C13, C18): length 0.96 (faster than median attacker 1.00), 30 - 41% damage share, 100% wins; Aura Eliminate 86% worth, Geo Break and Frost Slash 100% (CLASS_STAGES.md, SKILL_WORTH.md)',
        'Proposal': 'None: she leads intermediate physical damage',
        'Decision': '2026-10-10: no change; Berserker is one of the strongest intermediate carries',
        'Why': 'False: high damage and 100% wins at her stage'
    },
    'J20': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Vanilla_Result': 'Category error: Paladin is an intermediate class (levels 15 - 29), not advanced. In intermediate (C10, C13, C18) she has 100% wins as tank; at level 30 she promotes to Valkyrie or Samurai (CLASS_STAGES.md)',
        'Proposal': 'None: fulfills her role at intermediate tier',
        'Decision': '2026-10-10: no change; Paladin functions as intended in her tier',
        'Why': 'False: intermediate class compared to an advanced class'
    },
    'J21': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Vanilla_Result': 'Dark Lord is the top physical attacker in advanced and late game: 54 - 74% damage share, 100% wins at C40 and P24; Attract serves (+1% wins, SKILL_VALUE.md); Megaton Press 72% worth, Geo Break and Life Bring 100% (CLASS_STAGES.md, SKILL_WORTH.md)',
        'Proposal': 'None: Dark Lord is already top tier',
        'Decision': '2026-10-10: no change; Dark Lord is top tier',
        'Why': 'False: premier physical carry in advanced and postgame'
    },
    'J22': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Vanilla_Result': 'Samurai uses Katana (Katana Mastery, Flowing Water Stance); Dual Wield belongs to Kunoichi. Advanced/late: length 0.91 - 0.92, 28 - 42% damage share, 90 - 100% wins; Moonlight Purgatory and Onikagura 100% worth (CLASS_STAGES.md, SKILL_WORTH.md)',
        'Proposal': 'None: premier physical attacker',
        'Decision': '2026-10-10: no change; Samurai is in the top band',
        'Why': 'False: premise confuses Samurai with Kunoichi; Samurai performs at 90 - 100% wins'
    },
    'J24': {
        'Status': 'open',
        'Verdict': 'partly true',
        'Vanilla_Result': 'Basic Heal capped in vanilla. Bishop learns Full Heal (100% HP restore, no cap) and Circle Nurse; in advanced/late Bishop has 92 - 100% wins (+9% on median, CLASS_STAGES.md)',
        'Proposal': 'Solved by class identity (Full Heal) and H036/H037 (HP% scaling on Heal and Circle Heal)',
        'Decision': '2026-10-10: solved by Full Heal and H rows (H036, H037)',
        'Why': 'Bishop has Full Heal, and basic heals now scale via effect 106 in H rows'
    },
    'F01': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Vanilla_Result': 'A dedicated healer and buff management remain standard in hardcore DRPGs, but TP flattening (J03) and Auto-Refresh (J05) reduce Maid lock during exploration, while items and debuffers allow alternative party comps',
        'Proposal': 'Solved by systemic changes in T03 and T05',
        'Decision': '2026-10-10: no further change; party flexibility improved by T03 and class balance',
        'Why': 'Overstated: Maid and healer requirements softened by T03 and item viability'
    },
    'F02': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Vanilla_Result': 'Aura Eliminate (Berserker, Dark Lord) dispels on hit (100% rate, worth 86% in SKILL_WORTH); Dispel (Sage) adds +3% wins and -24% fight length (serves, SKILL_VALUE.md); Etoile dances mitigate stat buffs',
        'Proposal': 'None: physical and magical options exist',
        'Decision': '2026-10-10: no change; Aura Eliminate and Dispel cover both playstyles',
        'Why': 'Both archetypes possess reliable dispel tools'
    },
    'F07': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Vanilla_Result': 'Valkyrie provides top AoE safety with Spell/Material Guard, but Etoile debuffs and Paladin defense also clear late bosses (CLASS_STAGES.md)',
        'Proposal': 'None: class identity kept without mandatory status',
        'Decision': '2026-10-10: no change',
        'Why': 'As J04: top stability, but alternative strategies exist'
    },
    'F85': {
        'Status': 'out of the balance',
        'Verdict': None,
        'Vanilla_Result': 'Joint skills require specific party rosters and TP/gauge investment; parked by user decision',
        'Proposal': 'None: joint skills parked',
        'Decision': None,
        'Why': 'Joint skills parked by the user on 2026-10-04; DLC ignored'
    },
    'F87': {
        'Status': 'open',
        'Verdict': 'true',
        'Vanilla_Result': 'In vanilla, guns disabled Spiral Arrow, Divine Cannon, Warning Shot and Support Shot, shutting down Sniper exclusive kit',
        'Proposal': 'Solved by W036 - W039 (guns may use Spiral Arrow, Divine Cannon, Warning Shot, Support Shot)',
        'Decision': '2026-10-09: solved as J06/F10 (W036 - W039)',
        'Why': 'Guns enabled for Sniper signature skills in W rows'
    },
    'F94': {
        'Status': 'not a problem',
        'Verdict': 'true',
        'Vanilla_Result': 'Guest allies act by fixed enemy-like AI scripts with some flavor skills (e.g. Gold Steal); they are optional temporary helpers',
        'Proposal': 'None: keep flavor',
        'Decision': '2026-10-10: no change; mercenary flavor preserved',
        'Why': 'Cosmetic/flavor quirk with minimal progression impact'
    },
    'S07': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Vanilla_Result': 'Witch excels in story (length 0.63, 46 - 68% damage) due to elemental AoE coverage; normalizes in postgame (length 0.98, 39 - 52% damage) where boss elemental resistances rise (CLASS_STAGES.md)',
        'Proposal': 'None: fulfills magic attacker identity without breaking late game',
        'Decision': '2026-10-10: no change; class identity preserved',
        'Why': 'Strong in story by elemental coverage, balanced against physical burst in postgame'
    },
    'S17': {
        'Status': 'not a problem',
        'Verdict': 'true',
        'Vanilla_Result': 'Enchantment damage reductions have a hardcoded floor at 0.75 (max 25% reduction) in executable formula. Defensive builds rely on stats, guard skills, and shields instead (FORMULAS.md)',
        'Proposal': 'None: changing floor requires .text executable patch (principle 11, benign levers first)',
        'Decision': '2026-10-10: no change; defenses operate via stats and guard skills',
        'Why': 'Executable constant; defensive viability is provided through skills and stats'
    },
    'F81': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Vanilla_Result': 'Sharp Stab is worth 100% at base tier (C03) for Scout (SKILL_WORTH.md). It is a level 1 skill naturally superseded by intermediate/advanced techniques (Stings, Rain, Pursuit Stab)',
        'Proposal': 'None: natural skill tree progression',
        'Decision': '2026-10-10: no change; natural progression of base skills',
        'Why': 'Expected progression from base to advanced skills'
    },
    'J13': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Vanilla_Result': 'Guest ally scripts are data-driven weighted sets not surfaced in UI (BALANCE_BASELINE.md)',
        'Proposal': 'None: documented in baseline',
        'Decision': '2026-10-10: close without change; documented in baseline',
        'Why': 'Minor informational point; guest allies are optional'
    },
    'J34': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Vanilla_Result': 'Berserk lowers DEF by 11 - 25% (effect 28), not to zero; ATK +20 to +50% (effect 26). Displayed accurately in QoL variant',
        'Proposal': 'None: already shown in QoL',
        'Decision': '2026-10-10: close as false; figures verified',
        'Why': 'False: DEF drops by 11 - 25%, not to zero'
    }
}

for row in ws_crit.iter_rows(min_row=2):
    rid = row[0].value
    if rid in updates:
        u = updates[rid]
        for col_name, val in u.items():
            col_idx = headers_crit.index(col_name)
            row[col_idx].value = val

wb_crit.save('mod_project/data_tables/balance_criteria.xlsx')
print("Updated balance_criteria.xlsx successfully.")

# 2. Update balance_changes.xlsx to add F87 to W036-W039
wb_ch = openpyxl.load_workbook('mod_project/data_tables/balance_changes.xlsx')
ws_ch = wb_ch['Changes']
headers_ch = [c.value for c in ws_ch[1]]
serves_idx = headers_ch.index('Serves')

w_ids = {'W036', 'W037', 'W038', 'W039'}
for row in ws_ch.iter_rows(min_row=2):
    cid = row[0].value
    if cid in w_ids:
        cur_serves = str(row[serves_idx].value or '')
        if 'F87' not in cur_serves:
            row[serves_idx].value = cur_serves + ', F87'
            print(f"Updated {cid} Serves to: {row[serves_idx].value}")

wb_ch.save('mod_project/data_tables/balance_changes.xlsx')
print("Updated balance_changes.xlsx successfully.")
