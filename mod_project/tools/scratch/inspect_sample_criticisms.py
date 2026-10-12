import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', data_only=True)
ws = wb['Criticisms']
headers = [c.value for c in ws[1]]

sample_ids = ['J01', 'J02', 'J03', 'J05', 'J06', 'J07', 'U01', 'U02', 'F09', 'F10', 'F11', 'F15', 'S03']

for row in ws.iter_rows(min_row=2, values_only=True):
    rd = dict(zip(headers, row))
    if rd['ID'] in sample_ids:
        print(f"ID: {rd['ID']}")
        print(f"  Status: {rd.get('Status')} | Verdict: {rd.get('Verdict')}")
        print(f"  Vanilla_Result: {rd.get('Vanilla_Result')}")
        print(f"  Balance_Result: {rd.get('Balance_Result')}")
        print(f"  Proposal: {rd.get('Proposal')}")
        print(f"  Decision: {rd.get('Decision')}")
        print("-" * 50)
