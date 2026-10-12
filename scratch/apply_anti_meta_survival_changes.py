import openpyxl
import os
import sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

BC_PATH = os.path.join(BASE_DIR, 'data_tables', 'balance_changes.xlsx')
SETUP_PATH = os.path.join(BASE_DIR, 'data_tables', 'sim_setup.xlsx')

def update_balance_changes():
    wb = openpyxl.load_workbook(BC_PATH)
    ws = wb['Changes']
    
    # Clean up accidental D0100 at the end if present
    if ws.cell(ws.max_row, 1).value == 'D0100':
        ws.delete_rows(ws.max_row)
        print("Removed temporary D0100 row.")
    
    # 1. Update X056 - X060 (Life Bringer Effect 26)
    # Minimal tweak: [0.05, 0.10, 0.15, 0.20, 0.25]
    lb_vals = {1: 0.05, 2: 0.10, 3: 0.15, 4: 0.20, 5: 0.25}
    for row in ws.iter_rows(min_row=2):
        rid = row[0].value
        if rid in ('X056', 'X057', 'X058', 'X059', 'X060'):
            lvl = row[4].value
            row[6].value = lb_vals[lvl]
            row[10].value = '2026-10-10'
            row[11].value = 'decided'
            print(f"Updated {rid}: Life Bringer Lv{lvl} -> {lb_vals[lvl]}")

    # 2. Update D001 - D010 (Spartan Effect 147 turns: starts at 3, ends at 6)
    # [3, 3, 4, 4, 4, 5, 5, 5, 6, 6]
    spartan_turns = {1: 3, 2: 3, 3: 4, 4: 4, 5: 4, 6: 5, 7: 5, 8: 5, 9: 6, 10: 6}
    for row in ws.iter_rows(min_row=2):
        rid = row[0].value
        if rid in [f'D0{i:02d}' if i < 10 else f'D{i:03d}' for i in range(1, 11)]:
            lvl = row[4].value
            row[6].value = spartan_turns[lvl]
            row[10].value = '2026-10-10'
            row[11].value = 'decided'
            print(f"Updated {rid}: Spartan turns Lv{lvl} -> {spartan_turns[lvl]}")

    # Check existing IDs
    existing_ids = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            existing_ids.add(str(row[0]))

    # 3. Add Magic Barrier Effect 183: S022 - S026
    # ([-0.25, -0.28, -0.32, -0.36, -0.40])
    mb_vals = [(1, -0.18, -0.25), (2, -0.20, -0.28), (3, -0.22, -0.32), (4, -0.24, -0.36), (5, -0.26, -0.40)]
    for lvl, van, new_val in mb_vals:
        sid = f"S0{21+lvl}"
        if sid not in existing_ids:
            ws.append([
                sid, 'skills', "Technique 58 Magic Barrier", 'effect 183', lvl, van, new_val,
                'L13', 'F09', "Magic Barrier party magic damage reduction buff", '2026-10-10', 'decided'
            ])
            existing_ids.add(sid)
            print(f"Added {sid}: Magic Barrier Lv{lvl} ({van} -> {new_val})")

    # 4. Add Catalina's Action Effect 103 (Heal power): S027 - S031
    # [6, 8, 10, 13, 16]
    cat_heals = [(1, 4, 6), (2, 6, 8), (3, 8, 10), (4, 10, 13), (5, 12, 16)]
    for lvl, van, new_val in cat_heals:
        sid = f"S0{26+lvl}"
        if sid not in existing_ids:
            ws.append([
                sid, 'skills', "Technique 104 Catalina's Action", 'effect 103', lvl, van, new_val,
                'L13', 'F09', "Catalina modest heal buff for party sustain", '2026-10-10', 'decided'
            ])
            existing_ids.add(sid)
            print(f"Added {sid}: Catalina heal Lv{lvl} ({van} -> {new_val})")

    # 5. Add Equipment class bitmask for Kite Shield and Fluted Armor: W040, W041
    if 'W040' not in existing_ids:
        ws.append([
            'W040', 'equipment', '538 Kite Shield', 'class bitmask', None, '0x3f', '0x2003f',
            'L14', 'F09', "Magical Princess equipment compatibility: enable bit 17", '2026-10-10', 'decided'
        ])
        existing_ids.add('W040')
        print("Added W040: 538 Kite Shield class bitmask -> 0x2003f")

    if 'W041' not in existing_ids:
        ws.append([
            'W041', 'equipment', '666 Fluted Armor', 'class bitmask', None, '0x3f', '0x2003f',
            'L14', 'F09', "Magical Princess equipment compatibility: enable bit 17", '2026-10-10', 'decided'
        ])
        existing_ids.add('W041')
        print("Added W041: 666 Fluted Armor class bitmask -> 0x2003f")

    wb.save(BC_PATH)
    print("Saved balance_changes.xlsx successfully.")

def update_sim_setup():
    wb = openpyxl.load_workbook(SETUP_PATH)
    
    # 1. Update Roles sheet for magical princess
    ws_roles = wb['Roles']
    for row in ws_roles.iter_rows(min_row=2):
        if row[0].value == 'magical princess':
            row[5].value = 'Kite Shield|Leather Shield'
            row[7].value = 'Fluted Armor|Leather Armor'
            row[10].value = 'Magical Change; Sword Mastery; Magical Slash; Magic Barrier; Accelerator; Rainbow Flash; Double Attack; Magical Spark; Cover; Fireball; Ice Edge; Fire Blast'
            print("Updated sim_setup.xlsx Roles: magical princess shield, armor, and skill priority.")
            break

    # 2. Update Policies sheet for magical princess to upkeep Magic Barrier against bosses
    ws_policies = wb['Policies']
    has_mb = False
    for row in ws_policies.iter_rows(min_row=2, values_only=True):
        if row[0] == 'magical princess' and row[3] == 'status:183':
            has_mb = True
            break
            
    if not has_mb:
        mp_rows = []
        for r_idx, row in enumerate(ws_policies.iter_rows(min_row=2), start=2):
            if row[0].value == 'magical princess':
                mp_rows.append(r_idx)
        
        if mp_rows:
            dmg_row_idx = mp_rows[-1]
            ws_policies.cell(row=dmg_row_idx, column=2, value=3)
            ws_policies.insert_rows(dmg_row_idx)
            ws_policies.cell(row=dmg_row_idx, column=1, value='magical princess')
            ws_policies.cell(row=dmg_row_idx, column=2, value=2)
            ws_policies.cell(row=dmg_row_idx, column=3, value='boss')
            ws_policies.cell(row=dmg_row_idx, column=4, value='status:183')
            ws_policies.cell(row=dmg_row_idx, column=5, value='Magic Barrier kept up against bosses')
            print("Inserted Magic Barrier (status:183) into magical princess policy.")

    wb.save(SETUP_PATH)
    print("Saved sim_setup.xlsx successfully.")

if __name__ == '__main__':
    update_balance_changes()
    update_sim_setup()
