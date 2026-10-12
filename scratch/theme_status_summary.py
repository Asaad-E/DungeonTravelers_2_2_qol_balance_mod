import openpyxl
from collections import Counter

wb = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', read_only=True)
ws = wb['Criticisms']
headers = [c for c in next(ws.iter_rows(values_only=True))]
themes = {}
for r in ws.iter_rows(values_only=True):
    d = dict(zip(headers, r))
    if not d.get('ID'): continue
    t = d.get('Theme', 'None')
    themes.setdefault(t, Counter())[d.get('Status', 'None')] += 1

for t, c in sorted(themes.items()):
    print(f"Theme '{t}': {dict(c)}")
