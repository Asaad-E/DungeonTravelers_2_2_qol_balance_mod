import os
import openpyxl

XLSX_PATH = 'mod_project/data_tables/qol_skill_text.xlsx'

test_lines = {
    'TechniqueSkillData_0095.bin': (
        'ATK: 144% -> 147 -> 150 -> 154 -> 158 -> 162 -> 166 -> 170 -> 175 -> 180% (L10)\n'
        'Turns: 3 (Lv1-3) -> 4 (Lv4-6) -> 5 (Lv7-9) -> 6 (L10)    |    TP Cost: 20 -> 42'
    ),
    'TechniqueSkillData_0096.bin': (
        'ATK: 126% (+2-5%/Lv) -> 157%   |   Paralysis: 30% (+3-5%/Lv) -> 64% (L10)\n'
        'Turns: 3 (Lv1-4) -> 4 (Lv5-8) -> 5 (Lv9-10)  |  TP Cost: 24 -> 48  |  Speed: 0.9x'
    ),
    'TechniqueSkillData_0097.bin': (
        'Heal Power: 4 -> 6 -> 8 -> 10 -> 12 (L1-5 Max)\n'
        'Turns: 3 (Lv1-2) -> 4 (Lv3-4) -> 5 (L5 Max)   |   TP Cost: 30 -> 48   |   Speed: 0.9x'
    ),
    'TechniqueSkillData_0098.bin': (
        'ATK: 100% -> 104 -> 108 -> 112 -> 118 -> 123 -> 128 -> 133 -> 138 -> 143% (L10)\n'
        'Turns: 2 (Lv1-4) -> 3 (Lv5-7) -> 4 (Lv8-10)  |  TP Cost: 42 -> 70  |  Speed: 0.75x'
    ),
    'TechniqueSkillData_0100.bin': (
        'Pwr: 180 -> 240 -> 310 -> 390 -> 480 -> 585 -> 700 (L7 Max)\n'
        'Turns: 4    |    TP Cost: 38 (+4-6/Lv) -> 68    |    Speed: 0.82x -> 0.92x'
    ),
}

wb = openpyxl.load_workbook(XLSX_PATH)
ws = wb['Skills']
header = [c.value for c in ws[1]]
col_fn = header.index('Filename')
col_lvup = header.index('QoL_LvUp')

count = 0
for row in ws.iter_rows(min_row=2):
    fn = row[col_fn].value
    if fn in test_lines:
        row[col_lvup].value = test_lines[fn]
        count += 1
        print(f"Updated {fn}")

wb.save(XLSX_PATH)
print(f"Saved {XLSX_PATH}: {count} doll skills updated with clean spacing.")
