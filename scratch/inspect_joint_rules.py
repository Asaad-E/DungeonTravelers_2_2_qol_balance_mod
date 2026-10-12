import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/qol_rules.xlsx', data_only=True)
ws_sentences = wb['JointSkillSentences']
ws_effects = wb['JointEffects']

print("=== JOINT EFFECTS ===")
for r in ws_effects.iter_rows(values_only=True):
    print(r)

print("\n=== JOINT SKILL SENTENCES (first 15) ===")
for r in list(ws_sentences.iter_rows(values_only=True))[:15]:
    print(r)
