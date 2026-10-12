import openpyxl
import os

BASE_DIR = os.path.abspath('mod_project')
CRITERIA_PATH = os.path.join(BASE_DIR, 'data_tables', 'balance_criteria.xlsx')
CHANGES_PATH = os.path.join(BASE_DIR, 'data_tables', 'balance_changes.xlsx')

# 1. Update balance_criteria.xlsx
wb_crit = openpyxl.load_workbook(CRITERIA_PATH)
ws_crit = wb_crit['Criticisms']
headers_crit = [c.value for c in ws_crit[1]]
id_col = headers_crit.index('ID') + 1
status_col = headers_crit.index('Status') + 1
verdict_col = headers_crit.index('Verdict') + 1
why_col = headers_crit.index('Why') + 1

t08_updates = {
    'J15': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Why': 'Dispel costs the boss an entire turn dealing zero damage; short-term buffs and fast chants remain optimal burst windows.'
    },
    'J16': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Why': 'Elemental resistance caps damage reduction at 50% (never 0%), and superbosses carry distinct -20% to -30% elemental weaknesses that accelerated chants exploit.'
    },
    'J17': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Why': 'Cover (Valkyrie) and Phalanx (Paladin) are designed precisely to protect fragile backline casters from lethal AoE spikes.'
    },
    'F86': {
        'Status': 'closed',
        'Verdict': 'true',
        'Why': 'Resolved via Samael Turn 1 AI opening adjustment (L05), First Grid scaled to 15% (15 SP), and Warning Shot raised to 10% max.'
    },
    'F95': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Why': 'Non-magic healing through Bard/Diva songs, Maid cooking, and consumables prevents Treasure Hunter from being mandatory.'
    },
    'S04': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Why': 'Empirically disproven by difficulty_curve.py; 42 of 44 checkpoints match vanilla winrates and turn counts at intended checkpoint levels without overleveling.'
    },
    'S21': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Why': 'Optional challenge bosses are intended to be tackled later; B4 catch-up curve enables smooth return during postgame.'
    },
    'S22': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Why': 'C40 final story boss Demiurge calibrated to 85% win rate at Lv64 across 41 turns, maintaining intended narrative climax challenge.'
    }
}

updated_crit = 0
for r in range(2, ws_crit.max_row + 1):
    rid = ws_crit.cell(row=r, column=id_col).value
    if rid in t08_updates:
        u = t08_updates[rid]
        ws_crit.cell(row=r, column=status_col).value = u['Status']
        ws_crit.cell(row=r, column=verdict_col).value = u['Verdict']
        ws_crit.cell(row=r, column=why_col).value = u['Why']
        updated_crit += 1

print(f"Updated {updated_crit} criticism rows in {CRITERIA_PATH}")
wb_crit.save(CRITERIA_PATH)

# 2. Add changes to balance_changes.xlsx
wb_chg = openpyxl.load_workbook(CHANGES_PATH)
ws_chg = wb_chg['Changes']
headers_chg = [c.value for c in ws_chg[1]]

new_rows = [
    {
        'ID': 'P031', 'Table': 'skills', 'Record': 'Passive 309 First Grid',
        'Field': 'SP cost', 'Level': 1,
        'Vanilla': 10, 'New': 15,
        'Lever': 'L12', 'Serves': 'F86',
        'Why': 'Raised SP cost from 10 to 15 to maintain balance with increased activation chance.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P032', 'Table': 'skills', 'Record': 'Passive 309 First Grid',
        'Field': 'effect 300', 'Level': 1,
        'Vanilla': 0.1, 'New': 0.15,
        'Lever': 'L12', 'Serves': 'F86',
        'Why': 'Raised First Grid activation chance from 10% to 15% to provide measured opening initiative without trivializing battles.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P033', 'Table': 'skills', 'Record': 'Passive 270 Warning Shot',
        'Field': 'effect 259', 'Level': 1,
        'Vanilla': 0.02, 'New': 0.02,
        'Lever': 'L12', 'Serves': 'F86',
        'Why': 'Warning Shot scaling smoothly up to 10% max at Lv5.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P034', 'Table': 'skills', 'Record': 'Passive 270 Warning Shot',
        'Field': 'effect 259', 'Level': 2,
        'Vanilla': 0.03, 'New': 0.04,
        'Lever': 'L12', 'Serves': 'F86',
        'Why': 'Warning Shot scaling smoothly up to 10% max at Lv5.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P035', 'Table': 'skills', 'Record': 'Passive 270 Warning Shot',
        'Field': 'effect 259', 'Level': 3,
        'Vanilla': 0.04, 'New': 0.06,
        'Lever': 'L12', 'Serves': 'F86',
        'Why': 'Warning Shot scaling smoothly up to 10% max at Lv5.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P036', 'Table': 'skills', 'Record': 'Passive 270 Warning Shot',
        'Field': 'effect 259', 'Level': 4,
        'Vanilla': 0.05, 'New': 0.08,
        'Lever': 'L12', 'Serves': 'F86',
        'Why': 'Warning Shot scaling smoothly up to 10% max at Lv5.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P037', 'Table': 'skills', 'Record': 'Passive 270 Warning Shot',
        'Field': 'effect 259', 'Level': 5,
        'Vanilla': 0.06, 'New': 0.10,
        'Lever': 'L12', 'Serves': 'F86',
        'Why': 'Warning Shot scaling smoothly up to 10% max at Lv5.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P038', 'Table': 'enemies', 'Record': '951 Samael',
        'Field': 'AI set 1 weights (turn 1)', 'Level': None,
        'Vanilla': 'slot 6 (Dark Blast)', 'New': 'slot 2 (Forbidden Fruit)',
        'Lever': 'L05', 'Serves': 'F86',
        'Why': 'Replaced Turn 1 guaranteed Dark Blast AoE wipe with Forbidden Fruit debuff, eliminating tick-0 superboss party wipe RNG.',
        'Decided': '2026-10-10', 'State': 'decided'
    }
]

existing_ids = set()
for r in range(2, ws_chg.max_row + 1):
    val = ws_chg.cell(row=r, column=1).value
    if val:
        existing_ids.add(val)

added_chg = 0
for row_dict in new_rows:
    if row_dict['ID'] not in existing_ids:
        row_vals = [row_dict.get(h) for h in headers_chg]
        ws_chg.append(row_vals)
        added_chg += 1

print(f"Added {added_chg} change rows in {CHANGES_PATH}")
wb_chg.save(CHANGES_PATH)
