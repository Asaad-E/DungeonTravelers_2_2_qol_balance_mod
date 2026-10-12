import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', read_only=True)
ws = wb['Criticisms']
headers = [c for c in next(ws.iter_rows(values_only=True))]
for r in ws.iter_rows(values_only=True):
    d = dict(zip(headers, r))
    theme = str(d.get('Theme'))
    if any(t in theme for t in ['T10', 'T13', 'T14']):
        print(f"ID={d.get('ID')} | Theme={theme} | Subject={d.get('Subject')} | Status={d.get('Status')} | Verdict={d.get('Verdict')} | Why={d.get('Why')}")
