import openpyxl, os, struct, sys
sys.path.insert(0, os.path.abspath('tools'))
from pck_handler import PckArchive

base_dir = r'c:\Users\user\Personal\Physics\DungeonTraveler2-2\mod_project'
game_param = r'c:\Users\user\Personal\Physics\DungeonTraveler2-2\DungeonTravelers2-2_EN\Data\ENG\Parameter'

arc_enc = PckArchive(os.path.join(game_param, 'EnchantData.pck'))
for target_id in [23, 66, 270, 275]:
    fn = f'EnchantData_{target_id:04d}.bin'
    for f_name, b in arc_enc.files:
        if f_name == fn:
            p0 = struct.unpack('<I', b[8:12])[0]
            p1 = struct.unpack('<I', b[16:20])[0]
            s0 = b[p0:b.find(b'\x00', p0)].decode('latin-1', errors='ignore') if p0 < len(b) else ''
            s1 = b[p1:b.find(b'\x00', p1)].decode('latin-1', errors='ignore') if p1 < len(b) else ''
            print(f'{fn}: s0="{s0}", s1="{s1}"')

wb_enc = openpyxl.load_workbook(os.path.join(base_dir, 'data_tables', 'enchants.xlsx'))
ws = wb_enc.active
print('Header:', [c for c in next(ws.iter_rows(values_only=True))])
for row in ws.iter_rows(values_only=True):
    for c in row:
        if c and any(k in str(c) for k in ['275', 'Double EXP', '270', 'Dragon Papi', 'Higher-Level']):
            print('Row:', row)
            break
