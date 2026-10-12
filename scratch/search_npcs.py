import os
import glob
import openpyxl

targets = ['rin-rin', 'rinrin', 'rin_rin', 'mary conductor', 'maryconductor', 'conductor']

print("=== SEARCHING TEXT & SOURCE FILES ===")
for root, dirs, files in os.walk('.'):
    if '.git' in root or 'extracted_pcks' in root:
        continue
    for f in files:
        if f.endswith(('.py', '.txt', '.md', '.bin')):
            path = os.path.join(root, f)
            try:
                content = open(path, 'r', encoding='latin1', errors='ignore').read().lower()
                for t in targets:
                    if t in content:
                        print(f'Found "{t}" in {path}')
            except Exception as e:
                pass

print("\n=== SEARCHING EXCEL TABLES ===")
for x in glob.glob('mod_project/data/*.xlsx'):
    try:
        wb = openpyxl.load_workbook(x, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for r in ws.iter_rows(values_only=True):
                for cell in r:
                    if isinstance(cell, str):
                        for t in targets:
                            if t in cell.lower():
                                print(f'Found "{t}" in {os.path.basename(x)} [{sname}]')
    except Exception as e:
        pass
