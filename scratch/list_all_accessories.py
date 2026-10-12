import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/vanilla_source_of_truth.xlsx', data_only=True)
ws = wb['Equipment']
headers = [c.value for c in ws[1]]
accessories = []
for row in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(headers, row))
    if d.get('Category') == 'Accessory':
        accessories.append(d)

print(f"Total accessories: {len(accessories)}")
for a in accessories:
    enchs = [str(a[f'Enchant{i}_Name']) for i in range(1, 5) if a.get(f'Enchant{i}_Name')]
    ench_str = ', '.join(enchs) if enchs else 'None'
    desc = str(a.get('Vanilla_Desc', ''))[:45]
    print(f"{a['Equip_ID']:3d} | {a['Name']:26} | DEF:{a.get('DEF',0):2} | {ench_str:40} | {desc}")
