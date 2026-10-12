import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/vanilla_source_of_truth.xlsx', data_only=True)

print("=== 1. SKILLS FILENAMES & IDS ===")
skills = wb['Skills']
targets = [
    'Provoke', 'Aura Eliminate', 'Attract', 'Self-Sacrifice', 
    'Cover', 'Backward Swing', 'Trump Card', 'Concentrate', 
    'Auto-Guard', 'Imperial Guard', 'Parry', 'Tinker Squash'
]
for row in skills.iter_rows(values_only=True):
    if row[3] in targets:
        print(f"{row[3]:<18} | Filename: {row[1]:<28} | ID: {row[2]}")

print("\n=== 2. ITEMS (Resurrection Stone, Mandra Root, etc.) ===")
items = wb['Items']
for row in items.iter_rows(values_only=True):
    if row[1] in [26, 52, 514, 835]:
        print(f"Item {row[1]}: {row[2]:<26} | type: {row[6]} | val: {row[7]} | type2: {row[8]} | val2: {row[9]} | desc: {row[21]}")

print("\n=== 3. EFFECTS FROM COMMAND DATA ===")
effects = wb['Effects']
for row in effects.iter_rows(values_only=True):
    if row[0] in [0, 26, 48, 54, 170, 171, 264]:
        print(f"Effect {row[0]}: Game Name = '{row[1]}' | Inferred: '{row[2]}'")

print("\n=== 4. SKILL LEVELS: TRUMP CARD, ATTRACT, PROVOKE, MASTER SPELLS ===")
sl = wb['SkillLevels']
for row in sl.iter_rows(values_only=True):
    if row[3] in ['Trump Card', 'Attract', 'Provoke', 'Crimson Flare', 'Cocytus', 'Muon Plasma']:
        print(f"{row[3]:<15} | Field: {row[5]:<10} | Eff: {row[6]} | Vals: {row[10:17]}")
