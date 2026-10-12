import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_changes.xlsx', read_only=True)
print("Sheets:", wb.sheetnames)
ws = wb['Changes']
headers = [c for c in next(ws.iter_rows(values_only=True))]
print("Headers:", headers)
rows = list(ws.iter_rows(values_only=True))
print(f"Total rows in Changes: {len(rows)}")

# Print a few sample rows
for r in rows[:5]:
    print(dict(zip(headers, r)))
