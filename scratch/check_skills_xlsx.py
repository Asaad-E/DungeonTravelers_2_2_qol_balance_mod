import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/skills.xlsx', read_only=True)
print("Sheets in skills.xlsx:", wb.sheetnames)
ws = wb.active
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i < 15:
        print(f"Row {i}: {row}")
    else:
        break
