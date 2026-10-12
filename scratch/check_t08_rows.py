import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', read_only=True)
ws = wb['Criticisms']
headers = [c for c in next(ws.iter_rows(values_only=True))]
t08_rows = []
for r in ws.iter_rows(values_only=True):
    d = dict(zip(headers, r))
    if d.get('Theme') == 'T08 Bosses':
        t08_rows.append(d)

print(f"Total T08 Bosses rows: {len(t08_rows)}")
for r in t08_rows:
    print(f"\nID={r.get('ID')} | Subject={r.get('Subject')}")
    print(f"  Claim: {r.get('Claim')}")
    print(f"  Status: {r.get('Status')}")
    print(f"  Verdict: {r.get('Verdict')}")
    print(f"  Why: {r.get('Why')}")
