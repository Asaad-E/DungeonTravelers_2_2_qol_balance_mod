import sys, os
sys.path.insert(0, os.path.abspath('mod_project/tools'))
import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/vanilla_source_of_truth.xlsx', data_only=True)
print("Sheets in vanilla_source_of_truth:", wb.sheetnames)
if 'Equipment' in wb.sheetnames:
    ws = wb['Equipment']
    headers = [c.value for c in ws[1]]
    print("Headers in Equipment:", headers[:10])
