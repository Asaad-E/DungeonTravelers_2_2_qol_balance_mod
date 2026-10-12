import sys, os
sys.path.insert(0, os.path.abspath('mod_project/tools'))
import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/vanilla_source_of_truth.xlsx', data_only=True)
if 'Raw_Class' in wb.sheetnames:
    ws = wb['Raw_Class']
    for r in ws.iter_rows(min_row=2, values_only=True):
        print(f"Class {r[1]}: {r[2]}")
