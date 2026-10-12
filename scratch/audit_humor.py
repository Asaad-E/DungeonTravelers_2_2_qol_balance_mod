import openpyxl, os

tables_dir = 'mod_project/data_tables'

def check_table(filename, sheetname, orig_col, mod_col, name_col=None):
    path = os.path.join(tables_dir, filename)
    if not os.path.exists(path):
        return
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheetname] if sheetname in wb.sheetnames else wb.active
    
    total = 0
    preserved = 0
    replaced = 0
    examples_replaced = []
    
    for r in range(2, ws.max_row + 1):
        orig = ws.cell(row=r, column=orig_col).value or ""
        mod = ws.cell(row=r, column=mod_col).value or ""
        name = ws.cell(row=r, column=name_col).value if name_col else f"Row {r}"
        
        if not mod or mod == orig:
            continue
            
        total += 1
        # Check if original text (or first line of it) is preserved in mod
        orig_clean = orig.strip()
        first_line_orig = orig_clean.split('\n')[0].strip() if orig_clean else ""
        
        if orig_clean in mod or (first_line_orig and first_line_orig in mod):
            preserved += 1
        else:
            replaced += 1
            examples_replaced.append((name, orig_clean, mod))
            
    print(f"\n==========================================")
    print(f"TABLE: {filename} ({total} active modified rows)")
    print(f"Preserved original text: {preserved}")
    print(f"Completely replaced original text: {replaced}")
    if replaced > 0:
        print(f"--- First 5 Replaced Examples ---")
        for n, o, m in examples_replaced[:5]:
            print(f"[{n}]")
            print(f"  ORIG: {repr(o)}")
            print(f"  MOD : {repr(m)}")

# Let's inspect each table:
# 1. items.xlsx: Name=col 3, Orig=col 5, Mod=col 6
check_table('items.xlsx', 'Items', 5, 6, 3)

# 2. equipment.xlsx: Name=col 3, Orig=col 5, Mod=col 6
check_table('equipment.xlsx', 'Equipment', 5, 6, 3)

# 3. grand_seals.xlsx: Name=col 2, Orig=col 4, Mod=col 5
check_table('grand_seals.xlsx', 'GrandSeals', 4, 5, 2)

# 4. joint_skills.xlsx: Name=col 2, Orig=col 3, Mod=col 4
check_table('joint_skills.xlsx', 'JointSkills', 3, 4, 2)

# 5. skills.xlsx: Name=col 4, Orig=col 7, Mod=col 9
check_table('skills.xlsx', 'Skills', 7, 9, 4)

# 6. food_menu.xlsx: Name=col 2, Orig=col 4, Mod=col 5
check_table('food_menu.xlsx', 'FoodMenu', 4, 5, 2)

# 7. classes.xlsx: Name=col 3, Orig=col 4, Mod=col 5
check_table('classes.xlsx', 'Classes', 4, 5, 3)

# 8. party_skills.xlsx: Name=col 2, Orig=col 4, Mod=col 5
check_table('party_skills.xlsx', 'PartySkills', 4, 5, 2)
