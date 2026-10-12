import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', read_only=True)
ws = wb['Criticisms']
headers = [c for c in next(ws.iter_rows(values_only=True))]
target_ids = ['J04', 'J19', 'J20', 'J21', 'J22', 'J24', 'F01', 'F02', 'F07', 'F85', 'F87', 'F94', 'S07', 'S17', 'F81', 'J13', 'J34']
for r in ws.iter_rows(values_only=True):
    d = dict(zip(headers, r))
    if d.get('ID') in target_ids:
        print(f"ID={d.get('ID')} | Status={d.get('Status')} | Verdict={d.get('Verdict')}")
