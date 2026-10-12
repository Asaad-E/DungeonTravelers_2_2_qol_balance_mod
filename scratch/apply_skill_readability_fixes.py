import openpyxl

TABLE_PATH = 'mod_project/data_tables/qol_skill_text.xlsx'

wb = openpyxl.load_workbook(TABLE_PATH)
ws = wb['Skills']

header = [c.value for c in ws[1]]
name_idx = header.index('Skill_Name') + 1
fn_idx = header.index('Filename') + 1
lvup_idx = header.index('QoL_LvUp') + 1
edited_idx = header.index('Edited') + 1

updates = {
    'TechniqueSkillData_0036.bin': (
        'Fleet Foot',
        'Speed increase: 10% -> 11 -> 12 -> 13 -> 14 -> 15 -> 16 -> 17 -> 18 -> 20% (L10)\n'
        'Turns: 3 (Lv1-3) -> 4 (Lv4-6) -> 5 (Lv7-9) -> 6 (L10)  |  TP Cost: 14 (+2/Lv) -> 32'
    ),
    'TechniqueSkillData_0041.bin': (
        'Chain Spike',
        'Number of attacks: 3 (Lv1-4) -> 4 (Lv5-8) -> 5 (L9-10)   |   ATK: 104% -> 148%\n'
        'Power decrease: 25% -> 14% | TP Cost: 30 (+4-6/Lv) -> 72 | Speed: 0.75x -> 0.85x'
    ),
    'TechniqueSkillData_0018.bin': (
        'Megaton Press',
        'ATK: 120% -> 130 -> 140 -> 148 -> 156 -> 164 -> 170 -> 176 -> 180 -> 184% (L10)\n'
        'Targets: 2 (Lv1-3) -> 3 (Lv4-6) -> 4 (Lv7-8) -> 5 (L9-10) | TP: 26 (+4-6/Lv) -> 68'
    ),
    'MagicSkillData_0219.bin': (
        'Fire Scroll',
        'Max HP damage: 4% -> 5% -> 6% -> 7% -> 8% (L1-5 Max)   |   Accuracy: 64% (+4%/Lv)\n'
        'ATK: -4% (-1%/Lv) (3 Turns) | Damage cap: 250 -> 500 | TP Cost: 20 (+5/Lv) -> 40'
    ),
    'MagicSkillData_0222.bin': (
        'Wind Scroll',
        'Max HP damage: 4% -> 5% -> 6% -> 7% -> 8% (L1-5 Max)   |   Accuracy: 64% (+4%/Lv)\n'
        'Evasion: -7% (-2%/Lv) (3 Turns) | Damage cap: 250 -> 500 | TP Cost: 20 (+5/Lv) -> 40'
    ),
    'MagicSkillData_0223.bin': (
        'Earth Scroll',
        'Max HP damage: 4% -> 5% -> 6% -> 7% -> 8% (L1-5 Max)   |   Accuracy: 64% (+4%/Lv)\n'
        'DEF: -4% (-1%/Lv) (3 Turns) | Damage cap: 250 -> 500 | TP Cost: 20 (+5/Lv) -> 40'
    ),
    'MagicSkillData_0231.bin': (
        'Carrot and Stick',
        'DEF: -5% -> -6% -> -7% -> -8% -> -10% (L1-5 Max)   |   Magic defense: -10% -> -15%\n'
        'Accuracy: 80% -> 90% | HP recovery: 35 -> 25 | Turns: 3 | TP: 18 (+2-4/Lv) -> 32'
    )
}

applied = 0
for row_idx in range(2, ws.max_row + 1):
    fn = ws.cell(row=row_idx, column=fn_idx).value
    if fn in updates:
        expected_name, new_lvup = updates[fn]
        actual_name = ws.cell(row=row_idx, column=name_idx).value
        assert actual_name == expected_name, f"Name mismatch: {actual_name} != {expected_name}"
        ws.cell(row=row_idx, column=lvup_idx, value=new_lvup)
        ws.cell(row=row_idx, column=edited_idx, value='yes')
        print(f"Updated row {row_idx}: {actual_name} ({fn})")
        applied += 1

assert applied == len(updates), f"Expected {len(updates)} updates, but applied {applied}"

wb.save(TABLE_PATH)
print("Saved qol_skill_text.xlsx successfully!")
