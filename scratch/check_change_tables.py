import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_changes.xlsx', read_only=True)
ws = wb['Changes']
headers = [c for c in next(ws.iter_rows(values_only=True))]
tables = set()
for r in ws.iter_rows(values_only=True):
    d = dict(zip(headers, r))
    if d.get('Table'):
        tables.add(d['Table'])
print("Tables in Changes:", tables)

# Print a few rows from each non-skill table
for r in ws.iter_rows(values_only=True):
    d = dict(zip(headers, r))
    if d.get('Table') in ['sp_curve', 'equipment', 'enemies']:
        if d['ID'] in ['B001', 'W001', 'W036', 'A001', 'E001', 'G001']:
            print(f"{d['ID']}: Table={d['Table']} | Record={d['Record']} | Field={d['Field']} | Vanilla={d['Vanilla']} | New={d['New']} | Lever={d['Lever']} | Serves={d['Serves']}")
