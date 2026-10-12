import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/skills.xlsx', read_only=True)
ws = wb.active
headers = [c for c in next(ws.iter_rows(values_only=True))]
print("Headers:", headers)

for row in ws.iter_rows(values_only=True):
    cat = row[1]
    name = row[3]
    max_lv = row[4]
    tp = row[5]
    mod_lvup = row[9]
    if cat == 'Magic':
        print(f"Magic: {name:<20} | MaxLv: {max_lv:<2} | TP: {tp}")
        if mod_lvup:
            print(f"   ModLvUp: {mod_lvup.splitlines()[0] if mod_lvup else ''}")
