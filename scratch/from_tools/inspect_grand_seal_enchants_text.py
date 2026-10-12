import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/qol_enchant_text.xlsx', data_only=True)
ws = wb.active
headers = [c.value for c in ws[1]]

target_ench_ids = {212, 214, 217, 219, 220, 222, 223, 224, 225, 246, 248, 257, 258, 259, 260, 262, 263, 264, 265, 266, 267, 268, 270, 271, 272, 280, 281, 290, 291, 292, 293, 294, 296, 297, 298}

for row in ws.iter_rows(min_row=2, values_only=True):
    eid = row[0]
    if eid in target_ench_ids:
        vname = row[2]
        qol_name = row[6]
        note = row[11] if len(row)>11 else ''
        print(f"ID {eid:3d}: Vanilla='{vname}' | QoL='{qol_name}' | Note='{note}'")
