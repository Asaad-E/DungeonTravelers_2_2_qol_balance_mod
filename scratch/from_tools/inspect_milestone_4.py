import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', data_only=True)
ws = wb['Criticisms']
headers = [c.value for c in ws[1]]

milestone_4_themes = ['T10', 'T13', 'T14']
m4_rows = []

for row in ws.iter_rows(min_row=2, values_only=True):
    rd = dict(zip(headers, row))
    theme = str(rd.get('Theme') or '')
    rid = rd.get('ID')
    if any(theme.startswith(t) for t in milestone_4_themes) or rid == 'F14':
        m4_rows.append(rd)

print(f"Total rows in Milestone 4 (T10, T13, T14): {len(m4_rows)}")
for r in m4_rows:
    print(f"ID: {r['ID']} | Theme: {r['Theme']} | Subject: {r.get('Subject')}")
    print(f"  Claim: {r.get('Claim')}")
    print(f"  Status: {r.get('Status')} | Verdict: {r.get('Verdict')}")
    print(f"  Vanilla_Result: {r.get('Vanilla_Result')}")
    print(f"  Balance_Result: {r.get('Balance_Result')}")
    print(f"  Proposal: {r.get('Proposal')}")
    print(f"  Decision: {r.get('Decision')}")
    print("-" * 60)
