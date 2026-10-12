import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', data_only=True)
ws = wb['Criticisms']
headers = [c.value for c in ws[1]]

for r in ws.iter_rows(min_row=2, values_only=True):
    rd = dict(zip(headers, r))
    if rd['ID'] == 'S20':
        for k, v in rd.items():
            print(f"{k}: {v}")
