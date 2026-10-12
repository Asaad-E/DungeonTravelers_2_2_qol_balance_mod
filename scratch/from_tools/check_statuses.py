import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', data_only=True)
ws = wb['Criticisms']
headers = [c.value for c in ws[1]]

for row in ws.iter_rows(min_row=2, values_only=True):
    rd = dict(zip(headers, row))
    if rd['Status'] in ['not a problem', 'out of the balance']:
        print(f"{rd['ID']}: Status={rd['Status']} | Verdict={rd.get('Verdict')} | Why={rd.get('Why')} | Decision={rd.get('Decision')}")
