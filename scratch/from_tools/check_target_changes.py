import openpyxl

wb_crit = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', data_only=True)
ws_crit = wb_crit['Criticisms']
headers_crit = [c.value for c in ws_crit[1]]

target_ids = ['J04', 'J19', 'J20', 'J21', 'J22', 'J24', 'F01', 'F02', 'F07', 'F85', 'F87', 'F94', 'S07', 'S17', 'F81', 'J13', 'J34']

crit_data = {}
for row in ws_crit.iter_rows(min_row=2, values_only=True):
    rd = dict(zip(headers_crit, row))
    if rd['ID'] in target_ids:
        crit_data[rd['ID']] = rd

wb_ch = openpyxl.load_workbook('mod_project/data_tables/balance_changes.xlsx', data_only=True)
ws_ch = wb_ch['Changes']
headers_ch = [c.value for c in ws_ch[1]]

changes_by_row = {}
for row in ws_ch.iter_rows(min_row=2, values_only=True):
    rd = dict(zip(headers_ch, row))
    serves = str(rd.get('Serves_Row') or '')
    for tid in target_ids:
        if tid in serves:
            changes_by_row.setdefault(tid, []).append(rd['Change_ID'])

print("Changes associated with target rows:")
for tid in target_ids:
    print(f"{tid}: {changes_by_row.get(tid, [])}")
