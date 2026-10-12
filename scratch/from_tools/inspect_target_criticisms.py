import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', data_only=True)
ws = wb['Criticisms']
headers = [c.value for c in ws[1]]

target_ids = {'J04', 'J19', 'J20', 'J21', 'J22', 'J24', 'F01', 'F02', 'F07', 'F85', 'F87', 'F94', 'S07', 'S17', 'F81', 'J13', 'J34'}

for row in ws.iter_rows(min_row=2, values_only=True):
    row_dict = dict(zip(headers, row))
    rid = row_dict.get('ID')
    if rid in target_ids:
        print(f"ID: {rid} | Theme: {row_dict.get('Theme')} | Subject: {row_dict.get('Subject')}")
        print(f"  Claim: {row_dict.get('Claim')}")
        print(f"  Status: {row_dict.get('Status')} | Verdict: {row_dict.get('Verdict')}")
        print(f"  Proposal: {row_dict.get('Proposal')}")
        print(f"  Decision: {row_dict.get('Decision')}")
        print(f"  Vanilla_Result: {row_dict.get('Vanilla_Result')}")
        print(f"  Balance_Result: {row_dict.get('Balance_Result')}")
        print("-" * 60)
