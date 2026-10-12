import openpyxl

wb_ch = openpyxl.load_workbook('mod_project/data_tables/balance_changes.xlsx', data_only=True)
print("Sheet names:", wb_ch.sheetnames)

ws_ch = wb_ch['Changes']
headers = [c.value for c in ws_ch[1]]
print("Headers:", headers)

from collections import Counter
prefixes = Counter()
for row in ws_ch.iter_rows(min_row=2, values_only=True):
    cid = str(row[0] or '')
    if cid:
        prefixes[cid[0]] += 1

print("Prefixes count:", prefixes)
