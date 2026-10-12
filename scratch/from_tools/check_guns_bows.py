import openpyxl

wb_ch = openpyxl.load_workbook('mod_project/data_tables/balance_changes.xlsx', data_only=True)
ws_ch = wb_ch['Changes']
headers = [c.value for c in ws_ch[1]]

matches = []
for row in ws_ch.iter_rows(min_row=2, values_only=True):
    rd = dict(zip(headers, row))
    desc = str(rd.get('Description') or '')
    serves = str(rd.get('Serves_Row') or '')
    target = str(rd.get('Target') or '')
    field = str(rd.get('Field') or '')
    text_all = f"{desc} {serves} {target} {field}".lower()
    if any(k in text_all for k in ['gun', 'bow', 'j06', 'f10', 'f87', 'weapon_flag', 'weapon flag']):
        matches.append((rd.get('Change_ID'), serves, target, field, rd.get('Vanilla_Value'), rd.get('Modified_Value'), desc))

print(f"Total matches in balance_changes.xlsx: {len(matches)}")
for m in matches:
    print(m)

print("\n--- Check balance_criteria.xlsx for J06 and F10 ---")
wb_crit = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', data_only=True)
ws_crit = wb_crit['Criticisms']
headers_crit = [c.value for c in ws_crit[1]]
for row in ws_crit.iter_rows(min_row=2, values_only=True):
    rd = dict(zip(headers_crit, row))
    if rd.get('ID') in ['J06', 'F10', 'F87']:
        print(rd.get('ID'), rd.get('Subject'), rd.get('Claim'), "Verdict:", rd.get('Verdict'), "Proposal:", rd.get('Proposal'), "Decision:", rd.get('Decision'))
