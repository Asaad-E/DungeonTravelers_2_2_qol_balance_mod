import sys, os
sys.path.insert(0, os.path.abspath('mod_project/tools'))
import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/vanilla_source_of_truth.xlsx', data_only=True)

# 1. Skills sheet
ws_skills = wb['Skills']
headers_skills = [c.value for c in ws_skills[1]]
print("=== SKILLS SHEET ===")
for r in ws_skills.iter_rows(min_row=2, values_only=True):
    if r[2] in [270, 309]:
        d = dict(zip(headers_skills, r))
        print(f"\nSkill ID {d.get('Skill_ID')} ({d.get('Name')}):")
        for k, v in d.items():
            if v is not None:
                print(f"  {k}: {v}")

# 2. SkillLevels sheet
ws_levels = wb['SkillLevels']
headers_levels = [c.value for c in ws_levels[1]]
print("\n=== SKILL LEVELS SHEET ===")
for r in ws_levels.iter_rows(min_row=2, values_only=True):
    if r[2] in [270, 309]:
        d = dict(zip(headers_levels, r))
        print(f"ID {d.get('Skill_ID')} | Field={d.get('Field')} | Effect={d.get('Effect_ID')} ({d.get('Effect_Name')}) | Type={d.get('Data_Type')}")
        vals = [d.get(f'Lv{k}') for k in range(1, 11) if d.get(f'Lv{k}') is not None]
        print(f"  Levels: {vals}")
