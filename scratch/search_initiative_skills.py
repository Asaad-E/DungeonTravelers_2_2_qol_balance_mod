import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/vanilla_source_of_truth.xlsx', data_only=True)
ws = wb['Skills']
headers = [c.value for c in ws[1]]

found = []
for r in ws.iter_rows(min_row=2, values_only=True):
    name = str(r[3])
    if any(k in name.lower() for k in ['warning', 'first', 'quick', 'speed', 'delay', 'initiative', 'preemptive', 'slow', 'haste', 'fast', 'action']):
        found.append((r[1], str(r[2]), r[3]))

print(f"Skills found: {len(found)}")
for fn, sid, name in found:
    print(f"ID {sid:5s} ({fn}): {name}")
