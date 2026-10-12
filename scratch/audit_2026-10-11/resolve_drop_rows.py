"""Rows P011 - P025 of balance_changes.xlsx after the reading of the drop (scratch/drops_reading.md, checked against
0x558517 - 0x5586fa and the enemy records), and the decision of the user of 2026-10-11 on catching up (B4 alone, the
kept curve not built) in the register.

    python resolve_drop_rows.py
"""
import os
import sys
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(PROJECT, 'tools'))
import enemy_data                                            # noqa: E402

CHANGES = os.path.join(PROJECT, 'data_tables', 'balance_changes.xlsx')
CRITERIA = os.path.join(PROJECT, 'data_tables', 'balance_criteria.xlsx')
FLOOR = 2.0            # percent: the lowest chance of a fixed piece among the common enemies once the rows are built
TODAY = '2026-10-11'
# (row, enemy, piece): P024 is the row Garasha already had; the others are the enemies the same rule finds
PIECES = (('P024', 6, 339), ('P079', 212, 656), ('P080', 732, 761), ('P081', 236, 223), ('P082', 236, 764),
          ('P083', 933, 643), ('P084', 65, 340))
GUARANTEED = ('Dropped on 2026-10-11: the piece of this boss is already certain in vanilla. Its record holds a drop '
              'chance of 100 (0x7C), no item, and one slot (0x80 = 1) that names the piece with weight 100 (0x8C); the '
              'word the row changed is that count of slots, not a chance (0x558517; scratch/drops_reading.md).')


def chance(enemy, weights):
    """Percent of each slot: the chance of the equipment drop times the share of its weight (0x558331, 0x5585a8)."""
    total = float(sum(weights))
    return [enemy['equipment_drop'][0] * w / total for w in weights]


def main():
    enemies = {e['id']: e for e in enemy_data.enemies()}
    wb = openpyxl.load_workbook(CHANGES)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    col = {name: head.index(name) for name in head}
    by_id = {row[0].value: row for row in ws.iter_rows(min_row=2) if row[0].value}
    for number in range(11, 24):
        row = by_id[f'P{number:03d}']
        row[col['State']].value, row[col['Why']].value = 'dropped', GUARANTEED
    ryoma = by_id['P025']
    ryoma[col['State']].value = 'dropped'
    ryoma[col['Why']].value = ('Dropped on 2026-10-11: the word the row changed is the count of slots. The Sudarshana of '
                               'Ryoma falls 3.2% of the time (chance 16, weight 20 of 100), which is what a fixed piece '
                               'of a common enemy falls at across the game (2 to 4.5%).')
    wanted = {}
    for _, eid, piece in PIECES:
        wanted.setdefault(eid, []).append(piece)
    new_weight = {}
    for eid, pieces in wanted.items():
        enemy = enemies[eid]
        slots = enemy['equipment_slots']
        weight = 1
        while True:
            weights = [weight if zone == 0 and pid in pieces else w for zone, pid, w in slots]
            if min(c for c, (zone, pid, _) in zip(chance(enemy, weights), slots) if zone == 0 and pid in pieces) >= FLOOR:
                break
            weight += 1
        new_weight[eid] = (weight, weights)
    for rid, eid, piece in PIECES:
        enemy = enemies[eid]
        slots = enemy['equipment_slots']
        k = next(i for i, (zone, pid, _) in enumerate(slots) if zone == 0 and pid == piece)
        weight, weights = new_weight[eid]
        before, after = chance(enemy, [w for _, _, w in slots])[k], chance(enemy, weights)[k]
        values = {'ID': rid, 'Table': 'enemies', 'Record': f"{eid} {enemy['name']}",
                  'Field': f'equipment slot {k + 1} weight (piece {piece})', 'Level': None, 'Vanilla': slots[k][2],
                  'New': weight, 'Lever': 'L01', 'Serves': 'F90',
                  'Why': f'The fixed pieces of common enemies fall 2 to 4.5% of the time; seven fall under {FLOOR:g}%, '
                         f'this one {before:.2f}%: chance of the equipment drop {enemy["equipment_drop"][0]} (0x7C) times '
                         f'its weight {slots[k][2]} of {sum(w for _, _, w in slots)} (0x558570 - 0x5585a8). Raised to the '
                         f'lowest weight that gives {FLOOR:g}%: {after:.2f}%. The weights of the other slots are not '
                         f'touched. Reading: scratch/drops_reading.md, checked on 2026-10-11.',
                  'Decided': TODAY, 'State': 'decided'}
        if rid in by_id:
            for name, value in values.items():
                by_id[rid][col[name]].value = value
        else:
            ws.append([values[name] for name in head])
        print(rid, values['Record'], values['Field'], values['Vanilla'], '->', weight, f'{before:.2f}% -> {after:.2f}%')
    wb.save(CHANGES)

    wb = openpyxl.load_workbook(CRITERIA)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    col = {name: head.index(name) for name in head}
    rows = {row[0].value: row for row in ws.iter_rows(min_row=2) if row[0].value}

    def put(rid, **cells):
        for name, value in cells.items():
            rows[rid][col[name]].value = value

    alone = ('B4 alone (the user, 2026-10-11): the factors by level gap are in the program, the EXP of a level stays as in '
             'vanilla. The kept curve is not built: it would put a party 3 to 9 levels above the level the bosses were '
             'fitted at, and can be added later as data')
    put('U01', Proposal=alone + '. 10 levels behind: 20 - 47 battles (vanilla 48 - 144); 20 behind: 24 - 82 (65 - 311); '
        'from level 1: 24 - 107 (65 - 519) (docs/generated/CATCH_UP_OPTIONS.md, option B4)',
        Decision='2026-10-05: kept curve and factors B4; 2026-10-11: B4 alone, installed; joining level may be raised')
    put('J37', Proposal='Decided: B4 alone and one more point on the four upper steps of the reset (L25, L27), installed '
        '2026-10-11: 400 battles for one crown, 17 resets from level 70 (vanilla 4,393)',
        Decision='2026-10-05: B4 and one more point per reset; 2026-10-11: B4 alone, the step of 59 levels keeps 1 point')
    put('F25', Proposal='Decided: as J37, 400 battles, best by 17 resets from 70',
        Decision='2026-10-05: as J37; 2026-10-11: as J37')
    put('S11', Why='The growth tables give the same per level before and after 70, so as stated it is not in the data. '
        'What rises is the EXP each level costs; the gap factors answer that for whoever is below the enemies (a level '
        'costs 25 battles at 70 instead of 35, 11 at 90 instead of 18)',
        Proposal='Nothing beyond the factors B4: the kept curve is not built (the user, 2026-10-11)',
        Decision='2026-10-09: nothing beyond the kept curve; 2026-10-11: nothing beyond B4, the curve not built')
    why = rows['F92'][col['Why']].value
    put('F92', Why=why.replace('(option Kept + B4)', '(option B4, without the kept curve)')
        .replace('about 257 battles', 'about 400 battles'))
    put('F90', Vanilla_Result='The piece of a postgame boss is certain (drop chance 100, one slot of weight 100). A fixed '
        'piece of a common enemy falls 2 to 4.5% of the time, of a very powerful enemy 10 to 50%; seven pieces of six '
        'common enemies fall under 2% (0.8 to 1.5%). What is stolen and how often is another rule (FORMULAS 12.16)',
        Why='Read on 2026-10-11 (scratch/drops_reading.md, checked against 0x558517 - 0x5586fa): rows P011 - P023 and '
        'P025 of balance_changes.xlsx dropped, they changed the count of slots of drops that are certain or usual. '
        'Rows P024, P079 - P084 raise the seven pieces under 2% to 2%.')
    wb.save(CRITERIA)


if __name__ == '__main__':
    main()
