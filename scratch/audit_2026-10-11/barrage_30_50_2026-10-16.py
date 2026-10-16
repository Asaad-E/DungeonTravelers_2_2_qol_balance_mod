"""One-off edit: Barrage (passive 302, effect 251) from 30% at level 1 to 50% at level 10, rows W026 - W035 of
balance_changes.xlsx. A passive for one situation (the plain attack of the Kunoichi): at 25% - 40% it changed no fight
(passive_worth_2026-10-16.py, beside this file)."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
NEW = [0.3, 0.32, 0.34, 0.36, 0.38, 0.4, 0.42, 0.44, 0.47, 0.5]
WHY = ('A passive for one situation, the plain attack of the Kunoichi: at 25% - 40% (+25% over vanilla, 2026-10-09) it '
       'changed no fight, her plain attacks being 0.7 a fight at C38 (scratch/audit_2026-10-11/passive_worth_2026-10-16.py); '
       '30% at level 1 to 50% at level 10 (decided 2026-10-16)')

wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
done = 0
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 3).value == 'Passive 302 Barrage' and ws.cell(r, 4).value == 'effect 251':
        level = int(ws.cell(r, 5).value)
        ws.cell(r, 7).value = NEW[level - 1]
        ws.cell(r, 10).value = WHY
        ws.cell(r, 11).value = '2026-10-16'
        done += 1
assert done == 10, done
wb.save(PATH)
print('rows', done)
