"""One-off edit (2026-10-13): a turn of a status takes 20 steps of the map in the place of 5 (row X15 of the sheet
ExePatches, lever L47), so that a meal lasts; row J36 of the register closed; docs/FORMULAS.md 12.18 says which
statuses the count belongs to. Run once; kept as the record of the edit.
"""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CRITERIA = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
FORMULAS = os.path.join(ROOT, 'docs', 'FORMULAS.md')
X15 = ('X15', 'balance', 'L47', '0x48d5f3', '05', '14', None, None,
       'Steps of the map that count as one turn of a status with a number of turns (cmp dword ptr [esi + 0x2c], 5 at 0x48d5f0, '
       'reached only when the tick is not of a battle: the word at +0x40 of its second argument is 0). 5 -> 20: a meal lasts '
       '60 to 200 steps in the place of 15 to 50. Decided 2026-10-13')
L47 = ('L47', 'Steps of the map that count as one turn of a status with a number of turns', 'executable, 0x48d5f3', 'one byte',
       'Executable', 'A byte from 1 to 127; the same instruction, the same length',
       'docs/FORMULAS.md 12.18: the tick of the 97 effects of the class CCommandTurnBase (0x48d590). The ailments that end on '
       'a roll (0x48d180, its own count at 0x48d202) and petrification (0x483010) have theirs and are left')
J36 = {'Status': 'closed',
       'Balance_Result': 'Row X15 of the sheet ExePatches: a turn of a status takes 20 steps of the map in the place of 5, so a '
                         'meal lasts 60 to 200 steps, two to eight random battles of walking. Battles count turns as before',
       'Why': 'The buff of a meal was gone a battle later. The count of steps is one byte that only the map uses; with it the '
              'buffs and the lowered stats a member carries out of a battle last four times as long on the map too, while '
              'poison and the other ailments that end on a roll keep their own count',
       'Decision': '2026-10-13: approved, 20 steps a turn: a buff with its drawback'}
OLD = """* On the map a turn is **five steps**: the method runs on every step but only every fifth one counts (`0x48d5ed`).
"""
NEW = """* On the map a turn is **five steps**: the method runs on every step but only every fifth one counts (`0x48d5ed`;
  the count is skipped when the tick comes from a battle, the word at `+0x40` of its second argument). This is the
  tick of the 97 effects whose class is `CCommandTurnBase`: the buffs and the lowered stats with a number of turns,
  the buff of a meal among them (`measures/DESIGN_ROWS.md`). An ailment that ends on a roll has its own count of
  steps (`0x48d1fc - 0x48d208`: 5, or 2 with a piece of enchantment 265), and so has petrification (`0x483038`). The
  Balance variant makes the first one 20 (row X15 of the sheet `ExePatches`), which leaves battles and those ailments
  as they are.
"""


def main():
    wb = openpyxl.load_workbook(CRITERIA)
    ws = wb['ExePatches']
    assert not any(r[0].value == 'X15' for r in ws.iter_rows())
    at = next(r[0].row for r in ws.iter_rows() if r[0].value == 'X14')
    ws.insert_rows(at + 1)
    for k, value in enumerate(X15):
        ws.cell(at + 1, k + 1).value = value
    ws = wb['Levers']
    at = next(r[0].row for r in ws.iter_rows() if r[0].value == 'L46')
    ws.insert_rows(at + 1)
    for k, value in enumerate(L47):
        ws.cell(at + 1, k + 1).value = value
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == 'J36')
    for column, value in J36.items():
        row[head.index(column)].value = value
    wb.save(CRITERIA)
    with open(FORMULAS, encoding='utf-8') as f:
        text = f.read()
    assert text.count(OLD) == 1
    with open(FORMULAS, 'w', encoding='utf-8', newline='\n') as f:
        f.write(text.replace(OLD, NEW))
    print('X15, L47, J36')


if __name__ == '__main__':
    main()
