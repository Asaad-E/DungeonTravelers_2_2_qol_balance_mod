"""Moves the tutorial texts of scratch/update_euthys_and_teaching.py into a workbook,
data_tables/balance_tutorial_text.xlsx, so that they stop living in code and in the game folder.

One row per replaced box: the pack, the entry, the whole original box, the text, whether the line is voiced, and a
state: checked (every statement set against the executable or the data, with where), unverified (it states something
nothing has read yet: not to be built until read), held (withdrawn: the original line stays). The corrections of the
audit of 2026-10-11 are applied here; the reason of each is in the column Note.

    python tutorials_to_workbook.py
"""
import os
import re
import sys
import importlib.util
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, os.path.join(PROJECT, 'tools'))
from pck_handler import PckArchive                          # noqa: E402

SOURCE = os.path.join(PROJECT, 'scratch', 'update_euthys_and_teaching.py')
OUTPUT = os.path.join(PROJECT, 'data_tables', 'balance_tutorial_text.xlsx')
MAX_LINES, MAX_WIDTH = 3, 70

# entry -> (state, text that takes the place of the one of the script or None, note)
T, L = 'CheatsTeaching_%s.adv.bin', 'BattleTeacher_%s.adv.bin'
REVIEW = {
    T % '01': ('checked', None, 'WeaponRankData: twelve ranks, 1.00 to 2.90. Defend: gauge value x1.15 (0x44a326)'),
    T % '02': ('checked', None, 'General advice; TP cost and chant speed by level are in the skill files'),
    T % '03': ('checked', None, 'FORMULAS 5: F1 of the action, F2 of the chant; dolls act on their own turns (11.1)'),
    T % '04': ('checked', 'Each Sealbook offers different enchantments. A character who\ncarries one gains them as if they were on her equipment,\nso make plenty and use them to strengthen your party.',
               'Was "permanently unlocks monster Level, HP and Resistances" and "forge them onto weapons (+1 to +9)": '
               'neither is read. Read: the enchantments of a carried sealbook join the list of her pieces (0x4f7b5d)'),
    T % '05': ('checked', 'Carry one Grand Sealbook at a time: its enchantment works for\nthe whole party, and the mutant it holds fights beside you\nas an ally you do not command.',
               'Was "drop with 100% guarantee": the drop is not read. FORMULAS 16.1 and 11: one carried, its '
               'enchantment for the party, its summon in the party acting by its own script'),
    T % '06': ('checked', 'Each stat of a summon follows the best member of the party in\nthat stat; its form follows '
               'the average level of the party.\nSummons act on their own turns and cannot be commanded.',
               'Was "party averages (level, ATK, INT)": 0x4ffc90 takes the highest of the members for each stat and '
               'only the form by the mean level. "Heals and buffs do not affect summons" was not read: taken out'),
    T % '07': ('held', None, 'What a sealbook reveals is not read: the original line stays'),
    T % '08': ('held', None, 'The contents of the battle log are not read: the original line stays'),
    T % '09': ('checked', None, 'As Teaching 01'),
    T % '10': ('checked', 'Chant interrupt chance scales with damage taken vs Max HP.\nAny Critical Hit also has a '
               'flat 1 in 3 chance to interrupt!\nOne heavy hit or a high-crit skill stops a spell best.',
               'Was "multihit attacks to reliably stop spells": the rate is the damage of each hit against max HP '
               '(0x44d001), so many small hits interrupt less. The second roll is 33 in 100 (0x44d238)'),
    T % '11': ('checked', None, '0x44a7d0: short 0.75 and 0.5, mid 0.75 only when both stand behind, long full'),
    T % '12': ('checked', None, '0x44cd3b, 0x44eff9: x1.5 from blows and spells; 0x47ea6b: any damage wakes; 0x44af7e: '
               'hit for certain'),
    T % '13': ('checked', 'Each point of resistance cuts elemental damage by 0.5%, so\n100 resistance halves it. At '
               '100 resistance you are also\ncompletely immune to the status ailments of that element.',
               'Was "200 res = 100% immune": the enchantment sum stops at 100 (0x4f8490), so 200 is not a figure a '
               'player can plan on. Half a percent per point and the block at 100 are read (0x652b64, 0x481fc2)'),
    T % '14': ('checked', None, 'QuestRankData: bag +5 at ranks 8, 11, 15 and 20; the text names no rank'),
    T % '15': ('checked', None, 'FORMULAS 16.3: rank = level / 10, EX once the reset bonus reaches 50'),
    T % '16': ('held', None, 'No reading gives "+1% per 5 awards": PartySkillData holds no amount (QOL_ROADMAP)'),
    T % '17': ('checked', 'Joint skills unleash devastating power!\nEach joint skill can only be used once per battle, so time\ntheir activation carefully to break through dangerous phases.',
               'Was "pierce defenses": not read (a joint skill goes through the magic formula with its Power, '
               'FORMULAS 15). Once per battle is the original line'),
    T % '18': ('held', None, 'The levels a reset may stop at are not read: the original line stays'),
    T % '19': ('checked', None, "True on the Balance variant only: row W042 gives the Slave's Torc Double EXP; the gap "
               'factors are 0x450739'),
    T % '20': ('checked', None, 'FORMULAS 16.2: dolls are techniques and act in a zone that forbids magic'),
    T % '21': ('checked', None, 'QuestRankData: rank 12 opens a mercenary, rank 19 raises their levels'),
    T % '22': ('checked', None, 'General advice, nothing stated'),
    T % '23': ('checked', None, 'EnemyAIData: scripts switch sets on the HP share of the enemy (FORMULAS 9, source 0)'),
    T % '24': ('checked', None, 'Barrage Rondo, Provoke with Dark Aura and the critical second roll all exist as named'),
    T % '25': ('checked', None, 'FORMULAS 10: a reset of 59 levels or more adds 1 (up to 5 at 98); resistances too '
               '(0x4f5f04)'),
    L % '04': ('checked', None, 'As Teaching 01'),
    L % '07': ('checked', None, '0x44a7d0'),
    L % '08': ('checked', None, 'Resistance 100 halves (0x652b64)'),
    L % '09': ('checked', None, 'QuestRankData; the text names no rank'),
    L % '11-12': ('checked', None, 'True on the Balance variant only (row W042)'),
    L % '13': ('checked', '"I visit often. A hearty meal restores you before a journey,\nand the finest dishes send '
               'you off with Brave, Protect or\nthe three physical guards already cast. Exquisite and practical."',
               'Was "elemental guards that even Dispel traps cannot erase": FoodMenuData gives Brave (Mega-Meaty '
               'Ramen), Protect (Ramen Carbonara) and the Stab, Slash and Bash Guards (Ultimate Champon); nothing read '
               'says a trap cannot dispel them'),
    L % 'Sp01': ('checked', None, 'Circle Protect is the spell of the Enchantress (MagicSkillData 227)'),
    L % 'Sp02': ('checked', None, 'Dark Aura answers a short-range hit of any attacker without Anti-Dark Aura (12.9)'),
    L % 'Sp03': ('checked', None, 'FORMULAS 10 and 16.3'),
}
# Lesson 10 replaces two boxes: the first one states the awards figure, the second is advice
LESSON_10 = {'awards': ('held', None, 'No reading gives "+1% per 5 awards"'),
             'advice': ('checked', '"*giggle* Don\'t be so timid, Fried. Examine their requirements\nin the menu and coordinate your turns to store Pep. A well-timed\nJoint Skill can turn the tide in an instant."',
                        'Was "pierce tough defenses": not read. Pep needed and spent: FORMULAS 13')}


def boxes(blob):
    return [m.group()[:-1].decode('latin-1') for m in re.finditer(rb'[\x20-\x7e\n]{6,}\x00', blob)]


def voiced(blob, original):
    """The id of the voice clip that plays with the box, when the command before it is one."""
    found = boxes(blob)
    if original not in found:
        return None
    before = found[max(0, found.index(original) - 2):found.index(original)]
    return next((s for s in reversed(before) if s.startswith('VC_')), None)


def main():
    spec = importlib.util.spec_from_file_location('tutorial_script', SOURCE)
    script = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(script)
    rows, problems = [], []
    for pack, table in (('Lesson.pck', script.EUTHYS_REPLACEMENTS), ('Teaching.pck', script.TEACHING_REPLACEMENTS)):
        archive = dict(PckArchive(os.path.join(PROJECT, 'backups', pack)).files)
        for entry, replacements in table.items():
            for original, text in replacements.items():
                key = entry
                if entry == L % '10':
                    state, corrected, note = LESSON_10['awards' if 'awards you amass' in original else 'advice']
                else:
                    state, corrected, note = REVIEW[key]
                text = corrected or text
                whole = next((b for b in boxes(archive[entry]) if original in b), None)
                if whole is None:
                    problems.append(f'{entry}: the original line is not in the backup')
                    continue
                box = whole.replace(original, text)
                lines = box.strip('"').split('\n')
                if len(lines) > MAX_LINES or max(len(x) for x in lines) > MAX_WIDTH:
                    problems.append(f'{entry}: {len(lines)} lines, widest {max(len(x) for x in lines)}')
                rows.append((pack, entry, whole, box if state != 'held' else None, voiced(archive[entry], whole) or '',
                             state, note))
    if problems:
        raise SystemExit('\n'.join(problems))
    wb = openpyxl.Workbook()
    about = wb.active
    about.title = 'About'
    for line in (
            ('Topic', 'Detail'),
            ('What', 'The tutorial texts of the Balance variant: lines of the lessons of Euthys (Lesson.pck) and of the '
                     'library (Teaching.pck). They state figures of the Balance variant, so they belong to that build '
                     'and never to the QoL one. Nothing builds them yet: tools/build_variant.py has no Balance variant.'),
            ('State', 'checked: every statement of the text was set against the executable or the data (column Note says '
                      'where). unverified: it states something nothing has read; not to be built until it is read. held: '
                      'withdrawn, the original line stays.'),
            ('Limits', f'A box is the whole string of the script, not the part that changes: at most {MAX_LINES} lines and '
                       f'{MAX_WIDTH} characters a line (the originals reach 75 in Lesson.pck and 69 in Teaching.pck).'),
            ('Voiced', 'The clip that plays with the box. A voiced box whose text no longer says what the Japanese line '
                       'says is a choice to take before building it.'),
            ('Source', 'scratch/update_euthys_and_teaching.py as of 2026-10-11 14:03, reviewed in the audit of that day '
                       '(scratch/audit_2026-10-11/AUDIT_REPORT.md, section 5).')):
        about.append(line)
    limits = wb.create_sheet('Settings')
    for line in (('Key', 'Value', 'Note'),
                 ('max_lines', MAX_LINES, 'Lines of a box of a lesson or of the library'),
                 ('max_width', MAX_WIDTH, 'Characters of a line: the originals reach 75 in Lesson.pck and 69 in Teaching.pck')):
        limits.append(line)
    ws = wb.create_sheet('Tutorials')
    ws.append(('Pack', 'Entry', 'Original box', 'Text', 'Voiced', 'State', 'Note'))
    for row in rows:
        ws.append(row)
    wb.save(OUTPUT)
    count = {}
    for row in rows:
        count[row[5]] = count.get(row[5], 0) + 1
    print(len(rows), 'boxes:', count, '| voiced:', sum(1 for r in rows if r[4]))


if __name__ == '__main__':
    main()
