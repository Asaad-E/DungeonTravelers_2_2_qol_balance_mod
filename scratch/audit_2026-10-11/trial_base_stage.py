"""A trial for the base stage, written into balance_changes.xlsx as decided rows so that every tool measures it:
the TP cost of Danger Shoot at each level and the ATK of the Dagger. Undo with git checkout of the workbook.

    python trial_base_stage.py "15,18,21,25,29,33,37,41,45,50" 20
"""
import os
import sys
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, PROJECT)
sys.path.insert(0, os.path.join(PROJECT, 'tools'))
from sim.data import game_data                                # noqa: E402

CHANGES = os.path.join(PROJECT, 'data_tables', 'balance_changes.xlsx')
WHY_TP = ('Decided 2026-10-11 (the Spieler at the base stage): Danger Shoot is her one attack and its first level cost '
          '27 TP, more than the 20 - 25 TP she has at levels 3 - 5, so she could not use it once (the level a skill is '
          'used at is the level it was raised to)')
WHY_ATK = ('Decided 2026-10-11 (the Scout with a dagger at the base stage, 9 - 13% of the damage of the party): the '
           'Dagger holds 14 ATK where the Shortsword holds 20 and the Short Bow 27; at levels 5 - 9 the forge adds little')


def main():
    costs = [int(x) for x in sys.argv[1].split(',')]
    attack = int(sys.argv[2])
    data = game_data()
    (category, number), skill = next((k, s) for k, s in data.skills.items() if s.name == 'Danger Shoot')
    record = f'{category} {number} Danger Shoot'
    wb = openpyxl.load_workbook(CHANGES)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    col = {name: head.index(name) for name in head}
    done = set()
    for row in ws.iter_rows(min_row=2):
        if row[col['Record']].value == record and row[col['Field']].value == 'TP cost' and row[col['State']].value != 'dropped':
            level = int(row[col['Level']].value)
            row[col['New']].value, row[col['State']].value = costs[level - 1], 'decided'
            row[col['Why']].value = str(row[col['Why']].value).split(' | ')[0] + ' | ' + WHY_TP
            done.add(level)
    for level in range(1, 11):
        if level not in done and costs[level - 1] != skill.tp[level - 1]:
            ws.append([{'ID': f'Y{10 + level:03d}', 'Table': 'skills', 'Record': record, 'Field': 'TP cost', 'Level': level,
                        'Vanilla': skill.tp[level - 1], 'New': costs[level - 1], 'Lever': 'L12',
                        'Serves': 'T05 (every class at its stage)', 'Why': WHY_TP, 'Decided': '2026-10-11',
                        'State': 'decided'}.get(name) for name in head])
    piece = data.equipment[53]
    if attack != piece['stats']['atk']:
        ws.append([{'ID': 'W046', 'Table': 'equipment', 'Record': f"53 {piece['name']}", 'Field': 'ATK (base)',
                    'Vanilla': piece['stats']['atk'], 'New': attack, 'Lever': 'L14', 'Serves': 'T05 (every class at its stage)',
                    'Why': WHY_ATK, 'Decided': '2026-10-11', 'State': 'decided'}.get(name) for name in head])
    wb.save(CHANGES)
    print('Danger Shoot TP', skill.tp, '->', costs, '| Dagger ATK', piece['stats']['atk'], '->', attack)


if __name__ == '__main__':
    main()
