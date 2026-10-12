"""Eccentrics (Lilian, Lizerietta, Souffle) changed on request of 2026-10-12: its DEX bonus added nothing that
could be measured. The size of a party skill bonus is the program's, by effect (0x4f9000 - 0x4f94e6), so what a row
can change is which effect the skill gives: it takes the one of Board Members (DEX and AGI, 20 each at every proof)
in the place of its own (DEX, 32). Row P091 of data_tables/balance_changes.xlsx, lever L42.

    python change_eccentrics.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TABLES = os.path.join(PROJECT, 'data_tables')
ROW = {'ID': 'P091', 'Table': 'party_skills', 'Record': 'PartySkillData_0393 (Eccentrics)', 'Field': 'effect id',
       'Vanilla': 371, 'New': 415, 'Lever': 'L42', 'Serves': 'R02',
       'Why': 'Decided 2026-10-12: change it. Measured (scratch/audit_2026-10-11/PARTY_SKILLS.md, the full bonus on the '
              'members the skill needs, four late bosses): DEX +32 on three members leaves the wins and the length of '
              'the fights where they were, while DEX +20 with AGI +20, the bonus of Board Members, makes them 13% shorter. '
              'Read in the program (0x4f9268 - 0x4f94b3): each party effect is asked once of a member ("has it?") and '
              'its bonus added to a sum by stat, so a member who holds the same effect from two party skills gets it '
              'once: Lizerietta and Souffle gain nothing over Board Members, Lilian gains what they have',
       'Decided': '2026-10-12', 'State': 'decided'}
LEVER = ('L42', 'Which effect a party skill gives its members', 'PartySkillData.pck', '0x4C, one int', 'Benign',
         'The id is one of the party effects the program knows (360 - 379, 415)',
         'tools/text_checks.py PARTY_EFFECT; the bonus of each effect is worked out at 0x4f9000 - 0x4f94e6. Not seen '
         'working in the game yet')


def main():
    wb = openpyxl.load_workbook(os.path.join(TABLES, 'balance_changes.xlsx'))
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    if ROW['ID'] not in {r[0].value for r in ws.iter_rows(min_row=2)}:
        ws.append([ROW.get(c) for c in head])
    wb.save(os.path.join(TABLES, 'balance_changes.xlsx'))
    wb = openpyxl.load_workbook(os.path.join(TABLES, 'balance_criteria.xlsx'))
    ws = wb['Levers']
    if LEVER[0] not in {r[0].value for r in ws.iter_rows(min_row=2)}:
        ws.append(LEVER)
    wb.save(os.path.join(TABLES, 'balance_criteria.xlsx'))


if __name__ == '__main__':
    main()
