"""The last decisions of 2026-10-12 written into the workbooks. Run once.

* The chance to hold on with 1 HP (effect 254: Guts of the Etoile, the unique skill of Lizerietta on the variant) is
  no longer rolled for a unit that is at 1 HP already (decided: no loop of never falling): row X15 of the sheet
  ExePatches, lever L46, seven bytes of the same length.
* Eccentrics gives the effect of Vagabonds (AGI) in the place of the one of Board Members: left to the model by the
  user. Two of its three members form Board Members with one more girl, and the same effect counts once.

    python register_night.py
"""
import os
import sys

import openpyxl

TABLES = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables')
BOX = ('When a hit would fell her, she may hold on with 1 HP, unless she is at 1 HP already.' + chr(10) +
       'Activation rate: 10% -> 12 -> 14 -> 16 -> 18 -> 20 -> 22 -> 24 -> 26%   |   Crown EX: 35%')


def main():
    path = os.path.join(TABLES, 'balance_criteria.xlsx')
    wb = openpyxl.load_workbook(path)
    wb['ExePatches'].append(['X15', 'balance', 'L46', '0x4238a2', '2b454085c00f8f', '48483b45400f83', None, None,
                             'Effect 254 (Guts; the unique skill of Lizerietta on the variant) is not rolled for a unit '
                             'at 1 HP: "HP - damage > 0, skip" becomes "HP - 2 >= damage as unsigned, skip" (dec eax twice, '
                             'cmp eax with the damage, jae to the same place). A hit that leaves exactly 1 HP now rolls '
                             'too, with the same result. Decided 2026-10-12: no loop of never falling'])
    wb['Levers'].append(['L46', 'The roll of effect 254 is skipped at 1 HP', 'executable, 0x4238a2', 'seven bytes',
                         'Executable', 'The seven bytes are the ones the row names; the jump keeps its place and its target',
                         'docs/FORMULAS.md 12.7; eax is loaded again at both ends of the jump (0x4238ad, 0x423a18). No '
                         'enemy has the effect. Not seen working in the game yet'])
    wb.save(path)

    path = os.path.join(TABLES, 'balance_changes.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb['Changes']
    col = {c.value: k for k, c in enumerate(ws[1])}
    for row in ws.iter_rows(min_row=2):
        if row[0].value == 'P091':
            row[col['New']].value = 366
            row[col['Why']].value = (str(row[col['Why']].value) + '. Changed the same night to the effect of Vagabonds '
                                     '(366, AGI +20), the choice left to the model: Lizerietta and Souffle form '
                                     'Board Members with any one of four other girls, and the same effect from two party '
                                     'skills counts once, so the first choice gave them nothing there; Vagabonds shares '
                                     'only Souffle and needs two girls more. Measured worth of AGI +20: fights 11% '
                                     'shorter, against 13% for DEX and AGI')
            row[col['State']].value = 'decided'
    wb.save(path)

    path = os.path.join(TABLES, 'balance_skill_text.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb['Skills']
    col = {c.value: k for k, c in enumerate(ws[1])}
    for row in ws.iter_rows(min_row=2):
        if row[col['Filename']].value == 'UniqueSkillData_0370.bin':
            row[col['Balance_Desc']].value = BOX
    wb.save(path)
    print('written')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    main()
