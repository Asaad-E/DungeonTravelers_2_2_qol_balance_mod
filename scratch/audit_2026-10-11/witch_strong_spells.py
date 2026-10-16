"""Sheet Roles of sim_setup.xlsx: the Witch of the model raises her strong fire spells (Crimson Flare, Inferno) right after Fireball and
her first passives (2026-10-16): Crimson Flare right after Fireball, Inferno after Concentrate. With them late in her skill priority she kept them at level 1 and spammed Fireball (C30: Fireball 10,
every strong spell 1); at their top Crimson Flare has 1300 power against 520 of Fireball, about twice the damage per time."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'sim_setup.xlsx')
FIRST = ['Crimson Flare', 'Inferno']
wb = openpyxl.load_workbook(PATH)
ws = wb['Roles']
head = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(head)}
for r in range(2, ws.max_row + 1):
    if ws.cell(r, col['Role']).value == 'mage':
        cell = ws.cell(r, col['Skill priority'])
        skills = [s.strip() for s in cell.value.split(';')]
        for name in FIRST:
            skills.remove(name)
        # Crimson Flare (5 levels, 30 skill points) right after Fireball; Inferno (55 points) after her first passives: at level 44
        # the Witch holds 121 points, and Fireball and Magic Mastery alone took 110 of them
        skills.insert(skills.index('Fireball') + 1, 'Crimson Flare')
        skills.insert(skills.index('Concentrate') + 1, 'Inferno')
        cell.value = '; '.join(skills)
        print(cell.value)
wb.save(PATH)
