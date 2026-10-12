"""The rows of the program patch become the set of the Balance variant, and the crown rule as the patch gives it gets
its own row in CrownOptions. Writes data_tables/balance_enchant_text.xlsx: the one enchantment name the patch makes
false.

    python exe_set_balance.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CRITERIA = os.path.join(PROJECT, 'data_tables', 'balance_criteria.xlsx')
ENCHANTS = os.path.join(PROJECT, 'data_tables', 'balance_enchant_text.xlsx')
KEPT = 'One more point, lowest step kept'


def main():
    wb = openpyxl.load_workbook(CRITERIA)
    ws = wb['ExePatches']
    head = [c.value for c in ws[1]]
    renamed = 0
    for row in ws.iter_rows(min_row=2):
        if row[head.index('Set')].value == 'test1':
            row[head.index('Set')].value = 'balance'
            renamed += 1
        first = row[0].value
        if isinstance(first, str) and first.startswith('The step of 59 to 68 levels'):
            row[0].value = ('The step of 59 to 68 levels stays at 1 point (accepted by the user, 2026-10-05): its instruction '
                            'is an "inc", with no byte that holds the amount. Giving it 2 means rewriting the end of the '
                            'function. With the other four steps one point higher the cheapest crown no longer uses that '
                            'step (docs/generated/CATCH_UP_OPTIONS.md, rule "' + KEPT + '").')
        if isinstance(first, str) and first.startswith('Factors with the piece'):
            row[0].value = ('Factors with the piece of enchantment 270 (X06 to X08: 1.5, 4, 6) and the last band at 10 with '
                            'or without the piece: accepted by the user on 2026-10-05, after the test in the game. The name '
                            'of the enchantment in the Balance variant states them (balance_enchant_text.xlsx).')
        if isinstance(first, str) and first.startswith('VA: address'):
            row[0].value = ('VA: address of the first byte changed. Old_Hex must be found there before anything is written, '
                            'in both program files of the game folder (DT22_en.exe and DT22_en_unwrapped.exe): the set is '
                            'written into both or into neither (tools/exe_patch.py).')
    crown = wb['CrownOptions']
    if not any(r[0].value == KEPT for r in crown.iter_rows(min_row=2)):
        crown.insert_rows(4)
        for k, value in enumerate((KEPT, '98:6, 89:5, 79:4, 69:3, 59:1', 50,
                                   'What the set balance of ExePatches writes: four bytes; the step of 59 levels is an '
                                   '"inc" and keeps its point')):
            crown.cell(row=4, column=k + 1, value=value)
    wb.save(CRITERIA)
    print(renamed, 'rows renamed to the set balance')

    out = openpyxl.Workbook()
    about = out.active
    about.title = 'About'
    for line in (('Topic', 'Detail'),
                 ('What', 'Enchantment names of the Balance variant worded by hand: the names of the QoL variant '
                          '(qol_enchant_text.xlsx) whose figure a change of the Balance variant makes false.'),
                 ('Columns', 'Balance_Name takes the place of QoL_Name in the build of the Balance variant. Depends_On: '
                             'the rows of ExePatches (balance_criteria.xlsx) or of balance_changes.xlsx the name states.')):
        about.append(line)
    ws = out.create_sheet('Enchants')
    ws.append(('Enchant_ID', 'Filename', 'QoL_Name', 'Balance_Name', 'Depends_On', 'Note'))
    ws.append((270, 'EnchantData_0270.bin', 'EXP from higher-level enemies +15%',
               'EXP from enemies 1-19 levels higher +20-33%', 'X01 - X08',
               'Factor with the piece over the factor without it: 1.5 / 1.25 one to four levels below the enemy, 4 / 3 '
               'five to nine below, 6 / 5 ten to nineteen below; 10 / 10 from twenty below, where the piece adds nothing'))
    out.save(ENCHANTS)
    print('written', ENCHANTS)


if __name__ == '__main__':
    main()
