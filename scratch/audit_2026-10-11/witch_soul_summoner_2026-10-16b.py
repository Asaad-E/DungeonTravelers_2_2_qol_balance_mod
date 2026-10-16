"""Decided 2026-10-16 (second pass, with the band of the classes): the Witch INT x0.8 at every level (rows V062 - V160 of growth
row 18), and the Soul Summoner (growth row 34) AGI x1.1, TP x1.2, HP x1.1 at every level, in place of help to her doll."""
import os
import sys
import openpyxl

HERE = os.path.dirname(__file__)
ROOT = os.path.join(HERE, '..', '..')
sys.path.insert(0, os.path.join(ROOT, 'tools'))
import class_data                                   # noqa: E402

PATH = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
DECIDED = '2026-10-16'
WITCH_WHY = ('The Witch into the band of the other classes (S07, decided 2026-10-16): in vanilla she was in 97% of the fastest '
             'final-class parties (the next class 65%). Her INT x0.8 (growth row 18) and the basic spells at levels 6 - 10 x0.85; '
             'measured by the class-balance measure of the suite (docs/generated/CLASS_BALANCE.md).')
SOUL_WHY = ('The Soul Summoner helped through her own stats rather than her doll (S07, decided 2026-10-16): AGI x1.1, TP x1.2, HP x1.1 '
            'at every level (growth row 34); her passives scaled up barely moved her. Measured by the class-balance measure of the '
            'suite (docs/generated/CLASS_BALANCE.md).')
SOUL = {'agi': 1.1, 'tp': 1.2, 'hp': 1.1}

table = class_data.growth()
wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
head = [c.value for c in ws[1]]
assert head[:12] == ['ID', 'Table', 'Record', 'Field', 'Level', 'Vanilla', 'New', 'Lever', 'Serves', 'Why', 'Decided', 'State'], head
done = 0
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 2).value == 'class_growth' and ws.cell(r, 3).value == '18 Witch' and ws.cell(r, 4).value == 'int':
        ws.cell(r, 7).value = round(ws.cell(r, 6).value * 0.8)
        ws.cell(r, 10).value = WITCH_WHY
        done += 1
assert done == class_data.LEVELS, done
ids = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
assert not any(ws.cell(r, 3).value == '34 Soul Summoner' for r in range(2, ws.max_row + 1))
first = 1 + max(int(k[1:]) for k in ids if k and str(k).startswith('V'))
row, k = ws.max_row + 1, 0
for stat, f in SOUL.items():
    for level in range(1, class_data.LEVELS + 1):
        old = table[34][level - 1][stat]
        values = ('V%03d' % (first + k), 'class_growth', '34 Soul Summoner', stat, level, old, round(old * f), 'L10', 'S07', SOUL_WHY,
                  DECIDED, 'built')
        for c, value in enumerate(values, start=1):
            ws.cell(row, c).value = value
        row += 1
        k += 1
wb.save(PATH)
print('Witch rows', done, '; Soul Summoner V%03d - V%03d' % (first, first + k - 1))
