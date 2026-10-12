"""Two pieces of preparation the Japanese pages state outright, written into data_tables/sim_setup.xlsx:

* the three superbosses are fought with every member crowned (page 231: Lolo and Tenerezza are hard "even with every
  member crowned"; page 333: players report Marduk, a lesser fight, as not winnable without crowns): column Crown
  bonus of the sheet Checkpoints, 50 for P49, P50 and P51;
* the fifth Aradia cannot be played without instant death immunity or high earth resistance (page 207): a row of the
  sheet Preparation, earth 100 for P16 (the checkpoints from P17 on already hold it).

    python postgame_crowns.py
"""
import os
import sys
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, PROJECT)
from sim import scenario, setup as setup_mod                 # noqa: E402

CROWNED = ('P49', 'P50', 'P51')
CROWN_SOURCE = (' | Crown bonus 50 since 2026-10-12: atwiki dt2-2 page 231 (the superboss of the casino is hard even '
                'with every member crowned) and page 333 (Marduk reported as not winnable without crowns)')
EARTH = ('P16', 6, 'earth', 100, 'atwiki dt2-2 page 207: Murder License kills one member before each action of Aradia; '
         'instant death immunity (two accessories, two sealbooks) or high earth resistance is the premise of the fight')


def main():
    setup = setup_mod.load()
    wb = openpyxl.load_workbook(setup_mod.SETUP)
    ws = wb['Checkpoints']
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value in CROWNED and int(row[head.index('Crown bonus')].value or 0) != 50:
            row[head.index('Crown bonus')].value = 50
            row[head.index('Source')].value = str(row[head.index('Source')].value) + CROWN_SOURCE
    ws = wb['Preparation']
    cid, index, name, value, source = EARTH
    cp = setup.checkpoint(cid)
    have = {(r[0].value, str(r[1].value), str(r[3].value)) for r in ws.iter_rows(min_row=2) if r[0].value}
    for number in scenario.battle_numbers(cp):
        if (cp['Floor'], str(number), str(index)) not in have:
            ws.append((cp['Floor'], number, f"{cid} {cp.get('Boss (as the story page names it)') or ''}".strip(), index,
                       name, value, source))
    wb.save(setup_mod.SETUP)


if __name__ == '__main__':
    main()
