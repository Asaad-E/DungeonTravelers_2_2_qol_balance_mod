"""Policies: the maid role plays by the rows of the diva role (decided 2026-10-16). Her path ends in the Diva; with her own rows (heal at half
HP, TP before the songs, no speed song) the Diva of that role won 28 - 30% at C40 in class-stages against 71% for the diva role, with
the same skill plan, shield and character. A row whose skill the classes so far do not have is skipped, so at the start she is still a
Maid that revives, heals and gives TP."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'sim_setup.xlsx')
NOTE = ' | 2026-10-16: the maid role plays by the rows of the diva role (her path ends in the Diva; class-stages C40 28% with her own rows)'

wb = openpyxl.load_workbook(PATH)
ws = wb['Policies']
head = [c.value for c in ws[1]]
assert head[:5] == ['Role', 'Order', 'Condition', 'Action', 'Source'], head
rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
diva = [r for r in rows if r[0] == 'diva']
first = next(i for i, r in enumerate(rows) if r[0] == 'maid')
kept = [r for r in rows if r[0] != 'maid']
new = [['maid'] + r[1:4] + [str(r[4]) + NOTE] + r[5:] for r in diva]
at = sum(1 for r in rows[:first] if r[0] != 'maid')
out = kept[:at] + new + kept[at:]
width = len(head)
for i in range(len(rows) + 1, len(out) + 1 + 1):
    pass
for k, r in enumerate(out, start=2):
    for c in range(1, width + 1):
        ws.cell(k, c).value = r[c - 1] if c - 1 < len(r) else None
for k in range(len(out) + 2, ws.max_row + 1):
    for c in range(1, width + 1):
        ws.cell(k, c).value = None
wb.save(PATH)
print('maid rows', len(new), 'total rows', len(out))
