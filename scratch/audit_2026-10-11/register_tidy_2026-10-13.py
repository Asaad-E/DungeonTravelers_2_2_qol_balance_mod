"""One-off edit (2026-10-13) of balance_criteria.xlsx and balance_changes.xlsx: rows of the register closed with no
figure get the figure read from the data, texts that went stale are brought up to date, wording made neutral, the
levers seen working are marked so, and the rows of the two dances say what the corrected model measures.
Run once; kept as the record of the edit.
"""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CRITERIA = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
CHANGES = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')

# sheet Criticisms: (row id, column): new text
SET = {
    ('J15', 'Vanilla_Result'): 'An attack that removes buffs (effect 74: after the hit lands, a roll under its value) is carried by an '
                               'enemy of 8 of the 40 boss fights of the story (C22, C26, C29, C31, C33, C34, C39, C40) and of 40 of '
                               'the 55 of the postgame (the attacks of each enemy, data of the game, counted on 2026-10-13)',
    ('J15', 'Why'): 'True of the postgame, not of the story. The removal rides on a hit that also deals its damage, so it costs the '
                    'boss nothing; what keeps a buff worth its turn is its length against how often that attack is picked, and the '
                    'buffs measured still pay (docs/generated/SKILL_VALUE.md). Left as the game has it: it is what makes the '
                    'postgame bosses different from the story ones',
    ('J15', 'Decision'): '2026-10-13: no change',
    ('F96', 'Vanilla_Result'): 'The enemy with most HP of a boss fight has a resistance below 0 in 26 of the 40 fights of the story '
                               '(down to -75 at C17, -50 at C08, C20 and C22) and in 4 of the 55 of the postgame (P36 -50, P07 and P22 '
                               '-30, P39 -25): data of the game, counted on 2026-10-13',
    ('F96', 'Why'): 'A weakness is common early and rare where the claim places it: 4 of the 55 postgame fights. Each point below 0 '
                    'adds half a percent of damage (docs/FORMULAS.md 2.4), so -50 is a quarter more, not a trivial fight',
    ('F96', 'Decision'): '2026-10-13: no change',
    ('F65', 'Verdict'): 'partly true: songs compete for two slots with Carol and March; measured +0% of length and of wins at the '
                        'bosses of the reference party (docs/generated/SKILL_VALUE.md), where little damage is magic',
    ('S25', 'Why'): 'Content and maps are out of the balance. The conditions of side events, which are only text, were looked at as '
                    'an addition to the QoL variant and dropped on 2026-10-12',
    ('J09', 'Decision'): '2026-10-09: no change (the early misses are kept)',
    ('F85', 'Vanilla_Result'): 'Joint skills ask for specific members in the party and for TP and gauge; parked (decided 2026-10-04)',
    ('S23', 'Why'): 'An idea for later (2026-10-07): complex and interesting. Not evaluated and nothing is opened for it. The QoL '
                    'variant already shows the speed and chant of every skill',
    ('S26', 'Why'): 'Not for now: very specific, and almost nobody reports it (2026-10-07). Reports from the console versions and '
                    'early patches; none reproduced here',
}
SEEN = ('Not seen working in the game yet', 'Seen working in the game (play-tested, 2026-10-12)')
DANCES = ('Fascinating Waltz', 'Dumbing Mazurka')
DANCE_WHY = (' | 2026-10-13 : measured again with the dance bonus taken away as the program does (docs/FORMULAS.md 12.8; '
             'measures/dance_trial.py, 150 battles). Kept: a strong party gains at most 8 points of wins over the values of the '
             'game (C36 77% > 85%, C38 76% > 83%), and they are what holds the Juggernaut party up (C30 16% > 63%, C38 10% > 36%)')


def main():
    wb = openpyxl.load_workbook(CRITERIA)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    done = 0
    for row in ws.iter_rows(min_row=2):
        for (rid, column), text in SET.items():
            if row[0].value == rid:
                row[head.index(column)].value = text
                done += 1
    assert done == len(SET), done
    marked = 0
    for name in ('Levers', 'About'):
        for row in wb[name].iter_rows():
            for cell in row:
                value = str(cell.value or '')
                if SEEN[0] in value:
                    cell.value = value.replace(SEEN[0], SEEN[1])
                    marked += 1
                elif value.startswith('Decision of decided:'):
                    cell.value = value.replace('Decision of decided:', 'Decided:')
                elif value.startswith('Definition of the user (2026-10-05):'):
                    cell.value = value.replace('Definition of the user (2026-10-05):', 'Defined on 2026-10-05:')
    assert marked == 4, marked
    wb.save(CRITERIA)

    wb = openpyxl.load_workbook(CHANGES)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    noted = 0
    for row in ws.iter_rows(min_row=2):
        record, field = str(row[head.index('Record')].value), str(row[head.index('Field')].value)
        if row[head.index('State')].value == 'built' and any(d in record for d in DANCES) and field in ('effect 195', 'effect 201'):
            row[head.index('Why')].value = str(row[head.index('Why')].value) + DANCE_WHY
            noted += 1
    assert noted == 20, noted
    wb.save(CHANGES)
    print(done, 'cells of the register,', marked, 'levers,', noted, 'rows of the dances')


if __name__ == '__main__':
    main()
