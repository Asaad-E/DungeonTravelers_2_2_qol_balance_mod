"""One-off edit: the usage readings of the Etoile after her second dance became Debilitating Salsa
(etoile_salsa_2026-10-15.py): Dumbing Mazurka is now the dance left out, and every 'slots' reading names the new pair."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'sim_setup.xlsx')
SLOTS = 'Her two dance slots hold Fascinating Waltz and Debilitating Salsa (rules status:195, 199).'

wb = openpyxl.load_workbook(PATH)
ws = wb['AuditReadings']
for r in range(2, ws.max_row + 1):
    key = ws.cell(r, 1).value
    if key == 'usage etoile | Debilitating Salsa':
        ws.cell(r, 1).value = 'usage etoile | Dumbing Mazurka'
        ws.cell(r, 3).value = SLOTS + (' Waltz and Salsa win as often as Waltz and Mazurka and end the fights 5 - 33% '
                                       'sooner at C22 - C40 and P24 (measures/etoile_dances.py, 2026-10-15).')
    elif str(key).startswith('usage etoile |') and ws.cell(r, 2).value == 'slots':
        ws.cell(r, 3).value = SLOTS
wb.save(PATH)
print('done')
