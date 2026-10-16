"""One-off edit: rows F11 and F63 of the sheet Criticisms follow Barrage at 30% - 50% (barrage_30_50_2026-10-16.py)."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
TEXT = {
    'F11': ('W rows: Double Attack and Triple Attack +25%, built and read back (BALANCE_BUILD.md); Barrage 30% at level 1 to 50% at '
            'level 10 since 2026-10-16 (rows W026 - W035)'),
    'F63': ('Barrage acts on plain attacks only: the Kunoichi of the model makes 0.7 a fight at C38 (3% of her actions). At 25% - 40% '
            'the wins at C30 - C40 did not move with it or without it (300 battles; at a 100% chance C38 stays at 85%); at 30% - 50% '
            '(2026-10-16) the full suite moves by chance only, the Kunoichi inside the band of CLASS_STAGES.md'),
}

wb = openpyxl.load_workbook(PATH)
ws = wb['Criticisms']
col = [c.value for c in ws[1]].index('Balance_Result') + 1
done = 0
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 1).value in TEXT:
        ws.cell(r, col).value = TEXT[ws.cell(r, 1).value]
        done += 1
assert done == 2
wb.save(PATH)
print('rows', done)
