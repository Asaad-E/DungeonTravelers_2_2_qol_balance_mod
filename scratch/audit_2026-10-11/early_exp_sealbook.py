"""The effect "more EXP from higher-level enemies" (enchantment 270) earlier in the story, asked on
2026-10-12: in vanilla only the Grand Sealbook of Dragon Papi holds it (level 44). The Grand Sealbook of Molle, a boss
of the story at level 23 (chapter 5), gains it as a second enchantment: one book per boss, never dropped nor stolen.
Row P089 of data_tables/balance_changes.xlsx.

    python early_exp_sealbook.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CHANGES = os.path.join(PROJECT, 'data_tables', 'balance_changes.xlsx')
ROW = {'ID': 'P089', 'Table': 'sealbooks', 'Record': 'Grand Sealbook [Molle] (0829)', 'Field': 'secondary enchantment',
       'Vanilla': 'ench_cnt=1 (262 only)',
       'New': 'ench_cnt=2 (262 Escape probability UP, 270 Gain more EXP from higher-level enemies)',
       'Lever': 'L19', 'Serves': 'U01',
       'Why': 'Decided 2026-10-12: the piece that raises the EXP factors of B4 comes late (the Grand Sealbook of Dragon '
              'Papi, level 44); he wants the same effect around level 20 - 30, on one unique thing that is neither farmed '
              'nor stolen, without filling the game with such pieces. Molle is a boss of the story at level 23 and its '
              'grand sealbook, made once from it, held only "Escape probability UP". The book of Dragon Papi keeps its own',
       'Decided': '2026-10-12', 'State': 'decided'}


def main():
    wb = openpyxl.load_workbook(CHANGES)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    assert ROW['ID'] not in {r[0].value for r in ws.iter_rows(min_row=2)}
    ws.append([ROW.get(name) for name in head])
    wb.save(CHANGES)


if __name__ == '__main__':
    main()
