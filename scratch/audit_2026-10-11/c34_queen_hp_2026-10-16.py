"""C34 Queen of Darkness: HP x1.30 -> x1.15 (rows G114, G115), decided 2026-10-16 after dtcli class-search over every final-class party."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'balance_changes.xlsx')

NEW = {'G114': 9200, 'G115': 11500}
WHY = ('Harder than vanilla by a little, to make up for what the variant eased (R10): HP x1.15 (decided 2026-10-16). x1.30 of 2026-10-11 '
       'made a cliff for parties without a Maid-line member: dtcli class-search final C34 boss, every final-class party (14,835), 20 runs and '
       'the candidates at 200: x1.30 left 6.7% of the parties 10 or more points under vanilla (60 drops confirmed, the worst 74 points; parties '
       'without Diva, Mistress or Etoile won 38%); x1.15 leaves 1.1% (21 confirmed, the worst 27 points; those parties 56%), mean wins 82% '
       '(vanilla 65%), median length 18.1, no party shorter than the shortest of vanilla.')

wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
head = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(head)}
done = set()
for row in range(2, ws.max_row + 1):
    rid = ws.cell(row, col['ID']).value
    if rid in NEW:
        ws.cell(row, col['New']).value = NEW[rid]
        ws.cell(row, col['Why']).value = WHY
        ws.cell(row, col['Decided']).value = '2026-10-16'
        done.add(rid)
assert done == set(NEW), done
wb.save(PATH)
print('written', sorted(done))
