"""Restores, in balance_changes.xlsx, the immunity bits commit d25c8ed closed together with sleep: poison, silence and
slow go back to the values decided on 2026-10-09 (read from the workbook of commit f8cbd16); sleep stays as that
commit left it, closed as in vanilla. The Why of every boss row then says what the row does.

    python restore_boss_flags.py            shows what it would change
    python restore_boss_flags.py --write    writes the workbook
"""
import os
import sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
WORKBOOK = os.path.join(os.path.dirname(os.path.dirname(HERE)), 'data_tables', 'balance_changes.xlsx')
BEFORE = os.path.join(HERE, 'balance_changes_at_f8cbd16.xlsx')
SLEEP = 1 << 3
OLD_NOTE = ', and to sleep in the story'
NEW_NOTE = '; sleep as in vanilla (the opening of 2026-10-09 was taken back on 2026-10-11)'


def rows(path, read_only):
    wb = openpyxl.load_workbook(path, read_only=read_only, data_only=read_only)
    ws = wb['Changes']
    head = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    return wb, ws, {name: k for k, name in enumerate(head)}


def main(write):
    _, ws_old, col = rows(BEFORE, True)
    before = {r[col['ID']]: r[col['New']] for r in ws_old.iter_rows(min_row=2, values_only=True) if r[col['ID']]}
    wb, ws, col = rows(WORKBOOK, False)
    flags = notes = 0
    for row in ws.iter_rows(min_row=2):
        rid, field = row[col['ID']].value, str(row[col['Field']].value)
        if not rid or not field.startswith('immunity flags') or row[col['State']].value == 'dropped':
            continue
        now, old = int(str(row[col['New']].value), 16), int(str(before[rid]), 16)
        wanted = (old & ~SLEEP) | (now & SLEEP)          # the bits of 2026-10-09, sleep as it stands today
        if wanted != now:
            flags += 1
            print(f"{rid} {row[col['Record']].value}: {hex(now)} -> {hex(wanted)}")
            row[col['New']].value = hex(wanted)
        why = str(row[col['Why']].value or '')
        if OLD_NOTE in why:
            notes += 1
            row[col['Why']].value = why.replace(OLD_NOTE, NEW_NOTE)
    print(f'{flags} flag values restored, {notes} notes brought up to date')
    if write:
        wb.save(WORKBOOK)
        print('written:', WORKBOOK)


if __name__ == '__main__':
    main('--write' in sys.argv)
