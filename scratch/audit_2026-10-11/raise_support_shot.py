"""Support Shot (row F52 of the register), raised at the user's word on 2026-10-11: rows W047 - W051 of
data_tables/balance_changes.xlsx.

    python raise_support_shot.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CHANGES = os.path.join(PROJECT, 'data_tables', 'balance_changes.xlsx')
VANILLA = (0.04, 0.06, 0.08, 0.1, 0.12)
NEW = (0.1, 0.15, 0.2, 0.25, 0.3)
WHY = ('Decided 2026-10-11 (the user: raise it to a reasonable level): Support Shot fires only after a plain attack of '
       'another member, and at 4 - 12% the reference party with and without it was the same inside the noise (C22, C30, '
       'C34, C40). At 10 - 30% a Sniper beside members who strike plainly adds a shot to about one of their attacks in '
       'three; beside a party that spends TP every turn it stays what it was')


def main():
    wb = openpyxl.load_workbook(CHANGES)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    have = {r[0].value for r in ws.iter_rows(min_row=2)}
    for level, (old, new) in enumerate(zip(VANILLA, NEW), 1):
        rid = f'W{46 + level:03d}'
        assert rid not in have, rid
        ws.append([{'ID': rid, 'Table': 'skills', 'Record': 'Passive 271 Support Shot', 'Field': 'effect 260',
                    'Level': level, 'Vanilla': old, 'New': new, 'Lever': 'L12', 'Serves': 'F52', 'Why': WHY,
                    'Decided': '2026-10-11', 'State': 'decided'}.get(name) for name in head])
    wb.save(CHANGES)


if __name__ == '__main__':
    main()
