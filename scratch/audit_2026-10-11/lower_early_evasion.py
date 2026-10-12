"""The early ranks of the unique skill of Lizerietta lowered, the high ones kept (the user, 2026-10-12: it is the one
unique skill that is always at work, the others are situational): 4 - 26% by rank, 35% with the crown. Run once."""
import os
import openpyxl

TABLES = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables')
VALUES = [0.04, 0.06, 0.08, 0.11, 0.14, 0.17, 0.20, 0.23, 0.26, 0.35]
BOX = ('Enemy attacks and spells are less likely to hit her.' + chr(10) +
       'Evasion: 4% -> 6 -> 8 -> 11 -> 14 -> 17 -> 20 -> 23 -> 26%   |   Crown EX: 35%')
NOTE = ('. Lowered on its early ranks the same day (the user: the one unique skill always at work): 4 - 26%. Measured '
        '(scratch/audit_2026-10-11/curve_ablation.py, 200 battles a checkpoint): with these values, with 10 - 26%, with '
        'none and with the skill of the game the eight checkpoints give the same figures within their noise')

path = os.path.join(TABLES, 'balance_changes.xlsx')
wb = openpyxl.load_workbook(path)
ws = wb['Changes']
col = {c.value: k for k, c in enumerate(ws[1])}
for row in ws.iter_rows(min_row=2):
    if str(row[col['Record']].value) == 'Unique 370 Usurper of Souls' and row[col['Field']].value == 'value':
        row[col['New']].value = VALUES[int(row[col['Level']].value) - 1]
        row[col['Why']].value = str(row[col['Why']].value) + NOTE
        row[col['State']].value = 'decided'
wb.save(path)
path = os.path.join(TABLES, 'balance_skill_text.xlsx')
wb = openpyxl.load_workbook(path)
ws = wb['Skills']
col = {c.value: k for k, c in enumerate(ws[1])}
for row in ws.iter_rows(min_row=2):
    if row[col['Filename']].value == 'UniqueSkillData_0370.bin':
        row[col['Balance_Desc']].value = BOX
wb.save(path)
print('written')
