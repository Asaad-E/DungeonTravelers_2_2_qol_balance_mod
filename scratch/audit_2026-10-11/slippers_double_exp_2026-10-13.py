"""One-off edit (2026-10-13): the Bathroom Slippers carry Double EXP (275) in the place of the factor enchantment
(270): row P090 of balance_changes.xlsx and row F92 of the register. Run once; kept as the record of the edit.
"""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
path = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
wb = openpyxl.load_workbook(path)
ws = wb['Changes']
head = [c.value for c in ws[1]]
row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == 'P090')
assert row[head.index('New')].value == 270
row[head.index('New')].value = 275
row[head.index('Serves')].value = 'U01, F92'
row[head.index('Why')].value = ('One piece that speeds EXP up, early and for any class: the Bathroom Slippers, the reward of quest '
                                'rank 3, given once, neither farmed nor stolen. They keep "Nullify poison" and give up "Physical '
                                'defense UP" (enchantment 33) for Double EXP (275), which the game already puts on Lord '
                                "Mashiro's Fan (a drop of Kiyohime, level 38, for the classes of the fan) and Therion's Crown, "
                                'and reads at 0x445a1e; the count of innate enchantments stays as it is (lever L15). Decided '
                                '2026-10-12 with the factor enchantment (270), which only the Grand Sealbook of Dragon Papi has '
                                '(level 44); changed to Double EXP on 2026-10-13: the factors by level gap already catch '
                                'everyone up without a piece')
wb.save(path)
path = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
wb = openpyxl.load_workbook(path)
ws = wb['Criticisms']
head = [c.value for c in ws[1]]
row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == 'F92')
row[head.index('Why')].value = ('The DLC is ignored (2026-10-06). What it eased is covered by T01: the factors by level gap '
                                '(option B4, without the kept curve) catch a member up, and a crown costs about 400 battles '
                                'instead of 4,393 (docs/generated/CATCH_UP_OPTIONS.md). One piece gains Double EXP, as the DLC '
                                'gave: the Bathroom Slippers, the reward of quest rank 3 (row P090 of balance_changes.xlsx, '
                                '2026-10-13). The Slave\'s Torc carried it from 2026-10-11 to 2026-10-13 (row W042, dropped)')
wb.save(path)
print('P090 -> 275')
