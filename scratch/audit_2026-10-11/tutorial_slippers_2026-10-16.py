"""One-off edit: the two tutorial boxes of the Balance variant that named the Slave's Torc (row W042, dropped) as the piece
that doubles EXP name the Bathroom Slippers (row P090), which do; and the sheet About no longer says nothing builds them."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'balance_tutorial_text.xlsx')
TEXTS = {
    'BattleTeacher_11-12.adv.bin': (
        '"Simply equip her with the Bathroom Slippers to double her gains.\n'
        'She will soak up so much experience from stronger monsters that\n'
        'she\'ll rejoin the vanguard before you even have time to complain."',
        'True on the Balance variant only (row P090)'),
    'CheatsTeaching_19.adv.bin': (
        'Equip the Bathroom Slippers to double experience gains,\n'
        'rapidly catching up allies who underwent resets.\n'
        'Facing higher-level monsters also grants native bonus EXP.',
        'True on the Balance variant only: row P090 gives the Bathroom Slippers Double EXP; the gap factors are 0x450739'),
}

wb = openpyxl.load_workbook(PATH)
ws = wb['Tutorials']
header = [c.value for c in ws[1]]
entry, text, note = header.index('Entry') + 1, header.index('Text') + 1, header.index('Note') + 1
done = 0
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, entry).value
    if name in TEXTS and "Slave's Torc" in str(ws.cell(r, text).value):
        ws.cell(r, text).value, ws.cell(r, note).value = TEXTS[name]
        done += 1
assert done == 2, done
about = wb['About']
for row in about.iter_rows():
    for cell in row:
        if isinstance(cell.value, str) and 'Nothing builds them yet' in cell.value:
            cell.value = cell.value.replace(' Nothing builds them yet: tools/build_variant.py has no Balance variant.',
                                            ' tools/build_variant.py balance writes them into Lesson.pck and Teaching.pck.')
wb.save(PATH)
print('boxes', done)
