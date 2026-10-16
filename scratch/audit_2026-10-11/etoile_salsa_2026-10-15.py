"""One-off edit: the Etoile of the model keeps Fascinating Waltz and Debilitating Salsa up instead of Waltz and Dumbing
Mazurka (sheet Policies of sim_setup.xlsx). Measured by measures/etoile_dances.py at C22 - C40 and P24, 200 battles a
pair: the same wins, fights 5 - 33% shorter."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'sim_setup.xlsx')

wb = openpyxl.load_workbook(PATH)
ws = wb['Policies']
changed = 0
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 1).value == 'etoile' and ws.cell(r, 4).value == 'status:201':
        ws.cell(r, 4).value = 'status:199'
        ws.cell(r, 5).value = 'Debilitating Salsa on the enemies'
        changed += 1
assert changed == 1, changed
wb.save(PATH)
print('changed', changed)
