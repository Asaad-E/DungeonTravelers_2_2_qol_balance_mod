"""One-off edit: the readings of 34 Cleanse and 142 Heaven or Hell Effect in the sheet AuditReadings of sim_setup.xlsx,
from probe_two_actives_2026-10-15.py (beside this file)."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'sim_setup.xlsx')
ROWS = [
    ('34 Cleanse', 'context',
     'A cure aimed at one ally: the prepared battle aims it at a member with no ailment. Aimed at the member poisoned '
     'and asleep it removes both (probe_two_actives_2026-10-15.py, recheck 2026-10-15).'),
    ('142 Heaven or Hell Effect', 'context',
     'The heal of a random move, aimed at a foe: the prepared foes are at full HP. A foe at half HP regains 10% of '
     'its max HP, under the cap of 555 (probe_two_actives_2026-10-15.py, recheck 2026-10-15).'),
]

wb = openpyxl.load_workbook(PATH)
ws = wb['AuditReadings']
have = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
for row in ROWS:
    if row[0] not in have:
        ws.append(list(row))
wb.save(PATH)
print('rows', ws.max_row - 1)
