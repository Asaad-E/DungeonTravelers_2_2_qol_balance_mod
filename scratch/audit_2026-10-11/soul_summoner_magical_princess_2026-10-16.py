"""Soul Eater (the Soul Summoner's own attack) and the Magical Princess (Magical Change, Magical Spark), decided 2026-10-16.
Measured before building with dtcli class-search final C30, C34, C38, C40 boss, compare=variant, the parties that hold the class."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'balance_changes.xlsx')
DECIDED = '2026-10-16'
WHY_SE = ('The Soul Summoner herself dealt 1 - 2% of the damage of her party, her dolls the rest (C30 - C40); Soul Eater, her own '
          'attack, 20% cheaper and 15% stronger (decided 2026-10-16), and the model learns it (sheet Roles). class-search over the '
          '14,804 final-class parties with a Soul Summoner: wins 53.2% > 54.0%, no dominant party; it helps her, not the stacks of dolls.')
WHY_MP = ('The Magical Princess transformed, as the guides play her, dealt 19 - 35% of the damage of her party where a Witch deals 53 - '
          '57%, short of TP (Magical Change 69, Magical Spark up to 34 a cast). Magical Change two turns longer and 25% stronger, '
          'Magical Spark 20% cheaper (decided 2026-10-16); each alone added one point, together: class-search over the 15,289 '
          'final-class parties with her, wins 47.5% > 50.6%, the fights 9% shorter, no dominant party.')

NEW = []
for lv, (old, new) in enumerate(zip([32, 35, 38, 41, 44, 47, 50, 52, 54, 56], [26, 28, 30, 33, 35, 38, 40, 42, 43, 45]), start=1):
    NEW.append(('skills', 'Magic 233 Soul Eater', 'TP cost', lv, old, new, 'L12', 'J18, S08', WHY_SE))
for lv, old in enumerate([360, 400, 440, 480, 520, 560, 600, 640, 690, 740], start=1):
    NEW.append(('skills', 'Magic 233 Soul Eater', 'effect 1', lv, old, round(old * 1.15), 'L12', 'J18, S08', WHY_SE))
for lv, old in enumerate([4, 4, 5, 5, 6], start=1):
    NEW.append(('skills', 'Technique 59 Magical Change', 'turns', lv, old, old + 2, 'L12', 'J25, F05', WHY_MP))
for lv, (old, new) in enumerate(zip([0.2, 0.23, 0.26, 0.29, 0.32], [0.25, 0.29, 0.33, 0.36, 0.40]), start=1):
    NEW.append(('skills', 'Technique 59 Magical Change', 'effect 189', lv, old, new, 'L12', 'J25, F05', WHY_MP))
NEW.append(('skills', 'Technique 60 Magical Spark', 'TP cost', 1, 13, 10, 'L12', 'J25, F05', WHY_MP))
SPARK = {'C098': 11, 'C099': 13, 'C100': 14, 'C101': 15, 'C102': 18, 'C103': 20, 'C104': 22, 'C105': 25, 'C106': 27}

wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
head = [c.value for c in ws[1]]
assert head[:12] == ['ID', 'Table', 'Record', 'Field', 'Level', 'Vanilla', 'New', 'Lever', 'Serves', 'Why', 'Decided', 'State'], head
col = {name: i + 1 for i, name in enumerate(head)}
ids = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
for rid, new in SPARK.items():
    r = ids[rid]
    assert 'Magical Spark' in ws.cell(r, col['Record']).value and ws.cell(r, col['Field']).value == 'TP cost'
    ws.cell(r, col['New']).value = new
    ws.cell(r, col['Why']).value = (ws.cell(r, col['Why']).value or '').rstrip() + ' Then 20% cheaper again (2026-10-16): ' + WHY_MP
    ws.cell(r, col['Decided']).value = DECIDED
row = ws.max_row + 1
for k, values in enumerate(NEW, start=31):
    rid = 'V%03d' % k
    target = ids.get(rid)
    if target is None:
        target = row
        row += 1
    for c, value in enumerate((rid,) + values + (DECIDED, 'built'), start=1):
        ws.cell(target, c).value = value
wb.save(PATH)
print('written V031 - V%03d and' % (30 + len(NEW)), len(SPARK), 'rows of Magical Spark')
