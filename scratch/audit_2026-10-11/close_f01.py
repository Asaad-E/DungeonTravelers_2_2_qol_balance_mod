"""Row F01 of the register closed on the user's decision of 2026-10-12 (a fifth more TP from the fruit, rows
Y012 - Y014), with the figures measured on the build (no_maid_line.py and tp_items_trial.py beside this file).

    python close_f01.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
WORKBOOK = os.path.join(PROJECT, 'data_tables', 'balance_criteria.xlsx')
FIGURE = ('Y012 - Y014 (Magical Fruit 25 to 30 TP, Fairy Apple 50 to 60, Golden Apple 100 to 120), built and read back '
          '(1,401 of 1,401 rows). The Meta party with a Bishop in the place of its Diva, nobody of the Maid line in it '
          '(scratch/audit_2026-10-11/tp_items_trial.py, 60 battles): with the 8 fruit of the sheets 3% / 3% / 18% / 68% at '
          'C30 / C34 / C36 / C38; with 24 fruit in the bag 67% / 82% / 45% / 98% (47% / 52% / 38% / 97% before the raise). '
          'What stops such a party at C30 and C34 is TP (100% with TP at will); C36 does not turn on it, and the last boss '
          'is not won without the Maid line (0%; 45% with 99 fruit). Up to C28 the party wins 78 - 100% without her. '
          'Nothing else moved: curve 37 of 44, battery 75 of 85, classes 74 of 76, anchors 6 of 6')
VERDICT = ('true from C30 on in vanilla; on the variant a party without the Maid line wins C30 to C38 with a bag of '
           'fruit, and needs her at the last boss')
DECISION = ('; the user, 2026-10-12: a fifth more TP from the fruit and no new cut of TP costs; the Maid keeps her part '
            'and a party without her gets through with effort')


def main():
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    col = {name: head.index(name) for name in head}
    for row in ws.iter_rows(min_row=2):
        if row[0].value == 'F01':
            row[col['Status']].value = 'closed'
            row[col['Verdict']].value = VERDICT
            row[col['Balance_Result']].value = FIGURE
            row[col['Proposal']].value = 'Y012 - Y014: the three TP fruits of the shop x1.2'
            row[col['Decision']].value = str(row[col['Decision']].value) + DECISION
    wb.save(WORKBOOK)


if __name__ == '__main__':
    main()
