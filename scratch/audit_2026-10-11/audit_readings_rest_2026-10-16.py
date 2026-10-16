"""One-off edit: the readings of the last unread rows of the audit (sheet AuditReadings of sim_setup.xlsx), from the
lines of docs/generated/sim_audit/ and the rules of the sheet Policies, read on 2026-10-16."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'sim_setup.xlsx')
SAGE_BUFFS = ('Left out of her rules on 2026-10-13 (sheet Policies, rule status:161): each added nothing that could be '
              'measured at C30, C36, C38, C40 and P24 (150 battles, with and without each), and with all five kept up she '
              'never attacked. She keeps Circle Protect and Circle Brave, then casts her spells.')
ROWS = [
    ('276 Absolute Attack', 'context',
     'Effect 307 cancels the answers of the target to a blow (Avenger 249, Counter Whip 303, effect 250, the unique '
     'skill 447). Samael (C38) has none of them, so nothing changes there.'),
    ('302 Barrage', 'context',
     'Effect 251 adds two hits to a plain attack only. At C38 the skills of the Kunoichi always score above her plain '
     'attack, so she never makes one.'),
    ('usage soul summoner | Spartan', 'doll order',
     'Her rule fills both doll slots in the order Mercidarian, Spartan, Musiy, Catalina, Oracles: Mercidarian takes '
     'both slots.'),
    ('usage sage | Accelerator', 'role choice', SAGE_BUFFS),
    ('usage sage | Amplify', 'role choice', SAGE_BUFFS),
    ('usage sage | Sun Veil', 'role choice', SAGE_BUFFS),
    ('usage sage | Bible Thump', 'role choice', 'A tome skill: the Sage of the model holds no tome.'),
    ('usage kunoichi | Phoenix', 'role choice', 'A projectile skill: the Kunoichi of the model holds a Tantou.'),
]

wb = openpyxl.load_workbook(PATH)
ws = wb['AuditReadings']
have = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
for row in ROWS:
    if row[0] not in have:
        ws.append(list(row))
wb.save(PATH)
print('rows', ws.max_row - 1)
