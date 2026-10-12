"""Writes the pairs tools/boss_fit.py gentle kept into balance_changes.xlsx: one row per enemy and stat (lever L01).
An enemy stat a live row already changes keeps its row: only its new value and its reason move. Run after
"boss_fit.py gentle" and "gentle-report".

    python apply_gentle_fit.py            shows the rows
    python apply_gentle_fit.py --write    writes the workbook
"""
import os
import sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, PROJECT)
sys.path.insert(0, os.path.join(PROJECT, 'tools'))
import boss_fit                                           # noqa: E402
from sim import setup as setup_mod                        # noqa: E402
from sim.data import game_data                            # noqa: E402
from sim.variant import balance_data                      # noqa: E402

WORKBOOK = os.path.join(PROJECT, 'data_tables', 'balance_changes.xlsx')
TODAY = '2026-10-11'


def figures(f):
    return f"{f['won']:.0%} / {f['time']:.1f} / {f['lowest_hp']:.0%}" if f.get('time') else f"{f['won']:.0%}"


def main(write):
    data, setup = game_data(), setup_mod.load()
    variant = balance_data(data)
    found = boss_fit.gentle_choice(setup, data)
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    col = {name: k for k, name in enumerate(head)}
    live = {}
    numbers = []
    for row in ws.iter_rows(min_row=2):
        rid = row[col['ID']].value
        if not rid:
            continue
        if str(rid).startswith('G') and str(rid)[1:].isdigit():
            numbers.append(int(str(rid)[1:]))
        if row[col['Table']].value == 'enemies' and row[col['State']].value != 'dropped':
            live[(int(str(row[col['Record']].value).split()[0]), row[col['Field']].value)] = row
    next_id = max(numbers) + 1
    changed = added = 0
    for r in found:
        kept = r['kept']
        if not kept:
            print(f"{r['id']}: {r['why']}")
            continue
        base = r['tried'][0]
        factors = f"HP x{kept['hp']:.2f}" + (f", ATK and INT x{kept['attack']:.2f}" if kept['attack'] != 1.0 else '')
        why = (f"A little harder, to make up for what the variant eased (R10; the user, {TODAY}; tools/boss_fit.py gentle, "
               f"docs/generated/BOSS_FIT.md): {r['id']}, reference party, wins / length / lowest HP: vanilla "
               f"{figures(r['vanilla'])}, variant {figures(base['figures'])}, with {factors} {figures(kept['figures'])}. "
               f"Every archetype of the battery still wins at least the floor of the boss ({r['floor']:.0%}): "
               + ', '.join(f"{a} {v:.0%}" for a, v in kept['archetypes'].items()))
        for eid in r['enemies']:
            for field, factor in (('hp', kept['hp']), ('atk', kept['attack']), ('int', kept['attack'])):
                if factor == 1.0:
                    continue
                new = int(round(variant.enemies[eid][field] * factor))
                old_row = live.get((eid, field))
                if old_row is not None:
                    print(f"{r['id']} {old_row[col['ID']].value} {data.enemies[eid]['name']} {field}: "
                          f"{old_row[col['New']].value} -> {new} (row kept)")
                    old_row[col['New']].value = new
                    old_row[col['Why']].value = f"{old_row[col['Why']].value}. Then, {TODAY}: {why}"
                    changed += 1
                else:
                    rid = f'G{next_id:03d}'
                    next_id += 1
                    values = {'ID': rid, 'Table': 'enemies', 'Record': f"{eid} {data.enemies[eid]['name']} ({r['id']})",
                              'Field': field, 'Level': None, 'Vanilla': data.enemies[eid][field], 'New': new,
                              'Lever': 'L01', 'Serves': 'R08, R10', 'Why': why, 'Decided': TODAY, 'State': 'decided'}
                    ws.append([values.get(name) for name in head])
                    print(f"{r['id']} {rid} {data.enemies[eid]['name']} {field}: {values['Vanilla']} -> {new}")
                    added += 1
    print(f'{added} rows added, {changed} rows moved')
    if write:
        wb.save(WORKBOOK)
        print('written:', WORKBOOK)


if __name__ == '__main__':
    main('--write' in sys.argv)
