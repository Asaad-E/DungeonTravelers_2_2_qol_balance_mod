"""Easy fixes of the audit of 2026-10-17 in the workbooks: M-04, F87 names only W036 - W037 (W038 and W039 were dropped on 2026-10-11, a
passive record holds one weapon code), R06 and R07 carry the counts of the build of 2026-10-17; L-01, the note of enchantment 270 has
the factor of X01 (1.15 since 2026-10-13)."""
import os
import openpyxl

ROOT = os.path.join(os.path.dirname(__file__), '..', '..')

p = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
wb = openpyxl.load_workbook(p)
ws = wb['Criticisms']
head = [c.value for c in ws[1]]
col = {n: i + 1 for i, n in enumerate(head)}
for r in range(2, ws.max_row + 1):
    if ws.cell(r, col['ID']).value == 'F87':
        for name in ('Why', 'Proposal', 'Decision'):
            cell = ws.cell(r, col[name])
            v = str(cell.value or '')
            v = v.replace('W036-W039 making bow combat skills fully compatible with guns.',
                          'W036 - W037 opening Spiral Arrow and Divine Cannon to guns (W038 and W039, Warning Shot and Support Shot, '
                          'dropped on 2026-10-11: a passive record holds one weapon code).')
            v = v.replace('Solved by W036 - W039 (guns may use Spiral Arrow, Divine Cannon, Warning Shot, Support Shot)',
                          'Solved by W036 - W037 (guns may use Spiral Arrow and Divine Cannon; W038 - W039 dropped, as J06 says)')
            v = v.replace('2026-10-09: solved as J06/F10 (W036 - W039)', '2026-10-09: solved as J06/F10 (W036 - W037; W038 - W039 dropped 2026-10-11)')
            cell.value = v
ws = wb['Regression']
for r in range(2, ws.max_row + 1):
    rid = ws.cell(r, 1).value
    if rid == 'R06':
        ws.cell(r, 5).value = ('2026-10-17. BALANCE_BUILD.md: 2,238 of 2,238 live rows built, each naming the row it serves (the build '
                               'refuses a row that does not)')
    if rid == 'R07':
        cell = ws.cell(r, 5)
        cell.value = str(cell.value).replace('(99 dropped)', '(539 dropped on 2026-10-17)')
wb.save(p)

p = os.path.join(ROOT, 'data_tables', 'balance_enchant_text.xlsx')
wb = openpyxl.load_workbook(p)
done = 0
for ws in wb:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and '1.5 / 1.25 one to four levels below the enemy' in cell.value:
                cell.value = cell.value.replace('1.5 / 1.25 one to four levels below the enemy',
                                                '1.5 / 1.15 one to four levels below the enemy (X01 is 1.15 since 2026-10-13)')
                done += 1
assert done == 1, done
wb.save(p)
print('written')
