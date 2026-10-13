"""One-off edit (2026-10-13): C15 is accepted as easier (row R08 of the sheet Regression), and the HP of C30 stays."""
import os

import openpyxl

PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables',
                    'balance_criteria.xlsx')
wb = openpyxl.load_workbook(PATH)
ws = wb['Regression']
head = [c.value for c in ws[1]]
result = next(i for i, h in enumerate(head) if str(h).startswith('Result'))
row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == 'R08')
row[result].value = str(row[result].value) + ('. Accepted the same day: a little faster and within what was expected. The HP '
                                              'of C30 (x1.265 of the game) stays')
row[head.index('Verdict')].value = 'met; 10 bosses easier, accepted'
wb.save(PATH)
print('R08')
