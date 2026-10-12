"""The final pass of 2026-10-11 written into the sheet Regression of data_tables/balance_criteria.xlsx (columns Result
and Verdict), and the two rows of the register that waited for it (S27, F22). The figures are those of the generated
reports of the pass (scratch/audit_2026-10-11/regression_pass.log).

    python write_regression.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
WORKBOOK = os.path.join(PROJECT, 'data_tables', 'balance_criteria.xlsx')
RESULT, VERDICT = 'Result (final pass, 2026-10-11)', 'Verdict'
PASS = {
    'R01': ('met in what is measured',
            'PARTY_BATTERY.md: no party wins everywhere more than the best of vanilla did; the weakest party of vanilla '
            '(Anti-Meta-Plus, 0 - 14% at eight bosses) reaches 29 - 100%, still the weakest. CLASS_STAGES.md: 74 of 76 rows '
            'within the band. Limit: the band has no top, so a class that is too fast is not flagged (the Witch fights at '
            '0.59 of the median attacker at the advanced stage; row J11, left)'),
    'R02': ('met',
            'PARTY_BATTERY.md: 64 of 85 pairs at or above the floor in vanilla, 76 of 85 on the variant; no pair that met '
            'its floor in vanilla lost it. CLASS_STAGES.md: the two slow rows (Assassin and Treasure Hunter of the treasure '
            'hunter role) were slow before the changes (row J26)'),
    'R03': ('met',
            'CHECKPOINT_LEVELS.md: 95 checkpoints, 19 jumps; no row of balance_changes.xlsx changes the level of an enemy, '
            'so every jump stands where vanilla has it'),
    'R04': ('partly measured',
            'Not run as written (each class left out in turn). The nearest measure: five parties of different make at 17 '
            'bosses, 76 of 85 pairs above the floor, and no boss where only one of them wins except C29 (Mage-Heavy 100%, '
            'the others 6 - 37%: Calliope, which the model does not play as the guide describes)'),
    'R05': ('met by construction',
            'A rows: every story and postgame boss is immune to turn delay; Magnificent Steps and Warning Shot at their new '
            'values take 19% of the turns of an enemy with three bearers of each (33% in vanilla)'),
    'R06': ('met',
            'BALANCE_BUILD.md: 1,386 of 1,386 live rows built, and each names the row it serves (column Serves, none empty)'),
    'R07': ('not reported as a figure',
            'The distance from vanilla is not summed up in one number; each row of balance_changes.xlsx holds the vanilla '
            'value beside the new one, and the rows the measures did not need were dropped (99 dropped)'),
    'R08': ('met; 7 bosses easier, accepted',
            'DIFFICULTY_CURVE.md: 37 of 44 checkpoints within the windows of vanilla, 7 easier (C14, C26, C29, C30, C32, '
            'C37, C38), 0 harder. Decided 2026-10-11: the bosses the variant leaves easier stay as they are. C40, raised '
            'as an exception, is within (wins 88% > 95%, length 35.5 > 34.2). Limit found the same day: Heretic Samsara of '
            'Demiurge is not played by the model, in vanilla or on the variant'),
    'R09': ('met',
            'VERDICTS.md: of 162 rows, 72 closed with the figure before and after, 77 not a problem with their figure, 12 '
            'out of the balance, 1 open (J02, left for the postgame)'),
    'R10': ('met',
            'DIFFICULTY_CURVE.md, columns Lowest HP and Pressure: inside the window of vanilla at 37 of 44 checkpoints; '
            'where they leave it the party stands higher (lowest HP 11 - 31% > 22 - 53% at C29, C30, C38), never lower. '
            'SIM_ANCHORS.md: 6 of 6 anchors met'),
}
ROWS = {
    'S27': ('not a problem',
            'partly true where vanilla was unfair, and eased there; elsewhere the variant keeps the ratio of vanilla',
            'DIFFICULTY_CURVE.md (final pass, Regression R08 and R10): the damage the enemies deal per turn against the '
            "party's max HP and its lowest HP stay inside the window of vanilla at 37 of 44 bosses; at the other 7 the "
            'party stands higher'),
    'F22': ('not a problem',
            'true for Roulette Circus only, and by its nature: Magic Box, Showtime and Heaven or Hell are worth 100% '
            'after the L rows',
            'L rows built; L002 dropped on 2026-10-11 : its damage draw strikes any unit of the battle, the '
            'party included, so it stays at level x 10. Its heal of an enemy is 10% of max HP instead of all of it '
            '(L001). Cast every turn, rows of eight random battles last 3 - 25% longer and end with -1 to +3 points of HP: '
            'the cheap gamble of the base class (15 TP, one level), which can no longer sink a fight'),
}


def main():
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb['Regression']
    head = [c.value for c in ws[1]]
    for name in (RESULT, VERDICT):
        if name not in head:
            ws.cell(row=1, column=len(head) + 1, value=name)
            head.append(name)
    for row in ws.iter_rows(min_row=2):
        if row[0].value in PASS:
            verdict, result = PASS[row[0].value]
            ws.cell(row=row[0].row, column=head.index(RESULT) + 1, value=result)
            ws.cell(row=row[0].row, column=head.index(VERDICT) + 1, value=verdict)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    col = {name: head.index(name) for name in head}
    for row in ws.iter_rows(min_row=2):
        if row[0].value in ROWS:
            state, verdict, figure = ROWS[row[0].value]
            row[col['Status']].value, row[col['Verdict']].value, row[col['Balance_Result']].value = state, verdict, figure
            row[col['Decision']].value = f"{row[col['Decision']].value or ''}; 2026-10-11: closed with the final pass".strip('; ')
    wb.save(WORKBOOK)


if __name__ == '__main__':
    main()
