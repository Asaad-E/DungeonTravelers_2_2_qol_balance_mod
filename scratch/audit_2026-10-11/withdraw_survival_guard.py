"""The user, 2026-10-12: effect 254 stays as the game has it (it may be rolled at 1 HP), since the guard of row X15
changed Guts of the Etoile too. Row X15 and lever L46 leave the workbook; the box of the unique skill of Lizerietta
goes back to its sentence without the clause. Run once."""
import os
import openpyxl

TABLES = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables')
BOX = ('When a hit would fell her, she may hold on with 1 HP.' + chr(10) +
       'Activation rate: 10% -> 12 -> 14 -> 16 -> 18 -> 20 -> 22 -> 24 -> 26%   |   Crown EX: 35%')
path = os.path.join(TABLES, 'balance_criteria.xlsx')
wb = openpyxl.load_workbook(path)
for sheet, key in (('ExePatches', 'X15'), ('Levers', 'L46')):
    ws = wb[sheet]
    for row in list(ws.iter_rows(min_row=2)):
        if row[0].value == key:
            ws.delete_rows(row[0].row)
wb.save(path)
path = os.path.join(TABLES, 'balance_skill_text.xlsx')
wb = openpyxl.load_workbook(path)
ws = wb['Skills']
col = {c.value: k for k, c in enumerate(ws[1])}
for row in ws.iter_rows(min_row=2):
    if row[col['Filename']].value == 'UniqueSkillData_0370.bin':
        row[col['Balance_Desc']].value = BOX
wb.save(path)
print('written')
