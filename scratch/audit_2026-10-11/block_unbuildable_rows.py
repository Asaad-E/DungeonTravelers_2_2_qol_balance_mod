"""Moves to the state "blocked" the rows of balance_changes.xlsx the value step refuses (tools/balance_values.py),
with the reason, and writes the state into the sheet About. A blocked row is not of the variant: neither the
simulator nor the build applies it until the decision is taken on what it becomes.

    python block_unbuildable_rows.py --write
"""
import os
import sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(os.path.dirname(HERE))
WORKBOOK = os.path.join(PROJECT, 'data_tables', 'balance_changes.xlsx')
TODAY = '2026-10-11'

WEAPON = ('a passive skill record holds one weapon code (PassiveSkillData 0x54), so "8, 15" cannot be stored; the only '
          'code that covers a gun is 19, "all weapons" (Weapon Mastery, Knock Down), which is wider than bow or gun. '
          'To decide: 19, or the passive stays a bow skill')
DROP = ('the word the row means is EnemyStatusData 0x80, which is not a chance: it is how many pieces the drop draws '
        '(0x558517 allocates that many; 0 to 3 over the 516 enemies), and the chance of the drop is 0x7C, already 100 '
        'for these bosses (16 for Garasha and Ryoma). 5 is outside what the game holds. To decide: the count (3 at most), '
        'or nothing')
QUEST = ('QuestData 0x38 is the kind of the reward (0 an item, 1 a piece of equipment, 2 and 3 an amount at 0x40) and 0x3C '
         'the id of the item or the piece; the figure the row took for gold is that id in quests 0, 5 and 100 (items 30, 42 '
         'and 585) and in quest 200 (the piece 599, Postapocalypse Helm), and in quest 50 an amount of kind 3. Written as '
         'it stands the row would change what the quest gives, or name a piece that does not exist. To decide: which '
         'quests pay gold (kind 2, 37 of them) and by how much')
REASONS = {**{f'W{n:03d}': WEAPON for n in (38, 39)},
           **{f'P{n:03d}': DROP for n in range(11, 26)},
           **{f'P{n:03d}': QUEST for n in range(26, 31)}}
ABOUT = ('State blocked', 'A row that cannot be written into the game files as it stands (tools/balance_values.py refuses '
         'it), with the reason at the start of its Why. It is not of the variant: sim/variant.py and the build leave it '
         'out until the decision is taken on what it becomes.')


def main(write):
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    col = {name: k for k, name in enumerate(head)}
    done = 0
    for row in ws.iter_rows(min_row=2):
        rid = row[col['ID']].value
        if rid in REASONS and row[col['State']].value == 'decided':
            row[col['State']].value = 'blocked'
            row[col['Why']].value = f"Blocked on {TODAY}: {REASONS[rid]}. Was: {row[col['Why']].value}"
            done += 1
    about = wb['About']
    if not any(str(r[0].value).startswith('State blocked') for r in about.iter_rows() if r[0].value):
        about.append(ABOUT)
    print(done, 'rows blocked')
    if write:
        wb.save(WORKBOOK)


if __name__ == '__main__':
    main('--write' in sys.argv)
