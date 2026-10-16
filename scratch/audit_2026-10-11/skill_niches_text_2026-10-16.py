"""balance_skill_text.xlsx: the hand-written boxes of Oracles and Shadow Clone follow the rows of 2026-10-16 (D061 - D070 x1.5, V016 - V020).
A doll's ATK in its box is (1 + the power of its action) * 100% (Spartan 1.88 -> 288%)."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'balance_skill_text.xlsx')
ORACLES_POWER = [2.28, 2.34, 2.43, 2.52, 2.61, 2.70, 2.82, 2.94, 3.06, 3.21]
atk = [round(100 * (1 + v)) for v in ORACLES_POWER]
steps = [b - a for a, b in zip(atk, atk[1:])]
EDITS = {
    'Oracles': ('ATK: 252% (+4-10%/Lv) -> 314%',
                'ATK: %d%% (+%d-%d%%/Lv) -> %d%%' % (atk[0], min(steps), max(steps), atk[-1]),
                'Rows D041 - D080: power of Oracles\' Action (x1.5 more on 2026-10-16), life, speed'),
    'Shadow Clone': ('Turns: 3 -> 3 -> 4 -> 4 -> 5', 'Turns: 4 -> 4 -> 5 -> 5 -> 6',
                     'Rows K001, V016 - V020 (one more turn): the generated box fell back to "Number of dodges UP"'),
}

wb = openpyxl.load_workbook(PATH)
ws = wb['Skills']
head = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(head)}
done = set()
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, col['Skill_Name']).value
    if name in EDITS:
        old, new, note = EDITS[name]
        text = ws.cell(r, col['Balance_LvUp']).value
        assert old in text, (name, text)
        ws.cell(r, col['Balance_LvUp']).value = text.replace(old, new)
        ws.cell(r, col['Note']).value = note
        done.add(name)
assert done == set(EDITS), done
wb.save(PATH)
print('written', sorted(done), EDITS['Oracles'][1])
