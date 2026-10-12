"""The chance of the unique skill of Lizerietta lowered (decided 2026-10-12: 40% and 55% made it the best unique
skill of the game, near certain where the others are occasional): 10 - 26% by rank, 35% with the crown. Run once."""
import os
import openpyxl

TABLES = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables')
VALUES = [0.10, 0.12, 0.14, 0.16, 0.18, 0.20, 0.22, 0.24, 0.26, 0.35]
BOX = ('When a hit would fell her, she may hold on with 1 HP.' + chr(10) +
       'Activation rate: 10% -> 12 -> 14 -> 16 -> 18 -> 20 -> 22 -> 24 -> 26%   |   Crown EX: 35%')
NOTE = ('. Lowered the same night (decided: at 40% and 55% it was near certain, the best unique skill of the game): 10 - '
        '26%, 35% with the crown, beside Alter Fate, which stops any hit at 10 - 22% and 30%, and Guts, 10 - 30%')

path = os.path.join(TABLES, 'balance_changes.xlsx')
wb = openpyxl.load_workbook(path)
ws = wb['Changes']
col = {c.value: k for k, c in enumerate(ws[1])}
for row in ws.iter_rows(min_row=2):
    if str(row[col['Record']].value) == 'Unique 370 Usurper of Souls' and row[col['Field']].value == 'value':
        row[col['New']].value = VALUES[int(row[col['Level']].value) - 1]
        row[col['Why']].value = str(row[col['Why']].value) + NOTE
        row[col['State']].value = 'decided'
wb.save(path)
path = os.path.join(TABLES, 'balance_skill_text.xlsx')
wb = openpyxl.load_workbook(path)
ws = wb['Skills']
col = {c.value: k for k, c in enumerate(ws[1])}
for row in ws.iter_rows(min_row=2):
    if row[col['Filename']].value == 'UniqueSkillData_0370.bin':
        row[col['Balance_Desc']].value = BOX
wb.save(path)
print('written')
