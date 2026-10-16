"""Weak tools of the Maid line raised (decided 2026-10-16), measured by dtcli skill-battery (20 runs a cell, the 3 best again at 200):
Dumbing Mazurka (effect 201) x1.3 on its raised value (A124 - A133), niche 22% -> good 30%; Emaciating Hopak (effect 202) x1.4, niche
22% -> good 27 - 31%; Absolute Obedience power (effect 26) and its two ailment rolls (117, 115) x1.4, niche 18% -> good 25%. In the
band of the classes (class-search preset=balance) the Etoile stays at 86.6% (the raised dances take the place of others in her two
slots) and the Mistress goes 84.7% -> 86.0% (+1.6 sd, in the band). First Aid is left as it is: a basic tool of the start."""
import json
import os
import openpyxl

ROOT = os.path.join(os.path.dirname(__file__), '..', '..')
PATH = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
DECIDED = '2026-10-16'
WHY_D = ('The dances of the Etoile made even (decided 2026-10-16): skill-battery, Dumbing Mazurka x1.3 niche 22% -> good 30% of the cells, '
         'Emaciating Hopak x1.4 niche 22% -> good 27 - 31%; her wins in the band of the classes unchanged (86.6%).')
WHY_M = ('A unique tool of the Mistress made worth its turn (decided 2026-10-16): skill-battery, power and both ailment rolls x1.4, niche '
         '18% -> good 25% of the cells; the Mistress 84.7% -> 86.0% wins in the band of the classes (+1.6 sd).')

snap = json.load(open(os.path.join(ROOT, 'build', 'engine_snapshot.json'), encoding='utf-8'))['sides']['vanilla']['skills']
NEW = []
for sid, name, effect, f, digits, why in ((75, 'Emaciating Hopak', 202, 1.4, 0, WHY_D),
                                          (78, 'Absolute Obedience', 26, 1.4, 2, WHY_M),
                                          (78, 'Absolute Obedience', 117, 1.4, 2, WHY_M),
                                          (78, 'Absolute Obedience', 115, 1.4, 2, WHY_M)):
    skill = next(s for s in snap if s['id'] == sid)
    power = next(e[2] for e in skill['effects'] if e[0] == effect)
    field = 'effect 26 (power)' if effect == 26 else f'effect {effect}'
    for level, old in enumerate(power, start=1):
        new = round(old * f) if digits == 0 else round(old * f, digits)
        NEW.append(('skills', f'Technique {sid} {name}', field, level, old, new, 'L12', 'S07', why))

wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
head = [c.value for c in ws[1]]
assert head[:12] == ['ID', 'Table', 'Record', 'Field', 'Level', 'Vanilla', 'New', 'Lever', 'Serves', 'Why', 'Decided', 'State'], head
records = {(ws.cell(r, 3).value, ws.cell(r, 4).value) for r in range(2, ws.max_row + 1)}
assert not any((v[1], v[2]) in records for v in NEW), 'rows already there'
done = 0
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 1).value in ['A%03d' % k for k in range(124, 134)]:
        assert ws.cell(r, 4).value == 'effect 201'
        ws.cell(r, 7).value = round(ws.cell(r, 7).value * 1.3, 2)
        ws.cell(r, 10).value = str(ws.cell(r, 10).value) + ' ' + WHY_D
        done += 1
assert done == 10, done
ids = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
first = 1 + max(int(k[1:]) for k in ids if k and str(k).startswith('V'))
row = ws.max_row + 1
for k, values in enumerate(NEW):
    for c, value in enumerate(('V%03d' % (first + k),) + values + (DECIDED, 'built'), start=1):
        ws.cell(row, c).value = value
    row += 1
wb.save(PATH)
print('Mazurka rows', done, '; written V%03d - V%03d' % (first, first + len(NEW) - 1))
