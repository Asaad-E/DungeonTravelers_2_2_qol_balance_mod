"""One-off edit: the Why of rows X056 - X060 (Life Bringer) named "+0.7 power at every level" and a worth of 76%, which
are not the values built (power 0.05 - 0.25, +0.15 - 0.17 a level over the game's -0.12 - 0.10) nor its worth now
(48% beside Showtime, docs/generated/SKILL_WORTH.md; left as it is on 2026-10-11, row F72 of Criticisms)."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
OLD = '+0.7 power at every level (it was -0.12 to +0.10) brings it to 76%, the max HP cut kept'
NEW = ('power 0.05 to 0.25 (it was -0.12 to +0.10, +0.15 - 0.17 a level), the max HP cut kept; it stands at 48% beside '
       'Showtime (docs/generated/SKILL_WORTH.md) and was left there on 2026-10-11 (F72)')

wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
done = 0
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 1).value in ('X056', 'X057', 'X058', 'X059', 'X060'):
        why = ws.cell(r, 10).value
        assert OLD in why, ws.cell(r, 1).value
        ws.cell(r, 10).value = why.replace(OLD, NEW)
        done += 1
assert done == 5
wb.save(PATH)
print('rows', done)
