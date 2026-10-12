"""The EXP enchantment (270) on one unique accessory instead of a sealbook, the user's word of 2026-10-12: row P089
(the Grand Sealbook of Molle) is dropped and the Bathroom Slippers take it in the place of their second enchantment.
The Bathroom Slippers are given once, on reaching quest rank 3 (QuestRankData_03; atwiki dt2-2 page 48): never
dropped, never sold, never stolen. The book of Dragon Papi keeps its own. Rows of data_tables/balance_changes.xlsx.

    python early_exp_accessory.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CHANGES = os.path.join(PROJECT, 'data_tables', 'balance_changes.xlsx')
DROPPED = ('Dropped on 2026-10-12 (the user): the effect goes on an accessory, not on a sealbook; the book of Molle is as '
           'in vanilla. Was: ')
ROW = {'ID': 'P090', 'Table': 'equipment', 'Record': '732 Bathroom Slippers', 'Field': 'Enchant 2 (0x64)',
       'Vanilla': 33, 'New': 270, 'Lever': 'L15', 'Serves': 'U01',
       'Why': 'The user, 2026-10-12: the piece that raises the EXP factors of B4 comes late (the Grand Sealbook of Dragon '
              'Papi, level 44); he wants the same effect earlier on one unique accessory that is neither farmed nor '
              'stolen, without filling the game with such pieces. The Bathroom Slippers are the reward of quest rank 3, '
              'given once, wearable by every class; they keep "Nullify poison" and give up "Physical defense UP" '
              '(enchantment 33), the count of innate enchantments staying as it is (lever L15)',
       'Decided': '2026-10-12', 'State': 'decided'}


def main():
    wb = openpyxl.load_workbook(CHANGES)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    col = {name: head.index(name) for name in head}
    ids = set()
    for row in ws.iter_rows(min_row=2):
        ids.add(row[0].value)
        if row[0].value == 'P089' and row[col['State']].value != 'dropped':
            row[col['State']].value = 'dropped'
            row[col['Why']].value = DROPPED + str(row[col['Why']].value)
    if ROW['ID'] not in ids:
        ws.append([ROW.get(name) for name in head])
    wb.save(CHANGES)


if __name__ == '__main__':
    main()
