"""One-off edit (2026-10-13): row G118 (the HP of Samael at C38) says it is kept after the check of that day."""
import os

import openpyxl

PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables',
                    'balance_changes.xlsx')
wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
head = [c.value for c in ws[1]]
row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == 'G118')
row[head.index('Why')].value = str(row[head.index('Why')].value) + (
    ' | 2026-10-13 : kept after the check with the model as it is now (measures/BOSS_ROWS_CHECK.md, 200 battles): the '
    'reference party wins 90% with it and 94% without, 64% in the game')
wb.save(PATH)
print('G118')
