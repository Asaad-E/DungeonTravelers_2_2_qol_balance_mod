"""Two boxes of unique skills that did not say what the program does (docs/FORMULAS.md 12.19), corrected in
data_tables/qol_rules.xlsx, the user's word of 2026-10-12:

* Wreathed in Glory: the label "Turn delay" read the value as a cut of the delay; the program adds it to the speed
  of the skill or spell used (0x4f72a7).
* Tranquil Song: the box gave the chance and not what the song does: ATK and INT -0.1 on every opponent, asked for
  4 turns (0x43749f - 0x437650).

    python qol_unique_text.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RULES = os.path.join(PROJECT, 'data_tables', 'qol_rules.xlsx')
LABEL = ('Wreathed in Glory', 'Speed added',
         'Read in the program (docs/FORMULAS.md 12.19, 0x4f72a7): the value is added to the speed of the skill or spell '
         'she uses, the figure the boxes show as "Speed: 0.95x". The label "Turn delay" read it as a cut of the delay')
NOTE = ('Tranquil Song', "Enemy ATK and INT: -10% for 4 turns",
        'Read in the program (docs/FORMULAS.md 12.19, 0x43749f - 0x437650): on a success every unit of the other side '
        'receives an ATK debuff and an INT debuff of -0.1 (effects 41 and 43), asked for 4 turns')


def main():
    wb = openpyxl.load_workbook(RULES)
    ws = wb['UniqueSkillLabels']
    for row in ws.iter_rows(min_row=2):
        if row[0].value == LABEL[0]:
            row[1].value, row[2].value = LABEL[1], LABEL[2]
    ws = wb['UniqueSkillNotes']
    if NOTE[0] not in {r[0].value for r in ws.iter_rows(min_row=2)}:
        ws.append(NOTE)
    wb.save(RULES)


if __name__ == '__main__':
    main()
