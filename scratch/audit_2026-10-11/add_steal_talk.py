"""Two boxes of a short talk of Euthys (Lesson.pck, Talk_Scenario27, the one about gold) say what a steal takes and
from whom: rows of data_tables/balance_tutorial_text.xlsx. The wording was agreed on 2026-10-11.

    python add_steal_talk.py
"""
import os
import re
import sys
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(PROJECT, 'tools'))
from pck_handler import PckArchive                          # noqa: E402

WORKBOOK = os.path.join(PROJECT, 'data_tables', 'balance_tutorial_text.xlsx')
ENTRY = 'Talk_Scenario27.adv.bin'
NEW = (('While most things in this world',
        '"What gold cannot buy, a nimble thief may lift from a monster:\n'
        'one of the things it would drop, and only once from each."',
        'FORMULAS 12.16: a steal that succeeds draws among the drops of the enemy (0x4504bf calls 0x558180 with its '
        'steal flag), once per enemy (0x154). Not voiced. The first two boxes of the talk are the original ones'),
       ("You'll understand when you're older",
        '"Most bosses, mind you, guard their pockets far too well.\n'
        'You\'ll understand when you try... and then wish you hadn\'t."',
        'Bit 15 of the flags of an enemy (0x4503c5): 131 of the 157 enemies that only stand in fixed battles do not '
        'have it. The end of the original box is kept as the joke of the talk'))


def main():
    blob = bytes(dict(PckArchive(os.path.join(PROJECT, 'backups', 'Lesson.pck')).files)[ENTRY])
    boxes = [m.group()[:-1].decode('latin-1') for m in re.finditer(rb'[\x20-\x7e\n]{6,}\x00', blob)]
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb['Tutorials']
    limits = {r[0]: r[1] for r in wb['Settings'].iter_rows(min_row=2, values_only=True)}
    have = {(r[1].value, r[2].value) for r in ws.iter_rows(min_row=2) if r[0].value}
    for start, text, note in NEW:
        original = next(b for b in boxes if b.strip('"').startswith(start))
        lines = text.strip('"').split('\n')
        assert len(lines) <= limits['max_lines'] and max(len(x) for x in lines) <= limits['max_width'], text
        if (ENTRY, original) not in have:
            ws.append(('Lesson.pck', ENTRY, original, text, None, 'checked', note))
        print(repr(original), '->', repr(text))
    wb.save(WORKBOOK)


if __name__ == '__main__':
    main()
