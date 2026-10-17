"""The Diva's HP and RES rows (V493 - V690) dropped (decided 2026-10-16): her wins come from her survival, and at C40 (Demiurge) that
survival is what keeps a party with no other healer alive. Party battery, Juggernaut-Plus at C40 (floor 50%): 68% without the cut, 19%
with HP and RES x0.9, 36% with HP x0.9 alone; RES x0.9 alone broke nothing but left her at +2.4 sd in the band of the classes. No cut
put her in the band without breaking C40; she stays the best support, an outlier in wins kept by design (S07)."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'balance_changes.xlsx')
WHY = (' | Dropped 2026-10-16: at C40 (Demiurge) the Juggernaut-Plus party, which has no other healer, fell from 68% to 19% with HP and '
       'RES x0.9 (floor 50%), 36% with HP x0.9 alone; no cut put the Diva in the band without breaking C40. She stays the best support.')
wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
head = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(head)}
done = 0
for r in range(2, ws.max_row + 1):
    if ws.cell(r, col['Record']).value == '27 Diva' and ws.cell(r, col['State']).value == 'built':
        ws.cell(r, col['State']).value = 'dropped'
        ws.cell(r, col['Why']).value = str(ws.cell(r, col['Why']).value) + WHY
        done += 1
assert done == 198, done
wb.save(PATH)
print('dropped', done)
