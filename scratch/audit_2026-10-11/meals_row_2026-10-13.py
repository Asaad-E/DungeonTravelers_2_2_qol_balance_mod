"""One-off edit (2026-10-13): row J36 of the register (meals) takes its measure. Run once; kept as the record."""
import os

import openpyxl

PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables',
                    'balance_criteria.xlsx')
SET = {
    'Vanilla_Result': 'Of the 16 meals, 10 give a buff: the status of a skill at a fixed level, with the turns of that skill '
                      '(3 to 10). On the map a turn of a status is 5 steps (docs/FORMULAS.md 12.18), so a meal lasts 15 to 50 '
                      'steps, and in a battle each action of its bearer takes a turn off it. The mean walk between two random '
                      'battles is 18 to 28 steps (measures/DESIGN_ROWS.md, 2026-10-13)',
    'Verdict': 'true: a meal lasts 15 to 50 steps, about one battle',
    'Why': 'The shop appears where the roll of a step puts it, and its buff is gone a battle later: it reaches a boss only when '
           'the shop appears at its door. The turns are those of the skills themselves, and the five steps a turn are one byte '
           'of the program that every status of the map shares (ailments too)',
    'Decision': '2026-10-13: measured; the claim holds. No change so far: the state waits for a decision',
}


def main():
    wb = openpyxl.load_workbook(PATH)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == 'J36')
    for column, value in SET.items():
        row[head.index(column)].value = value
    wb.save(PATH)
    print('J36')


if __name__ == '__main__':
    main()
