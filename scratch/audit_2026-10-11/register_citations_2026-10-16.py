"""One-off edit: four rows of the sheet Criticisms (balance_criteria.xlsx) that cited rows of Changes taken back or named
for what they do not do, as the audit of 2026-10-16 found: J06 (W038, W039 dropped), J24 (H036, H037 are a cure, not Full
Heal), F17 (T031 - T040 dropped for S003 - S012), I11 (the values of No Miracles Needed lowered after the proposal)."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
ch = {r[0]: r for r in openpyxl.load_workbook(os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx'), read_only=True)['Changes']
      .iter_rows(values_only=True) if r[0]}
assert 'Checkmate' in str(ch['S003'][2]) and ch['S003'][11] == 'built', ch['S003'][:3]
assert ch['W038'][11] == ch['W039'][11] == 'dropped' and ch['W036'][11] == ch['W037'][11] == 'built'
EDITS = {
    'J06': {'Balance_Result': 'W036, W037: Spiral Arrow and Divine Cannon open to guns, built and read back (BALANCE_BUILD.md); W038, '
                              'W039 (Warning Shot, Support Shot) dropped on 2026-10-11, a passive record holds one weapon code; '
                              'shooter within the band at the four stages (0.85, 0.89, 0.88, 1.09; CLASS_STAGES.md)'},
    'J24': {'Why': 'Solved by class identity: the Bishop learns Full Heal (100% of max HP, no cap). H036 and H037 do not touch '
                   'the cap: they let Heal and Circle Heal cure poison too (and paralysis at the upper levels of Heal)',
            'Proposal': 'Solved by class identity (Full Heal); H036 / H037 add a cure to Heal and Circle Heal',
            'Decision': '2026-10-10: solved by Full Heal; H036, H037 (the cure on the heals) decided 2026-10-09'},
    'F17': {'Balance_Result': 'S003 - S012 (Checkmate, direct values 0.40 to 1.20; they replaced T031 - T040, x1.6, dropped to avoid '
                              'stacking two buffs) and Y rows, built and read back (BALANCE_BUILD.md); Trickster 1.08 - 1.13, Joker '
                              '1.00 advanced and 0.96 late, within the band (CLASS_STAGES.md)',
            'Proposal': 'T031 - T040: Checkmate x1.6 power; replaced by S003 - S012 (direct values)'},
    'I11': {'Proposal': 'U011 - U020, U051: the skill becomes the chance to hold on with 1 HP when a hit would fell her (effect 254, '
                        '10 - 26% by rank, 35% with the crown, lowered from the 15 - 40% / 55% first proposed), named No Miracles '
                        'Needed in the variant'},
}

wb = openpyxl.load_workbook(PATH)
ws = wb['Criticisms']
header = [c.value for c in ws[1]]
done = 0
for r in range(2, ws.max_row + 1):
    rid = ws.cell(r, 1).value
    for col, text in EDITS.get(rid, {}).items():
        ws.cell(r, header.index(col) + 1).value = text
        done += 1
assert done == sum(len(v) for v in EDITS.values()), done
wb.save(PATH)
print('cells', done)
