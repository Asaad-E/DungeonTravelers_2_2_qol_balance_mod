"""C19 Siesta: DEF x1.2 (new row G122), decided 2026-10-16: the fight had fallen under the length window of vanilla."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'balance_changes.xlsx')
WHY = ('The length of C19 back into the window of vanilla (R08): DEF x1.2 (decided 2026-10-16). Reference party, 400 runs: vanilla 8.68, '
       'the variant 6.43 (under the window, 6.53), with DEF x1.2 7.32, lowest HP 61% (vanilla 58%); ATK x1.2 did not lengthen it (6.46). '
       'dtcli class-search final C19 boss, the 12,344 final-class parties: wins vanilla 76.0%, variant with DEF x1.2 88.0%, every class 6 - 23 '
       'points above vanilla, 2 parties 10 or more points under it (the worst 11). Only Siesta: the two other enemies of the fight are met on the floor.')
ROW = ('G122', 'enemies', '795 Siesta (C19)', 'def', None, 85, 102, 'L01', 'R08, R10', WHY, '2026-10-16', 'built')

wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
ids = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
if 'G122' in ids:
    at = ids['G122']
else:
    at = ids['G121'] + 1
    ws.insert_rows(at)
for col, value in enumerate(ROW, start=1):
    ws.cell(at, col).value = value
wb.save(PATH)
print('written G122')
