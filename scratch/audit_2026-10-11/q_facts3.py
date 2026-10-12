"""Third pass of facts: Slave's Torc, who carries Double EXP, Circle Protect's class, meals, installed files vs backups."""
import os
import sys
import struct
import hashlib
import openpyxl
sys.path.insert(0, r'C:\Users\user\Personal\Physics\DungeonTraveler2-2\mod_project')
sys.path.insert(0, r'C:\Users\user\Personal\Physics\DungeonTraveler2-2\mod_project\tools')
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
ROOT = r'C:\Users\user\Personal\Physics\DungeonTraveler2-2'
SOT = os.path.join(ROOT, 'mod_project', 'data_tables', 'vanilla_source_of_truth.xlsx')
wb = openpyxl.load_workbook(SOT, read_only=True, data_only=True)


def table(name):
    rows = list(wb[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


eq = table('Equipment')
print("== Slave's Torc, every enchant column")
for d in eq:
    if d['Name'] == "Slave's Torc":
        print('   ', {k: v for k, v in d.items() if 'nchant' in k and v not in (None, '')})
print('== Equipment that carries enchant 275 (Double EXP) or 270')
for d in eq:
    ids = [v for k, v in d.items() if k.startswith('Enchant') and k.endswith('_ID')]
    for e in (275, 270):
        if e in ids:
            print(f"    {e}: {d['Equip_ID']} {d['Name']} ({d['Category']})")
seal = list(wb['Raw_SealScroll'].iter_rows(values_only=True))
for r in seal[1:]:
    if 275 in r[10:14] or 270 in r[10:14]:
        print('    sealbook:', r[1], r[8:14])

from sim import data as D
van = D.game_data()
s = van.skills.get(('Magic', 227))
print('== Circle Protect:', s.name, 'classes:', getattr(s, 'classes', None), [van.class_names.get(c) for c in (getattr(s, 'classes', None) or [])])


food = openpyxl.load_workbook(os.path.join(ROOT, 'mod_project', 'data_tables', 'qol_food_text.xlsx'), read_only=True, data_only=True)
for ws in food.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    print('== qol_food_text sheet', ws.title, len(rows))
    for r in rows[:1] + [x for x in rows[1:] if any('step' in str(c).lower() or 'ramen' in str(c).lower() for c in x)][:6]:
        print('   ', tuple(str(c)[:110] for c in r if c is not None))

print('== installed game files that differ from the backups, and whether build/qol explains them')
game = os.path.join(ROOT, 'DungeonTravelers2-2_EN', 'Data', 'ENG')
back = os.path.join(ROOT, 'mod_project', 'backups')
qol = os.path.join(ROOT, 'mod_project', 'build', 'qol', 'Data', 'ENG')


def md5(p):
    with open(p, 'rb') as fh:
        return hashlib.md5(fh.read()).hexdigest()


backup_by_name = {}
for base, _, files in os.walk(back):
    for fn in files:
        backup_by_name.setdefault(fn, os.path.join(base, fn))
for sub in ('Parameter', 'Script'):
    for fn in sorted(os.listdir(os.path.join(game, sub))):
        gp = os.path.join(game, sub, fn)
        if not os.path.isfile(gp):
            continue
        bp = backup_by_name.get(fn)
        qp = os.path.join(qol, sub, fn)
        g = md5(gp)
        same_b = bp is not None and md5(bp) == g
        same_q = os.path.isfile(qp) and md5(qp) == g
        if not same_b:
            import time
            print(f"    {sub}/{fn}: vanilla backup {'MISSING' if bp is None else 'differs'}; equals build/qol: {same_q}; "
                  f"modified {time.strftime('%Y-%m-%d %H:%M', time.localtime(os.path.getmtime(gp)))}")
