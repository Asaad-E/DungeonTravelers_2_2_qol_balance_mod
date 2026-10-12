"""What the Japanese guide says a player brings to the postgame, written into data_tables/sim_setup.xlsx: the
resistances of the party from the Plains of Light on (sheet Preparation), the zone of the weakened Crom Cruach (sheet
ZoneChanges) and, beside what GameFAQs says of each boss, what the Japanese pages say (sheet PostgameBosses).
Read on 2026-10-11 from w.atwiki.jp/dt2-2; paraphrased, never copied.

    python postgame_preparation.py
"""
import os
import sys
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, PROJECT)
sys.path.insert(0, os.path.join(PROJECT, 'tools'))
from sim import scenario, setup as setup_mod                 # noqa: E402

FIRST = 'P17'                    # the first boss of the Plains of Light in the order of the guide
ELEMENTS = ((3, 'fire'), (4, 'water'), (5, 'wind'), (6, 'earth'), (7, 'lightning'))
PREMISE = ('atwiki dt2-2 page 326 (Plains of Light, overview): by then the party wears leather pieces and rings forged '
           'with AGI, speed and elemental resistance, seven resistance enchantments of 15 fill it, and every '
           'resistance at 100 is the premise of the dungeon; taken for every boss from there on. The enchantment '
           'slots it takes are not charged (the sheet EnchantPlan fills them apart)')
NOTES = {
    'P03': 'atwiki page 37: high physical attack, raise physical defence; stronger than the first time',
    'P10': 'atwiki page 29 (not read yet)',
    'P12': 'atwiki page 219: with bash resistance raised and a physical barrier the fight is stable',
    'P13': 'atwiki page 219: fought on a cell that forbids magic; poison, freeze, silence and slow work on it',
    'P17': 'atwiki page 326: a strong enemy; high defence in a zone that raises magic attack and physical defence',
    'P20': 'atwiki page 271: the fire of the quest removes the water zone and the second form and does not weaken '
           'its stats; no chance on arrival; with it a party of mean level 85 or more can just fight it',
    'P22': 'atwiki page 206: a party of physical attackers cannot scratch it (high defence, voids ranged blows, '
           'Parry kept up); earth magic on its earth cell, or fixed damage',
    'P23': 'atwiki pages 206 and 317: Chaos Eye pierces and petrifies for sure without earth resistance 100 or a '
           'piece against petrification',
    'P29': 'atwiki page 326: easier to fight than the strong enemies before it as long as healing is kept up',
    'P30': 'atwiki page 231: with Lolo present she always acts first; the weakened fight is the one of level 90',
    'P31': 'atwiki page 219: the incense weakens the boss and removes its two escorts',
    'P50': 'atwiki page 231: the strongest enemy of the game, hard even with every member crowned',
}


def main():
    setup = setup_mod.load()
    wb = openpyxl.load_workbook(setup_mod.SETUP)
    order = [cp['ID'] for cp in setup.checkpoints if cp['ID'].startswith('P')]
    late = order[order.index(FIRST):]
    ws = wb['Preparation']
    have = {(r[0].value, str(r[1].value), str(r[3].value)) for r in ws.iter_rows(min_row=2) if r[0].value}
    added = 0
    for cid in late:
        cp = setup.checkpoint(cid)
        for number in scenario.battle_numbers(cp):
            for index, name in ELEMENTS:
                if (cp['Floor'], str(number), str(index)) not in have:
                    ws.append((cp['Floor'], number, f"{cid} {cp.get('Boss (as the story page names it)') or ''}".strip(),
                               index, name, 100, PREMISE))
                    added += 1
    zones = wb['ZoneChanges']
    crom = setup.checkpoint('P20')
    number = scenario.battle_numbers(crom)[0]
    if not any(r[0].value == crom['Floor'] and str(r[1].value) == str(number) for r in zones.iter_rows(min_row=2)):
        zones.append((crom['Floor'], number, 'Crom Cruach, weakened (P20)', 8, None, 'atwiki dt2-2 page 271 (Hell\'s Hole): '
                      'the fire of the quest, used at the crater, removes the zone that strengthens water for the fight '
                      'and keeps the second form away; the stats of the boss do not change'))
    bosses = wb['PostgameBosses']
    head = [c.value for c in bosses[1]]
    says = head.index('What the guide says')
    for row in bosses.iter_rows(min_row=2):
        note = NOTES.get(row[0].value)
        if note and note not in str(row[says].value or ''):
            row[says].value = (str(row[says].value) + ' | ' if row[says].value else '') + note
    wb.save(setup_mod.SETUP)
    print(added, 'rows of preparation for', len(late), 'checkpoints from', FIRST)


if __name__ == '__main__':
    main()
