"""Easy fixes of the audit of 2026-10-17 in data_tables/sim_setup.xlsx: M-02, Shadow Sew gets a floor of its own (Settings
ailment_min_chance_123: its roll is a fixed chance, 30 - 40%, under the general floor of 50%, so the model never cast it); M-03, a reading
for 299 Song Mastery (sheet AuditReadings)."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'sim_setup.xlsx')
wb = openpyxl.load_workbook(PATH)

ws = wb['Settings']
head = [c.value for c in ws[1]]
keys = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
if 'ailment_min_chance_123' not in keys:
    row = ws.max_row + 1
    values = ['ailment_min_chance_123', 0.3,
              'Shadow Sew (effect 123) is cast when its chance is at least this (the floor of ailment_min_chance for this effect alone): its roll is a '
              'fixed 30 - 40% by level, no resistance, and only on a target of a lower level than the user (docs/FORMULAS.md), so the general floor of '
              '50% kept the model from ever casting it',
              'model choice 2026-10-17 (audit M-02): the lowest chance of the skill, so it is cast where the program lets it land']
    for c, v in enumerate(values[:len(head)], start=1):
        ws.cell(row, c).value = v

ws = wb['AuditReadings']
ids = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
if '299 Song Mastery' not in ids:
    row = ws.max_row + 1
    for c, v in enumerate(['299 Song Mastery', 'context',
                           'At C18 the songs of the maid role (a Bard, level 29) already keep the party at full HP and TP, so a larger song bonus '
                           'changes no fight there. It serves where HP or TP run short: the skill battery reads it niche (best cell C38 boss, '
                           'Mage-Heavy, fights 16% shorter; 2026-10-17). Unread since the maid role took the plan and policy of the diva '
                           '(2026-10-16).'], start=1):
        ws.cell(row, c).value = v
wb.save(PATH)
print('written')
