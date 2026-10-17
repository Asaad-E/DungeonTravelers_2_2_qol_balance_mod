"""Roles: the maid role gets a skill plan and the Leather Shield of the diva role (decided 2026-10-16). Her path ends in the Diva, and with her
old priority list (Self Session last, no shield) the Diva of that role won 28% at C40 in class-stages against 71% for the diva role. A plan
skips what the classes so far cannot learn, so the plan below keeps the Maid of the start (First Aid, Cook, Generous Heart, Maid Mastery)
and gives the Bard and the Diva the songs of the guide first (GameFAQs faqs/80834, Hentaicheg: the heal song maxed first, Self Session
level 1 as soon as possible)."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'sim_setup.xlsx')
PLAN = ('Replenishing Carol:max; Self Session:1; Song Mastery:max; First Aid:3; Cook:3; Generous Heart:3; Maid Mastery:3; Lightning March:3; '
        'Blissful Chorale:3; Maid Mastery:max; Generous Heart:max; Cook:max; First Aid:max; Lightning March:max; Concentration Prelude:max; '
        'Auto-Recovery:max; Fast Talk:max')
SOURCE = (' | 2026-10-16: the skill plan and the Leather Shield of the diva role (GameFAQs faqs/80834, Hentaicheg), with the Maid skills of '
          'the start kept: with the priority list alone the Diva of this role won 28% at C40 (class-stages) against 71%.')

wb = openpyxl.load_workbook(PATH)
ws = wb['Roles']
head = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(head)}
for r in range(2, ws.max_row + 1):
    if ws.cell(r, col['Role']).value == 'maid':
        assert ws.cell(r, col['Lv30 class']).value == 'Diva'
        ws.cell(r, col['Skill plan']).value = PLAN
        ws.cell(r, col['Off hand']).value = 'Leather Shield'
        ws.cell(r, col['Source']).value = str(ws.cell(r, col['Source']).value) + SOURCE
        break
else:
    raise SystemExit('no maid role')
wb.save(PATH)
print('maid role written')
