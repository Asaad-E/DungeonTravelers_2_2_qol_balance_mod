"""Criticisms S07 (the Witch so dominant she trivialises the story) opened again on 2026-10-16, after dtcli class-search: the Witch is
in 87% of the fastest parties of the final classes, in vanilla and on the variant alike. The model's Witch spammed Fireball with
her strong spells at level 1 (skill priority of the sheet Roles), so her real margin is measured again before any decision."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'balance_criteria.xlsx')
NOTE = (' 2026-10-16, opened again: class-search over every final-class party puts the Witch in 87% of the fastest 5% of the parties '
        '(vanilla 87%); the model leveled Fireball and left her strong spells at level 1, so her margin is measured again with the '
        'spells a player raises.')
wb = openpyxl.load_workbook(PATH)
ws = wb['Criticisms']
head = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(head)}
for r in range(2, ws.max_row + 1):
    if ws.cell(r, col['ID']).value == 'S07':
        ws.cell(r, col['Status']).value = 'open'
        why = ws.cell(r, col['Why']).value or ''
        if NOTE.strip() not in why:
            ws.cell(r, col['Why']).value = (why.rstrip() + NOTE).strip()
        print('S07 open; was:', why[:120])
wb.save(PATH)
