"""One-off edit (2026-10-13) of the sheet Criticisms: four rows closed on an argument get the figure read from the
data and a reason that holds (J16, J17, S18, S21), so that they are not opened again. Run once; kept as the record.
"""
import os

import openpyxl

PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables',
                    'balance_criteria.xlsx')
SET = {
    'J16': {'Vanilla_Result': 'The enemy with most HP of each boss fight, its five elemental resistances (data of the game, '
                              '2026-10-13): at 100 or more in all five in 1 of the 55 fights of the postgame and in none of '
                              'the 40 of the story; in three or more elements in 4 of the 55; at least one element below 100 '
                              'in 54 of the 55',
            'Verdict': 'false as stated: one postgame boss in 55 resists every element at 100',
            'Why': 'A resistance of 100 halves the damage and never voids it (docs/FORMULAS.md 2.4), and every postgame boss '
                   'but one leaves an element below 100 to strike. The wall the claim describes is one fight, not the postgame',
            'Decision': '2026-10-13: closed with the count; no change'},
    'J17': {'Status': 'not a problem',
            'Why': 'True from C26 on, and meant: one enemy action can fell the frailest member from full HP in most battles of '
                   'the last third of the story, with Cover and the guards played. It is the reason the Valkyrie and the Paladin '
                   'have a place (rows J04, F07), and the parties without them still win those fights (measures/WIN_RATES.md: '
                   'Dual-Support 91% and Mage-Heavy 96% over the 17 bosses of the battery on the variant). Taking the spikes '
                   'away would take the guard classes with them',
            'Decision': '2026-10-13: measured; the claim holds and nothing is changed: it is the back and forth of a boss fight '
                        '(Regression R10)'},
    'S18': {'Vanilla_Result': 'The bag starts at 30 slots and gains 5 at quest ranks 8, 11, 15 and 20, up to 50 (QuestRankData, '
                              'reward kind 5; 200, 340, 560 and 1,000 trust points)',
            'Verdict': 'partly true: 30 slots until quest rank 8, 50 at the end',
            'Why': 'The 30 is a constant of the program and the four enlargements are data of the quest ranks; a player who does '
                   'quests has 35 slots by the middle of the story. Choosing what to carry back is part of an expedition, and '
                   'the variant leaves it: a larger bag is a patch of the program for a matter of comfort',
            'Decision': '2026-10-13: closed with the figure; no change'},
    'S21': {'Vanilla_Result': 'Of the 194 fixed battles on floors with random battles, 41 hold an enemy 20 levels or more above '
                              'the common enemies of their floor, and 31 one 30 or more above (data of the game, 2026-10-13): a '
                              'boss of level 98 on a floor of the story whose common enemies are of level 27. 34 of the 41 are '
                              'checkpoints of the model, nearly all of them of the postgame',
            'Verdict': 'true: a fifth of the fixed battles stand 20 levels or more above their floor',
            'Why': 'They are the bosses of the postgame, placed on the floors of the story: the game sends the party back to old '
                   'floors for them, and nothing on the map tells their level. The variant does not move them (R03: the jumps of '
                   'level of the game are kept); what it changes is how fast a party that comes back catches up (the EXP factors '
                   'by level gap)',
            'Decision': '2026-10-13: closed with the count; the bosses stay where the game has them'},
}


def main():
    wb = openpyxl.load_workbook(PATH)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    done = 0
    for row in ws.iter_rows(min_row=2):
        for column, value in SET.get(row[0].value, {}).items():
            row[head.index(column)].value = value
            done += 1
    assert done == sum(len(v) for v in SET.values())
    wb.save(PATH)
    print(done, 'cells')


if __name__ == '__main__':
    main()
