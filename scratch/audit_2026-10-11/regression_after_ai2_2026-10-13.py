"""One-off edit (2026-10-13): the sheet Regression after the pass with the fourth habit of the model (no punished blow
that is worth nothing). Run once; kept as the record of the edit.
"""
import os

import openpyxl

PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables',
                    'balance_criteria.xlsx')
ROWS = {
    'R01': ('2026-10-13, after the model was taught four habits of a player (sim/policy.py). measures/WIN_RATES.md (200 '
            'battles): the three parties that work in the game end level, Meta 93%, Dual-Support 91%, Mage-Heavy 96% of wins '
            'over the 17 bosses of the battery (88%, 90%, 92% in the game); two cells of them gain 15 points or more (Meta at '
            'C38, Mage-Heavy at C36). The weakest party of the game (Anti-Meta, 46% on the mean) reaches 78%. CLASS_STAGES.md: '
            '74 of 76 rows within the band. Limit: the band has no top (row J11, left). The Witch is the case: far ahead at C30 '
            'alone, a boss that resists physical damage at 90 - 100 (row S07, decided 2026-10-13: not a problem there)',
            'met in what is measured'),
    'R02': ('2026-10-13. PARTY_BATTERY.md: 70 of 85 pairs at or above the floor in vanilla, 77 of 85 on the variant; no pair '
            'that met its floor in vanilla lost it. CLASS_STAGES.md: the two slow rows (Assassin and Treasure Hunter of the '
            'treasure hunter role) were slow before the changes (row J26)', 'met'),
    'R08': ('2026-10-13, after the model was taught four habits of a player. DIFFICULTY_CURVE.md (40 battles a boss): 34 of 44 '
            'checkpoints within the windows of vanilla, 10 easier, 0 harder. Nine are accepted (C14, C26, C29, C32, C37, C38 on '
            '2026-10-11; C17, C19 and C40 on 2026-10-13). C30 is within: the reference party wins 98% of it in the game and on '
            'the variant in 39 and 40 turns (200 battles; what the model saw as hard there was its own members striking a boss '
            'with Dark Aura). C15 is new with the pass: 32% shorter (16.1 > 11.0, wins 99% both, 200 battles)',
            'met but for 10 bosses easier: 9 accepted, C15 not decided'),
}
wb = openpyxl.load_workbook(PATH)
ws = wb['Regression']
head = [c.value for c in ws[1]]
result = next(i for i, h in enumerate(head) if str(h).startswith('Result'))
for row in ws.iter_rows(min_row=2):
    if row[0].value in ROWS:
        row[result].value, row[head.index('Verdict')].value = ROWS[row[0].value]
wb.save(PATH)
print(len(ROWS), 'rows')
