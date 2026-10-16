"""Criticisms S07 closed again on 2026-10-16 with the figures of the change (rows V062 - V185)."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'balance_criteria.xlsx')
RESULT = ('2026-10-16: the Witch INT x0.9 (growth row 18) and the five basic spells at levels 6 - 10 x0.85 (V062 - V185). dtcli '
          'class-search final C30, C34, C38, C40 boss, 3,000 parties a boss: in vanilla she is in 97% of the fastest 5% of the parties '
          '(the next class 65%); now 75% (the next 74%), her parties win 49.9% (median of the other classes 52.4%).')
wb = openpyxl.load_workbook(PATH)
ws = wb['Criticisms']
head = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(head)}
for r in range(2, ws.max_row + 1):
    if ws.cell(r, col['ID']).value == 'S07':
        ws.cell(r, col['Status']).value = 'closed'
        ws.cell(r, col['Balance_Result']).value = RESULT
wb.save(PATH)
print('S07 closed')
