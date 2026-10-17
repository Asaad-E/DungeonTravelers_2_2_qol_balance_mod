"""Criticisms J13 and F94: the figures of dtcli mercenary-battery written in the morning of 2026-10-17 set each mercenary in the place of
the summon and at the stage of the party's level; the program fields both and offers the stage its story flags open (measures/
MERCENARY_RESEARCH.md, checked in the program). That text is replaced by the figures of the corrected measure."""
import json
import os
import openpyxl

ROOT = os.path.join(os.path.dirname(__file__), '..', '..')
PATH = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
rows = json.load(open(os.path.join(ROOT, 'docs', 'generated', 'mercenary_battery', 'part_1.json'), encoding='utf-8'))
count = {}
for r in rows:
    count[r['verdict']] = count.get(r['verdict'], 0) + 1
up = sum(1 for r in rows if r['mean_wins'] > 0)
TEXT = ('2026-10-17, dtcli mercenary-battery (docs/generated/MERCENARY_BATTERY.md): each mercenary hired beside the summon of the grand '
        'sealbook (the program builds both, units 3000 and 3001), at the stage the tavern offers at that point of the story (sheet '
        'MercenaryUnlocks, from the flag at 0xB4 of each stage; measures/MERCENARY_RESEARCH.md), in the reference party and the five battery '
        'parties, at the bosses of the curve where she is offered and on rows of random battles; 20 runs a cell, the 3 best at 200: '
        '%d mercenaries, %s; %d of them raise the wins of the party on the mean. A guest ally has no HP and is never a target.'
        % (len(rows), ', '.join('%d %s' % (n, v) for v, n in sorted(count.items())), up))
wb = openpyxl.load_workbook(PATH)
ws = wb['Criticisms']
head = [c.value for c in ws[1]]
col = {n: i + 1 for i, n in enumerate(head)}
for r in range(2, ws.max_row + 1):
    if ws.cell(r, col['ID']).value in ('J13', 'F94'):
        cell = ws.cell(r, col['Balance_Result'])
        v = str(cell.value or '')
        i = v.find('2026-10-17, dtcli mercenary-battery')
        v = (v[:i].rstrip() if i >= 0 else v.rstrip())
        cell.value = (v + ' ' if v else '') + TEXT
wb.save(PATH)
print(TEXT)
