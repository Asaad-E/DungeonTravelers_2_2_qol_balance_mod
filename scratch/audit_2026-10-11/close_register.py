"""Closes the rows of the sheet Criticisms that were reviewed on 2026-10-11: the state of each and, in the column
Balance_Result, the figure it is closed on. The figures are those of the generated reports of that day
(CLASS_STAGES.md, SKILL_WORTH.md, PARTY_BATTERY.md, BALANCE_BUILD.md). Rows not named here stay as they are.

    python close_register.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
WORKBOOK = os.path.join(PROJECT, 'data_tables', 'balance_criteria.xlsx')
DAY = '2026-10-11'
BUILT = 'built and read back (BALANCE_BUILD.md, 1,382 of 1,382 rows)'
CLASSES = 'CLASS_STAGES.md'
WORTH = 'SKILL_WORTH.md (78 of 84 active skills at 60% or more of the best strike of their class)'
USER_KEEPS = f'{DAY}, decided: the change stays, it is justified in its row'

CLOSED = {
    'J03': f'C001 - C168, the TP cost of 22 damage skills, {BUILT}',
    'J05': f'R rows, Auto-Refresh 3 - 9 to 4 - 11 TP per turn, {BUILT}; every role spends 15% of its TP on an ally under 15%',
    'J06': f'W rows, four bow skills open to guns, {BUILT}; shooter within the band at the four stages (0.85, 0.89, 0.88, 1.09; {CLASSES})',
    'J07': f'W rows, two-handed base ATK -10%, {BUILT}',
    'J10': 'as J03',
    'J14': f'M rows, top single spells chant at 0.85, {BUILT}',
    'J18': f'D rows, {BUILT}; soul summoner within the band at the four stages (1.12, 0.98, 0.87, 1.02; {CLASSES})',
    'J23': f'T066 - T075, {BUILT}; Sage wins 12% under the median support at the advanced stage and 2% late, within the band ({CLASSES})',
    'J25': f'Q001 - Q015, {BUILT}; Magical Princess 1.06 advanced, 1.01 late, within the band ({CLASSES})',
    'J26': f'T041 - T050, {BUILT}; Treasure Hunter 1.21 advanced (slow), 1.00 late; the user, {DAY}: what is left is not a problem',
    'J27': f'K001 - K025 and W045 (Tantou forge rate x2), {BUILT}; Kunoichi 1.08 advanced, 1.01 late, within the band ({CLASSES})',
    'J29': f'T051 - T065, {BUILT}; Mistress 0.95 advanced, 1.02 late, within the band ({CLASSES})',
    'J30': f'Y001 - Y011 (Danger Shoot +0.10, from 20 TP) and W046 (Dagger 20 ATK), {BUILT}; Spieler 1.11 - 1.12 at the base stage, '
           f'Trickster 1.08 - 1.13, within the band ({CLASSES})',
    'J31': f'A rows, every story and postgame boss immune to turn delay, {BUILT}',
    'J33': f'M rows, {BUILT}; Witch 0.59 advanced, 0.94 late ({CLASSES})',
    'J37': 'ExePatches set balance in both program files (12 rows, 17 bytes each); seen working in the game',
    'F03': 'as J29',
    'F04': 'as J27',
    'F05': 'as J25',
    'F08': f'W rows, two-handed base ATK -10%, {BUILT}',
    'F09': f'W rows, whip upgrade rate 2.0 and Whip Mastery x2, {BUILT}; Hard Whipping worth 100% ({WORTH})',
    'F10': 'as J06',
    'F11': f'W rows, Double Attack, Triple Attack and Barrage +25%, {BUILT}',
    'F12': f'E rows, Dark Aura of common enemies 4 stats instead of 6, {BUILT}',
    'F13': f'M rows, magic defence buffs and debuffs x1.5, {BUILT}',
    'F15': f'A001 - A239, {BUILT}; SKILL_AILMENTS.md',
    'F17': f'T031 - T040 and Y rows, {BUILT}; Trickster 1.08 - 1.13, Joker 1.00 advanced and 0.96 late, within the band ({CLASSES})',
    'F23': f'A rows, Magnificent Steps 2 - 4% and Warning Shot 1 - 3%, {BUILT}',
    'F24': f'looked at skill by skill in T06: {WORTH}; the chants and speeds that held a skill under the line are in the M, K and X rows',
    'F25': 'as J37',
    'F26': 'as J05',
    'F31': f'X rows, {BUILT}; Blood Drain worth 76% ({WORTH})',
    'F32': f'X rows, {BUILT}; Soul Drain worth 78%, Soul Absorb 68% ({WORTH})',
    'F34': f'X088 and its levels, Blood Rage 43 to 32 TP, {BUILT}. {USER_KEEPS}',
    'F35': f'X rows, Brave and Protect x2, {BUILT}. {USER_KEEPS}',
    'F44': f'X rows, Illusion 3 / 4 / 5 hits, {BUILT}. {USER_KEEPS}',
    'F45': f'Q rows, {BUILT}; Magical Spark worth 100% ({WORTH})',
    'F46': f'Q001 and its levels, Magical Change 52 to 39 TP, {BUILT}',
    'F48': f'T066 - T070, {BUILT}',
    'F57': f'T056 - T065, {BUILT}; Hard Whipping worth 100% ({WORTH})',
    'F59': f'K rows, {BUILT}; Ninja Piledriver worth 100% ({WORTH})',
    'F60': f'K rows (TP -25%) and X rows (chant 0.75 at every level), {BUILT}. {USER_KEEPS}',
    'F61': f'X rows, Shadow Sew -25% TP, {BUILT}. {USER_KEEPS}',
    'F62': f'X rows, {BUILT}; the scrolls are worth 48 - 60% beside Ninja Piledriver, which the Tantou (W045) raised ({WORTH}). '
           f'{DAY}, decided: they stay as they are, they are debuffs late (J27)',
    'F68': f'T051 - T055, {BUILT}; Enticing Love worth 100% ({WORTH})',
    'F69': f'M rows, Carrot and Stick x1.5, {BUILT}',
    'F71': f'D rows, {BUILT}; dolls within the band at the four stages (1.12, 0.97, 1.04, 1.13; {CLASSES})',
    'F72': f'X rows, {BUILT}; Life Bringer worth 48% beside Showtime ({WORTH}). {DAY}, decided: it stays as it is',
    'F74': f'Limit Field TP x0.55, {BUILT}',
    'F75': f'S rows, {BUILT}; Graceful Rose worth 100% ({WORTH})',
    'F78': f'W rows, +25%, {BUILT}',
    'F79': f'W rows, +25%, {BUILT}',
    'F88': f'K001 - K025, T001 - T030, T051 - T065, W045, {BUILT}; Kunoichi and Mistress within the band at the four stages ({CLASSES})',
    'F90': f'P024, P079 - P084 (seven pieces under 2% raised to 2%) and P085 - P088 (Item Steal open on Odin, Myrddin, Asmodeus and '
           f'Carmilla), {BUILT}; their pieces come out of 6 - 9% of the steals that succeed; a talk of Euthys says it in the game',
    'F93': 'as U01',
    'U01': 'ExePatches set balance in both program files; CATCH_UP_OPTIONS.md, row B4; seen working in the game',
    'S01': 'as U01',
    'S02': f'E rows, common enemies of level 80 or more lose 10 points of physical resistance, {BUILT}',
    'S03': f'E rows, four enemy passives capped at the party top, {BUILT}',
    'S05': f'A rows, {BUILT}; SKILL_AILMENTS.md',
    'S06': f'A rows, debuffs x1.5, {BUILT}',
    'S11': 'as U01',
    'S12': f'B001 - B040, +1 skill point per level from 31 to 70, {BUILT}',
    'S20': f'41 P rows of the table quests, gold reward x1.40, {BUILT}; EXP for the members behind comes from the factors B4 (U01)',
    'U02': f'D001 - D243, {BUILT}; dolls 1.12, 0.97, 1.04, 1.13 and soul summoner 1.12, 0.98, 0.87, 1.02 at the four stages, '
           f'within the band ({CLASSES})',
}
DECIDED = {
    'F34': f'{DAY}: approved, the change stays',
    'F35': f'{DAY}: approved, the change stays',
    'F44': f'{DAY}: approved, the change stays',
    'F60': f'{DAY}: approved, the change stays',
    'F61': f'{DAY}: approved, the change stays',
    'F90': f'{DAY}: approved (seven drop weights to the 2% floor; Item Steal open on the four immune commons)',
    'S20': f'{DAY}: approved (quest gold x1.40)',
}
NOT_A_PROBLEM = ('J11', 'J32', 'F18', 'F19', 'F20', 'F21', 'F27', 'F28', 'F29', 'F30', 'F33', 'F36', 'F37', 'F38', 'F39',
                 'F40', 'F41', 'F43', 'F50', 'F51', 'F53', 'F58', 'F63', 'F64', 'F65', 'F66', 'F67', 'F70', 'F76', 'F77',
                 'F80', 'F82', 'S10')
OUT = ('F49', 'F54', 'F55', 'F56', 'F83', 'F84')


def main():
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    col = {name: head.index(name) for name in head}
    seen = set()
    for row in ws.iter_rows(min_row=2):
        rid = row[0].value
        if rid in CLOSED:
            state, figure = 'closed', CLOSED[rid]
        elif rid in NOT_A_PROBLEM:
            state, figure = 'not a problem', None      # the figure of these rows is their verdict
        elif rid in OUT:
            state, figure = 'out of the balance', None
        else:
            continue
        assert row[col['Status']].value in ('open', state), (rid, row[col['Status']].value)
        row[col['Status']].value = state
        if figure:
            row[col['Balance_Result']].value = figure
        if rid in DECIDED:
            old = row[col['Decision']].value
            row[col['Decision']].value = DECIDED[rid] if not old or old == 'proposed' else f'{old}; {DECIDED[rid]}'
        seen.add(rid)
    missing = (set(CLOSED) | set(NOT_A_PROBLEM) | set(OUT)) - seen
    assert not missing, missing
    wb.save(WORKBOOK)
    print(len(CLOSED), 'closed,', len(NOT_A_PROBLEM), 'not a problem,', len(OUT), 'out of the balance')


if __name__ == '__main__':
    main()
