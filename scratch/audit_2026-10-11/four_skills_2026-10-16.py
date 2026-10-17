"""Four weak skills of the battery (decided 2026-10-16), measured by dtcli skill-battery with variants (every rule of use, 20 runs a cell,
the 3 best at 200). Each change is applied on the value the variant has now (its row if there is one, else vanilla):
Panic Voice TP x0.8 (dead -> good, 29% of the cells; a small raise: the Diva is among the best classes, and its stun is left as it is so
the Mistress keeps her part); Accelerator 5 turns and its speed x1.67, up to +25% (hurts -> niche); Blood Rage TP x0.5 and +5 turns
(hurts -> niche); Illusion TP x0.5 and its chant speed x1.6 (dead -> niche)."""
import json
import os
import openpyxl

ROOT = os.path.join(os.path.dirname(__file__), '..', '..')
PATH = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
DECIDED = '2026-10-16'
WHY = {
    'Panic Voice': ('Panic Voice was dead in the skill battery (2026-10-16): 96 TP, the Diva hardly cast it. TP x0.8 made it good (29% of the '
                    'cells); a small raise, the Diva being among the best classes, and the stun left so the Mistress keeps her part.'),
    'Accelerator': 'Accelerator still hurt in the skill battery (2026-10-16): 5 turns and up to +25% speed made it niche.',
    'Blood Rage': ('Blood Rage hurt in the skill battery (2026-10-16): it raises the max HP, not the HP, so it needs a heal behind it; TP x0.5 '
                   'and 5 more turns made it niche.'),
    'Illusion': ('Illusion was dead in the skill battery (2026-10-16): its clones guard the caster, who stands at the back; TP x0.5 and a '
                 'chant x1.6 faster made it niche.'),
}
SERVES = {'Panic Voice': 'T06', 'Accelerator': 'T06', 'Blood Rage': 'T06', 'Illusion': 'T06'}
# (record, field, array of the snapshot or effect id, op, by, digits, name)
PLAN = [
    ('Magic 225 Panic Voice', 'TP cost', 'tp', '*', 0.8, 0, 'Panic Voice'),
    ('Magic 197 Accelerator', 'turns', 'turns', '+', 1, 0, 'Accelerator'),
    ('Magic 197 Accelerator', 'effect 168', 168, '*', 1.67, 2, 'Accelerator'),
    ('Magic 165 Blood Rage', 'TP cost', 'tp', '*', 0.5, 0, 'Blood Rage'),
    ('Magic 165 Blood Rage', 'turns', 'turns', '+', 5, 0, 'Blood Rage'),
    ('Magic 200 Illusion', 'TP cost', 'tp', '*', 0.5, 0, 'Illusion'),
    ('Magic 200 Illusion', 'chant speed factor', 'cast', '*', 1.6, 2, 'Illusion'),
]
snap = {s['id']: s for s in json.load(open(os.path.join(ROOT, 'build', 'engine_snapshot.json'), encoding='utf-8'))['sides']['vanilla']['skills']}

wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
head = [c.value for c in ws[1]]
assert head[:12] == ['ID', 'Table', 'Record', 'Field', 'Level', 'Vanilla', 'New', 'Lever', 'Serves', 'Why', 'Decided', 'State'], head
existing = {(ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 5).value): r for r in range(2, ws.max_row + 1) if ws.cell(r, 12).value == 'built'}
ids = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
nxt = 1 + max(int(k[1:]) for k in ids if k and str(k).startswith('V'))
row = ws.max_row + 1
scaled = added = 0


def apply(value, op, by, digits):
    out = value * by if op == '*' else value + by
    return round(out) if digits == 0 else round(out, digits)


for record, field, key, op, by, digits, name in PLAN:
    sid = int(record.split()[1])
    values = snap[sid][key] if isinstance(key, str) else next(e[2] for e in snap[sid]['effects'] if e[0] == key)
    for level, old in enumerate(values, start=1):
        at = existing.get((record, field, level))
        if at:
            ws.cell(at, 7).value = apply(ws.cell(at, 7).value, op, by, digits)
            ws.cell(at, 10).value = str(ws.cell(at, 10).value) + ' ' + WHY[name]
            scaled += 1
            continue
        vals = ('V%03d' % nxt, 'skills', record, field, level, old, apply(old, op, by, digits), 'L12', SERVES[name], WHY[name], DECIDED, 'built')
        for c, value in enumerate(vals, start=1):
            ws.cell(row, c).value = value
        row += 1
        nxt += 1
        added += 1
wb.save(PATH)
print('scaled', scaled, 'added', added)
