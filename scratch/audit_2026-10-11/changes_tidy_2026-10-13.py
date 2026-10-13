"""One-off edit (2026-10-13), after the second audit: rows of balance_changes.xlsx that change nothing are dropped,
reasons that quote a value since changed are brought up to date, and the gold of quests serves the row of the
register it answers (S19, the early economy, closed with it). Run once; kept as the record of the edit.
"""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CHANGES = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
CRITERIA = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
SAME = ' | 2026-10-13 : dropped: the value of the game and the new one are the same at this level, the row changed nothing'
STALE = {
    'X06': (' | 2026-10-13 : the party buffs it is compared with are now 0.31 at their top level (rows T066 - T075)',
            lambda rid: 'X061' <= rid <= 'X070'),
    'G02': (' | 2026-10-13 : the HP of this boss is now x1.265 of the game (rows G110, G111), not x1.30; its ATK and INT stay '
            'at x0.92, checked that day (measures/BOSS_ROWS_CHECK.md: 98% of wins in 39.6 turns against 98% in 38.9 in the game)',
            lambda rid: rid in ('G029', 'G030', 'G032', 'G033')),
}


def same(a, b):
    try:
        return abs(float(a) - float(b)) < 1e-9
    except (TypeError, ValueError):
        return a == b


def main():
    wb = openpyxl.load_workbook(CHANGES)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    col = {k: head.index(k) for k in ('ID', 'Table', 'Vanilla', 'New', 'Why', 'State', 'Serves')}
    dropped = moved = noted = 0
    for row in ws.iter_rows(min_row=2):
        rid = str(row[col['ID']].value)
        if row[col['State']].value != 'built':
            continue
        if same(row[col['Vanilla']].value, row[col['New']].value):
            row[col['State']].value = 'dropped'
            row[col['Why']].value = str(row[col['Why']].value) + SAME
            dropped += 1
            continue
        if row[col['Table']].value == 'quests' and row[col['Serves']].value == 'S20':
            row[col['Serves']].value = 'S19'
            moved += 1
        for note, mine in STALE.values():
            if mine(rid):
                row[col['Why']].value = str(row[col['Why']].value) + note
                noted += 1
    wb.save(CHANGES)

    wb = openpyxl.load_workbook(CRITERIA)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == 'S19':
            row[head.index('Status')].value = 'closed'
            row[head.index('Balance_Result')].value = ('The gold of the 41 quests that pay gold x1.40 (P rows of the table '
                                                       'quests): quests are where the early gold comes from')
            row[head.index('Decision')].value = '2026-10-11: quest gold x1.40 approved; 2026-10-13: it is this row those rows answer'
        if row[0].value == 'S20':
            row[head.index('Balance_Result')].value = ('The EXP of a quest is not changed. The gold of quests was raised for '
                                                       'row S19; the members who are behind catch up through the EXP factors '
                                                       'by level gap (sheet ExePatches)')
    wb.save(CRITERIA)
    print(dropped, 'rows dropped,', moved, 'moved to S19,', noted, 'reasons brought up to date')


if __name__ == '__main__':
    main()
