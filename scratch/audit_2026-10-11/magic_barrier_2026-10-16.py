"""One-off edit: Magic Barrier (Technique 58, effect 183) from 25% at level 1 to 45% at level 10 (decided 2026-10-16).
Rows S022 - S026 held levels 1 - 5 only (25% - 40%), so levels 6 - 10 kept the game's 28% - 36% and the skill fell from
40% to 28% at level 6; rows S043 - S047 give levels 6 - 10. Lever L12 (a value of a skill per level)."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
VANILLA = [-0.18, -0.2, -0.22, -0.24, -0.26, -0.28, -0.3, -0.32, -0.34, -0.36]
NEW = [-0.25, -0.27, -0.29, -0.32, -0.34, -0.36, -0.38, -0.4, -0.43, -0.45]
WHY = ('Party magic damage reduction buff: 25% at level 1 to 45% at level 10, every level a little more (decided '
       '2026-10-16). Rows S022 - S026 held levels 1 - 5 alone (25% - 40%), so levels 6 - 10 kept the game\'s 28% - 36% and '
       'the skill fell from 40% to 28% at level 6')
IDS = {1: 'S022', 2: 'S023', 3: 'S024', 4: 'S025', 5: 'S026', 6: 'S043', 7: 'S044', 8: 'S045', 9: 'S046', 10: 'S047'}

wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
by_id = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
template = [ws.cell(by_id['S022'], c).value for c in range(1, ws.max_column + 1)]
for level, rid in IDS.items():
    if rid in by_id:
        r = by_id[rid]
    else:
        ws.append(list(template))
        r = ws.max_row
        ws.cell(r, 1).value = rid
        ws.cell(r, 5).value = level
    assert ws.cell(r, 3).value == 'Technique 58 Magic Barrier' and ws.cell(r, 5).value == level
    ws.cell(r, 6).value = VANILLA[level - 1]
    ws.cell(r, 7).value = NEW[level - 1]
    ws.cell(r, 8).value = 'L12'
    ws.cell(r, 10).value = WHY
    ws.cell(r, 11).value = '2026-10-16'
    ws.cell(r, 12).value = 'built'
wb.save(PATH)
print('rows', len(IDS))
