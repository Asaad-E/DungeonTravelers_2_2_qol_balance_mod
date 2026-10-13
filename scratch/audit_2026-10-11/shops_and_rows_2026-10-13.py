"""One-off edit (2026-10-13): docs/FORMULAS.md 14 says what the five special chances of a zone start (the wandering
shops), and the sheets Criticisms and Regression of balance_criteria.xlsx take the measures of that day (Sacred Hymn,
Ecstasy, meals, the two bosses accepted). Run once; kept as the record of the edit.
"""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FORMULAS = os.path.join(ROOT, 'docs', 'FORMULAS.md')
CRITERIA = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
OLD = """After that roll the same step rolls five special chances of the zone (`zone[0x20]` to `zone[0x30]`), each one only
while its game flag (601 to 605) is on and it has not fired yet, and only once the dungeon counters allow it
(`0x556380`: a counter of at least 5 and at least 50 steps, 25 with a piece of enchantment 259); the same four
enchantments shift them. What each of the five starts was not followed.
"""
NEW = """After that roll the same step rolls five more chances of the zone, in this order, and the first that falls under
its chance fires (`0x4bcb60`): each makes a wandering shop appear, the class the program builds for it named by its
type information (`0x4b28b7`, jump table `0x4b36cc`).

| Chance | Field of the zone | Shop | Game flag | Vanilla values |
|---|---|---|---|---|
| 1 | `0x20` | Peddler (`CShopPeddler`) | 601 | 3 - 10 |
| 2 | `0x24` | Blacksmith (`CShopBlacksmith`) | 602 | 2 - 10 |
| 3 | `0x28` | Ice cream shop (`CShopIce`) | 603 | 0, or 5 - 10 |
| 4 | `0x2C` | Ramen shop (`CShopRamen`) | 604 | 0, or 6 - 12 |
| 5 | `0x30` | Dark peddler (`CShopDarkPeddler`) | 605 | 0 in every zone |

A shop is rolled only when all of this holds (`0x4bc8a1 - 0x4bca94`, `0x556380`):

* its game flag is on (the story has opened that shop);
* it has not appeared yet: a byte of the dungeon for each shop (`+0x781` to `+0x785`), set when it appears
  (`0x4bcbad`) and cleared with the other counters of the dungeon (`0x4b613f - 0x4b6175`);
* the counter of battles of the dungeon (`+0x770`, one more at the start of each battle, `0x4b2730`) is at least 5;
* at least 50 steps were taken since the last shop appeared (`+0x200`, one more a step at `0x4b92b3`, set to 0 when a
  shop appears, `0x4b289b`); 25 steps with a piece of enchantment 259.

Then `chance + shift of the four enchantments above`, at least 1, against a number below 100. The battle is rolled
first and the peddler before the blacksmith, and a shop that appears sets the 50 steps back for the others, so the
blacksmith of a zone at 5 with a peddler at 6 and battles at 6 is the next shop a little under half of the times.
An item sets a byte (`+0x7A4`) that puts the ice cream and ramen chances of the zone at 100 for the next step
(`0x4bcaa9 - 0x4bcb52`). The five chances are ints of the floor file, as the chance of a battle: data, not code.
"""
SOURCE = ' (measures/one_skill_trial.py, 150 battles, 2026-10-13)'
SET = {
    'F65': {'Verdict': 'partly true: it serves at the last boss alone',
            'Vanilla_Result': 'The Maid of the reference party without the song and with it at its top level: wins 84% > 91% '
                              'and HP left 82% > 88% at C40, 93% > 95% at C26, no change at C30, C36 and C38; the length of '
                              'the fights does not move' + SOURCE,
            'Why': 'A song of RES is worth its slot where the boss casts: the last one. Elsewhere Carol and March are better '
                   'uses of the two slots',
            'Decision': '2026-10-13: measured; no change'},
    'F70': {'Verdict': 'true, for another reason: it stacks with War Cry but adds nothing that can be measured',
            'Vanilla_Result': 'The passive gives ATK and DEX +10% to +15% for 3 turns, once a battle, the first time its '
                              'bearer falls below half her HP (EffectRules 312); its statuses (41, 45) are not those of War '
                              'Cry (161), so they add. The Mistress without it and with it at its top level: wins within 2 '
                              'points and length within 2% at C26, C30, C36, C38 and C40' + SOURCE,
            'Why': 'Three turns once a battle are too little at that size. At 30% to 50% the fights of C30 and C40 are 5% to 8% '
                   'shorter; at 60% to 100%, 8% to 10% (the same measure with the values changed)',
            'Decision': '2026-10-13: measured; a dead passive as the game has it. Whether to raise it waits for a decision'},
    'J36': {'Why': 'A quarter of the meals have nothing that could run out. For the other twelve the buff is the skill itself, '
                   'performed at a fixed level by the ordinary skill action (0x50b030, 0x41d2b0); what ends it outside a battle '
                   'was looked for on 2026-10-13 and not found (the list emptied every 100 steps at 0x4b92ed holds the shops '
                   'that appeared, not the buffs)',
            'Decision': '2026-10-13: read in part; the length of the buff is still not known'},
}
R08 = ('2026-10-13. DIFFICULTY_CURVE.md (40 battles a boss): 33 of 44 checkpoints within the windows of vanilla, 11 easier, 0 '
       'harder. Six were accepted on 2026-10-11 (C14, C26, C29, C32, C37, C38); C30 is back within with its HP x1.15. The five '
       'that crossed since then, with 200 battles (measures/curve_recheck.py): C15, C28 and C39 are within, it was chance; C17 '
       'is 28% shorter (49.1 > 35.4, wins 100% both) and C40 is won more (85% > 97%). Both accepted on 2026-10-13',
       'met; 8 bosses easier, accepted')


def main():
    with open(FORMULAS, encoding='utf-8') as f:
        text = f.read()
    assert text.count(OLD) == 1
    with open(FORMULAS, 'w', encoding='utf-8', newline='\n') as f:
        f.write(text.replace(OLD, NEW))
    wb = openpyxl.load_workbook(CRITERIA)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    done = 0
    for row in ws.iter_rows(min_row=2):
        for column, value in SET.get(row[0].value, {}).items():
            row[head.index(column)].value = value
            done += 1
    assert done == sum(len(v) for v in SET.values())
    ws = wb['Regression']
    head = [c.value for c in ws[1]]
    result = next(i for i, h in enumerate(head) if str(h).startswith('Result'))
    for row in ws.iter_rows(min_row=2):
        if row[0].value == 'R08':
            row[result].value, row[head.index('Verdict')].value = R08
    wb.save(CRITERIA)
    print(done, 'cells')


if __name__ == '__main__':
    main()
