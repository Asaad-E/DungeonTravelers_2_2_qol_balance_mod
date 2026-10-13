"""One-off edit of balance_changes.xlsx (2026-10-13): Alluring Samba back to 40% at its top level, and the bosses of
C30 and C38 with 15% more HP than the variant had. Run once; kept as the record of the edit.
"""
import os

import openpyxl

PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables',
                    'balance_changes.xlsx')
SAMBA = dict(zip(('A094', 'A095', 'A096', 'A097', 'A098', 'A099', 'A100', 'A101', 'A102', 'A103'),
                 (-0.09, -0.12, -0.16, -0.19, -0.23, -0.26, -0.3, -0.33, -0.37, -0.4)))
SAMBA_WHY = (' | 2026-10-13 : back to 40% at its top level. The program takes the dance bonus of the performer away from '
             'a dance that lowers a stat (docs/FORMULAS.md 12.8); the model added it, so the trial of 2026-10-12 weighed '
             'the dance too light. With Dance Mastery and enchantments 60% is 77% in play, and with Armor Penetration the '
             'DEF of the enemy fell below zero. With the model corrected (measures/def_debuff_trial.py, 150 battles) the '
             'dance at 60% made the fights of C30 - C40 25 - 35% shorter than the usual dances, at the values of the game '
             '9 - 20%')
HP = {'G110': 10752, 'G111': 8602}
HP_WHY = (' | 2026-10-13 : HP x1.15 on top (x1.265 of the game). The party that was already the strongest in the game won '
          'far more here on the variant (measures/win_rates.py, 200 battles: Meta 66% in the game, 97% on the variant); '
          'with it Meta wins 86% and the parties that lag keep 64 - 72% (measures/boss_trim_trial.py). Putting the ATK and '
          'INT of the game back instead cost those parties too much (Anti-Meta-Plus 80% to 32%)')
NEW = ('G118', 'enemies', '950 Samael (C38)', 'hp', None, 18500, 21275, 'L01', 'R08, R10',
       'The party that was already the strongest in the game won far more here on the variant (measures/win_rates.py, 200 '
       'battles: Meta 67% in the game, 96% on the variant). HP x1.15: Meta 90%, Dual-Support 95%, Mage-Heavy 98%, and the '
       'parties that lag keep theirs but Juggernaut without a Diva (56% to 40%; measures/boss_trim_trial.py). x1.25 cost '
       'them more (Juggernaut-Plus 82% to 59%)', '2026-10-13', 'built')


def main():
    wb = openpyxl.load_workbook(PATH)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    new, why = head.index('New') + 1, head.index('Why') + 1
    done = 0
    for row in ws.iter_rows(min_row=2):
        rid = row[0].value
        if rid in SAMBA:
            ws.cell(row[0].row, new).value = SAMBA[rid]
            ws.cell(row[0].row, why).value += SAMBA_WHY
            done += 1
        elif rid in HP:
            ws.cell(row[0].row, new).value = HP[rid]
            ws.cell(row[0].row, why).value += HP_WHY
            done += 1
    assert done == 12 and not any(r[0].value == NEW[0] for r in ws.iter_rows(min_row=2))
    ws.append(NEW)
    wb.save(PATH)
    print(done + 1, 'rows')


if __name__ == '__main__':
    main()
