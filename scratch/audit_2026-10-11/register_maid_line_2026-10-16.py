"""Criticisms F01, J05, F26 (the Maid line as the TP source), J11 (the magic curve) and S07 (the Witch) given the figures of 2026-10-16
(dtcli class-search and class-stages on the variant of commit baab113)."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'balance_criteria.xlsx')
NO_MAID = ('2026-10-16, dtcli class-search preset=no_maid_line (every final-class party without a Diva, Mistress or Etoile, 1,000 a boss; '
           'measures/CLASS_SEARCH_FINAL_REG_NOMAID.md, _REG_NOMAID24.md): with the 8 fruit of the sheets the best such party wins C30, C34 '
           'and C38 100% and C40 20%; with 24 fruit in the bag all four 100% (C40: Bishop, Papillon, two Soul Summoners, Valkyrie, 18.2 '
           'turns). The Maid line stays the TP source by design: a party without her gets through the story with a bag of fruit.')
SET = {
    'F01': {'add': NO_MAID},
    'J05': {'add': NO_MAID},
    'F26': {'set': 'as J05. ' + NO_MAID},
    'J11': {'set': ('2026-10-16, docs/generated/CLASS_STAGES.md: the mage role against the median attacker, length of the won fights / '
                    'share of the party damage: base (Magic User) 0.77 / 38 - 50%, intermediate (Sorceress) 0.91 / 27 - 51%, advanced '
                    '(Witch) 0.70 / 44 - 55%, late (Witch) 0.96 / 3 - 54%: the middle peak is softer than in vanilla (0.57 advanced) and '
                    'there is no cliff; no change to the curve of magic itself.')},
    'S07': {'set': ('2026-10-16: the Witch INT x0.8 (growth row 18) and the five basic spells at levels 6 - 10 x0.85 (V062 - V185). '
                    'dtcli suite, class-balance (docs/generated/CLASS_BALANCE.md, every boss of the final classes): her parties win '
                    '73.4% (-0.7 sd against the classes; mean 76.6%). Speed: one Witch puts a party among the fastest 5% x1.85 as often '
                    'as the mean (+2.5 sd, the next class x1.59), two or more x6.6; in vanilla she was in 97% of the fastest parties. '
                    'Lower INT only takes wins from her, not speed: the speed of a stack of mages is what is left, and it is kept. '
                    'CLASS_STAGES.md: her fights last 0.70 of those of the median attacker (vanilla 0.57).')},
}
wb = openpyxl.load_workbook(PATH)
ws = wb['Criticisms']
head = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(head)}
done = []
for r in range(2, ws.max_row + 1):
    rid = ws.cell(r, col['ID']).value
    if rid in SET:
        cell = ws.cell(r, col['Balance_Result'])
        how = SET[rid]
        cell.value = how['set'] if 'set' in how else (str(cell.value or '').rstrip() + ' ' + how['add']).strip()
        done.append(rid)
assert sorted(done) == sorted(SET), done
wb.save(PATH)
print('written', ', '.join(done))
