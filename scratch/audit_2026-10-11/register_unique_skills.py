"""The unique skills of the eighteen girls as rows of the register (sheet Criticisms of
data_tables/balance_criteria.xlsx), asked by the user on 2026-10-12: what the Japanese pages of each girl say
(w.atwiki.jp/dt2-2, read that day) and what the GameFAQs guide says (faqs/80834, page "Girl Unique Skills + Party
Skills"), paraphrased, beside the values of the game. Theme T15, state open: to be measured before anything is decided.

    python register_unique_skills.py
"""
import os
import sys
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, PROJECT)
sys.path.insert(0, os.path.join(PROJECT, 'tools'))
from sim.data import game_data                            # noqa: E402

WORKBOOK = os.path.join(PROJECT, 'data_tables', 'balance_criteria.xlsx')
THEME = 'T15 Unique skills'
TEST = 'The reference party with the girl, with and without her unique skill: it adds wins, speed or HP left, and not more than the strongest of the others'
# girl: (atwiki page, kind, what the Japanese page says, what GameFAQs says)
ROWS = {
    'Irena': (44, 'Dominant', 'called a strong unique outright, above all with attack magic of long chant; with the crown said to fire one time in two',
              'the most powerful passive; the chance felt very high, close to 40 - 50% with the crown'),
    'Yae': (109, 'Dominant', 'two actions are strong and easy to use in any class; with Power Charge the damage explodes; it costs more TP',
            'monstrous value as a Dark Lord when it fires during Power Charge; the extra turn does not spend buffs and cannot fire twice in a row'),
    'Ist': (260, 'Dominant', 'a very good unique, useful in any class; with the crown First Aid heals about half the HP at once',
            'improves every maid skill by 1 - 13%, Clean included; best as an Etoile'),
    'Melvy': (261, 'Dominant', 'brutal figures with the crown, up to +20% in all; as a Magical Princess her ATK passes 1,000',
              'one of the few ways to make the Princess somewhat viable'),
    'Mefmera': (253, 'Mixed', 'depends on the enemy but a strong unique; at a high level it cuts a lot',
                'the cut is noticeable, but few endgame bosses are of those races'),
    'Eltricia': (229, 'Mixed', 'crowned and with high AGI it is always her turn; worse for casters of long chants since it no longer shortens the chant',
                 'the less time spent chanting the more it gives; pairs with the best grand sealbook of casters'),
    'Alisia': (230, 'Mixed', 'always on, so her damage is stable; weak against bosses that take Pep, and escaping lowers it',
               'about +15% damage at level 90 from the default Pep to the maximum; on par with Yae and Luco in mean power'),
    'Luco': (43, 'Mixed', 'the shadow strikes with about twice her attack for three actions; low chance (about 5% at rank 1); '
                          'a fault: if the shadow drains HP or TP its attack falls back to hers',
             'very powerful for a long while, falls off a little in the very late game; chance about 5% +1% a rank'),
    'Grishna': (257, 'Mixed', 'useful for any attacker; best with strikes of many hits (Samurai, attacking Valkyrie)',
                'a slight buff from the first game that hardly matters; best with Daedalus Attack'),
    'Monica': (254, 'Mixed', 'the highest weapon damage of the Scout line; wasted as a Kunoichi, who turns to ninjutsu',
               'felt like a chance of +20 - 25% damage with no pattern he could find; pure damage scouts are weak anyway'),
    'Lizerietta': (107, 'Mixed', 'a good skill of narrow use: it takes turns from bosses that summon',
                   'nice against bosses with familiars or dolls, useless outside of them; the chance is pretty high'),
    'Souffle': (255, 'Mixed', 'only a Treasure Hunter gets the most of it',
                'its line damage item is massive when she joins and next to nothing late; her worth is her party skills'),
    'Conette': (111, 'Mixed', 'a little stronger than in the first game (the speed cut lands often); only in the front row',
                'a great effect even when the slow is blocked; needs her in front taking hits'),
    'Fiora': (259, 'Mixed', 'an insurance', 'the common 5 - 20% chance; nothing more said'),
    'Yuni': (97, 'Dead', 'the debuff now ends after a few turns: a fix or a plain nerf; fine for a Maid line girl',
             'a slight indirect nerf since debuffs no longer last for ever; it also overwrites the ATK and INT buffs of the enemy'),
    'Lilian': (258, 'Dead', 'the ailment may be blocked by resistances, and the dispel does not fire as of version 1.05 (a bug)',
               'the dispel is useful at the right moment but that is very rare; less useful late since most enemies are immune'),
    'Tsurara': (41, 'Dead', 'she rises with 1 HP at a time nobody can tell, so it is lost unless healed at once, and it gets in the way of Resurrect',
                'kind of pointless: a revive item costs almost no time and is more reliable'),
    'Siena': (42, 'Dead', 'never useless in a class, never of real use either; a comment asks why such a pointless unique was left',
              'chance always about 15%, the share reflected from 50% to 200%; true damage; best in the front row'),
}


def main():
    data = game_data()
    girls = {c['name']: c for c in data.characters}
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    have = {r[0].value for r in ws.iter_rows(min_row=2)}
    added = 0
    for number, (girl, (page, kind, japanese, english)) in enumerate(ROWS.items(), 1):
        rid = f'I{number:02d}'
        if rid in have:
            continue
        unique = data.unique[girls[girl]['unique_skill']]
        values = ', '.join(f'{v:g}' for v in unique['values'][:9]) + f"; crown {unique['values'][9]:g}"
        ws.append([{'ID': rid, 'Source': 'JP community; GameFAQs DT2-2', 'Scope': 'Skill',
                    'Subject': f"{unique['name']} ({girl})",
                    'Claim': f'atwiki dt2-2 page {page}: {japanese} | GameFAQs faqs/80834, Girl Unique Skills: {english}',
                    'Kind': kind, 'Test': TEST,
                    'Vanilla_Result': f"Effect {unique['effect']}, values by rank {values} (UniqueSkillData; rule in "
                                      f'docs/FORMULAS.md 12.19; rank = level / 10, the crown value once the reset bonus reaches 50)',
                    'Status': 'open', 'Theme': THEME, 'Verdict': 'to measure'}.get(name) for name in head])
        added += 1
    wb.save(WORKBOOK)
    print(added, 'rows added')


if __name__ == '__main__':
    main()
