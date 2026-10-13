"""One-off edit (2026-10-13): the sheet Regression after the pass with the three habits of the members who heal
(sim/policy.py) and the buffs of the Sage line. Run once; kept as the record of the edit.
"""
import os

import openpyxl

PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables',
                    'balance_criteria.xlsx')
ROWS = {
    'R01': ('2026-10-13, after the model was taught three habits of a healer. measures/WIN_RATES.md (200 battles): the three '
            'parties that work in the game end level, Meta 93%, Dual-Support 92%, Mage-Heavy 96% of wins over the 17 bosses of '
            'the battery (86%, 90%, 91% in the game). The weakest party of the game (Anti-Meta-Plus, 50% on the mean) reaches '
            '84%. CLASS_STAGES.md: 74 of 76 rows within the band. Limit: the band has no top (row J11, left). The Witch is the '
            'case: far ahead at C30 alone, a boss that resists physical damage at 90 - 100 (row S07, decided 2026-10-13: not a '
            'problem there)', 'met in what is measured'),
    'R02': ('2026-10-13. PARTY_BATTERY.md: 68 of 85 pairs at or above the floor in vanilla, 75 of 85 on the variant; no pair '
            'that met its floor in vanilla lost it. CLASS_STAGES.md: the two slow rows (Assassin and Treasure Hunter of the '
            'treasure hunter role) were slow before the changes (row J26)', 'met'),
    'R08': ('2026-10-13, after the model was taught three habits of a healer. DIFFICULTY_CURVE.md (40 battles a boss): 34 of 44 '
            'checkpoints within the windows of vanilla, 10 easier, 0 harder. Eight are accepted (C14, C26, C29, C32, C37, C38 on '
            '2026-10-11; C17 and C40 on 2026-10-13). Two are new with the pass, checked with 200 battles '
            '(measures/curve_recheck.py): C19 is 27% shorter (8.0 > 5.8, wins 96% > 98%) and C30 is won more (76% > 94%, its '
            'HP x1.15 since that day)', 'met but for 10 bosses easier: 8 accepted, C19 and C30 not decided'),
}
wb = openpyxl.load_workbook(PATH)
ws = wb['Regression']
head = [c.value for c in ws[1]]
result = next(i for i, h in enumerate(head) if str(h).startswith('Result'))
for row in ws.iter_rows(min_row=2):
    if row[0].value in ROWS:
        row[result].value, row[head.index('Verdict')].value = ROWS[row[0].value]
wb.save(PATH)
print(len(ROWS), 'rows')
