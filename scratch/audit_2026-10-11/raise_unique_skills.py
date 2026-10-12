"""Three unique skills raised on request of 2026-10-12, after the trials of unique_trial.py beside this file:

* Third Eye Wide (Lilian): its chance times 2 (rows U001 - U010 of data_tables/balance_changes.xlsx).
* Usurper of Souls (Lizerietta): its chance times 1.6, never over 1 (rows U011 - U020): "very effective where it
  belongs", against the enemies that summon.
* Arcane Shield (Siena): the roll written in the program from 15 to 25 in 100 (row X13 of the sheet ExePatches, set
  balance, lever L38), and the box of the skill worded for it (data_tables/balance_skill_text.xlsx).

    python raise_unique_skills.py
"""
import os
import sys
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, PROJECT)
sys.path.insert(0, os.path.join(PROJECT, 'tools'))
from sim.data import game_data                            # noqa: E402

TABLES = os.path.join(PROJECT, 'data_tables')
RAISES = (('Third Eye Wide', 2.0, 'I15',
           'Decided 2026-10-12: times 2. Without it the reference party won 78.6% of five late bosses, with it 78.0%, '
           'at twice the chance 80.2% (500 battles a row); rows of eight random battles at four floors do not run away '
           'with it. Its seventh result, the dispel, still removes nothing: that is a fault of the program (0x47ee6f)'),
          ('Usurper of Souls', 1.6, 'I11',
           'Decided 2026-10-12: raise its chance so that it is very effective where it belongs, against the enemies '
           'that summon; it does nothing anywhere else, so nothing else moves. Times 1.6 and never over 1: the crown '
           'value becomes certain'))
PATCH = ('X13', 'balance', 'L38', '0x424768', '0f', '19', None, None,
         'Roll of Arcane Shield, a number below 100 under: 15 -> 25 (decided 2026-10-12; cmp eax, imm8 at 0x424766)')
ARCANE_FILE = 'UniqueSkillData_0379.bin'
ARCANE_OLD, ARCANE_NEW = 'Activation rate: 15%', 'Activation rate: 25%'


def main():
    data = game_data()
    wb = openpyxl.load_workbook(os.path.join(TABLES, 'balance_changes.xlsx'))
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    have = {r[0].value for r in ws.iter_rows(min_row=2)}
    number = 0
    for name, factor, serves, why in RAISES:
        uid, skill = next((k, s) for k, s in data.unique.items() if s['name'] == name)
        for rank, value in enumerate(skill['values'], 1):
            number += 1
            rid = f'U{number:03d}'
            if rid in have:
                continue
            ws.append([{'ID': rid, 'Table': 'unique', 'Record': f'Unique {uid} {name}', 'Field': 'value', 'Level': rank,
                        'Vanilla': value, 'New': min(1.0, round(value * factor, 2)), 'Lever': 'L41', 'Serves': serves,
                        'Why': why, 'Decided': '2026-10-12', 'State': 'decided'}.get(c) for c in head])
    wb.save(os.path.join(TABLES, 'balance_changes.xlsx'))

    wb = openpyxl.load_workbook(os.path.join(TABLES, 'balance_criteria.xlsx'))
    ws = wb['ExePatches']
    if PATCH[0] not in {r[0].value for r in ws.iter_rows(min_row=2)}:
        last = max(r[0].row for r in ws.iter_rows(min_row=2) if str(r[0].value or '').startswith('X') and r[1].value == 'balance')
        ws.insert_rows(last + 1)
        for k, value in enumerate(PATCH, 1):
            ws.cell(row=last + 1, column=k, value=value)
    ws = wb['Levers']
    if 'L41' not in {r[0].value for r in ws.iter_rows(min_row=2)}:
        ws.append(('L41', 'The ten values of a unique skill', 'UniqueSkillData.pck', '0x1C, ten floats', 'Benign',
                   'A chance stays between 0 and 1', 'tools/extract_vanilla.py reads the ten values; docs/FORMULAS.md 12.19 and 16.3'))
    wb.save(os.path.join(TABLES, 'balance_criteria.xlsx'))

    qol = openpyxl.load_workbook(os.path.join(TABLES, 'qol_skill_text.xlsx'), read_only=True)['Skills']
    rows = qol.iter_rows(values_only=True)
    qhead = list(next(rows))
    box = next(r[qhead.index('QoL_Desc')] for r in rows if r and r[qhead.index('Filename')] == ARCANE_FILE)
    assert ARCANE_OLD in box, box
    wb = openpyxl.load_workbook(os.path.join(TABLES, 'balance_skill_text.xlsx'))
    ws = wb['Skills']
    bhead = [c.value for c in ws[1]]
    if ARCANE_FILE not in {r[bhead.index('Filename')].value for r in ws.iter_rows(min_row=2)}:
        ws.append([{'Category': 'Unique', 'Filename': ARCANE_FILE, 'Balance_Desc': box.replace(ARCANE_OLD, ARCANE_NEW)}.get(c)
                   for c in bhead])
    wb.save(os.path.join(TABLES, 'balance_skill_text.xlsx'))
    print(number, 'value rows, patch', PATCH[0])


if __name__ == '__main__':
    main()
