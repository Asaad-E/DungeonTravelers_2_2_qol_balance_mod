"""Sheet Roles of sim_setup.xlsx: the Soul Summoner of the model learns Soul Eater, her own attack (2026-10-16).
It was missing from her skill priority, so she never cast it; dtcli skill-battery rates it good."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'sim_setup.xlsx')
wb = openpyxl.load_workbook(PATH)
ws = wb['Roles']
head = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(head)}
for r in range(2, ws.max_row + 1):
    if ws.cell(r, col['Role']).value == 'soul summoner':
        cell = ws.cell(r, col['Skill priority'])
        skills = [s.strip() for s in cell.value.split(';')]
        if 'Soul Eater' not in skills:
            skills.insert(skills.index('Musiy') + 1, 'Soul Eater')
        cell.value = '; '.join(skills)
        print(cell.value)
wb.save(PATH)
