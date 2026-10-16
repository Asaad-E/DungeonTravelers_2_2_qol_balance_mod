"""One-off edit (decided 2026-10-16): Circle Heal's heal power rises by 0.5 at every level, 10 to 14.5 (5x to 7.25x RES on
screen): levels 1 - 9 back to 10 - 14 (rows H001 - H009, as built before the even curve of 10 - 14 that left
two-decimal figures), level 10 14.5 (new row H038; the game has 14). Its text in balance_skill_text.xlsx follows."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CHANGES = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
TEXT = os.path.join(ROOT, 'data_tables', 'balance_skill_text.xlsx')
NOTE = 'Circle Heal rises evenly to its top level (14, as the game): level 9 healed as much as level 10'
NEW_NOTE = 'Circle Heal rises by 0.5 at every level, 10 to 14.5 (5x to 7.25x): level 9 healed as much as level 10 (14)'
POWER = [10, 10.5, 11, 11.5, 12, 12.5, 13, 13.5, 14, 14.5]

wb = openpyxl.load_workbook(CHANGES)
ws = wb['Changes']
rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
for level in range(1, 10):
    r = rows[f'H{level:03d}']
    assert ws.cell(r, 3).value == 'Magic 188 Circle Heal' and ws.cell(r, 5).value == level
    ws.cell(r, 7).value = POWER[level - 1]
    ws.cell(r, 10).value = ws.cell(r, 10).value.replace(NOTE, NEW_NOTE)
if 'H038' not in rows:
    template = [ws.cell(rows['H009'], c).value for c in range(1, ws.max_column + 1)]
    ws.append(template)
    r = ws.max_row
    ws.cell(r, 1).value = 'H038'
    ws.cell(r, 5).value = 10
    ws.cell(r, 6).value = 14
    ws.cell(r, 7).value = POWER[9]
wb.save(CHANGES)

wb = openpyxl.load_workbook(TEXT)
ws = wb['Skills']
header = [c.value for c in ws[1]]
col = header.index('Balance_LvUp') + 1
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 2).value == 'MagicSkillData_0188.bin':
        old = ws.cell(r, col).value
        new = old.replace('5x -> 5.25 -> 5.5 -> 5.75 -> 6 -> 6.25 -> 6.5 -> 6.75 -> 7 -> 7x (L10)',
                          '5x -> 5.25 -> 5.5 -> 5.75 -> 6 -> 6.25 -> 6.5 -> 6.75 -> 7 -> 7.25x (L10)')
        assert new != old
        ws.cell(r, col).value = new
wb.save(TEXT)
print('done')
