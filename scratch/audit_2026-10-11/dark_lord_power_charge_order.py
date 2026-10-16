"""Sheet Policies of sim_setup.xlsx: the Dark Lord keeps Berserk up before she charges Power Charge (2026-10-16).

With Power Charge first, the turn after the charge went to Berserk and the charge was lost (trace of C29, Meta party): Power Charge
adds its value to the other bonus of the next physical blow (docs/FORMULAS.md 2.4: final = damage C * (1 + other bonuses)), so the
blow after the charge must be an attack."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'sim_setup.xlsx')
SOURCE = ('Berserk kept up before Power Charge, so that the turn after the charge is a blow (2026-10-16: with the charge first, the '
          'charged turn went to Berserk; dtcli trace C29 Meta)')

wb = openpyxl.load_workbook(PATH)
ws = wb['Policies']
head = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(head)}
found = {}
for row in range(2, ws.max_row + 1):
    if ws.cell(row, col['Role']).value == 'dark lord':
        found[ws.cell(row, col['Action']).value] = row
charge, berserk = found['status:191'], found['status:26']
order_charge = ws.cell(charge, col['Order']).value
order_berserk = ws.cell(berserk, col['Order']).value
assert (order_charge, order_berserk) == (4, 5), (order_charge, order_berserk)
ws.cell(charge, col['Order']).value = order_berserk
ws.cell(berserk, col['Order']).value = order_charge
ws.cell(charge, col['Source']).value = SOURCE
wb.save(PATH)
print('dark lord: Berserk', order_charge, 'Power Charge', order_berserk)
