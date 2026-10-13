"""One-off edit (2026-10-13) of the sheets Criticisms and Regression of balance_criteria.xlsx: the seven rows closed on
a reading of the design get the figures of measures/DESIGN_ROWS.md, and R08 the check of five bosses with 200
battles. Where the figure says the claim holds the verdict changes; the state of a row is left for a decision.
Run once; kept as the record of the edit.
"""
import os

import openpyxl

PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables',
                    'balance_criteria.xlsx')
SOURCE = ' (measures/DESIGN_ROWS.md, 2026-10-13)'
SET = {
    'J17': {'Vanilla_Result': 'The reference party at the 40 boss fights of the story, 30 battles each: the largest damage of one '
                              'enemy action on one member is above the whole HP of the frailest member in the median battle of '
                              '12 fights (C26, C29, C30, C31, C32, C34 to C40; 119% to 201% of her HP) and at least once in 18. '
                              'Before C26 it stays below (40% to 94%)' + SOURCE,
            'Verdict': 'true from C26 on',
            'Why': 'From C26 on one enemy action can fell the frailest member from full HP in most battles, with Cover and the '
                   'guards played. It is what makes the Valkyrie and the Paladin worth their place (rows J04, F07); nothing '
                   'was changed for it',
            'Decision': '2026-10-13: measured; the claim holds. No change so far: the state waits for a decision'},
    'J35': {'Vanilla_Result': 'Each upgrade level adds int(level * rate) to every stat of the piece: for the 17 weapons that '
                              'take levels, 5.4% to 9.7% of base ATK a level (median 6.3%), not one percent. Only 29 of the 551 '
                              'pieces take levels at all. What a level costs in gold is in the program and was not read' + SOURCE,
            'Verdict': 'false on the size of the bonus; its cost not measured',
            'Why': 'A level is worth about 6% of the base ATK of the weapon, and the levels add up (a Claymore goes from 37 to '
                   '312 ATK at the top level). The half of the claim about the price is not answered',
            'Decision': '2026-10-13: no change'},
    'S16': {'Vanilla_Result': 'The best weapon that takes levels, at the highest level a zone gives (99), has less ATK than the '
                              'best fixed piece in each of the 10 weapon types (Claymore 312 against Amenomurakumo 433, Heavy '
                              'Axe 361 against Adamas Greatscythe 462, Kukri 137 against Kali Kukri 279). The fixed pieces '
                              'carry 1 to 4 enchantments of their own; the chosen enchantments of a forged piece are not '
                              'weighed here' + SOURCE,
            'Verdict': 'false on ATK; the enchantments not compared',
            'Why': 'On attack the unique pieces stay ahead to the end, by 20% to 140%. Whether three or four chosen enchantments '
                   'make up for it is not measured',
            'Decision': '2026-10-13: no change'},
    'S14': {'Vanilla_Result': 'Of the 103 floors, 20 have cells that forbid magic and 26 cells that lower physical or magic '
                              'attack; where a floor has them they cover a median of 10% and 16% of the cells of its map (walls '
                              'counted), at most 36% and 44%. A lowering zone multiplies attack by 0.5 to 0.75, for both sides'
                              + SOURCE,
            'Verdict': 'a matter of taste; the figure: a fifth to a quarter of the floors, a tenth to a sixth of their cells',
            'Why': 'The zones weigh on both sides alike and cover a small part of the floors that have them. Maps are out of '
                   'what the variant changes',
            'Decision': '2026-10-13: no change'},
    'S19': {'Vanilla_Result': 'A random battle of the first two dungeons gives 5 to 18 gold and of the next two 27 to 88, '
                              'against 90 for the cheapest healing item (Croissant), 180 for an Egg Sandwich and 450 for a '
                              'Danish; a piece sells for a third of its price. The prices of equipment are not decoded' + SOURCE,
            'Verdict': 'partly true: five to eighteen battles pay one basic heal at the start',
            'Why': 'Battles pay little at the start; quests are where the early gold comes from, and the variant raises the gold '
                   'of every quest by 40% (row S20)',
            'Decision': '2026-10-13: answered in part by the rows of S20; nothing more'},
    'J36': {'Vanilla_Result': 'Of the 16 meals, 4 only restore HP or TP or cure, at once, and 12 also give a buff (a skill at a '
                              'fixed level). How many steps that buff lasts is in the program and was not read, so the test of '
                              'this row is not answered' + SOURCE,
            'Verdict': 'not measured',
            'Why': 'A quarter of the meals have nothing that could run out. For the other twelve the length of the buff is not '
                   'known here',
            'Decision': '2026-10-13: left; it needs a reading of the program'},
    'F95': {'Vanilla_Result': 'Boss fights on a cell that forbids magic: C29a in the story, P07, P13, P25 and P47 in the '
                              'postgame. C29a is won without her (Dual-Support 100%). In the postgame the same party with and '
                              'without the Treasure Hunter, 100 battles: Meta 57% > 18% at P13 and 96% > 76% at P25; '
                              'Dual-Support with her in the place of the shooter 14% > 43% and 52% > 75%; at P07 she changes '
                              'nothing (97% > 95%, 3% > 3%) and nobody wins P47' + SOURCE,
            'Verdict': 'partly true: she adds 20 to 40 points of wins at two of the four postgame fights, and is not needed '
                       'in the story',
            'Why': 'She helps where magic is shut, as the guide says, but a party wins P07 and P25 without her (Meta 95% and 76%)',
            'Decision': '2026-10-13: no change'},
}
R08 = ('2026-10-13. DIFFICULTY_CURVE.md (40 battles a boss): 33 of 44 checkpoints within the windows of vanilla, 11 easier, 0 '
       'harder. Six were accepted on 2026-10-11 (C14, C26, C29, C32, C37, C38); C30 is back within with its HP x1.15. The five '
       'that crossed since then, with 200 battles (measures/curve_recheck.py): C15, C28 and C39 are within, it was chance; C17 '
       'is 28% shorter (49.1 > 35.4, wins 100% both) and C40 is won more (85% > 97%)',
       'met but for 8 bosses easier: 6 accepted, C17 and C40 not decided')


def main():
    wb = openpyxl.load_workbook(PATH)
    ws = wb['Criticisms']
    head = [c.value for c in ws[1]]
    done = 0
    for row in ws.iter_rows(min_row=2):
        for column, text in SET.get(row[0].value, {}).items():
            row[head.index(column)].value = text
            done += 1
    assert done == sum(len(v) for v in SET.values()), done
    ws = wb['Regression']
    head = [c.value for c in ws[1]]
    result = next(i for i, h in enumerate(head) if str(h).startswith('Result'))
    for row in ws.iter_rows(min_row=2):
        if row[0].value == 'R08':
            row[result].value, row[head.index('Verdict')].value = R08
    wb.save(PATH)
    print(done, 'cells')


if __name__ == '__main__':
    main()
