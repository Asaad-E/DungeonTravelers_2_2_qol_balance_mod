"""One-off edit (2026-10-13): row S17 of the register carries its decision. Run once; kept as the record."""
import os

import openpyxl

PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables',
                    'balance_criteria.xlsx')
wb = openpyxl.load_workbook(PATH)
ws = wb['Criticisms']
head = [c.value for c in ws[1]]
row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == 'S17')
row[head.index('Decision')].value = ('2026-10-10: no change; 2026-10-13, confirmed: the cap of a quarter stays, raising it would '
                                     'make every party stronger')
wb.save(PATH)
print('S17')
