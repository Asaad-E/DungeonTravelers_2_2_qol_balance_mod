"""Decided 2026-10-16: the Etoile and the Diva lowered into the band of the classes (dtcli class-search preset=balance): Debilitating
Salsa's damage cap (effect 296) x0.7 and Fascinating Waltz (effect 195) x0.85 on its raised value; Magical Spark power x1.15; the Diva (growth row 27) HP and RES x0.85
at every level. With all of it: 15 of 16 classes in the band; the Diva 91.6% -> 87.2% wins (+2.0 sd), the Etoile 89.3% -> 86.4% wins
and x1.98 -> x1.52 among the fastest (+1.5 sd). Her songs do not carry the Diva's wins (Carol x0.6, Chorale x0.8, Self Session without
effect changed nothing): her own survival does."""
import json
import os
import sys
import openpyxl

HERE = os.path.dirname(__file__)
ROOT = os.path.join(HERE, '..', '..')
sys.path.insert(0, os.path.join(ROOT, 'tools'))
import class_data                                   # noqa: E402

PATH = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
DECIDED = '2026-10-16'
WHY_E = ('The Etoile into the band of the classes (S07, decided 2026-10-16): in the final classes she was an outlier in speed (x1.98 '
         'among the fastest 5%, +2.3 sd), carried by Debilitating Salsa, and second in wins, carried by Fascinating Waltz. Salsa cap '
         'x0.7 and Waltz x0.85: x1.52 (+1.5 sd), wins 89.3% -> 86.4% (docs/generated/CLASS_BALANCE.md).')
WHY_D = ('The Diva into the band of the classes (S07, decided 2026-10-16): first in wins of the final classes (91.6%, +2.3 sd), carried '
         'by her own survival, not by her songs (Carol x0.6, Chorale x0.8 or Self Session without effect changed nothing). HP and RES '
         'x0.85: 87.2% (+2.0 sd) (docs/generated/CLASS_BALANCE.md).')

snap = json.load(open(os.path.join(ROOT, 'build', 'engine_snapshot.json'), encoding='utf-8'))['sides']['vanilla']['skills']
NEW = []
for sid, name, effect, f, digits in ((73, 'Debilitating Salsa', 296, 0.7, 0),):
    skill = next(s for s in snap if s['id'] == sid)
    power = next(e[2] for e in skill['effects'] if e[0] == effect)
    for level, old in enumerate(power, start=1):
        new = round(old * f) if digits == 0 else round(old * f, digits)
        NEW.append(('skills', f'Technique {sid} {name}', f'effect {effect}', level, old, new, 'L12', 'S07', WHY_E))
table = class_data.growth()
for stat in ('hp', 'res'):
    for level in range(1, class_data.LEVELS + 1):
        old = table[27][level - 1][stat]
        NEW.append(('class_growth', '27 Diva', stat, level, old, round(old * 0.85), 'L10', 'S07', WHY_D))

wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
head = [c.value for c in ws[1]]
assert head[:12] == ['ID', 'Table', 'Record', 'Field', 'Level', 'Vanilla', 'New', 'Lever', 'Serves', 'Why', 'Decided', 'State'], head
records = {(ws.cell(r, 3).value, ws.cell(r, 4).value) for r in range(2, ws.max_row + 1)}
assert not any((v[1], v[2]) in records for v in NEW), 'rows already there'
# rows already in the sheet: Fascinating Waltz (A084 - A093, raised 2026-10-09) x0.85 on its New; Magical Spark power (Q006 - Q015)
# x1.15 on its New (decided 2026-10-16, with the Soul Summoner and the Witch)
WHY_S = ('The Magical Princess helped (S07, decided 2026-10-16): Magical Spark power x1.15 on its raised value, with its TP x0.8; she '
         'was last but one in wins of the final classes (docs/generated/CLASS_BALANCE.md).')
for r in range(2, ws.max_row + 1):
    rid = ws.cell(r, 1).value
    if rid in ['A%03d' % k for k in range(84, 94)]:
        assert ws.cell(r, 4).value == 'effect 195'
        ws.cell(r, 7).value = round(ws.cell(r, 7).value * 0.85, 2)
        ws.cell(r, 10).value = str(ws.cell(r, 10).value) + ' ' + WHY_E
    if rid in ['Q%03d' % k for k in range(6, 16)]:
        assert ws.cell(r, 4).value == 'effect 26 (power)'
        ws.cell(r, 7).value = round(ws.cell(r, 7).value * 1.15, 2)
        ws.cell(r, 10).value = str(ws.cell(r, 10).value) + ' ' + WHY_S
ids = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
first = 1 + max(int(k[1:]) for k in ids if k and str(k).startswith('V'))
row = ws.max_row + 1
for k, values in enumerate(NEW):
    for c, value in enumerate(('V%03d' % (first + k),) + values + (DECIDED, 'built'), start=1):
        ws.cell(row, c).value = value
    row += 1
wb.save(PATH)
print('written V%03d - V%03d' % (first, first + len(NEW) - 1))
