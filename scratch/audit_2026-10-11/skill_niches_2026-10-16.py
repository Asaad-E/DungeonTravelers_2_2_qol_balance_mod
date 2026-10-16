"""Skills without a niche in dtcli skill-battery (2026-10-16): Power Charge, Accelerator, Shadow Clone, Prism Dawn, Oracles' Action.

Measured before building: skill-battery NAMES 20 200 variants TRIAL (every active under five rules, boss / random / walk then boss) and, for
Power Charge and Oracles, class-search final C34, C38 boss with the parties that hold the class, the variant against the variant with the
value (compare=variant)."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'balance_changes.xlsx')
DECIDED = '2026-10-16'

WHY_PC = ('Power Charge a little stronger, for flavour (decided 2026-10-16), the charged blow at most x2.4: its value adds to the other bonus '
          'of the next physical blow (docs/FORMULAS.md 2.4), so 1.05 - 1.30 gave x2.05 - x2.3 and 1.13 - 1.40 gives x2.13 - x2.4. skill-battery: '
          'niche to good (serves in 26% of the cells); class-search C34, C38, the 7,632 final-class parties with a Dark Lord: wins 69.9% > 70.6%, '
          'length 15.2 > 15.0, no dominant party. The Dark Lord of the model charges after Berserk (sheet Policies, 2026-10-16).')
WHY_ACC = ('Accelerator a little stronger (decided 2026-10-16): a fun skill whose numbers were thin, +8 - 12% speed to one ally for 3 turns '
           'against a turn of the Sage. 10 - 15% for 4 turns; skill-battery: niche, its best cell +12% wins (C19, a walk then the boss). The '
           'model gives it to the ally who hits hardest (Policy.cs, 2026-10-16).')
WHY_SC = ('Shadow Clone one turn longer (decided 2026-10-16): with K001 - K005 and X075 - X078 it still cost the Kunoichi more than it saved '
          '(skill-battery: hurts, 51% of the cells worse). One more turn at every level: niche; 0.7 times the TP did not help.')
WHY_PD = ('Prism Dawn x1.4 power (decided 2026-10-16): even forced into use it added nothing to the Sage (skill-battery: dead, also at x1.2); '
          'at x1.4 it serves in some fights (niche).')
WHY_OR = (' Then x1.5 (decided 2026-10-16): Oracles took the place of a better doll and made the party worse (skill-battery: hurts at x1.0, '
          'x1.3 and x1.4, niche at x1.5); the magic doll against foes of low RES, a little above Spartan in power. class-search C34, C38, the '
          '7,402 final-class parties with a Soul Summoner: wins 70.2% > 70.5%, no dominant party.')

NEW = []
for level, (old, new) in enumerate([(1.05, 1.13), (1.1, 1.19), (1.16, 1.25), (1.22, 1.32), (1.3, 1.40)], start=1):
    NEW.append(('skills', 'Technique 28 Power Charge', 'effect 191', level, old, new, 'L12', 'J21, F18', WHY_PC))
for level, (old, new) in enumerate([(0.08, 0.10), (0.09, 0.11), (0.10, 0.13), (0.11, 0.14), (0.12, 0.15)], start=1):
    NEW.append(('skills', 'Magic 197 Accelerator', 'effect 168', level, old, new, 'L12', 'J23', WHY_ACC))
for level in range(1, 6):
    NEW.append(('skills', 'Magic 197 Accelerator', 'turns', level, 3, 4, 'L12', 'J23', WHY_ACC))
for level, old in enumerate([3, 3, 4, 4, 5], start=1):
    NEW.append(('skills', 'Magic 216 Shadow Clone', 'turns', level, old, old + 1, 'L12', 'F60', WHY_SC))
for level, old in enumerate([100, 130, 160, 200, 240, 280, 320, 370, 420, 480], start=1):
    NEW.append(('skills', 'Magic 232 Prism Dawn', 'effect 1', level, old, round(old * 1.4), 'L12', 'J23', WHY_PD))
ORACLES = {'D%03d' % (60 + level): new for level, new in
           enumerate([2.28, 2.34, 2.43, 2.52, 2.61, 2.70, 2.82, 2.94, 3.06, 3.21], start=1)}

wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
head = [c.value for c in ws[1]]
assert head[:12] == ['ID', 'Table', 'Record', 'Field', 'Level', 'Vanilla', 'New', 'Lever', 'Serves', 'Why', 'Decided', 'State'], head
col = {name: i + 1 for i, name in enumerate(head)}
ids = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
for rid, new in ORACLES.items():
    r = ids[rid]
    assert 'Oracles' in ws.cell(r, col['Record']).value and ws.cell(r, col['Field']).value == 'effect 26'
    ws.cell(r, col['New']).value = new
    why = ws.cell(r, col['Why']).value or ''
    if WHY_OR.strip() not in why:
        ws.cell(r, col['Why']).value = why.rstrip() + WHY_OR
    ws.cell(r, col['Decided']).value = DECIDED
# the new rows: V001 on, at the end of the sheet; a row already written is overwritten
taken = [k for k in ids if str(k).startswith('V')]
row = ws.max_row + 1
for k, values in enumerate(NEW, start=1):
    rid = 'V%03d' % k
    target = ids.get(rid, None)
    if target is None:
        target = row
        row += 1
    for c, value in enumerate((rid,) + values + (DECIDED, 'built'), start=1):
        ws.cell(target, c).value = value
wb.save(PATH)
print('written', len(NEW), 'rows V001 - V%03d' % len(NEW), 'and', len(ORACLES), 'rows of Oracles\' Action')
