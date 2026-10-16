"""One-off edit: the levers whose changes were checked in the game on 2026-10-16 say so in the column Evidence of the sheet
Levers of balance_criteria.xlsx, as L42 - L45 did on 2026-10-12."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
SEEN = {
    'L12': 'Circle Brave, Circle Protect, Sun Veil and Accelerator at their new values and chant speed; Ecstasy 20% - 30%',
    'L15': 'Bathroom Slippers with Double EXP (enchantment 275)',
    'L25': 'the patched program files run, the one with its protector too',
    'L27': 'the patched program files run, the one with its protector too',
    'L38': 'the patched program files run, the one with its protector too',
    'L46': 'the wandering blacksmith comes more often',
    'L47': 'a meal lasts a couple of battles while poison still wears off as before',
}

wb = openpyxl.load_workbook(PATH)
ws = wb['Levers']
col = [c.value for c in ws[1]].index('Evidence') + 1
done = 0
for r in range(2, ws.max_row + 1):
    lever = ws.cell(r, 1).value
    if lever in SEEN:
        old = ws.cell(r, col).value or ''
        note = f'Seen working in the game (play-tested, 2026-10-16: {SEEN[lever]})'
        if note not in old:
            ws.cell(r, col).value = (old.rstrip('. ') + '. ' if old else '') + note
        done += 1
assert done == len(SEEN), done
wb.save(PATH)
print('levers', done)
