"""One-off edit: the figures of the nine rows of the register closed on an argument (sheet Criticisms of
balance_criteria.xlsx, column Balance_Result), measured or read on 2026-10-16. Avenger and Barrage:
passive_worth_2026-10-16.py beside this file; the rest from the generated files named in each text."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
FIGURES = {
    'J13': 'No text of the game or of the mod states what a mercenary does (0 of 31). What each adds is measured in '
           'measures/MERCENARIES.md: the reference party at C36, C40 and P24 with each guest against none (58%, 92%, 100% '
           'of wins without).',
    'F07': 'At every story boss of measures/WIN_RATES.md (17 of 17) a party without a Valkyrie wins on the variant: 92% to '
           '100% at 16 of them, 74% at C36 (Dual-Support). The superbosses of the postgame are outside the measure.',
    'F20': 'Each level buys a stated gain (skill records of the variant): Panic Voice +40 power a level (240 to 600) for 6 '
           'to 8 TP more, the same damage per TP (6.7 at level 1, 6.3 at level 10); Divine Shield 30% to 45% for 24 to 42 '
           'TP; Adaptation 6 to 15 and Preparation 15% to 20%, passives that cost nothing.',
    'F21': 'Nurse and Cleanse cost TP, not gold, and serve in a fight (docs/generated/SKILL_VALUE.md): Nurse at C38 wins +7 '
           'points, HP left +9; Cleanse at C40 wins +3 points, HP left +11. Return costs 1 TP at its level 5.',
    'F28': 'Attract on the Dark Lord at C38 (docs/generated/SKILL_VALUE.md): wins +0 points, fights 5% shorter, HP left +6 '
           'points: it serves.',
    'F38': 'Avenger on the tank at C30, C34, C36, C38 and C40, against the same tank without it (300 battles): wins within '
           '1 point at each, length within 0.2: no harm late, and no gain.',
    'F63': 'Barrage acts on plain attacks only: the Kunoichi of the model makes 0.7 a fight at C38 (3% of her actions). '
           'With it or without it the wins at C30 - C40 do not move (300 battles); at a 100% chance C38 stays at 85%.',
    'F82': 'Cleanse on the Shooter at C40 (docs/generated/SKILL_VALUE.md): wins +3 points, HP left +11, for TP instead of '
           'an item.',
    'F87': 'The Sniper is inside the band at every stage of docs/generated/CLASS_STAGES.md (length 0.91 and 1.04 of the '
           'median of the attackers, wins 85% - 100%).',
}

wb = openpyxl.load_workbook(PATH)
ws = wb['Criticisms']
header = [c.value for c in ws[1]]
col = header.index('Balance_Result') + 1
done = 0
for r in range(2, ws.max_row + 1):
    rid = ws.cell(r, 1).value
    if rid in FIGURES:
        old = ws.cell(r, col).value
        ws.cell(r, col).value = FIGURES[rid] if not old or old == 'None' else f'{old} | {FIGURES[rid]}'
        done += 1
assert done == len(FIGURES), done
wb.save(PATH)
print('rows', done)
