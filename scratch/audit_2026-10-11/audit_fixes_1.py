"""After the audit of 2026-10-13: Provoke is never below the game (its level 1 was), and its rows say what they do and
which complaint they serve. Run once."""
import os
import openpyxl

TABLES = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables')
WHY = ('Provoke grows with its level: in the game its enmity is 250 at every level, which draws about 77% of the attacks an '
       'enemy aims at one member (weight = share of the position + value, docs/FORMULAS.md 9.1), and a level only adds a '
       'turn. Levels 2 to 5 become 450, 750, 1100 and 1500, about 85%, 91%, 93% and 95%; level 1 stays at 250. Measured on '
       '2026-10-13 (measures/provoke_trial.py, the tank in the reference party, 150 battles at C18, C30, C36 and C40): the '
       'wins and the length of the fights are the same with 250, with these values and with a gentler curve, since the '
       'strong blows of the bosses fall on the whole party')
path = os.path.join(TABLES, 'balance_changes.xlsx')
wb = openpyxl.load_workbook(path)
ws = wb['Changes']
col = {c.value: k for k, c in enumerate(ws[1])}
for row in ws.iter_rows(min_row=2):
    if str(row[col['Record']].value) == 'Technique 0 Provoke' and str(row[col['Field']].value) == 'effect 170':
        row[col['Serves']].value = 'F27'
        if int(row[col['Level']].value) == 1:
            row[col['State']].value = 'dropped'
            row[col['Why']].value = ('Dropped on 2026-10-13: 200 was below the 250 of the game, and no change may leave a '
                                     'level worse than it was. ' + WHY)
        else:
            row[col['Why']].value = WHY
wb.save(path)
path = os.path.join(TABLES, 'balance_criteria.xlsx')
wb = openpyxl.load_workbook(path)
ws = wb['Criticisms']
col = {c.value: k for k, c in enumerate(ws[1])}
for row in ws.iter_rows(min_row=2):
    if row[0].value == 'F27':
        row[col['Status']].value = 'closed'
        row[col['Proposal']].value = 'S033 - S036: the enmity of Provoke grows with its level (250 at level 1, 1500 at level 5)'
        row[col['Decision']].value = (str(row[col['Decision']].value or '') + ' | 2026-10-13: the rows that raise it by level '
                                      'were already built under another complaint; they are filed here, level 1 put back to '
                                      'the value of the game, and measured (no boss moves)')
wb.save(path)
print('written')
