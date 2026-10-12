"""The three TP fruits of the shop give a fifth more TP, the user's word of 2026-10-12 (row F01 of the register):
rows Y012 - Y014 of data_tables/balance_changes.xlsx.

    python raise_tp_fruit.py
"""
import os
import openpyxl

PROJECT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CHANGES = os.path.join(PROJECT, 'data_tables', 'balance_changes.xlsx')
FRUIT = (('Y012', '19 Magical Fruit', 25, 30), ('Y013', '20 Fairy Apple', 50, 60), ('Y014', '21 Golden Apple', 100, 120))
WHY = ('The user, 2026-10-12 (row F01: no party without the Maid line wins from C30 on): a fifth more TP from the fruit, '
       'not a new cut of TP costs, which would move every party. Measured before it: the Meta party with a Bishop for '
       'its Diva wins 2% at C30 and C34 with 8 fruit of 50 TP, 50% with 24, 100% with TP at will; C36 and the last boss '
       'do not turn on TP. The Maid keeps her part; without her the fights are won with effort and a bag of fruit')


def main():
    wb = openpyxl.load_workbook(CHANGES)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    have = {r[0].value for r in ws.iter_rows(min_row=2)}
    for rid, record, old, new in FRUIT:
        assert rid not in have, rid
        ws.append([{'ID': rid, 'Table': 'items', 'Record': record, 'Field': 'effect 1 value (TP to one member)',
                    'Vanilla': old, 'New': new, 'Lever': 'L19', 'Serves': 'F01', 'Why': WHY, 'Decided': '2026-10-12',
                    'State': 'decided'}.get(name) for name in head])
    wb.save(CHANGES)


if __name__ == '__main__':
    main()
