"""The unique skill of Lizerietta, decided again on 2026-10-12 (night): not evasion, which is always at
work where every other unique skill is situational, but the chance to hold on with 1 HP when a hit would fell her
(effect 254, the one of Guts, which no class of hers has). Named after what she says of herself in her event
(sev_180: she does not believe in God and would not be glad to be saved by a miracle). Run once.

    python register_lize_survival.py
"""
import os
import sys

import openpyxl

TABLES = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'data_tables')
VALUES = [0.15, 0.18, 0.21, 0.24, 0.27, 0.30, 0.33, 0.36, 0.40, 0.55]
NAME = 'No Miracles Needed'
BOX = ('When a hit would fell her, she may hold on with 1 HP.' + chr(10) +
       'Activation rate: 15% -> 18 -> 21 -> 24 -> 27 -> 30 -> 33 -> 36 -> 40%   |   Crown EX: 55%')
WHY = ('Decided 2026-10-12 (night): evasion was always at work and made her simply better than any other Spieler; the '
       'skill becomes the chance to hold on with 1 HP when a hit would fell her (effect 254, read among the statuses of '
       'the target at 0x4238ad, where a unique skill is found as any status): situational as the others, it tells who she '
       'is (the smallest, who fights on without complaining and waits for no miracle: her event sev_180), and it serves '
       'her dolls without being tied to them, since they end when she falls. Looked at first: the DLC adds no playable '
       'girl; the unique skills the first game gave and this one dropped (GameFAQs faqs/72905) are not in the program. '
       'Tried (scratch/audit_2026-10-11/lize_ideas.py, 100 battles a cell, as papillon, doll master and joker at C36, '
       'C40, P27): at 0.20, 0.35 and 0.50 the wins move by 0 to 6 points, most at C40, where she falls most')


def main():
    path = os.path.join(TABLES, 'balance_changes.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb['Changes']
    col = {c.value: k for k, c in enumerate(ws[1])}
    for row in ws.iter_rows(min_row=2):
        if str(row[col['Record']].value) != 'Unique 370 Usurper of Souls':
            continue
        if row[col['Field']].value == 'value':
            row[col['New']].value = VALUES[int(row[col['Level']].value) - 1]
        elif row[col['Field']].value == 'effect id':
            row[col['New']].value = 254
        row[col['Why']].value = WHY
        row[col['State']].value = 'decided'
    wb.save(path)

    path = os.path.join(TABLES, 'balance_skill_text.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb['Skills']
    col = {c.value: k for k, c in enumerate(ws[1])}
    for row in ws.iter_rows(min_row=2):
        if row[col['Filename']].value == 'UniqueSkillData_0370.bin':
            row[col['Balance_Desc']].value = BOX
            row[col['Balance_Name']].value = NAME
            row[col['Note']].value = ('Rows U011 - U020, U051: in the variant the skill is the chance to hold on with 1 HP, '
                                      'with a name and a sentence of its own')
    wb.save(path)

    path = os.path.join(TABLES, 'qol_rules.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb['UniqueSkillLabels']
    for row in ws.iter_rows(min_row=2):
        if row[0].value == 'Small Target':
            row[0].value, row[1].value = NAME, 'Activation rate'
            row[2].value = ('The name the Balance variant gives the unique skill of Lizerietta (balance_skill_text.xlsx): its '
                            'value is the chance of effect 254, rolled when a hit would take her last HP')
    wb.save(path)

    path = os.path.join(TABLES, 'balance_criteria.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb['Criticisms']
    col = {c.value: k for k, c in enumerate(ws[1])}
    for row in ws.iter_rows(min_row=2):
        if row[0].value == 'I11':
            row[col['Proposal']].value = ('U011 - U020, U051: the skill becomes the chance to hold on with 1 HP when a hit '
                                          'would fell her (effect 254, 15 - 40% by rank, 55% with the crown), named again in '
                                          'the variant')
            row[col['Decision']].value = (str(row[col['Decision']].value or '') + ' | 2026-10-12 (night), decided: not '
                                          'evasion (always at work); the chance to survive a killing hit, situational as '
                                          'the others')
    wb.save(path)
    print('written')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    main()
