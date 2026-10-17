"""Criticisms: the rows that cite class-stages get its figures of 2026-10-17 (each role in the reference party and the five battery
parties, 240 fights a cell; docs/generated/CLASS_STAGES.md), for the classes or roles each row is about; F61 (Shadow Sew) gets the
reading of the skill battery."""
import os
import re
import openpyxl

ROOT = os.path.join(os.path.dirname(__file__), '..', '..')
PATH = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
TAG = '2026-10-17, class-stages'

# the stage tables of the report: (stage, role, class, measure, verdict)
rows, stage = [], None
for line in open(os.path.join(ROOT, 'docs', 'generated', 'CLASS_STAGES.md'), encoding='utf-8'):
    m = re.match(r'## (\w+) \(', line)
    if m:
        stage = m.group(1)
        continue
    if stage and line.startswith('| ') and not line.startswith('| Role') and '---' not in line:
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        rows.append((stage, cells[0], cells[1], cells[-2], cells[-1]))

# what each criticism is about: classes, or (role,) for a role
ABOUT = {
    'J06': [('shooter',)], 'J11': [('mage',)], 'J33': [('mage',)], 'F63': ['Kunoichi'], 'F71': [('dolls',)], 'F87': ['Sniper'],
    'F88': ['Kunoichi', 'Mistress'], 'J30': ['Spieler'], 'F17': ['Spieler', 'Trickster', 'Joker'], 'U02': ['Doll Master', 'Papillon', 'Soul Summoner'],
}


def figures(about):
    out = []
    for st, role, cls, measure, verdict in rows:
        hit = any((isinstance(a, tuple) and role == a[0]) or (isinstance(a, str) and cls == a) for a in about)
        if hit:
            out.append('%s %s (%s role): %s, %s' % (st, cls, role, measure, verdict))
    return out


wb = openpyxl.load_workbook(PATH)
ws = wb['Criticisms']
head = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(head)}
done, skipped = [], []
for r in range(2, ws.max_row + 1):
    rid = ws.cell(r, col['ID']).value
    text = ' '.join(str(ws.cell(r, c).value or '') for c in range(1, ws.max_column + 1))
    cell = ws.cell(r, col['Balance_Result'])
    if rid == 'F61':
        note = (' 2026-10-17, dtcli skill-battery with every rule of use: no effect. By the rule of the program (docs/FORMULAS.md, effect '
                '123) it takes only a target whose level is below the user\'s, and 67 of 101 story bosses are immune: a tool against weaker '
                'foes, left as the game has it.')
        if 'skill-battery with every rule of use: no effect' not in str(cell.value):
            cell.value = str(cell.value or '').rstrip() + note
            done.append(rid)
        continue
    if not ('CLASS_STAGES' in text or 'class_stages' in text or 'class-stages' in text) or TAG in str(cell.value):
        continue
    about = ABOUT.get(rid, [ws.cell(r, col['Subject']).value])
    if about is None:
        skipped.append(rid)
        continue
    figs = figures(about)
    if not figs:
        skipped.append(rid)
        continue
    cell.value = (str(cell.value or '').rstrip() + ' ' + TAG + ' (each role in the reference party and the five battery parties, 240 '
                  'fights a cell; docs/generated/CLASS_STAGES.md): ' + '; '.join(figs) + '.')
    done.append(rid)
wb.save(PATH)
print('written', len(done), ' '.join(done), '| no figure:', ' '.join(skipped))
