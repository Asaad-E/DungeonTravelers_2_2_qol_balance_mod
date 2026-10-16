"""Criticisms S09: the bound on stacking All stats, read from the drop code and measured (2026-10-16)."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'balance_criteria.xlsx')

WHY_ADD = (' The stacking cannot be farmed into a party-wide stack: a random drop rolls All stats only on the Ring (piece 710;'
           ' bit 8 of EnchantData 0x20, 0x556bf0), at most one per piece (group 3 at 0x28), from item level 10 / 30 / 50 / 75'
           ' for +3 / +6 / +9 / +12 (0x2C, 0x556cd3; +12 is postgame); the other sources are 15 fixed pieces with +1 to +9,'
           ' one per slot. About +30 on one or two members at most.')
RESULT = ('Upper bound, +30 on all five members of the reference party (impossible in play), 400 runs, median length:'
          ' C34 33.9 -> 16.3, C38 12.3 -> 6.6 (wins 91% -> 100%), C40 29.2 -> 15.5; vanilla with the same pieces counts'
          ' one: C38 18.4 -> 11.7. dtcli run, seed 7000.')

wb = openpyxl.load_workbook(PATH)
ws = wb['Criticisms']
head = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(head)}
for row in range(2, ws.max_row + 1):
    if ws.cell(row, col['ID']).value == 'S09':
        why = ws.cell(row, col['Why']).value or ''
        if WHY_ADD.strip() not in why:
            ws.cell(row, col['Why']).value = why.rstrip() + WHY_ADD
        ws.cell(row, col['Balance_Result']).value = RESULT
        break
else:
    raise SystemExit('S09 not found')
wb.save(PATH)
print('ok')
