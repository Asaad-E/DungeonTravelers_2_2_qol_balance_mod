"""One-off edit: row G119 of balance_changes.xlsx, the ATK of Demon God Fried (C36) x1.2, measured with
c36_atk_trial_2026-10-16.py (beside this file) on 2026-10-16."""
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
ROW = ['G119', 'enemies', '1015 Demon God Fried (C36)', 'atk', None, 265, 318, 'L01', 'R08, R10',
       'Difficulty kept (R08, R10) with attack, not HP: the classes and skills of the variant made C36 easier (wins / length / '
       'pressure of the reference party: vanilla 90% / 17.1 / 0.21, variant 98% / 13.3 / 0.22; Meta 88% > 98%, Mage-Heavy '
       '45% > 80%). ATK x1.20 brings the reference party (92% / 14.3 / 0.25), Meta (90%) and Dual-Support (74%, vanilla 75%) '
       'back to vanilla, inside every window of the curve; Mage-Heavy stays at 70%, the part of its gain that comes from the '
       'magic of the variant against a boss weak to it (ATK x1.40 would take it under +15 points but leave the reference '
       'party at 78% and Dual-Support at 56%, harder than vanilla). Anti-Meta-Plus 48% > 38% (vanilla 0%). 200 battles a '
       'cell (scratch/audit_2026-10-11/c36_atk_trial_2026-10-16.py). The curve of 40 battles puts its pressure on the edge '
       '(0.20 > 0.25); measures/curve_recheck.py C36 400: 89% > 91%, 16.8 > 14.2, pressure 0.21 > 0.25, within; 1,000 '
       'battles: pressure 0.214 > 0.246 (+15%, window 25%)',
       '2026-10-16', 'built']

wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
ids = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
assert ROW[0] not in ids
ws.append(ROW)
wb.save(PATH)
print('added', ROW[0])
