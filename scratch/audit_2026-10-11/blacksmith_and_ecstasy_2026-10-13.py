"""One-off edit (2026-10-13): the blacksmith appears half again as often in every zone that has her (rows Z of table
zones, lever L46, row S28 of the register), and Ecstasy gives twice as much (rows Y of its five levels, row F70).
Run once; kept as the record of the edit.
"""
import os
import struct
import sys
import zipfile

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(ROOT, 'tools'))
import floor_data                                         # noqa: E402
import floor_patch                                        # noqa: E402

CRITERIA = os.path.join(ROOT, 'data_tables', 'balance_criteria.xlsx')
CHANGES = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
FIELD, FACTOR, TODAY = 'blacksmith chance', 1.5, '2026-10-13'
ECSTASY = {'record': 'Passive 280 Ecstasy', 'field': 'effect 312', 'old': [0.1, 0.11, 0.12, 0.13, 0.15],
           'new': [0.2, 0.22, 0.24, 0.26, 0.3]}
S28 = ('S28', 'Play test', 'System', 'Blacksmith', 'The wandering blacksmith shows up too seldom to count on her', 'Tedium', None,
       'Chance that the blacksmith is the next shop to appear in a zone, and the steps it takes',
       'Floor files, zone 0x24 (docs/FORMULAS.md 14)',
       'Her chance a step is 2 to 10 in the 166 zones that have her (5 to 7 in most), rolled after the battle and after the '
       'peddler (3 to 10), only after 5 battles and 50 steps since the last shop, once until the counters of the dungeon are '
       'cleared. In a zone at 5 with a peddler at 6 she is the next shop 44% of the times',
       '166 Z rows of the table zones, her chance x1.5 rounded (3 to 15): in the same zone she is the next shop 56% of the '
       'times. The waits of 5 battles and 50 steps are the program\'s and stay',
       'closed', 'yes, small', 'More often everywhere, and not much: the peddler keeps his turn and the waits stay',
       'T10 Forge, enchantments and sealbooks', 'true', 'Her chance x1.5 in every zone (L46)',
       f'{TODAY}: approved (more often in general, not too high)')
L46 = ('L46', 'Chance a step of a wandering shop of a zone (peddler, blacksmith, ice cream, ramen, dark peddler)',
       'Dungeon/stNN/Data.pck floor file', 'zone 0x20 - 0x30, five ints', 'Benign',
       'A whole number from 1 to 100; a zone with 0 keeps 0. Written in place and put back the same way '
       '(tools/floor_patch.py): the packs are not copied',
       'docs/FORMULAS.md 14: each is rolled against a number below 100 on a step (0x4bc8a1 - 0x4bcb77), as the chance of a '
       'battle (L29)')


def main():
    wb = openpyxl.load_workbook(CRITERIA)
    ws = wb['Criticisms']
    if any(r[0].value == 'S28' for r in ws.iter_rows(min_row=2)):       # the first sheets were written by a run that stopped
        return changes()
    ws.append(S28)
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == 'F70':
            row[head.index('Status')].value = 'closed'
            row[head.index('Balance_Result')].value = ('5 Y rows, ATK and DEX +20% to +30%: the fights of C30 and C40 a little '
                                                       'shorter (measures/one_skill_trial.py)')
            row[head.index('Decision')].value = f'{TODAY}: raised to 30% at its top level'
    levers = wb['Levers']
    at = next(r[0].row for r in levers.iter_rows() if r[0].value == 'L45')
    levers.insert_rows(at + 1)
    for k, value in enumerate(L46):
        levers.cell(at + 1, k + 1).value = value
    wb.save(CRITERIA)
    changes()


def changes():
    wb = openpyxl.load_workbook(CHANGES)
    ws = wb['Changes']
    head = [c.value for c in ws[1]]
    ids = {r[0].value for r in ws.iter_rows(min_row=2)}
    last_y = max(int(str(i)[1:].split('.')[0]) for i in ids if str(i).startswith('Y'))
    for k, (old, new) in enumerate(zip(ECSTASY['old'], ECSTASY['new'])):
        ws.append((f'Y{last_y + 1 + k:03d}', 'skills', ECSTASY['record'], ECSTASY['field'], k + 1, old, new, 'L12', 'F70',
                   'A passive that added nothing that could be measured: ATK and DEX +10% to +15% for 3 turns, once a battle '
                   '(measures/one_skill_trial.py, 150 battles: wins within 2 points, length within 2%). Twice as much',
                   TODAY, 'built'))
    blocks, count = floor_patch.zone_blocks(), 0
    offset = floor_patch.FIELDS[FIELD]
    with zipfile.ZipFile(floor_data.FLOOR_ZIP) as archive:
        for label, (dungeon, entry, block) in sorted(blocks.items()):
            held = struct.unpack_from('<i', archive.read(f'{dungeon}/{entry}'), block + offset)[0]
            if held > 0:
                count += 1
                ws.append((f'Z{count:03d}', 'zones', label, FIELD, None, held, int(held * FACTOR + 0.5), 'L46', 'S28',
                           'The blacksmith more often in every zone that has her, by half (S28): the waits of the program '
                           '(5 battles, 50 steps since the last shop) and the turn of the peddler before her stay',
                           TODAY, 'built'))
    wb.save(CHANGES)
    print(count, 'zones, 5 levels of Ecstasy')


if __name__ == '__main__':
    main()
