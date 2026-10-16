"""The Diva's HP and RES x0.9 instead of x0.85 (rows V493 - V690, decided 2026-10-16): with both at x0.85 the Diva of the reference party
of C40 won 60% of her fights (class-stages, weak; 92% without the change); at x0.9 she wins 88% and her parties 88.8% (+2.2 sd in the
band of the classes, from 91.5%)."""
import os
import openpyxl

PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'data_tables', 'balance_changes.xlsx')
WHY = ('The Diva lowered a little (S07, decided 2026-10-16): first in wins of the final classes (91.5%, +2.4 sd), carried by her own '
       'survival, not by her songs (Carol x0.6, Chorale x0.8 or Self Session without effect changed nothing). HP and RES x0.9: 88.8% '
       '(+2.2 sd, docs/generated/CLASS_BALANCE.md); x0.85 on both took the Diva of the reference party of C40 from 92% to 60% wins '
       '(docs/generated/CLASS_STAGES.md), x0.9 keeps 88%. The best support in wins is her identity.')
wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
done = 0
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 3).value == '27 Diva':
        ws.cell(r, 7).value = round(ws.cell(r, 6).value * 0.9)
        ws.cell(r, 10).value = WHY
        done += 1
assert done == 198, done
wb.save(PATH)
print('Diva rows', done)
