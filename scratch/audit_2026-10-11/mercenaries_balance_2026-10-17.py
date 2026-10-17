import json
import os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PATH = os.path.join(ROOT, 'data_tables', 'balance_changes.xlsx')
DECIDED = '2026-10-17'

PLAN = {
    'Euphemia': {
        'stats': {'atk': 1.25, 'agi': 1.20},
        'why': 'Euphemia had only +0.1% mean wins and 60% serve due to sluggish spd=0.5-0.6 skills; ATK x1.25 and AGI x1.20 give heavy-hitter initiative, lifting serve to 84% and mean wins to +0.6%.'
    },
    'Rebecca': {
        'stats': {'dex': 1.15, 'atk': 1.15},
        'why': 'Rebecca is an exploration thief; DEX x1.15 and ATK x1.15 improve combat presence in dungeon randoms (+28% length reduction) while keeping her steal identity intact.'
    },
    'Riesco': {
        'stats': {'atk': 1.30, 'int': 1.25, 'agi': 1.25},
        'why': 'Riesco was the least-serving story mercenary (44% serve); ATK x1.3, INT x1.25 and AGI x1.25 lift serve to 59% and mean wins to +1.1%, giving frost attacks punch.'
    },
    'Carol': {
        'stats': {'int': 1.30},
        'why': 'Carol keeps her comedic 22-ATK slaps while INT x1.30 makes her elemental spells devastate on impact, raising serve from 60% to 73% and eliminating hurts.'
    },
    'Three Fates': {
        'stats': {'int': 1.25, 'atk': 1.15},
        'why': 'Three Fates keeps the Take Break meme and costly contract; INT x1.25 and ATK x1.15 ensure that when the sisters coordinate, strikes deal divine damage (serve 58%, mean wins +1.7%).'
    },
    'Lolo': {
        'stats': {'agi': 0.80, 'atk': 0.80, 'dex': 0.85},
        'why': 'Lolo had excessive speed (AGI 122-138) and +4.9% mean wins overshadowing other postgame mercenaries; AGI x0.8, ATK x0.8, DEX x0.85 moderates her to +4.0% wins and 58% serve.'
    }
}

npcs = json.load(open(os.path.join(ROOT, 'build', 'engine_snapshot.json'), encoding='utf-8'))['sides']['vanilla']['npcs']

wb = openpyxl.load_workbook(PATH)
ws = wb['Changes']
head = [c.value for c in ws[1]]
assert head[:12] == ['ID', 'Table', 'Record', 'Field', 'Level', 'Vanilla', 'New', 'Lever', 'Serves', 'Why', 'Decided', 'State'], head

# Existing npcs rows
existing = {(ws.cell(r, 3).value, ws.cell(r, 4).value): r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == 'npcs'}
ids = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
nxt = 1 + max(int(str(k)[1:]) for k in ids if k and str(k).startswith('N'))
row = ws.max_row + 1
added = updated = 0

for name, cfg in PLAN.items():
    st_npcs = sorted((x for x in npcs if x['name'] == name), key=lambda x: x['level'])
    for n in st_npcs:
        record = f"{n['id']} {n['name']} (level {n['level']})"
        for stat, factor in cfg['stats'].items():
            vanilla_val = n[stat]
            new_val = round(vanilla_val * factor)
            at = existing.get((record, stat))
            if at:
                ws.cell(at, 7).value = new_val
                ws.cell(at, 10).value = cfg['why']
                ws.cell(at, 11).value = DECIDED
                ws.cell(at, 12).value = 'decided'
                updated += 1
            else:
                change_id = f'N{nxt:03d}'
                ws.append([change_id, 'npcs', record, stat, None, vanilla_val, new_val, 'L44', 'F94', cfg['why'], DECIDED, 'decided'])
                nxt += 1
                added += 1

wb.save(PATH)
print(f'Done: {added} rows added, {updated} rows updated in balance_changes.xlsx')
