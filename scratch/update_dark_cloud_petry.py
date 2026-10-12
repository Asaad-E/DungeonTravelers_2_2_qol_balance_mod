import openpyxl

rules_path = 'mod_project/data_tables/qol_rules.xlsx'
wb = openpyxl.load_workbook(rules_path)
ws_effects = wb['ItemEffects']
ws_short = wb['ShortItems']

# Update Type 56 and Type 17 in ItemEffects
for row in ws_effects.iter_rows(min_row=2):
    if row[0].value == 56:
        row[1].value = '{v}% HP dmg'
    if row[0].value == 17:
        row[1].value = '{v}% petry'

# Update Dark Cloud in ShortItems
found = False
for row in ws_short.iter_rows(min_row=2):
    if row[0].value == 'Dark Cloud':
        row[1].value = 'Roiling mass. Is that a green lady in there?'
        found = True
        break
if not found:
    ws_short.append(['Dark Cloud', 'Roiling mass. Is that a green lady in there?'])

wb.save(rules_path)
print("Updated ItemEffects and ShortItems for Dark Cloud in qol_rules.xlsx")
