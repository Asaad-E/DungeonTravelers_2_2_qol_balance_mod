import openpyxl
import os
import sys

BASE_DIR = os.path.abspath('mod_project')
CRITERIA_PATH = os.path.join(BASE_DIR, 'data_tables', 'balance_criteria.xlsx')
CHANGES_PATH = os.path.join(BASE_DIR, 'data_tables', 'balance_changes.xlsx')

# 1. Update balance_criteria.xlsx
wb_crit = openpyxl.load_workbook(CRITERIA_PATH)
ws_crit = wb_crit['Criticisms']
headers_crit = [c.value for c in ws_crit[1]]
id_col = headers_crit.index('ID') + 1
status_col = headers_crit.index('Status') + 1
verdict_col = headers_crit.index('Verdict') + 1
why_col = headers_crit.index('Why') + 1

crit_updates = {
    # From Step 1 closure:
    'J24': {
        'Status': 'closed',
        'Verdict': 'partly true',
        'Why': 'Addressed by Full Heal (H036) and Circle Full Heal (H037) additions to restore Bishop healing identity.'
    },
    'F87': {
        'Status': 'closed',
        'Verdict': 'true',
        'Why': 'Addressed by W036-W039 making bow combat skills fully compatible with guns.'
    },
    # Milestone 4 - T10:
    'J35': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Why': 'Forging costs scale quadratically as an endgame gold sink; S16 complains of the opposite (forged gear being too strong). Progression functions as intended.'
    },
    'F14': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Why': 'Reaching 100% resistance across all elements and ailments requires committing multiple equipment slots away from offensive/defensive multipliers; it is an intentional player build choice.'
    },
    'F89': {
        'Status': 'closed',
        'Verdict': 'true',
        'Why': 'Underperforming postgame Grand Sealbooks given complementary secondary enchantments (ench_cnt = 2 in SealScrollData.bin) boosting defense, critical rate, and TP sustainability.'
    },
    'F90': {
        'Status': 'closed',
        'Verdict': 'true',
        'Why': 'Rarest equipment drop rates on postgame/boss enemies gently bumped from 1% to 3% and 2%-3% to 5% in EnemyStatusData.bin.'
    },
    'S09': {
        'Status': 'closed',
        'Verdict': 'true',
        'Why': 'All stats enchantments (IDs 24-28) bit 0 enabled at offset 0x20 in EnchantData.pck (word 8 changed from 428 to 429), allowing gear piece stacking.'
    },
    'S16': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Why': 'Contradicts J35. Unique postgame weapons feature distinct innates and higher base stats; forged gear provides customization flexibility without invalidating unique top-tier drops.'
    },
    # Milestone 4 - T13:
    'F91': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Why': 'Elemental zone bonuses apply symmetrically to party and enemies, creating intentional positioning and build puzzles inherent to DT2-2 dungeon exploration.'
    },
    'F96': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Why': 'Exploiting enemy elemental affinities with zone synergy is the primary intended combat puzzle for magical and hybrid party setups.'
    },
    'S13': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Why': 'Encounter rate of 5-8% in story dungeons conforms to dungeon crawler genre norms. Repel items and Grand Sealbook Tres Durian provide encounter mitigation.'
    },
    'S14': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Why': 'Anti-magic, darkness, and damage tiles are core dungeon crawling hazards with dedicated counterplay (Light spell, levitate/cushion items, Grand Sealbook Sahohime).'
    },
    'S15': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Why': 'Scout class tree provides Trap Search and Trap Disarm; Grand Sealbook Deep Slapyon mitigates trap damage by 90%.'
    },
    # Milestone 4 - T14:
    'J36': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Why': 'Meals provide pre-dungeon or mid-dungeon prep; cooking and eating items near boss rooms or utilizing long-duration food recipes addresses boss preparation.'
    },
    'S18': {
        'Status': 'not a problem',
        'Verdict': 'partly true',
        'Why': 'Starting capacity of 30 is hardcoded in the engine executable; Quest Ranks naturally expand it to 50 slots. Preserves expedition resource management without unstable executable patches.'
    },
    'S19': {
        'Status': 'not a problem',
        'Verdict': 'false',
        'Why': 'Early game resource scarcity is intentional survival tension in dungeon crawlers; early story quests and enemy drops provide adequate funds for baseline gear.'
    },
    'S20': {
        'Status': 'closed',
        'Verdict': 'partly true',
        'Why': 'Quest gold rewards increased by +30% to +50% in QuestData, allowing players to purchase better gear and seals, while underleveled party members catch up rapidly via the B4 EXP curve on quest monsters.'
    }
}

updated_crit = 0
for r in range(2, ws_crit.max_row + 1):
    rid = ws_crit.cell(row=r, column=id_col).value
    if rid in crit_updates:
        u = crit_updates[rid]
        ws_crit.cell(row=r, column=status_col).value = u['Status']
        ws_crit.cell(row=r, column=verdict_col).value = u['Verdict']
        ws_crit.cell(row=r, column=why_col).value = u['Why']
        updated_crit += 1

print(f"Updated {updated_crit} criticism rows in {CRITERIA_PATH}")
wb_crit.save(CRITERIA_PATH)

# 2. Add change rows to balance_changes.xlsx
wb_chg = openpyxl.load_workbook(CHANGES_PATH)
ws_chg = wb_chg['Changes']
headers_chg = [c.value for c in ws_chg[1]]

new_rows = [
    # F89: Grand Sealbooks (P001 - P005)
    {
        'ID': 'P001', 'Table': 'sealbooks', 'Record': 'Grand Sealbook [Melek Taus] (0955)',
        'Field': 'secondary enchantment', 'Level': None,
        'Vanilla': 'ench_cnt=1 (222 only)', 'New': 'ench_cnt=2 (222 Front line ATK, 268 Critical UP)',
        'Lever': 'L19', 'Serves': 'F89',
        'Why': 'Postgame Grand Sealbook given complementary offensive secondary enchantment (Critical rate UP) via native engine loop.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P002', 'Table': 'sealbooks', 'Record': 'Grand Sealbook [Samael] (0951)',
        'Field': 'secondary enchantment', 'Level': None,
        'Vanilla': 'ench_cnt=1 (224 only)', 'New': 'ench_cnt=2 (224 Back line INT, 248 Spell TP Down)',
        'Lever': 'L19', 'Serves': 'F89',
        'Why': 'Postgame Grand Sealbook given magic sustain secondary enchantment (TP Cost for spells DOWN) via native engine loop.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P003', 'Table': 'sealbooks', 'Record': 'Grand Sealbook [Sukuna-Hikona] (0775)',
        'Field': 'secondary enchantment', 'Level': None,
        'Vanilla': 'ench_cnt=1 (281 only)', 'New': 'ench_cnt=2 (281 Back line DEX+AGI, 246 Tech TP Down)',
        'Lever': 'L19', 'Serves': 'F89',
        'Why': 'Postgame Grand Sealbook given technique sustain secondary enchantment (TP Cost for techniques DOWN) via native engine loop.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P004', 'Table': 'sealbooks', 'Record': 'Grand Sealbook [Genbu] (0801)',
        'Field': 'secondary enchantment', 'Level': None,
        'Vanilla': 'ench_cnt=1 (223 only)', 'New': 'ench_cnt=2 (223 Front line ATK+DEF, 257 Defend HP recovery)',
        'Lever': 'L19', 'Serves': 'F89',
        'Why': 'Postgame Grand Sealbook given defensive sustain secondary enchantment (HP recovers while defending) via native engine loop.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P005', 'Table': 'sealbooks', 'Record': 'Grand Sealbook [Murasaki Shikibu] (0722)',
        'Field': 'secondary enchantment', 'Level': None,
        'Vanilla': 'ench_cnt=1 (292 only)', 'New': 'ench_cnt=2 (292 Defend TP recovery, 265 Ailment dissipation UP)',
        'Lever': 'L19', 'Serves': 'F89',
        'Why': 'Postgame Grand Sealbook given ailment mitigation secondary enchantment (Rate of status ailments dissipation UP) via native engine loop.',
        'Decided': '2026-10-10', 'State': 'decided'
    },

    # S09: All Stats enchantments stacking (P006 - P010)
    {
        'ID': 'P006', 'Table': 'enchantments', 'Record': 'EnchantData_0024 (All stats +1)',
        'Field': 'stacking flag (0x20 bit 0)', 'Level': None,
        'Vanilla': 428, 'New': 429,
        'Lever': 'L19', 'Serves': 'S09',
        'Why': 'Bit 0 at 0x20 in EnchantData enabled, matching IDs 0-23 so All Stats +1 stacks across multiple equipped pieces.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P007', 'Table': 'enchantments', 'Record': 'EnchantData_0025 (All stats +3)',
        'Field': 'stacking flag (0x20 bit 0)', 'Level': None,
        'Vanilla': 428, 'New': 429,
        'Lever': 'L19', 'Serves': 'S09',
        'Why': 'Bit 0 at 0x20 in EnchantData enabled so All Stats +3 stacks across multiple equipped pieces.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P008', 'Table': 'enchantments', 'Record': 'EnchantData_0026 (All stats +6)',
        'Field': 'stacking flag (0x20 bit 0)', 'Level': None,
        'Vanilla': 428, 'New': 429,
        'Lever': 'L19', 'Serves': 'S09',
        'Why': 'Bit 0 at 0x20 in EnchantData enabled so All Stats +6 stacks across multiple equipped pieces.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P009', 'Table': 'enchantments', 'Record': 'EnchantData_0027 (All stats +9)',
        'Field': 'stacking flag (0x20 bit 0)', 'Level': None,
        'Vanilla': 428, 'New': 429,
        'Lever': 'L19', 'Serves': 'S09',
        'Why': 'Bit 0 at 0x20 in EnchantData enabled so All Stats +9 stacks across multiple equipped pieces.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P010', 'Table': 'enchantments', 'Record': 'EnchantData_0028 (All stats +12)',
        'Field': 'stacking flag (0x20 bit 0)', 'Level': None,
        'Vanilla': 428, 'New': 429,
        'Lever': 'L19', 'Serves': 'S09',
        'Why': 'Bit 0 at 0x20 in EnchantData enabled so All Stats +12 stacks across multiple equipped pieces.',
        'Decided': '2026-10-10', 'State': 'decided'
    },

    # F90: Best-in-slot drop/steal rates (P011 - P025)
    {
        'ID': 'P011', 'Table': 'enemies', 'Record': '331 Overdue (P06)',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 1, 'New': 3,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame boss rarest equipment drop rate gently raised from 1% to 3% to alleviate excessive reload grinding.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P012', 'Table': 'enemies', 'Record': '751 Azathoth (P18)',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 1, 'New': 3,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame boss rarest equipment drop rate gently raised from 1% to 3% to alleviate excessive reload grinding.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P013', 'Table': 'enemies', 'Record': '862 Last Judgment (P24)',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 1, 'New': 3,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame superboss rarest equipment drop rate gently raised from 1% to 3%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P014', 'Table': 'enemies', 'Record': '921 Hell\'s Angel',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 1, 'New': 3,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame superboss rarest equipment drop rate gently raised from 1% to 3%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P015', 'Table': 'enemies', 'Record': '951 Samael',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 1, 'New': 3,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame superboss rarest equipment drop rate gently raised from 1% to 3%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P016', 'Table': 'enemies', 'Record': '955 Melek Taus',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 1, 'New': 3,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame superboss rarest equipment drop rate gently raised from 1% to 3%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P017', 'Table': 'enemies', 'Record': '775 Sukuna-Hikona',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 2, 'New': 5,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame rare equipment drop rate gently raised from 2% to 5%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P018', 'Table': 'enemies', 'Record': '801 Genbu',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 2, 'New': 5,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame rare equipment drop rate gently raised from 2% to 5%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P019', 'Table': 'enemies', 'Record': '712 Iktomi',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 2, 'New': 5,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame rare equipment drop rate gently raised from 2% to 5%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P020', 'Table': 'enemies', 'Record': '722 Murasaki Shikibu',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 2, 'New': 5,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame rare equipment drop rate gently raised from 2% to 5%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P021', 'Table': 'enemies', 'Record': '848 Qitian Dasheng',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 2, 'New': 5,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame rare equipment drop rate gently raised from 2% to 5%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P022', 'Table': 'enemies', 'Record': '887 Nuwa',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 2, 'New': 5,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame rare equipment drop rate gently raised from 2% to 5%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P023', 'Table': 'enemies', 'Record': '890 Kuchinawa',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 2, 'New': 5,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Postgame rare equipment drop rate gently raised from 2% to 5%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P024', 'Table': 'enemies', 'Record': '6 Garasha',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 2, 'New': 5,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Rare equipment drop rate gently raised from 2% to 5%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P025', 'Table': 'enemies', 'Record': '7 Ryoma',
        'Field': 'equipment drop rate', 'Level': None,
        'Vanilla': 2, 'New': 5,
        'Lever': 'L01', 'Serves': 'F90',
        'Why': 'Rare equipment drop rate gently raised from 2% to 5%.',
        'Decided': '2026-10-10', 'State': 'decided'
    },

    # S20: Quest gold rewards (P026 - P030)
    {
        'ID': 'P026', 'Table': 'quests', 'Record': 'QuestData_0000 (Fruit Fighter)',
        'Field': 'gold reward', 'Level': None,
        'Vanilla': 30, 'New': 45,
        'Lever': 'L19', 'Serves': 'S20',
        'Why': 'Early progression quest gold raised by +50% to ease early consumable and gear acquisition.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P027', 'Table': 'quests', 'Record': 'QuestData_0005 (Pest Control)',
        'Field': 'gold reward', 'Level': None,
        'Vanilla': 42, 'New': 60,
        'Lever': 'L19', 'Serves': 'S20',
        'Why': 'Early progression quest gold raised by +43% to ease early consumable and gear acquisition.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P028', 'Table': 'quests', 'Record': 'QuestData_0050 (Last-Minute Souvenir)',
        'Field': 'gold reward', 'Level': None,
        'Vanilla': 350, 'New': 500,
        'Lever': 'L19', 'Serves': 'S20',
        'Why': 'Midgame quest gold raised by +43% to incentivize quest completion for lagging parties.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P029', 'Table': 'quests', 'Record': 'QuestData_0100 (Midgame Hunting Quest)',
        'Field': 'gold reward', 'Level': None,
        'Vanilla': 585, 'New': 850,
        'Lever': 'L19', 'Serves': 'S20',
        'Why': 'Midgame quest gold raised by +45% to incentivize quest completion for lagging parties.',
        'Decided': '2026-10-10', 'State': 'decided'
    },
    {
        'ID': 'P030', 'Table': 'quests', 'Record': 'QuestData_0200 (Late Game Hunting Quest)',
        'Field': 'gold reward', 'Level': None,
        'Vanilla': 599, 'New': 900,
        'Lever': 'L19', 'Serves': 'S20',
        'Why': 'Late game quest gold raised by +50% to incentivize quest completion for lagging parties.',
        'Decided': '2026-10-10', 'State': 'decided'
    }
]

# Check existing IDs in ws_chg
existing_ids = set()
for r in range(2, ws_chg.max_row + 1):
    val = ws_chg.cell(row=r, column=1).value
    if val:
        existing_ids.add(val)

added_chg = 0
for row_dict in new_rows:
    if row_dict['ID'] not in existing_ids:
        row_vals = [row_dict.get(h) for h in headers_chg]
        ws_chg.append(row_vals)
        added_chg += 1

print(f"Added {added_chg} change rows in {CHANGES_PATH}")
wb_chg.save(CHANGES_PATH)
