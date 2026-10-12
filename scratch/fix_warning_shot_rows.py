import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_changes.xlsx')
ws = wb['Changes']

# Delete rows P033 to P037
rows_to_delete = []
for r in range(2, ws.max_row + 1):
    val = ws.cell(row=r, column=1).value
    if val in ['P033', 'P034', 'P035', 'P036', 'P037']:
        rows_to_delete.append(r)

for r in reversed(rows_to_delete):
    ws.delete_rows(r)

wb.save('mod_project/data_tables/balance_changes.xlsx')
print(f"Deleted {len(rows_to_delete)} rows from balance_changes.xlsx")
