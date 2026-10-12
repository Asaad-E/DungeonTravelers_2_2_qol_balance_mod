import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_changes.xlsx', read_only=True)
ws = wb['Changes']
headers = [c for c in next(ws.iter_rows(values_only=True))]
seen_prefixes = set()
for r in ws.iter_rows(values_only=True):
    d = dict(zip(headers, r))
    if not d.get('ID') or d['ID'] == 'ID':
        continue
    pref = d['ID'][0]
    if pref not in seen_prefixes:
        seen_prefixes.add(pref)
        print(f"Prefix {pref}: sample {d['ID']} | Table={d.get('Table')} | Record={d.get('Record')} | Field={d.get('Field')} | Lever={d.get('Lever')} | Serves={d.get('Serves')} | State={d.get('State')}")
