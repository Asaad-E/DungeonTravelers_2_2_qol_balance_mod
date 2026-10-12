import openpyxl

rules_path = 'mod_project/data_tables/qol_rules.xlsx'
wb = openpyxl.load_workbook(rules_path)

# Let's inspect current JointSkillSentences headers and rows
ws = wb['JointSkillSentences']
headers = [c.value for c in ws[1]]
print("Headers:", headers)
print("Count of rows:", ws.max_row)
