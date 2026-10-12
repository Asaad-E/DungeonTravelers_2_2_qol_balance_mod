import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_changes.xlsx')
ws = wb['Changes']
header = [c.value for c in ws[1]]

id_idx = header.index('ID')
rec_idx = header.index('Record')
field_idx = header.index('Field')
lvl_idx = header.index('Level')
van_idx = header.index('Vanilla')
new_idx = header.index('New')
why_idx = header.index('Why')
state_idx = header.index('State')

# 1. Update S003 - S012 (Checkmate 0.40 to 1.20)
cm_new_vals = {
    1: 0.40,
    2: 0.48,
    3: 0.56,
    4: 0.64,
    5: 0.72,
    6: 0.80,
    7: 0.90,
    8: 1.00,
    9: 1.10,
    10: 1.20
}

# 2. Update T046 - T050 (Pursuit Stab 1.00 to 1.55)
ps_new_vals = {
    1: 1.00,
    2: 1.15,
    3: 1.28,
    4: 1.42,
    5: 1.55
}

modified_count = 0
dropped_count = 0
rounded_count = 0

for row in ws.iter_rows(min_row=2):
    rid = row[id_idx].value
    if not rid: continue
    
    # S003 - S012
    if rid in ('S003', 'S004', 'S005', 'S006', 'S007', 'S008', 'S009', 'S010', 'S011', 'S012'):
        lvl = int(row[lvl_idx].value)
        row[new_idx].value = cm_new_vals[lvl]
        row[why_idx].value = f"Checkmate: 0.40 to 1.20 cleanly rounded for UI/balance (level {lvl} = {cm_new_vals[lvl]})."
        row[state_idx].value = 'decided'
        modified_count += 1
        
    # T031 - T040 (Drop second pass)
    elif rid in ('T031', 'T032', 'T033', 'T034', 'T035', 'T036', 'T037', 'T038', 'T039', 'T040'):
        row[state_idx].value = 'dropped'
        row[why_idx].value = "Dropped in favor of direct moderated values S003-S012 (0.40 to 1.20) to avoid double-buff stacking."
        dropped_count += 1
        
    # T046 - T050
    elif rid in ('T046', 'T047', 'T048', 'T049', 'T050'):
        lvl = int(row[lvl_idx].value)
        row[new_idx].value = ps_new_vals[lvl]
        row[why_idx].value = f"Pursuit Stab: power 1.00 to 1.55 (+35% to +43%), -25% TP cost kept (level {lvl} = {ps_new_vals[lvl]})."
        row[state_idx].value = 'decided'
        modified_count += 1
        
    # Other active rows with ugly floats (round to 2 decimals)
    elif row[state_idx].value != 'dropped' and isinstance(row[new_idx].value, float):
        val = row[new_idx].value
        s = str(val)
        if '.' in s and len(s.split('.')[1]) > 2:
            rounded_val = round(val, 2)
            row[new_idx].value = rounded_val
            rounded_count += 1

wb.save('mod_project/data_tables/balance_changes.xlsx')
print(f"Applied changes: {modified_count} skill levels updated, {dropped_count} double-buff rows dropped, {rounded_count} ugly floats rounded.")
