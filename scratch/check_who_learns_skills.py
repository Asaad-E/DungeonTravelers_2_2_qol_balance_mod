import sys, os
sys.path.insert(0, os.path.abspath('mod_project/tools'))
import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/vanilla_source_of_truth.xlsx', data_only=True)
ws_skills = wb['Skills']

# Find Warning Shot (270)
for r in ws_skills.iter_rows(min_row=2, values_only=True):
    if r[2] == 270:
        print("Warning Shot row:", r)

# Also check who learns First Grid (309)
for r in ws_skills.iter_rows(min_row=2, values_only=True):
    if r[2] == 309:
        print("First Grid row:", r)
