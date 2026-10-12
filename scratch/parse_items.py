# mod_project/scratch/parse_items.py
import openpyxl, re

wb = openpyxl.load_workbook('mod_project/data_tables/items.xlsx')
ws = wb['Items']

def get_concise_data(mod_str):
    m = mod_str.strip()
    
    # SP
    sp_match = re.search(r'Permanently grant \+(\d+) Skill Points', m)
    if sp_match:
        return f"+{sp_match.group(1)} SP"
        
    # Full restore
    if "Fully restore an ally's HP and TP" in m or "Fully restore all allies' HP and TP" in m:
        return "Full HP/TP"
    if "Fully restore an ally's HP" in m:
        return "Full HP"
    if "Fully restore an ally's TP" in m:
        return "Full TP"

    # Revive
    rev_match = re.search(r'Revive a fallen ally with (\d+)% HP', m)
    if rev_match:
        return f"Revive {rev_match.group(1)}% HP"
    if "Revive a fallen ally" in m:
        return "Revive"

    # Buffs
    if "Physical ATK +" in m and "Speed +" in m:
        atk = re.search(r'ATK \+(\d+)%', m).group(1)
        spd = re.search(r'Speed \+(\d+)%', m).group(1)
        return f"+{atk}% ATK, +{spd}% SPD"
    if "Physical ATK +" in m:
        atk = re.search(r'ATK \+(\d+)%', m).group(1)
        return f"+{atk}% ATK"
    if "Physical DEF +" in m:
        df = re.search(r'DEF \+(\d+)%', m).group(1)
        return f"+{df}% DEF"
    if "Evasion +" in m:
        ev = re.search(r'Evasion \+(\d+)%', m).group(1)
        return f"+{ev}% EVA"

    # Escape / Light
    esc_match = re.search(r'Illuminates dark dungeon floors for (\d+) steps', m)
    if esc_match:
        return f"Light {esc_match.group(1)} steps"
    if "escape from the dungeon" in m:
        return "Escape Dungeon"

    # Debuffs / Cure
    parts = []
    is_party = "all allies" in m.lower()
    
    # HP / TP
    hp_tp = re.search(r'HP and TP by ([\d,]+)', m)
    hp_val = re.search(r'HP by ([\d,]+)', m)
    tp_val = re.search(r'TP by ([\d,]+)', m)
    
    if hp_tp:
        val = hp_tp.group(1).replace(',', '')
        parts.append(f"+{val} HP/TP")
    else:
        if hp_val:
            val = hp_val.group(1).replace(',', '')
            parts.append(f"+{val} HP")
        if tp_val:
            val = tp_val.group(1).replace(',', '')
            parts.append(f"+{val} TP")

    # Ailment cure
    if "Cure all status ailments and stat debuffs" in m:
        parts.append("Cure Ailments & Debuffs")
    elif "Cure all status ailments" in m:
        parts.append("Cure All Ailments")
    elif "Cure 1 status ailment" in m:
        parts.append("Cure 1 Ailment")
    elif "Cure Shadow Sew and remove skill bind" in m:
        parts.append("Cure Shadow Sew & Bind")
    elif "Cure Shadow Sew" in m:
        parts.append("Cure Shadow Sew")
    elif "Cure Freeze" in m:
        parts.append("Cure Freeze")
    elif "Cure Sleep" in m:
        parts.append("Cure Sleep")
    elif "Cure Poison" in m:
        parts.append("Cure Poison")
    elif "Restore all debuffed stats" in m or "Restore debuffed stats" in m or "clear stat debuffs" in m:
        parts.append("Clear Debuffs")

    # Negative effects / Inflictions
    death = re.search(r'(?:Inflict Instant Death|risk of Instant Death) \(?(\d+)%', m)
    if death:
        parts.append(f"{death.group(1)}% Death")
    pois = re.search(r'(?:Inflict Poison|inflict Poison) \(?(\d+)%', m)
    if pois:
        parts.append(f"{pois.group(1)}% Poison")
    para = re.search(r'(?:Inflict Paralysis|inflict Paralysis) \(?(\d+)%', m)
    if para:
        parts.append(f"{para.group(1)}% Paralysis")
    slp = re.search(r'(?:Inflict Sleep|inflict Sleep) \(?(\d+)%', m)
    if slp:
        parts.append(f"{slp.group(1)}% Sleep")
    sil = re.search(r'(?:Inflict Silence|inflict Silence) \(?(\d+)%', m)
    if sil:
        parts.append(f"{sil.group(1)}% Silence")
    blnd = re.search(r'(?:Inflict Blind|inflict Blind) \(?(\d+)%', m)
    if blnd:
        parts.append(f"{blnd.group(1)}% Blind")
    frz = re.search(r'(?:Inflict Freeze|inflict Freeze) \(?(\d+)%', m)
    if frz:
        parts.append(f"{frz.group(1)}% Freeze")
    if "Deal 10 fixed damage" in m:
        parts.append("10 Fixed Dmg")
    if "Deal explosive damage" in m:
        parts.append("Explosive Dmg")

    res = ", ".join(parts)
    if is_party and res:
        res = "Party: " + res
    return res if res else m

# Test on all rows
results = []
for r in range(2, ws.max_row + 1):
    status = ws.cell(row=r, column=7).value
    mod = ws.cell(row=r, column=6).value
    orig = ws.cell(row=r, column=5).value
    name = ws.cell(row=r, column=3).value
    if status in ['Approved', 'Applied', 'Draft'] and mod:
        data = get_concise_data(mod)
        # Clean OG
        orig_clean = orig.strip() if orig else ""
        if orig_clean.endswith('.'):
            combined_1line = f"{orig_clean[:-1]} ({data})."
        else:
            combined_1line = f"{orig_clean} ({data})."
        results.append((r, name, orig_clean, mod, data, combined_1line))

print(f"Total processed: {len(results)}")
over_85 = [x for x in results if len(x[5]) > 85]
print(f"Items exceeding 85 chars on single line: {len(over_85)}")
for r, n, o, m, d, c in over_85[:10]:
    print(f"Row {r:03d} [{n}] len={len(c)}")
    print(f"   OG  : {o}")
    print(f"   DATA: {d}")
    print(f"   COMB: {c}")
