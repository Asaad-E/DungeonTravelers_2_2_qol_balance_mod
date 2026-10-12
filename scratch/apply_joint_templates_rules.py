import openpyxl
import os

rules_path = 'mod_project/data_tables/qol_rules.xlsx'
wb = openpyxl.load_workbook(rules_path)

# 1. Update JointSkillSentences in qol_rules.xlsx
from test_joint_templates_resolved import joint_templates

ws_joint = wb['JointSkillSentences']
# Clear existing rows (keep header)
header = [c.value for c in ws_joint[1]]
ws_joint.delete_rows(2, ws_joint.max_row)

for skill_name, sentence in joint_templates.items():
    if skill_name != "Genocide Rain's Action":
        ws_joint.append([skill_name, sentence, 'QoL with TP and Pep costs'])

wb.save(rules_path)
print(f"Updated {ws_joint.max_row - 1} JointSkillSentences in qol_rules.xlsx")
