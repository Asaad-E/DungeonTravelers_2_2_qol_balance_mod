import openpyxl
from collections import Counter

wb = openpyxl.load_workbook('mod_project/data_tables/balance_changes.xlsx', read_only=True)
ws = wb['Changes']
headers = [c for c in next(ws.iter_rows(values_only=True))]
rows = list(ws.iter_rows(values_only=True))
prefixes = Counter(r[0][0] for r in rows[1:] if r and r[0])
print("Prefixes in Changes:", prefixes)

# Check highest ID for each prefix
max_ids = {}
for r in rows[1:]:
    if r and r[0]:
        pref = r[0][0]
        try:
            num = int(r[0][1:])
            max_ids[pref] = max(max_ids.get(pref, 0), num)
        except ValueError:
            pass
print("Max IDs:", max_ids)
