import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', read_only=True)
ws = wb['Criticisms']
headers = [c for c in next(ws.iter_rows(values_only=True))]
target_ids = ['J35', 'F89', 'F90', 'S09', 'S16', 'F14', 'F91', 'F96', 'S13', 'S14', 'S15', 'J36', 'S18', 'S19', 'S20']
for r in ws.iter_rows(values_only=True):
    d = dict(zip(headers, r))
    if d.get('ID') in target_ids:
        print(f"\nID={d.get('ID')} | Subject={d.get('Subject')}")
        print(f"  Theme: {d.get('Theme')}")
        print(f"  Claim: {d.get('Claim')}")
        print(f"  Status: {d.get('Status')}")
        print(f"  Verdict: {d.get('Verdict')}")
        print(f"  Why: {d.get('Why')}")
