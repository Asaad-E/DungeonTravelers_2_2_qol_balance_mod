import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', read_only=True)
ws = wb['Criticisms']
headers = [c for c in next(ws.iter_rows(values_only=True))]
open_with_verdict = []
open_without_verdict = []
for r in ws.iter_rows(values_only=True):
    d = dict(zip(headers, r))
    if not d.get('ID'): continue
    if d.get('Status') == 'open':
        if d.get('Verdict'):
            open_with_verdict.append((d['ID'], d['Theme'], d['Subject'], d['Verdict']))
        else:
            open_without_verdict.append((d['ID'], d['Theme'], d['Subject']))

print(f"Open WITH verdict: {len(open_with_verdict)}")
print(f"Open WITHOUT verdict: {len(open_without_verdict)}")
if open_without_verdict:
    print("Open without verdict:", open_without_verdict)
