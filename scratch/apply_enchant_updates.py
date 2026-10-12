import openpyxl

RULES_PATH = 'mod_project/data_tables/qol_rules.xlsx'

wb = openpyxl.load_workbook(RULES_PATH)

# 1. Update EnchantSources
ws_sources = wb['EnchantSources']
sources = [r[0] for r in ws_sources.iter_rows(values_only=True)]
if 'exe-verified' not in sources:
    ws_sources.append(['exe-verified', 'Directly decompiled from DT2-2 unwrapped x86 executable', 'yes'])
    print('Added exe-verified to EnchantSources')

# 2. Update EnchantValues
ws_values = wb['EnchantValues']

# Map of existing rows by ID
row_idx_by_id = {}
for idx, r in enumerate(ws_values.iter_rows(values_only=True), start=1):
    if r[0] is not None and str(r[0]).strip() != '' and r[0] != 'Enchant_ID':
        row_idx_by_id[int(r[0])] = idx

# 43 Updated entries (DT1 rows verified in DT2-2 binary)
updates = {
    2: ('-10%', None, 'exe-verified', 'x86: 0x4f83b6 (-10% Max HP)'),
    29: ('+4%', None, 'exe-verified', 'x86: 0x44b69a (mulss 4.0)'),
    30: ('+2%', None, 'exe-verified', 'x86: 0x44b6d4 (addss xmm0, xmm0)'),
    37: ('+4%', None, 'exe-verified', 'x86: 0x44df30 (mulss 4.0)'),
    38: ('+2%', None, 'exe-verified', 'x86: 0x44df6a (addss xmm0, xmm0)'),
    49: ('+10', None, 'exe-verified', 'x86: 0x44a736 (flat +10 enmity points)'),
    50: ('-5', None, 'exe-verified', 'x86: 0x44a74d (flat -5 enmity points)'),
    51: ('+10%', None, 'exe-verified', 'x86: 0x44a114 (mulss 0.10)'),
    52: ('+5%', None, 'exe-verified', 'x86: 0x44a14d (mulss 0.05)'),
    54: ('+20%', None, 'exe-verified', 'x86: 0x449f6a (mulss 0.20)'),
    55: ('+10%', None, 'exe-verified', 'x86: 0x449fa3 (mulss 0.10)'),
    58: ('-6%', None, 'exe-verified', 'x86: 0x44d0bd (mulss 0.06)'),
    59: ('-3%', None, 'exe-verified', 'x86: 0x44d0f4 (mulss 0.03)'),
    60: ('+6%', None, 'exe-verified', 'x86: 0x44d12b (mulss 0.06)'),
    72: (None, 'Damage to TP: 5%', 'exe-verified', 'x86: 0x47e68d (mulss 0.05)'),
    78: ('+20%', None, 'exe-verified', 'x86: 0x44e04b (mulss 0.20)'),
    79: ('+20%', None, 'exe-verified', 'x86: 0x44e101 (mulss 0.20)'),
    80: ('+20%', None, 'exe-verified', 'x86: 0x44e1b7 (mulss 0.20)'),
    81: ('+20%', None, 'exe-verified', 'x86: 0x44e26d (mulss 0.20)'),
    82: ('+20%', None, 'exe-verified', 'x86: 0x44e37d (mulss 0.20)'),
    119: ('-10', None, 'exe-verified', 'x86: 0x4f8b86 (-10 flat resistance)'),
    161: (None, 'Attacks absorb 50% HP', 'exe-verified', 'x86: 0x42156f (movss 0.50, no undead)'),
    162: (None, 'Attacks absorb 15% HP', 'exe-verified', 'x86: 0x4215c4 (movss 0.15, no undead)'),
    163: (None, 'Attacks absorb 20% TP', 'exe-verified', 'x86: 0x421619 (movss 0.20, no undead)'),
    164: (None, 'Attacks absorb 5% TP', 'exe-verified', 'x86: 0x42166e (movss 0.05, no undead)'),
    170: ('+30%', None, 'exe-verified', 'x86: 0x44bfbb (mulss 0.30)'),
    171: ('+15%', None, 'exe-verified', 'x86: 0x44bff5 (mulss 0.15)'),
    172: ('+30%', None, 'exe-verified', 'x86: 0x44c06f (mulss 0.30)'),
    173: ('+15%', None, 'exe-verified', 'x86: 0x44c01d (mulss 0.15)'),
    174: ('+30%', None, 'exe-verified', 'x86: 0x44c0ce (mulss 0.30)'),
    175: ('+15%', None, 'exe-verified', 'x86: 0x44c01d (mulss 0.15)'),
    176: ('+60%', None, 'exe-verified', 'x86: 0x44c131 (mulss 0.60)'),
    177: ('+30%', None, 'exe-verified', 'x86: 0x44c16b (mulss 0.30)'),
    178: ('+15%', None, 'exe-verified', 'x86: 0x44c1a5 (mulss 0.15)'),
    179: ('+30%', None, 'exe-verified', 'x86: 0x44c1e5 (mulss 0.30)'),
    180: ('+15%', None, 'exe-verified', 'x86: 0x44c01d (mulss 0.15)'),
    181: ('+30%', None, 'exe-verified', 'x86: 0x44c244 (mulss 0.30)'),
    182: ('+15%', None, 'exe-verified', 'x86: 0x44c01d (mulss 0.15)'),
    183: ('+30%', None, 'exe-verified', 'x86: 0x44c2a3 (mulss 0.30)'),
    184: ('+15%', None, 'exe-verified', 'x86: 0x44c01d (mulss 0.15)'),
    185: ('+100%', None, 'exe-verified', 'x86: 0x44c32a (+1.0)'),
    186: ('+50%', None, 'exe-verified', 'x86: 0x44c352 (mulss 0.50)'),
    196: ('+100%', None, 'exe-verified', 'x86: 0x44c3ab (+1.0)'),
}

for eid, (fig, name, src, note) in updates.items():
    if eid in row_idx_by_id:
        idx = row_idx_by_id[eid]
        ws_values.cell(row=idx, column=2, value=fig)
        ws_values.cell(row=idx, column=3, value=name)
        ws_values.cell(row=idx, column=4, value=src)
        ws_values.cell(row=idx, column=5, value=note)
        print(f'Updated row {idx} for ID {eid}')
    else:
        print(f'WARNING: ID {eid} not found in existing rows!')

# 59 New entries
new_entries = [
    (31, '+4%', None, 'exe-verified', 'x86: 0x44ad18 (mulss 4.0)'),
    (32, '+2%', None, 'exe-verified', 'x86: 0x44ad52 (addss xmm0, xmm0)'),
    (33, '+4%', None, 'exe-verified', 'x86: 0x44b8eb (mulss 4.0)'),
    (34, '+2%', None, 'exe-verified', 'x86: 0x44b925 (addss xmm0, xmm0)'),
    (35, '+4%', None, 'exe-verified', 'x86: 0x44ad91 (mulss 4.0)'),
    (36, '+2%', None, 'exe-verified', 'x86: 0x44adcb (addss xmm0, xmm0)'),
    (39, '+3%', None, 'exe-verified', 'x86: 0x44d67b (mulss 3.0)'),
    (40, '+1%', None, 'exe-verified', 'x86: 0x44d6bd (addss 1.0 flat)'),
    (41, '+4%', None, 'exe-verified', 'x86: 0x44e5b1 (mulss 4.0)'),
    (42, '+2%', None, 'exe-verified', 'x86: 0x44e637 (addss xmm1, xmm1)'),
    (43, '+3%', None, 'exe-verified', 'x86: 0x44d6f0 (mulss 3.0)'),
    (44, '+1%', None, 'exe-verified', 'x86: 0x44d742 (addss 1.0 flat)'),
    (45, '+4%', None, 'exe-verified', 'x86: 0x44f357 (mulss 4.0)'),
    (46, '+1%', None, 'exe-verified', 'x86: 0x44f3f7 (addss 1.0 flat)'),
    (47, '-4%', None, 'exe-verified', 'x86: 0x44f3c0 (mulss 4.0)'),
    (48, '-2%', None, 'exe-verified', 'x86: 0x44f43f (addss xmm0, xmm0)'),
    (53, '-10%', None, 'exe-verified', 'x86: 0x44a185 (mulss 0.10)'),
    (56, '-20%', None, 'exe-verified', 'x86: 0x449fdb (mulss 0.20)'),
    (57, '-10%', None, 'exe-verified', 'x86: 0x44d083 (mulss 0.10)'),
    (61, '-20%', None, 'exe-verified', 'x86: 0x44d166 (mulss 0.20)'),
    (73, None, 'Damage to TP: 10%', 'exe-verified', 'x86: 0x47e6be (mulss 0.10)'),
    (203, '+60%', None, 'exe-verified', 'x86: 0x44c04d (mulss 0.60)'),
    (204, '+60%', None, 'exe-verified', 'x86: 0x44c02f (mulss 0.60)'),
    (205, '+60%', None, 'exe-verified', 'x86: 0x44c02f (mulss 0.60)'),
    (206, '+60%', None, 'exe-verified', 'x86: 0x44c02f (mulss 0.60)'),
    (207, '+60%', None, 'exe-verified', 'x86: 0x44c02f (mulss 0.60)'),
    (208, '+60%', None, 'exe-verified', 'x86: 0x44c02f (mulss 0.60)'),
    (209, '+100%', None, 'exe-verified', 'x86: 0x44c3ab (+1.0)'),
    (211, None, 'Encounter rate -1%', 'exe-verified', 'x86: 0x4bc812 (-1 flat)'),
    (212, None, 'Encounter rate -2%', 'exe-verified', 'x86: 0x4bc823 (-2 flat)'),
    (213, None, 'Encounter rate +1%', 'exe-verified', 'x86: 0x4bc834 (+1 flat)'),
    (214, None, 'Encounter rate +2%', 'exe-verified', 'x86: 0x4bc843 (+2 flat)'),
    (219, '+10%', None, 'exe-verified', 'x86: 0x44ca0b (mulss 0.10)'),
    (220, '+10%', None, 'exe-verified', 'x86: 0x44ca60 (mulss 0.10)'),
    (246, '-20%', None, 'exe-verified', 'x86: 0x4fd2a7 (mulss 0.80)'),
    (247, '-40%', None, 'exe-verified', 'x86: 0x4fd31e (mulss 0.60)'),
    (248, '-20%', None, 'exe-verified', 'x86: 0x4fd2f8 (mulss 0.80)'),
    (249, '-40%', None, 'exe-verified', 'x86: 0x4fd31e (mulss 0.60)'),
    (251, '+20%', None, 'exe-verified', 'x86: 0x44c485 (mulss 0.20)'),
    (252, '+20%', None, 'exe-verified', 'x86: 0x44c485 (mulss 0.20)'),
    (253, '+20%', None, 'exe-verified', 'x86: 0x44c485 (mulss 0.20)'),
    (254, '+20%', None, 'exe-verified', 'x86: 0x44c485 (mulss 0.20)'),
    (255, '+20%', None, 'exe-verified', 'x86: 0x44c485 (mulss 0.20)'),
    (259, None, 'Shop encounters: steps -50%', 'exe-verified', 'x86: 0x5563b3 (steps counter 50->25)'),
    (261, '+10%', None, 'exe-verified', 'x86: 0x44c8b4 (mulss 0.10)'),
    (265, None, 'Ailment recovery: 2 turns', 'exe-verified', 'x86: 0x48d202 (ailment threshold 2 turns)'),
    (268, None, 'Min Critical rate: 20%', 'exe-verified', 'x86: 0x44f4c3 (crit floor 20%)'),
    (272, None, 'Missed spell penalty: 25%', 'exe-verified', 'x86: 0x44ee4e (75% dmg on miss)'),
    (275, None, 'Double EXP (+100%)', 'exe-verified', 'x86: 0x445a1e (x2.0 EXP)'),
    (276, '+20%', None, 'exe-verified', 'x86: 0x558255 (+20 drop rate)'),
    (280, '+15%', None, 'exe-verified', 'x86: 0x44cd83 (mulss 0.15)'),
    (290, '+20%', None, 'exe-verified', 'x86: 0x44c8ee (mulss 0.20)'),
    (291, None, 'Mimic chest rate +20%', 'exe-verified', 'x86: 0x55577a (+20% mimic chance)'),
    (296, None, '4% chance: full HP on turn', 'exe-verified', 'x86: 0x437d1a (4% full HP on turn)'),
    (298, None, 'Chant speed +25%', 'exe-verified', 'x86: 0x44a380 (+25% chant speed)'),
    (302, '+2%', None, 'exe-verified', 'x86: 0x48c04a (mulss 0.02)'),
    (303, '+1%', None, 'exe-verified', 'x86: 0x48c07e (mulss 0.01)'),
    (304, '+2%', None, 'exe-verified', 'x86: 0x44b533 (mulss 0.02)'),
    (305, '+1%', None, 'exe-verified', 'x86: 0x44b58e (mulss 0.01)'),
]

for eid, fig, name, src, note in new_entries:
    if eid in row_idx_by_id:
        print(f'ERROR: ID {eid} already in sheet!')
    else:
        ws_values.append([eid, fig, name, src, note])
        print(f'Appended ID {eid}')

wb.save(RULES_PATH)
print('Successfully saved qol_rules.xlsx')
