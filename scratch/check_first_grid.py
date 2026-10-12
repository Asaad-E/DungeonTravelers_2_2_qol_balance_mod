import sys, os, struct
sys.path.insert(0, os.path.abspath('mod_project/tools'))
import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/vanilla_source_of_truth.xlsx', data_only=True)
ws = wb['Skills']
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[2] == 309 or r[3] == 'First Grid' or r[2] == 270 or r[3] == 'Warning Shot':
        print("Skill:", r[:8])

ws_lvl = wb['SkillLevels']
for r in ws_lvl.iter_rows(min_row=2, values_only=True):
    if r[1] in [270, 309]:
        print(f"Level row: ID={r[1]}, Level={r[2]}, TP={r[3]}, Effects={r[4:10]}")
