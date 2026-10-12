import openpyxl

wb = openpyxl.load_workbook('mod_project/data_tables/balance_criteria.xlsx', read_only=True)
ws = wb['Criticisms']
headers = [c for c in next(ws.iter_rows(values_only=True))]
open_rows = []
for row in ws.iter_rows(values_only=True):
    d = dict(zip(headers, row))
    if d.get('Status') == 'open':
        open_rows.append(d)

print(f"Total open rows: {len(open_rows)}")
for r in open_rows:
    print(f"ID={r.get('ID')} | Theme={r.get('Theme')} | Subject={r.get('Subject')} | Verdict={r.get('Verdict')}")
