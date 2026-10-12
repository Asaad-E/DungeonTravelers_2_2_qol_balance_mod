import glob
import os
import openpyxl

terms = ['patch', 'mod ', 'curve', 'calibrat', 'rebalance', 'vanilla', 'nerf', 'overhaul', 'our ']

print("--- SCANNING PYTHON TOOLS ---")
for py in glob.glob('mod_project/tools/*.py'):
    with open(py, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    for idx, l in enumerate(lines):
        for t in terms:
            if t in l.lower() and ('"' in l or "'" in l) and not l.strip().startswith('#'):
                if any(k in py for k in ['apply_mod', 'pck_handler']):
                    print(f'{py}:{idx+1}: [{t}] {l.strip()[:100]}')

print("\n--- SCANNING EXCEL DATA TABLES ---")
for x in glob.glob('mod_project/data/*.xlsx'):
    try:
        wb = openpyxl.load_workbook(x, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for r in ws.iter_rows(values_only=True):
                for cell in r:
                    if isinstance(cell, str):
                        for t in terms:
                            if t in cell.lower():
                                print(f'{os.path.basename(x)} [{sname}]: [{t}] {cell[:100]}')
    except Exception as e:
        print(f"Error {x}: {e}")
