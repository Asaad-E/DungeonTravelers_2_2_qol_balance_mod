"""
The choices of the model, read from data_tables/sim_setup.xlsx (see its sheet About). Nothing here holds a value.
"""
import os

import openpyxl

from gamedata.data import PROJECT_DIR

SETUP = os.path.join(PROJECT_DIR, 'data_tables', 'sim_setup.xlsx')


def _rows(wb, name):
    it = wb[name].iter_rows(values_only=True)
    header = next(it)
    return [dict(zip(header, r)) for r in it if any(v is not None for v in r)]


class Setup:
    def __init__(self, path=SETUP):
        wb = openpyxl.load_workbook(path, read_only=True)
        self.settings = {r['Key']: r['Value'] for r in _rows(wb, 'Settings')}
        self.policies = _rows(wb, 'Policies')
        self.roles = {r['Role']: r for r in _rows(wb, 'Roles')}
        self.composition = [r['Role'] for r in sorted(_rows(wb, 'Composition'), key=lambda r: r['Order'])]
        self.checkpoints = _rows(wb, 'Checkpoints')
        self.joins = [(r['Chapter'], [n.strip() for n in str(r['Characters who join in it']).split(';')])
                      for r in _rows(wb, 'Joins')]
        self.assumptions = _rows(wb, 'Assumptions')
        self.calibration = _rows(wb, 'Calibration')
        self.audit_readings = _rows(wb, 'AuditReadings') if 'AuditReadings' in wb.sheetnames else []
        self.grand_sealbooks = _rows(wb, 'GrandSealbooks') if 'GrandSealbooks' in wb.sheetnames else []
        self.zone_changes = _rows(wb, 'ZoneChanges') if 'ZoneChanges' in wb.sheetnames else []
        self.unique_gear = _rows(wb, 'UniqueGear') if 'UniqueGear' in wb.sheetnames else []
        self.personal_sealbooks = _rows(wb, 'PersonalSealbooks') if 'PersonalSealbooks' in wb.sheetnames else []
        self.kit = _rows(wb, 'Kit') if 'Kit' in wb.sheetnames else []
        self.preparation = _rows(wb, 'Preparation') if 'Preparation' in wb.sheetnames else []
        self.enchant_plan = _rows(wb, 'EnchantPlan') if 'EnchantPlan' in wb.sheetnames else []
        self.enchant_slots = _rows(wb, 'EnchantSlots') if 'EnchantSlots' in wb.sheetnames else []
        self.guide_levels = {r['Checkpoint']: r for r in _rows(wb, 'GuideLevels')}
        self.anchors = _rows(wb, 'Anchors') if 'Anchors' in wb.sheetnames else []
        # the class the guide gives each boss (easy, normal, hard): its floor of wins is Settings guide_<class>_min_won
        self.guide_difficulty = {r['Checkpoint']: r['Class'] for r in
                                 (_rows(wb, 'GuideDifficulty') if 'GuideDifficulty' in wb.sheetnames else [])}
        # the parties of the battery: {party: (archetype, [roles])} (tools/party_battery.py, tools/boss_fit.py gentle)
        self.battery = {r['Party']: (r['Archetype'], [x.strip() for x in str(r[next(k for k in r if str(k).startswith('Roles'))]).split(';')])
                        for r in (_rows(wb, 'BatteryParties') if 'BatteryParties' in wb.sheetnames else [])}

    def checkpoint(self, cid):
        return next(c for c in self.checkpoints if c['ID'] == cid)


_cache = {}


def load(path=SETUP):
    stamp = os.path.getmtime(path)
    if _cache.get('stamp') != stamp:
        _cache['setup'], _cache['stamp'] = Setup(path), stamp
    return _cache['setup']
